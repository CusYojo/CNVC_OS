import sys, os
sys.path.insert(0, os.path.abspath('tmp/pydeps'))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

src='尽调报告/智灵动力_投资尽职调查报告_20260806.md'
out='尽调报告/智灵动力_投资尽职调查报告_20260806.docx'
lines=open(src,encoding='utf-8').read().splitlines()
doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(2.3); sec.right_margin=Cm(2.3)
styles=doc.styles
for s in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    st=styles[s]
    st.font.name='Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'微软雅黑')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Normal'].paragraph_format.line_spacing=1.25
styles['Heading 1'].font.size=Pt(16); styles['Heading 1'].font.color.rgb=RGBColor(31,78,121)
styles['Heading 2'].font.size=Pt(13); styles['Heading 2'].font.color.rgb=RGBColor(47,84,150)
styles['Heading 3'].font.size=Pt(11.5); styles['Heading 3'].font.color.rgb=RGBColor(68,68,68)

# header/footer
header=sec.header.paragraphs[0]; header.text='浙江赛智伯乐股权投资管理有限公司｜内部投研材料'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(110,110,110)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('智灵动力投资尽职调查报告  ·  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

# cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(90)
r=p.add_run('智灵动力'); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor(31,78,121); r.font.name='微软雅黑'
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('投资尽职调查报告'); r.bold=True; r.font.size=Pt(24); r.font.name='微软雅黑'
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
r=p.add_run('初步版｜公开资料驱动'); r.font.size=Pt(13); r.font.color.rgb=RGBColor(100,100,100)
for text in ['项目ID：62e6c2f3-5d0a-4cd9-b49d-2919b8ba1249','报告日期：2026年8月6日','内部资料｜未经许可不得外传']:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(text).font.size=Pt(10.5)
doc.add_page_break()

# contents placeholder (Word can update fields)
p=doc.add_paragraph('目录',style='Heading 1')
toc=OxmlElement('w:fldSimple'); toc.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u'); p2=doc.add_paragraph(); p2._p.append(toc)
doc.add_page_break()

def clean(s):
    return s.replace('**','').replace('`','').strip()

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

i=0
while i < len(lines):
    line=lines[i].rstrip()
    if not line or line=='---': i+=1; continue
    if i<8 and (line.startswith('# ') or line.startswith('**项目') or line.startswith('**拟核验') or line.startswith('**报告') or line.startswith('**项目ID') or line.startswith('**报告属性')):
        i+=1; continue
    if line.startswith('> '):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.right_indent=Cm(.5)
        r=p.add_run(clean(line[2:])); r.italic=True; r.font.color.rgb=RGBColor(120,70,20)
        i+=1; continue
    if line.startswith('### '): doc.add_paragraph(clean(line[4:]),style='Heading 3'); i+=1; continue
    if line.startswith('## '): doc.add_paragraph(clean(line[3:]),style='Heading 2'); i+=1; continue
    if line.startswith('# '): doc.add_paragraph(clean(line[2:]),style='Heading 1'); i+=1; continue
    if line.startswith('|') and i+1<len(lines) and lines[i+1].startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            vals=[clean(x) for x in lines[i].strip('|').split('|')]
            if not all(set(v)<=set('-: ') for v in vals): rows.append(vals)
            i+=1
        if rows:
            cols=max(map(len,rows)); tbl=doc.add_table(rows=len(rows), cols=cols); tbl.style='Table Grid'; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
            for rr,row in enumerate(rows):
                for cc in range(cols):
                    cell=tbl.cell(rr,cc); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=row[cc] if cc<len(row) else ''
                    for para in cell.paragraphs:
                        for run in para.runs: run.font.size=Pt(8.5); run.font.name='微软雅黑'
                    if rr==0:
                        shade(cell,'D9EAF7')
                        for run in cell.paragraphs[0].runs: run.bold=True
            doc.add_paragraph()
        continue
    if line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); p.add_run(clean(line[2:])); i+=1; continue
    import re
    m=re.match(r'^(\d+)\.\s+(.*)',line)
    if m:
        p=doc.add_paragraph(style='List Number'); p.add_run(clean(m.group(2))); i+=1; continue
    p=doc.add_paragraph();
    # simple bold whole paragraph if wrapped
    text=clean(line); r=p.add_run(text)
    if line.startswith('**') and line.endswith('**'): r.bold=True
    i+=1

# metadata
props=doc.core_properties
props.title='智灵动力投资尽职调查报告（初步版）'; props.subject='投资尽职调查'; props.author='浙江赛智伯乐股权投资管理有限公司投资中台'
doc.save(out)
print(out)
