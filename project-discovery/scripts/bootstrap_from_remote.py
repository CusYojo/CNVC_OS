#!/usr/bin/env python3
"""Bootstrap a local Project Discovery Radar from its read-only HTTP API."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
import tempfile
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import Workbook


SOURCE_FILES = {
    "arxiv": "arxiv_candidates.jsonl",
    "wechat_985": "wechat_985_candidates.jsonl",
    "wechat_api": "wechat_api_candidates.jsonl",
    "wechat_chat": "wechat_chat_candidates.jsonl",
    "investment": "investment_candidates.jsonl",
}


def fetch_json(base_url: str, path: str, params: dict[str, Any] | None = None) -> Any:
    query = f"?{urlencode(params)}" if params else ""
    url = f"{base_url.rstrip('/')}{path}{query}"
    last_error: Exception | None = None
    for attempt in range(3):
        try:
            request = Request(url, headers={"User-Agent": "cybernaut-radar-bootstrap/1.0"})
            with urlopen(request, timeout=30) as response:
                return json.loads(response.read().decode("utf-8"))
        except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as exc:
            last_error = exc
            if attempt < 2:
                time.sleep(1 + attempt)
    raise RuntimeError(f"读取远程雷达失败: {url}: {last_error}")


def candidate_key(item: dict[str, Any]) -> str:
    for field in ("source_id", "fingerprint", "link", "title"):
        value = str(item.get(field, "") or "").strip()
        if value:
            return value
    return ""


def clean_candidate(item: dict[str, Any]) -> dict[str, Any]:
    cleaned = dict(item)
    cleaned.pop("_file", None)
    cleaned.pop("_line", None)
    return cleaned


def atomic_write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        stamp = datetime.now().strftime("%Y%m%d%H%M%S")
        shutil.copy2(path, path.with_suffix(f"{path.suffix}.bak-{stamp}"))
    fd, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as handle:
            for row in rows:
                handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
        os.replace(temp_name, path)
    except Exception:
        Path(temp_name).unlink(missing_ok=True)
        raise


def write_accounts_workbook(path: Path, accounts: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for account in accounts:
        sheet = str(account.get("sheet") or account.get("group") or "公众号")[:31]
        grouped.setdefault(sheet, []).append(account)
    for sheet, rows in sorted(grouped.items()):
        worksheet = workbook.create_sheet(sheet)
        worksheet.append(["公众号", "帐号名"])
        seen: set[tuple[str, str]] = set()
        for row in rows:
            key = (
                str(row.get("account_name", "") or "").strip(),
                str(row.get("wx_name", "") or "").strip(),
            )
            if not all(key) or key in seen:
                continue
            seen.add(key)
            worksheet.append(list(key))
        worksheet.freeze_panes = "A2"
        worksheet.column_dimensions["A"].width = 36
        worksheet.column_dimensions["B"].width = 30
    if not workbook.sheetnames:
        worksheet = workbook.create_sheet("公众号")
        worksheet.append(["公众号", "帐号名"])
    workbook.save(path)


def fetch_source(base_url: str, source: str) -> tuple[str, list[dict[str, Any]], int]:
    payload = fetch_json(
        base_url,
        "/api/candidates",
        {"source": source, "attention_only": "false", "limit": 500},
    )
    return source, payload.get("items", []), int(payload.get("total", 0))


def fetch_wechat_account(base_url: str, source_key: str) -> tuple[str, list[dict[str, Any]], int]:
    payload = fetch_json(
        base_url,
        "/api/candidates",
        {
            "source": "wechat_api",
            "source_key": source_key,
            "attention_only": "false",
            "limit": 500,
        },
    )
    return source_key, payload.get("items", []), int(payload.get("total", 0))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--remote-base", default="http://101.126.93.130:8121")
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--workers", type=int, default=12)
    args = parser.parse_args()

    project_dir = args.project_dir.resolve()
    data_dir = project_dir / "data"
    accounts_path = project_dir / "公众号来源.xlsx"
    remote_base = args.remote_base.rstrip("/")

    health = fetch_json(remote_base, "/api/health")
    summary = fetch_json(remote_base, "/api/summary")
    account_payload = fetch_json(remote_base, "/api/wechat-api/accounts")
    source_payload = fetch_json(remote_base, "/api/wechat/sources")
    accounts = account_payload.get("accounts", [])
    source_keys = sorted({
        str(account.get("wx_name", "") or "").strip()
        for account in accounts
        if str(account.get("wx_name", "") or "").strip()
    })

    rows_by_source: dict[str, dict[str, dict[str, Any]]] = {
        source: {} for source in SOURCE_FILES
    }
    warnings: list[str] = []

    generic_sources = ["arxiv", "wechat_985", "wechat_chat", "investment"]
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as executor:
        futures = [
            executor.submit(fetch_source, remote_base, source)
            for source in generic_sources
        ]
        # The endpoint has no offset and caps at 500, so recover WeChat per account.
        futures.extend(
            executor.submit(fetch_wechat_account, remote_base, source_key)
            for source_key in source_keys
        )
        # Also include the generic first page to catch historical rows whose account
        # was removed from the current workbook.
        futures.append(executor.submit(fetch_source, remote_base, "wechat_api"))

        completed = 0
        total_futures = len(futures)
        for future in as_completed(futures):
            key, items, total = future.result()
            completed += 1
            if completed % 50 == 0 or completed == total_futures:
                print(
                    f"已读取远程信源 {completed}/{total_futures}",
                    file=sys.stderr,
                    flush=True,
                )
            if total > len(items):
                warnings.append(f"{key}: total={total}, returned={len(items)}")
            for raw in items:
                item = clean_candidate(raw)
                source = str(item.get("source", "") or "")
                if source not in rows_by_source:
                    continue
                dedupe_key = candidate_key(item)
                if dedupe_key:
                    rows_by_source[source][dedupe_key] = item

    restored_counts: dict[str, int] = {}
    for source, filename in SOURCE_FILES.items():
        rows = sorted(
            rows_by_source[source].values(),
            key=lambda row: (
                str(row.get("published_at", "") or ""),
                int(row.get("attention_score", 0) or 0),
            ),
            reverse=True,
        )
        atomic_write_jsonl(data_dir / filename, rows)
        restored_counts[source] = len(rows)

    write_accounts_workbook(accounts_path, accounts)
    data_dir.mkdir(parents=True, exist_ok=True)
    (data_dir / "wechat_985_sources.json").write_text(
        json.dumps({"sources": source_payload.get("sources", [])}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    expected_counts = {
        str(source): int(count)
        for source, count in (summary.get("sources") or {}).items()
    }
    mismatches = {
        source: {"expected": expected, "restored": restored_counts.get(source, 0)}
        for source, expected in expected_counts.items()
        if restored_counts.get(source, 0) != expected
    }
    manifest = {
        "restored_at": datetime.now(timezone.utc).isoformat(),
        "remote_base": remote_base,
        "remote_health": health,
        "remote_summary": summary,
        "restored_counts": restored_counts,
        "account_count": len(accounts),
        "source_key_count": len(source_keys),
        "wechat_source_count": len(source_payload.get("sources", [])),
        "warnings": warnings,
        "count_mismatches": mismatches,
        "gsdata_credentials": "not_exported_by_remote_api",
    }
    (data_dir / "bootstrap_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    if mismatches:
        raise SystemExit("恢复数量与远程摘要不一致，请检查 bootstrap_manifest.json")


if __name__ == "__main__":
    main()
