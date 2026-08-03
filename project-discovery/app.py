from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import os
import re
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Literal
from urllib.parse import parse_qs, quote_plus, urljoin, urlparse
from zoneinfo import ZoneInfo

import feedparser
import httpx
from bs4 import BeautifulSoup
from fastapi import FastAPI, HTTPException, Query
from openpyxl import load_workbook
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field


def env_flag(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("RADAR_DATA_DIR", str(BASE_DIR / "data"))).expanduser().resolve()
STATIC_DIR = BASE_DIR / "static"
ARXIV_FILE = DATA_DIR / "arxiv_candidates.jsonl"
WECHAT_FILE = DATA_DIR / "wechat_985_candidates.jsonl"
WECHAT_API_FILE = DATA_DIR / "wechat_api_candidates.jsonl"
WECHAT_CHAT_CANDIDATES_FILE = DATA_DIR / "wechat_chat_candidates.jsonl"
WECHAT_CHAT_CANDIDATES_DIR = DATA_DIR / "wechat_chat_candidates_by_group"
WECHAT_CHAT_MESSAGES_DIR = DATA_DIR / "wechat_chat_messages"
INVESTMENT_FILE = DATA_DIR / "investment_candidates.jsonl"
WECHAT_SOURCES_FILE = DATA_DIR / "wechat_985_sources.json"
AUTO_STATUS_FILE = DATA_DIR / "auto_crawler_status.json"
WECHAT_DAILY_STATUS_FILE = DATA_DIR / "wechat_daily_status.json"
WECHAT_ACCOUNTS_XLSX = Path(
    os.getenv("RADAR_WECHAT_ACCOUNTS_XLSX", str(BASE_DIR / "公众号来源.xlsx"))
).expanduser().resolve()
GSDATA_CREDENTIALS_FILE = DATA_DIR / "gsdata_credentials.json"
GSDATA_API_URL = "http://databus.gsdata.cn:8888/api/service"
GSDATA_WECHAT_ROUTER = "/weixin/article/search1"
GSDATA_WECHAT_CONTENT_ROUTER = "/weixin/article/content"
PITCHHUB_FLOW_URL = "https://gateway.36kr.com/api/mis/nav/home/project/bulletin/flow"
ARXIV_API_URL = "https://export.arxiv.org/api/query"
ARXIV_RSS_URL = "https://arxiv.org/rss/{category}"
ARXIV_ID_RE = re.compile(r"arxiv\.org/abs/([0-9]{4}\.[0-9]{4,5}(?:v[0-9]+)?)")
WECHAT_ARTICLE_MAX_CHARS = 12000
WECHAT_ARTICLE_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.8",
}
CHAT_CONTEXT_BEFORE = 4
CHAT_CONTEXT_AFTER = 4
AUTO_CRAWL_INTERVAL_SECONDS = 30 * 60
AUTO_CRAWL_GROUPS = ["创投新闻", "海外项目", "论文"]
AUTO_CRAWL_ENABLED = env_flag("RADAR_AUTO_CRAWL_ENABLED", True)
WECHAT_DAILY_ENABLED = env_flag("RADAR_WECHAT_DAILY_ENABLED", True)
CHINA_TZ = ZoneInfo("Asia/Shanghai")
WECHAT_DAILY_RUN_HOUR = 8
WECHAT_DAILY_RUN_MINUTE = 30
WECHAT_API_MAX_WORKERS = 2


TOP_VENUES = (
    "NeurIPS", "ICML", "ICLR", "CVPR", "ICCV", "ECCV", "ACL", "EMNLP",
    "NAACL", "AAAI", "IJCAI", "KDD", "SIGIR", "WWW", "The Web Conference",
    "CHI", "SIGGRAPH", "OSDI", "SOSP", "NSDI", "SIGCOMM", "USENIX Security",
    "IEEE S&P", "CCS", "NDSS", "Nature", "Science", "Cell", "PNAS",
    "Lancet", "NEJM", "Nature Medicine", "Nature Biotechnology",
    "Nature Machine Intelligence", "Nature Communications", "Science Robotics",
    "Science Advances",
)

MAJOR_FUNDERS = (
    "National Key R&D Program", "National Natural Science Foundation of China",
    "NSFC", "国家自然科学基金", "国家重点研发计划", "科技创新2030",
    "National Science Foundation", "NSF", "National Institutes of Health",
    "NIH", "DARPA", "IARPA", "ARPA-E", "Department of Energy", "DOE",
    "Department of Defense", "DOD", "NASA", "European Research Council",
    "ERC", "Horizon Europe", "UKRI", "EPSRC", "Wellcome Trust",
    "Gates Foundation", "Bill & Melinda Gates Foundation", "Chan Zuckerberg",
    "HHMI", "DFG", "ANR", "JST", "JSPS", "KAKENHI", "AMED", "NEDO",
)

UNIVERSITY_985 = (
    "北京大学", "清华大学", "中国人民大学", "北京航空航天大学", "北京理工大学", "中国农业大学",
    "北京师范大学", "中央民族大学", "南开大学", "天津大学", "大连理工大学", "东北大学",
    "吉林大学", "哈尔滨工业大学", "复旦大学", "同济大学", "上海交通大学", "华东师范大学",
    "南京大学", "东南大学", "浙江大学", "中国科学技术大学", "厦门大学", "山东大学",
    "中国海洋大学", "武汉大学", "华中科技大学", "湖南大学", "中南大学", "国防科技大学",
    "中山大学", "华南理工大学", "四川大学", "电子科技大学", "重庆大学", "西安交通大学",
    "西北工业大学", "西北农林科技大学", "兰州大学",
    "Peking University", "Tsinghua University", "Renmin University of China",
    "Beihang University", "Beijing Institute of Technology", "China Agricultural University",
    "Beijing Normal University", "Nankai University", "Tianjin University",
    "Dalian University of Technology", "Northeastern University", "Jilin University",
    "Harbin Institute of Technology", "Fudan University", "Tongji University",
    "Shanghai Jiao Tong University", "East China Normal University", "Nanjing University",
    "Southeast University", "Zhejiang University", "University of Science and Technology of China",
    "Xiamen University", "Shandong University", "Ocean University of China",
    "Wuhan University", "Huazhong University of Science and Technology",
    "Hunan University", "Central South University", "National University of Defense Technology",
    "Sun Yat-sen University", "South China University of Technology", "Sichuan University",
    "University of Electronic Science and Technology of China", "Chongqing University",
    "Xi'an Jiaotong University", "Northwestern Polytechnical University", "Northwest A&F University",
    "Lanzhou University",
)

UNIVERSITIES_985_SOURCES = (
    {"school": "北京大学", "province": "北京"},
    {"school": "清华大学", "province": "北京"},
    {"school": "中国人民大学", "province": "北京"},
    {"school": "北京航空航天大学", "province": "北京"},
    {"school": "北京理工大学", "province": "北京"},
    {"school": "中国农业大学", "province": "北京"},
    {"school": "北京师范大学", "province": "北京"},
    {"school": "中央民族大学", "province": "北京"},
    {"school": "南开大学", "province": "天津"},
    {"school": "天津大学", "province": "天津"},
    {"school": "大连理工大学", "province": "辽宁"},
    {"school": "东北大学", "province": "辽宁"},
    {"school": "吉林大学", "province": "吉林"},
    {"school": "哈尔滨工业大学", "province": "黑龙江"},
    {"school": "复旦大学", "province": "上海"},
    {"school": "同济大学", "province": "上海"},
    {"school": "上海交通大学", "province": "上海"},
    {"school": "华东师范大学", "province": "上海"},
    {"school": "南京大学", "province": "江苏"},
    {"school": "东南大学", "province": "江苏"},
    {"school": "浙江大学", "province": "浙江"},
    {"school": "中国科学技术大学", "province": "安徽"},
    {"school": "厦门大学", "province": "福建"},
    {"school": "山东大学", "province": "山东"},
    {"school": "中国海洋大学", "province": "山东"},
    {"school": "武汉大学", "province": "湖北"},
    {"school": "华中科技大学", "province": "湖北"},
    {"school": "湖南大学", "province": "湖南"},
    {"school": "中南大学", "province": "湖南"},
    {"school": "国防科技大学", "province": "湖南"},
    {"school": "中山大学", "province": "广东"},
    {"school": "华南理工大学", "province": "广东"},
    {"school": "四川大学", "province": "四川"},
    {"school": "电子科技大学", "province": "四川"},
    {"school": "重庆大学", "province": "重庆"},
    {"school": "西安交通大学", "province": "陕西"},
    {"school": "西北工业大学", "province": "陕西"},
    {"school": "西北农林科技大学", "province": "陕西"},
    {"school": "兰州大学", "province": "甘肃"},
)

PROJECT_KEYWORDS = (
    "科研", "成果", "转化", "获批", "项目", "课题", "基金", "重点实验室", "工程中心",
    "技术", "论文", "专利", "团队", "合作", "突破", "发布", "入选", "重大",
    "Nature", "Science", "Cell", "顶刊", "产业化", "孵化", "揭榜挂帅", "国家重点",
    "青年科学家", "杰青", "优青", "院士", "国际合作", "临床", "样机", "平台",
)

INVESTMENT_KEYWORDS = (
    "融资", "投资", "天使轮", "种子轮", "Pre-A", "A轮", "B轮", "C轮", "战略投资",
    "并购", "上市", "IPO", "独角兽", "估值", "商业化", "量产", "订单", "客户",
    "产业化", "落地", "试点", "合作", "签约", "获批", "临床", "注册证", "FDA",
    "专利", "授权", "发明", "技术转移", "成果转化", "许可", "孵化", "产品发布",
    "developer tools", "AI", "machine learning", "robotics", "biotech", "medtech",
)

SKILL_SCORER_VERSION = "private-market-1.1.4"
PRIVATE_MARKET_RETAIN_SCORE = 60
UNIVERSITY_TECH_RETAIN_SCORE = 52

PRIMARY_MARKET_TERMS = (
    "天使轮", "种子轮", "Pre-A", "pre-A", "Pre A", "A轮", "B轮", "C轮", "D轮",
    "首轮融资", "新一轮融资", "完成融资", "轮融资", "融资", "股权融资", "战略融资", "战略投资", "私募",
    "PE", "成长轮", "投资方", "领投", "跟投", "资本", "基金", "未上市",
    "初创", "创业公司", "startup", "start-up", "venture", "growth equity",
)

PRIVATE_ENTITY_TERMS = (
    "初创公司", "创业公司", "未上市", "科技公司", "有限公司", "团队", "项目",
    "实验室", "研究院", "孵化", "校友创业", "成果转化", "技术转移", "产业化",
    "spinout", "spin-off", "startup", "founder", "co-founder", "创始人", "CEO", "CTO",
)

TECH_DIFFERENTIATION_TERMS = (
    "专利", "发明", "授权", "论文", "顶刊", "Nature", "Science", "Cell", "算法",
    "模型", "芯片", "机器人", "人工智能", "大模型", "AI", "机器学习", "合成生物",
    "生物医药", "医疗器械", "创新药", "材料", "传感器", "低空经济", "具身智能",
    "数据集", "开源", "样机", "原型", "注册证", "临床", "FDA", "NMPA",
)

UNIVERSITY_VALUE_TERMS = (
    "科技成果", "科研成果", "成果转化", "技术转移", "产业化", "转化落地", "孵化",
    "研究成果", "实验室成果", "项目成果", "重大成果", "创新成果",
    "专利", "发明专利", "授权专利", "论文", "顶刊", "高水平论文", "Nature", "Science", "Cell",
    "获奖", "国家科学技术奖", "科技奖", "科学技术奖", "自然科学奖", "技术发明奖",
    "科技进步奖", "专利奖",
    "实验室", "重点实验室", "工程中心", "研究中心", "课题组", "教授", "院士", "博士",
    "样机", "原型", "中试", "临床", "注册证", "获批", "试点", "示范应用", "产业应用",
    "核心技术", "关键技术", "突破", "首创", "国际领先", "国内首个", "填补空白",
    "人工智能", "大模型", "机器人", "芯片", "半导体", "光电", "量子", "低空经济",
    "合成生物", "生物医药", "医疗器械", "新材料", "新能源", "核聚变", "传感器",
)

UNIVERSITY_STRONG_VALUE_TERMS = (
    "科技成果", "科研成果", "成果转化", "技术转移", "产业化", "转化落地",
    "研究成果", "实验室成果", "项目成果", "重大成果", "创新成果",
    "专利", "发明专利", "授权专利", "顶刊", "高水平论文", "Nature", "Science", "Cell",
    "国家科学技术奖", "科技奖", "科学技术奖", "自然科学奖", "技术发明奖",
    "科技进步奖", "专利奖",
    "重点实验室", "工程中心", "研究中心", "课题组", "样机", "原型", "中试",
    "临床试验", "临床研究", "临床应用", "注册证", "医疗器械研发", "示范应用",
    "产业应用", "核心技术", "关键技术", "突破", "首创", "国际领先", "国内首个",
    "填补空白", "获批",
)

UNIVERSITY_RESEARCH_OUTCOME_TERMS = (
    "科技成果", "科研成果", "成果转化", "技术转移", "产业化", "转化落地",
    "研究成果", "实验室成果", "项目成果", "重大成果", "创新成果",
    "专利", "发明专利", "授权专利", "顶刊", "高水平论文", "Nature", "Science", "Cell",
    "取得重要进展", "突破", "首创", "核心技术",
    "关键技术", "样机", "原型", "中试", "临床试验", "临床研究", "临床应用",
    "医疗器械研发", "注册证", "示范应用", "产业应用", "国际领先", "国内首个",
    "填补空白", "获批", "获奖", "科技奖", "科学技术奖", "自然科学奖",
    "技术发明奖", "科技进步奖", "专利奖",
)

UNIVERSITY_WEAK_RESEARCH_TERMS = (
    "发文", "发表", "刊发", "学术发表", "最新发现", "揭示", "开发",
)

UNIVERSITY_CONCRETE_OUTCOME_TERMS = (
    "科技成果", "科研成果", "成果转化", "技术转移", "产业化", "转化落地",
    "专利", "发明专利", "授权专利", "样机", "原型", "中试", "临床试验",
    "临床研究", "临床应用", "医疗器械研发", "注册证", "示范应用", "产业应用",
    "新靶点", "治疗机制", "技术路线", "技术平台", "国家科学技术奖",
    "科学技术奖", "自然科学奖", "技术发明奖", "科技进步奖", "专利奖",
)

UNIVERSITY_TEAM_CONTEXT_TERMS = (
    "实验室", "重点实验室", "工程中心", "研究中心", "课题组", "院士", "教授",
    "PI", "研究团队", "科研团队",
)

UNIVERSITY_INVESTABLE_TECH_DOMAIN_TERMS = (
    "人工智能", "大模型", "算法", "数据集", "芯片", "半导体", "光电", "量子",
    "机器人", "具身智能", "低空经济", "无人机", "传感器", "新材料", "新能源",
    "储能", "核聚变", "合成生物", "生物医药", "医疗器械", "创新药", "药物",
    "靶点", "基因", "细胞", "蛋白", "RNA", "疫苗", "诊断", "检测", "脑机",
    "网络安全", "隐私计算", "工业软件",
)

UNIVERSITY_HARD_NOISE_TERMS = (
    "高考", "家长", "招生", "毕业季", "图书馆", "学术资源", "开通试用",
    "资源动态", "好书", "新书", "出版社", "课程预告", "课程", "讲座",
    "论坛", "研讨会", "活动预告", "会议通知", "邀请参加", "报名", "培训",
    "夏令营", "简章", "转载", "调查报告", "博士后科技服务团", "通知",
    "考试", "资格考试", "准考证", "打印准考证", "考生",
)

UNIVERSITY_HARD_NOISE_RESCUE_TERMS = (
    "发明专利", "授权专利", "专利授权", "技术许可", "样机", "原型", "中试",
    "临床试验", "临床应用", "注册证", "医疗器械研发", "新靶点", "治疗机制",
    "技术平台", "国家科学技术奖", "科学技术奖", "自然科学奖", "技术发明奖",
    "科技进步奖", "专利奖",
)

UNIVERSITY_NOISE_TERMS = (
    "直播", "预告", "讲座", "论坛", "研讨会", "沙龙", "培训", "课程", "招生",
    "毕业季", "展售", "好物", "好书", "AI头条", "盘点", "报告下载", "工具大全",
    "实战技巧", "报名", "活动", "会议通知", "招聘", "宣讲", "志愿", "换届",
    "理事会", "大会举行", "训练营", "夏令营", "简章", "新书", "发布会",
    "讲堂", "项目夏季路演日", "决赛名单", "知识星球", "扫码", "全文下载",
)

TEAM_INSTITUTION_TERMS = (
    "院士", "教授", "博士", "博士后", "实验室", "重点实验室", "工程中心", "研究中心",
    "杰青", "优青", "青年科学家", "校友", "创始人", "连续创业", "核心团队",
)

MARKET_URGENCY_TERMS = (
    "国产替代", "降本增效", "刚需", "痛点", "市场规模", "千亿", "百亿", "需求",
    "政策", "国家级", "揭榜挂帅", "采购", "招标", "供应链", "出海", "产业链",
    "商业化", "量产", "规模化", "标准化", "合规", "监管", "双碳", "新质生产力",
)

TRACTION_TERMS = (
    "客户", "头部客户", "签约", "合作", "试点", "示范", "订单", "交付", "量产",
    "投产", "落地", "收入", "营收", "付费", "预售", "采购", "中标", "获批",
    "注册证", "临床", "产品发布", "上线", "发布", "接入", "部署",
)

TIMING_TERMS = (
    "近日", "日前", "今日", "昨日", "昨天", "本周", "本月", "今年", "最新",
    "首次", "刚刚", "宣布", "发布", "启动", "获批", "完成", "突破", "入选",
)

SECONDARY_MARKET_TERMS = (
    "股价", "涨停", "跌停", "目标价", "券商", "研报", "评级", "买入", "卖出",
    "增持", "减持", "持有", "EPS", "每股收益", "市盈率", "市净率", "分红",
    "回购", "二级市场", "股票", "A股", "港股", "美股", "上市公司", "财报",
    "季报", "年报", "trading", "stock price", "analyst rating", "price target",
)

DISCLOSURE_AMOUNT_RE = re.compile(
    r"((?:超|近|约|逾|数)?\d+(?:\.\d+)?\s*(?:万亿元|亿元|万元|亿美元|万美元|亿|万|元|美元|人民币|美金|USD|RMB|million|billion|M|B)"
    r"|(?:超|近|约|逾|数)?(?:数十|数百|数千)\s*(?:万亿元|亿元|万元|亿美元|万美元|亿|万|元|美元|人民币|美金)"
    r"|(?:超|近|约|逾|数)?(?:千万|百万|十亿|百亿|千亿)(?:元|美元|人民币|美金)?级"
    r"|(?:超|近|约|逾|数)?(?:千万|百万|十亿|百亿|千亿)(?:元|美元|人民币|美金))",
    re.IGNORECASE,
)

ROUND_TERMS = (
    "种子轮", "天使轮", "Pre-A", "pre-A", "Pre A", "A轮", "B轮", "C轮", "D轮",
    "首轮融资", "新一轮融资", "战略融资", "战略投资", "股权融资", "私募", "PE", "成长轮",
)

INDUSTRY_KEYWORDS = (
    ("具身智能", "具身智能/机器人"),
    ("人形机器人", "具身智能/机器人"),
    ("机器人", "机器人"),
    ("外骨骼", "医疗康复/机器人"),
    ("人工智能", "人工智能"),
    ("大模型", "人工智能/大模型"),
    ("AI", "人工智能"),
    ("机器学习", "人工智能"),
    ("芯片", "半导体/芯片"),
    ("半导体", "半导体/芯片"),
    ("硅光", "半导体/光电子"),
    ("光通讯", "半导体/光电子"),
    ("核聚变", "新能源/核聚变"),
    ("新能源", "新能源"),
    ("储能", "新能源/储能"),
    ("生物医药", "生物医药"),
    ("创新药", "生物医药"),
    ("医疗器械", "医疗器械"),
    ("临床", "医疗健康"),
    ("材料", "新材料"),
    ("传感器", "智能硬件/传感器"),
    ("电子皮肤", "智能硬件/传感器"),
    ("低空经济", "低空经济"),
    ("农业", "农业科技"),
    ("数据中心", "算力基础设施"),
    ("开发者", "企业服务/开发者工具"),
    ("developer", "企业服务/开发者工具"),
    ("消费", "消费科技"),
)

SOURCE_DEFAULTS = (
    {
        "key": "arxiv_cs_ai",
        "name": "arXiv cs.AI",
        "group": "论文",
        "type": "arxiv_rss",
        "url": "https://rss.arxiv.org/rss/cs.AI",
        "frequency": "工作日",
        "enabled": True,
    },
    {
        "key": "arxiv_eess_sp",
        "name": "arXiv eess.SP",
        "group": "论文",
        "type": "arxiv_rss",
        "url": "https://rss.arxiv.org/rss/eess.SP",
        "frequency": "工作日",
        "enabled": True,
    },
    {
        "key": "wanfang_ai",
        "name": "万方论文搜索 人工智能",
        "group": "论文",
        "type": "wanfang_search",
        "url": "https://s.wanfangdata.com.cn/paper?q={keyword}",
        "keyword": "人工智能",
        "frequency": "每天",
        "enabled": False,
        "note": "公开页当前主要返回前端外壳，默认不自动抓取。",
    },
    {
        "key": "google_patents_ml",
        "name": "Google Patents machine learning",
        "group": "专利",
        "type": "rss",
        "url": "https://patents.google.com/patent/rss?q=machine+learning&after=20260701&language=ZH",
        "frequency": "每天",
        "enabled": True,
    },
    {
        "key": "google_patents_ai_cn",
        "name": "Google Patents 人工智能",
        "group": "专利",
        "type": "rss",
        "url": "https://patents.google.com/patent/rss?q=%E4%BA%BA%E5%B7%A5%E6%99%BA%E8%83%BD&after=20260701&language=ZH",
        "frequency": "每天",
        "enabled": True,
    },
    {
        "key": "tsinghua_otl",
        "name": "清华大学技术转移研究院 成果库",
        "group": "高校成果",
        "type": "html_list",
        "url": "https://otl.tsinghua.edu.cn/achievements/list.htm",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "pku_tech",
        "name": "北京大学科技开发部 技术成果",
        "group": "高校成果",
        "type": "html_list",
        "url": "https://kjkfb.pku.edu.cn/cgxx/jscg/index.htm",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "sjtu_aitri",
        "name": "上海交通大学先进产业技术研究院 科技成果",
        "group": "高校成果",
        "type": "html_list",
        "url": "https://aitri.sjtu.edu.cn/achievement/list.htm",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "zju_ttri",
        "name": "浙江大学工业技术转化研究院 成果展示",
        "group": "高校成果",
        "type": "html_list",
        "url": "https://ttri.zju.edu.cn/cgzs/list.htm",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "casip",
        "name": "中国科学院知识产权与产业化网",
        "group": "高校成果",
        "type": "html_list",
        "url": "http://www.casip.ac.cn/kjcg/index.html",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "hust_ttc",
        "name": "华中科技大学科技成果转化服务中心",
        "group": "高校成果",
        "type": "html_list",
        "url": "http://ttc.hust.edu.cn/kjcg/kjcg.htm",
        "frequency": "每周",
        "enabled": True,
    },
    {
        "key": "pedaily_quicknews",
        "name": "投资界 快讯",
        "group": "创投新闻",
        "type": "rss",
        "url": "https://feeds.pedaily.cn/n/quicknews",
        "frequency": "每小时",
        "enabled": False,
        "note": "2026-07-29 起源站 TLS 握手失败，暂时停用，避免每轮产生确定性错误。",
    },
    {
        "key": "36kr_feed",
        "name": "36氪 RSS",
        "group": "创投新闻",
        "type": "rss",
        "url": "https://36kr.com/feed",
        "frequency": "每小时",
        "enabled": True,
    },
    {
        "key": "36kr_pitchhub_financing_flash",
        "name": "36氪 PitchHub 融资快报",
        "group": "创投新闻",
        "type": "36kr_financing_flash",
        "url": "https://pitchhub.36kr.com/financing-flash",
        "frequency": "每天",
        "enabled": True,
        "max_pages": 8,
        "max_entries_per_run": 100,
        "note": "公开 HTML 页面，默认只保留页面标记为昨天的融资信息。",
    },
    {
        "key": "lieyunwang_feed",
        "name": "猎云网 RSS",
        "group": "创投新闻",
        "type": "rss",
        "url": "https://www.lieyunwang.com/feed",
        "frequency": "每小时",
        "enabled": False,
        "note": "2026-07-29 起源站 TLS 握手失败，暂时停用，待源站恢复后重新启用。",
    },
    {
        "key": "producthunt",
        "name": "Product Hunt 今日新品",
        "group": "海外项目",
        "type": "rss",
        "url": "https://www.producthunt.com/feed",
        "frequency": "每天",
        "enabled": True,
    },
    {
        "key": "producthunt_devtools",
        "name": "Product Hunt Developer Tools",
        "group": "海外项目",
        "type": "rss",
        "url": "https://www.producthunt.com/categories/developer-tools/feed",
        "frequency": "每天",
        "enabled": False,
        "note": "Product Hunt 已移除分类 RSS（当前返回 404），保留官方主 Feed。",
    },
    {
        "key": "cnki_reference",
        "name": "中国知网公开摘要页",
        "group": "论文",
        "type": "manual",
        "url": "https://kns.cnki.net/kns8s/AdvSearch?classid=YSTT4HG0",
        "frequency": "参考",
        "enabled": False,
        "note": "反爬严格，默认不自动抓取。",
    },
    {
        "key": "sogou_weixin_reference",
        "name": "搜狗微信搜索",
        "group": "微信生态",
        "type": "manual",
        "url": "https://weixin.sogou.com/weixin?type=2&query=%E5%A4%A9%E4%BD%BF%E8%BD%AE%E8%9E%8D%E8%B5%84&page=1",
        "frequency": "低频",
        "enabled": False,
        "note": "验证码和反爬明显，默认不自动抓取。",
    },
    {
        "key": "itjuzi_reference",
        "name": "IT 桔子公司库",
        "group": "创投新闻",
        "type": "manual",
        "url": "https://www.itjuzi.com/company?page=1",
        "frequency": "低频",
        "enabled": False,
        "note": "反爬中等，建议后续单独接 Playwright。",
    },
)


class ArxivRunRequest(BaseModel):
    categories: list[str] = Field(default_factory=lambda: ["cs.AI", "cs.CL", "cs.CV", "cs.LG"])
    keywords: list[str] = Field(default_factory=list)
    max_results: int = Field(default=20, ge=1, le=100)
    days: int = Field(default=14, ge=1, le=90)
    watch_authors: list[str] = Field(default_factory=list)


class WechatAccount(BaseModel):
    name: str = ""
    rss_url: str = ""


class WechatSource(BaseModel):
    school: str
    province: str = ""
    accounts: list[WechatAccount] = Field(default_factory=list)


class WechatSourcesRequest(BaseModel):
    sources: list[WechatSource]


class WechatRunRequest(BaseModel):
    max_entries_per_feed: int = Field(default=20, ge=1, le=100)


class InvestmentRunRequest(BaseModel):
    groups: list[str] = Field(default_factory=list)
    max_entries_per_source: int = Field(default=15, ge=1, le=60)
    keyword: str = "人工智能"


class WechatApiRunRequest(BaseModel):
    date: str = ""
    days: int = Field(default=7, ge=1, le=30)
    groups: list[str] = Field(default_factory=list)
    wx_names: list[str] = Field(default_factory=list)
    max_accounts: int = Field(default=0, ge=0, le=5000)
    limit_per_account: int = Field(default=100, ge=1, le=500)


class WechatChatFile(BaseModel):
    file_serial_no: str = ""
    file_name: str = ""
    file_url: str = ""


class WechatChatMessageIn(BaseModel):
    msg_key: str = ""
    group_name: str = ""
    group_serial_no: str = ""
    sender_name: str = ""
    sender_serial_no: str = ""
    cite_content: str = ""
    msg_time: str = ""
    send_time: str = ""
    message_time: str = ""
    content: str = ""
    msg_content: str = ""
    msg_content_decoded: str = ""
    raw_msg_content: str = ""
    file: WechatChatFile | dict[str, Any] | None = None
    msg_type: int | str | None = None


class WechatChatPushRequest(BaseModel):
    merchant_no: str = ""
    pushed_at: str = ""
    messages: list[WechatChatMessageIn] = Field(default_factory=list)


auto_crawler_task: asyncio.Task | None = None
wechat_daily_task: asyncio.Task | None = None
auto_crawler_running = False
wechat_daily_running = False


app = FastAPI(title="Project Discovery Radar")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def contains_phrase(text: str, phrase: str) -> bool:
    if not phrase:
        return False
    if re.search(r"[\u4e00-\u9fff]", phrase):
        return phrase in text
    pattern = r"(?<![A-Za-z0-9])" + re.escape(phrase) + r"(?![A-Za-z0-9])"
    return bool(re.search(pattern, text, flags=re.IGNORECASE))


def extract_arxiv_id(entry_id: str) -> str:
    match = ARXIV_ID_RE.search(entry_id or "")
    return match.group(1) if match else ""


def extract_entry_arxiv_id(entry: dict) -> str:
    candidates = [entry.get("id", ""), entry.get("link", "")]
    candidates.extend(link.get("href", "") for link in entry.get("links", []) if isinstance(link, dict))
    for value in candidates:
        arxiv_id = extract_arxiv_id(value)
        if arxiv_id:
            return arxiv_id
    return ""


def source_key(item: dict) -> str:
    for field in ("source_id", "fingerprint", "link", "title"):
        value = clean_text(str(item.get(field, "")))
        if value:
            return value
    return ""


def radar_source_key(item: dict) -> str:
    key = source_key(item)
    source = clean_text(str(item.get("source", ""))) or "unknown"
    return f"{source}:{key}" if key else ""


def parse_candidate_datetime(value: Any) -> datetime | None:
    text = clean_text(str(value or ""))
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        parsed = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
    if parsed is None:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=CHINA_TZ)
    return parsed.astimezone(timezone.utc)


def candidate_cursor_key(item: dict) -> tuple[int, str]:
    parsed = None
    for field in ("collected_at", "published_at", "updated_at"):
        parsed = parse_candidate_datetime(item.get(field))
        if parsed is not None:
            break
    micros = int(parsed.timestamp() * 1_000_000) if parsed is not None else 0
    identity = radar_source_key(item) or json.dumps(item, ensure_ascii=False, sort_keys=True)
    digest = hashlib.sha256(identity.encode("utf-8")).hexdigest()
    return micros, digest


def encode_candidate_cursor(key: tuple[int, str]) -> str:
    payload = json.dumps({"v": 1, "t": key[0], "k": key[1]}, separators=(",", ":"))
    return base64.urlsafe_b64encode(payload.encode("utf-8")).decode("ascii").rstrip("=")


def decode_candidate_cursor(value: str) -> tuple[int, str]:
    try:
        padded = value + "=" * (-len(value) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded).decode("utf-8"))
        if payload.get("v") != 1:
            raise ValueError("unsupported cursor version")
        timestamp = int(payload["t"])
        digest = str(payload["k"])
        if timestamp < 0 or not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise ValueError("invalid cursor fields")
        return timestamp, digest
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise ValueError("invalid candidate cursor") from exc


def entry_datetime(entry: dict, field: str) -> str:
    parsed = entry.get(f"{field}_parsed")
    if parsed:
        return datetime(*parsed[:6], tzinfo=timezone.utc).isoformat()
    return entry.get(field, "")


def normalize_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", name.casefold())


def phrase_hits(text: str, phrases: tuple[str, ...], limit: int = 8) -> list[str]:
    hits = []
    for phrase in phrases:
        if contains_phrase(text, phrase) and phrase not in hits:
            hits.append(phrase)
            if len(hits) >= limit:
                break
    return hits


def university_top_venue_hits(text: str, raw_hits: list[str], limit: int = 5) -> list[str]:
    if not raw_hits:
        return []
    paper_context_re = re.compile(
        r"(发文|发表|刊发|在线发表|接收|见刊|登上|收录|论文|研究成果|最新发现|封面|"
        r"发表于|发表于《|在《|《[^》]{0,80}》)",
        flags=re.IGNORECASE,
    )
    venue_context_re = re.compile(
        r"(Nature|Science|Cell|PNAS|Lancet|NEJM|CVPR|ICCV|ECCV|NeurIPS|ICML|ICLR|"
        r"ACL|EMNLP|AAAI|IJCAI|KDD|SIGIR|CHI|SIGGRAPH|OSDI|SOSP|NSDI|SIGCOMM|CCS|NDSS)"
        r"([A-Za-z .&-]{0,40})?"
        r"(\+|发文|发表|刊发|论文|封面|接收|见刊|在线发表|《|》|：|:|！|!|，|,|$)",
        flags=re.IGNORECASE,
    )
    hits = []
    for hit in raw_hits:
        if hit == "Science" and re.search(r"\bData Science\b|\bSpatial Data Science\b", text, flags=re.IGNORECASE):
            continue
        if contains_phrase(text, f"《{hit}》") or contains_phrase(text, f"{hit}+"):
            hits.append(hit)
        elif hit in {"Nature", "Science", "Cell", "PNAS", "Lancet", "NEJM"}:
            for match in re.finditer(re.escape(hit), text, flags=re.IGNORECASE):
                context = text[max(0, match.start() - 24): match.end() + 36]
                if paper_context_re.search(context) or venue_context_re.search(context):
                    hits.append(hit)
                    break
        else:
            hits.append(hit)
        if len(hits) >= limit:
            break
    return unique_keep_order(hits, limit)


def dimension_item(code: str, label: str, score: int, max_score: int, detail: str, hits: list[str] | None = None) -> dict:
    return {
        "code": code,
        "label": label,
        "score": max(0, min(max_score, score)),
        "max_score": max_score,
        "detail": detail,
        "hits": hits or [],
    }


def unique_keep_order(values: list[str], limit: int = 8) -> list[str]:
    rows = []
    for value in values:
        value = clean_text(value)
        if value and value not in rows:
            rows.append(value)
            if len(rows) >= limit:
                break
    return rows


def nearby_amounts(text: str, keywords: tuple[str, ...]) -> list[str]:
    values = []
    for match in DISCLOSURE_AMOUNT_RE.finditer(text):
        context = text[max(0, match.start() - 32): match.end() + 32]
        if any(contains_phrase(context, keyword) for keyword in keywords):
            values.append(match.group(1).strip())
    return unique_keep_order(values, 4)


def extract_round(text: str) -> str:
    hits = phrase_hits(text, ROUND_TERMS, 4)
    if hits:
        return hits[0]
    if contains_phrase(text, "完成融资") or contains_phrase(text, "融资"):
        return "融资轮次待核实"
    return "未披露/待核实"


def extract_investors(text: str) -> str:
    patterns = (
        r"(?:由|获|获得)([^。；，,.]{2,80}?)(?:领投|独家投资|投资|联合领投|参投|跟投)",
        r"([^。；，,.]{2,80}?)(?:领投|联合领投|参投|跟投)",
        r"(?:投资方|投资机构)(?:为|包括|有|：|:)([^。；]{2,120})",
    )
    hits = []
    stop_words = ("融资", "本轮", "近日", "完成", "宣布", "公司")
    for pattern in patterns:
        for match in re.finditer(pattern, text):
            value = clean_text(match.group(1))
            value = re.sub(r"^(本轮|本次|此次|近日|正式|已完成)", "", value)
            value = re.sub(r"(等多家机构|等机构|等)$", "", value)
            if value and not all(word in value for word in stop_words[:2]):
                hits.append(value)
    if not hits:
        hits = phrase_hits(text, ("资本", "基金", "创投", "投资", "金控", "产投", "险资"), 5)
    return "；".join(unique_keep_order(hits, 4)) or "未披露/待核实"


def extract_industries(text: str, explicit: list[str] | None = None) -> str:
    hits = list(explicit or [])
    for keyword, industry in INDUSTRY_KEYWORDS:
        if contains_phrase(text, keyword):
            hits.append(industry)
    return "、".join(unique_keep_order(hits, 5)) or "未识别/待核实"


def split_sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[。！？!?])|[；;]\s*", text)
    return [clean_text(part) for part in parts if len(clean_text(part)) >= 8]


def select_sentences(text: str, keywords: tuple[str, ...], limit: int = 3) -> list[str]:
    rows = []
    for sentence in split_sentences(text):
        if any(contains_phrase(sentence, keyword) for keyword in keywords):
            rows.append(sentence[:180])
            if len(rows) >= limit:
                break
    return rows


def extract_people(text: str) -> str:
    rows = select_sentences(text, ("创始人", "联合创始人", "CEO", "CTO", "教授", "博士", "博士后", "团队", "核心团队"), 3)
    return "；".join(rows) or "未披露/待核实"


def extract_contact(text: str, link: str = "") -> str:
    contacts = []
    for pattern in (
        r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+",
        r"(?:\+?86[-\s]?)?1[3-9]\d{9}",
        r"(?:电话|联系人|邮箱|微信|联系方式)[:： ]{0,3}([^。；\s]{3,60})",
    ):
        contacts.extend(match.group(0 if "@" in pattern or "\\d" in pattern else 1) for match in re.finditer(pattern, text))
    if not contacts and link:
        contacts.append(f"来源链接: {link}")
    return "；".join(unique_keep_order(contacts, 4)) or "未披露/待核实"


def extract_lab(text: str, fallback_school: str = "") -> str:
    patterns = (
        r"([\u4e00-\u9fffA-Za-z0-9]{2,30}(?:重点实验室|实验室|研究中心|工程中心|研究院|课题组|团队))",
        r"((?:National|State|Key|Lab|Laboratory|Institute|Center)[A-Za-z0-9 ,&-]{4,80})",
    )
    hits = []
    for pattern in patterns:
        for match in re.finditer(pattern, text, flags=re.IGNORECASE):
            candidate = clean_text(match.group(1))
            # 过滤被贪婪匹配吃入的谓语片段（例如"引导广大科研团队主动走出实验室"）
            if VAGUE_PROJECT_START_RE.search(candidate) or PROJECT_PREDICATE_RE.search(candidate):
                continue
            if re.search(r"[，。！？；;：:\n]", candidate):
                continue
            if candidate:
                hits.append(candidate)
    if fallback_school and not hits:
        hits.append(fallback_school)
    return "；".join(unique_keep_order(hits, 4)) or "未披露/待核实"


COMPANY_NAME_PLACEHOLDERS = {
    "",
    "待核验",
    "待核实",
    "未披露",
    "未披露/待核实",
    "未披露/待验证",
    "未识别/待核实",
    "不适用",
    "无",
    "-",
    "N/A",
    "null",
}
LEGAL_COMPANY_SUFFIX_RE = r"(?:股份有限公司|有限责任公司|有限公司)"
LEGAL_COMPANY_CHARS_RE = r"[\u4e00-\u9fffA-Za-z0-9（）()·&＋+\-]"


def is_meaningful_company_name(value: Any) -> bool:
    name = clean_text(str(value or "")).strip("“”\"' ")
    return bool(name) and name not in COMPANY_NAME_PLACEHOLDERS and len(name) <= 128


GENERIC_PROJECT_SUBJECTS = {
    "人工智能", "大模型", "机器人", "新材料", "新能源",
    "项目", "团队", "研究团队", "实验室", "课题组", "相关项目", "某项目", "作者", "负责人",
    "文章来源", "文章转载", "全新突破", "硬氪前线", "硬氪首发", "独家", "首发", "喜报",
    "来源", "清心新闻", "硬氪", "36氪", "十五五", "融资", "会赚钱", "保险科技",
    "消费级智能", "中国半导体", "入局物理智能", "X教授", "前瞻理论研究与创新平台",
    "AI软件", "卡脖子", "6氪", "定价权", "财务数据", "涌现新项目",
}
VAGUE_PROJECT_START_RE = re.compile(
    r"^(他|她|其|该|这|此|其中|上述|相关|目前|同时|此外|另|据|对于|关于|要求|需要|应当|必须|"
    r"支持|推动|加强|开展|主动|引导|持续|继续|曾|曾经|担任|联创|联合创始人?|成立|创始人|科研人员|参赛|"
    r"本次|全体|让|让更多|让我|使我|并使|并|基于|后两年|共享|未来|"
    r"作为|为|用于|以|从|在|将|把|被|由|于|联合|面对|通过|围绕|聚焦|落地|年初|年末|年底|月初|月末|算|"
    r"根据|我要|我们要|"
    r"了解|开拓|真实|价值|推荐阅读|背靠|赠礼环节|学员们|深刻|深刻认识|深刻体会|深刻理解|"
    r"购票观众|课题被|科创报国|促进|紧跟|对标|正是|也|过去|活动在|锤炼|需要|更多|不再|只有|至今|随后|共同|双方|各自|截止|"
    r"一是|二是|三是|四是)"
)
VAGUE_PROJECT_BODY_RE = re.compile(
    r"(岗位记录|关键证明|证明材料|要求主动|主动适应|走出实验室|为核心业务|核心业务的|系统梳理|以及团队|"
    r"进入导师课题组|共享两个学院|获颁|获评|荣获|获奖|荣誉|Award|"
    r"文章来源|论文合作者|合作者为|受试者|研究参与者|筛选期|完全开放|依托高校|顺利通过|"
    r"一行到访|成功举办|先后发言|按姓氏拼音排序|首先来到|带队|不是在实验室|"
    r"老师.*介绍|分别介绍了|介绍了其|第一城|早些时候|一批|多家|数个|能够替代|能够|"
    r"key observation|by the paper|(?:已经|已|正在)组建|已打通|过往所投公司的创业团队|"
    r"等信息|等材料|等证明|等方面|等工作|带来的变化)",
    flags=re.IGNORECASE,
)
VAGUE_PROJECT_END_RE = re.compile(r"(材料|记录|信息|情况|内容|要求|工作|方面|变化|问题|任务|路径|策略|证明|累计|责编|来源|再|消息|报道)$")
PROJECT_PREDICATE_RE = re.compile(
    r"(要求|适应|指出|表示|强调|认为|提出|推动|支持|引导|开展|打造|落地|实现|完成|获得|发布|宣布|拟|提供|形成|"
    r"建立|构建|促进|提升|加强|记录|证明|担任|任职|毕业|来自|师从|进入|共享|梳理|发表|结合|"
    r"参与|经历|合作者|申请|接受|走出|走访|到访|举办|发言|带队|来到|组建|体会|体会到了?|感受到|意识到|认识到|学到|了解到)"
)
NUMBERED_TECH_FRAGMENT_RE = re.compile(r"^[\u4e00-\u9fffA-Za-z]{1,8}[-—–][\u4e00-\u9fffA-Za-z]{1,8}\d{1,2}$")
GENERIC_INSTITUTION_TECH_RE = re.compile(
    r"^(?:清华|北大|北航|上交大|复旦|浙大|中科大|哈工大)(?:系)?"
    r"(?:机器人|芯片|人工智能|大模型)(?:团队|项目)?$"
)
LOW_VALUE_PUBLIC_LEAD_RE = re.compile(
    r"(?:院系之声.{0,30}(?:荣誉|获奖|Award)|"
    r"(?:教授|研究员|学者).{0,30}(?:获颁|获评|荣获|获奖|Award|荣誉|发文|发表文章)|"
    r"(?:获得|获评|入选|荣获|获).{0,24}(?:奖|荣誉|称号|教学团队|表彰|标兵)|"
    r"(?:科学技术奖|科技奖|自然科学奖|技术发明奖|科技进步奖).{0,40}(?:揭晓|获奖|表彰)|"
    r"\d+\s*项.{0,12}(?:获奖|获表彰)|"
    r"(?:国家级|省级|全国高校).{0,16}(?:教学团队|教学成果|荣誉|奖|标兵)|奖学金|"
    r"受试者招募|招募(?:研究参与者|受试者)|参与本研究|临床试验.{0,50}(?:招募|受试者|研究参与者)|"
    r"实践成果.{0,24}(?:申请|硕士学位)|学位答辩|专业学位培养改革|"
    r"论文.{0,40}(?:期刊|发表|刊发|接受|接收|accepted)|学术成果|研究论文|文章来源|转载全文|"
    r"毕业(?:季|典礼|致辞|生|倒计时|设计)|毕业生去哪儿|校友招聘|社会招聘|诚聘|实习生|招聘|"
    r"党支部|党员|党务|党建|革命先辈|校史|悼念|缅怀|研修班|训练营|课程|移动课堂|工作坊|"
    r"讲座(?:预告)?|活动(?:预告|抢先知)|Information Session|参访|探访|师生校友|院友沙龙|"
    r"创新大赛|参赛队伍|\d+\s*家.{0,24}(?:企业|公司).{0,30}(?:融资|投资)|"
    r"专场(?:科创)?路演|路演举办|加速计划.{0,20}(?:招募|启动)|"
    r"最前线|解码硬科技|罚单|行业进入强监管|"
    r"(?:\d+点\d*氪|氪星|创投|财经)(?:晚报|早报)?|为什么资本|什么样的.{0,20}(?:能|会)|"
    r"行业观察|赛道观察|赴港上市|登陆资本市场|IPO认购|上市获|"
    r"要报.{0,12}专业吗|我要不要学AI|招生(?:简章|宣传|咨询|专业|对象)?|"
    r"培养方案|课程介绍|实验班介绍|研修班|培训班|结业证书|能力提升计划|名家面对面|"
    r"学员企业|毕业典礼|毕业致辞|发表致辞|兼任|受聘|履新|任命|"
    r"(?:记者|人物)?专访|人物访谈|观点访谈|深度解读|系统剖析)",
    flags=re.IGNORECASE,
)
PURE_ACADEMIC_PUBLIC_LEAD_RE = re.compile(
    r"(?:(?:课题组|团队|实验室).{0,100}(?:发表|论文|研究|揭示|破解|开发|发现|成果)|"
    r"(?:学术成果|科研成果|研究进展|研究论文|最新研究|多项研究|两项研究|研究成果|合作论文).{0,100}"
    r"(?:课题组|团队|教授|研究员|实验室|突破|发现|揭示|开发|发表|刊发|接收)?|"
    r"(?:论文|研究成果).{0,80}(?:发表|刊发|接收|accepted|publication)|"
    r"(?:发表于|在线发表于|accepted by).{0,60}(?:期刊|journal|nature|science|IEEE)|"
    r"Science Publication|论文摘要|"
    r"(?:团队|课题组).{0,60}(?:算法|模型|数据|机制|通路|架构))",
    flags=re.IGNORECASE,
)
COMMERCIAL_PUBLIC_LEAD_RE = re.compile(
    r"(?:成果转化|技术转移|转化落地|产业化|中试|技术平台|工程化|技术许可|专利转让|孵化(?:成立|企业|公司)|"
    r"创办公司|成立公司|产品获批|注册证|临床应用|应用新场景|示范应用|产业应用|"
    r"客户验证|客户订单|采购|中标|签约|量产|营收|商业化)",
    flags=re.IGNORECASE,
)
REAL_INVESTMENT_EVENT_RE = re.compile(
    r"(?:(?:完成|宣布|获得|获).{0,16}(?:融资|投资)|(?:融资|投资).{0,16}(?:领投|跟投|交割)|"
    r"估值.{0,12}(?:亿元|万美元|亿美元|万元)|(?:天使轮|种子轮|Pre-?A|A轮|B轮|C轮|D轮|战略融资|战略投资))",
    flags=re.IGNORECASE,
)
VERIFIED_SOURCE_SUBJECT_RULES = (
    (re.compile(r"北航机器人所团队创业.{0,24}智能变刚度关节"), "航墨科技"),
    (re.compile(r"清华系初创完成数亿元种子轮融资.{0,30}世界模型"), "厘清智能"),
    (re.compile(r"前大疆科学家创业.{0,30}(?:四轮|耀途资本|锦秋基金)"), "硅羽科技"),
)


def is_specific_project_subject_name(value: Any) -> bool:
    name = clean_text(str(value or "")).strip("“”\"'「」『』 ,，。！？；;：:")
    if not name or name in COMPANY_NAME_PLACEHOLDERS or name == "未命名项目":
        return False
    if len(name) < 2 or len(name) > 60 or name in GENERIC_PROJECT_SUBJECTS:
        return False
    if NUMBERED_TECH_FRAGMENT_RE.fullmatch(name):
        return False
    if GENERIC_INSTITUTION_TECH_RE.search(name):
        return False
    if ("（" in name and "）" not in name) or ("(" in name and ")" not in name):
        return False
    if re.search(r"[，。！？；;：:、|｜丨\n]", name):
        return False
    if VAGUE_PROJECT_START_RE.search(name) or VAGUE_PROJECT_BODY_RE.search(name) or VAGUE_PROJECT_END_RE.search(name):
        return False
    if re.fullmatch(r"(?:(?:\d+|数|多)?人|个人|创始人?)创业团队", name):
        return False
    if re.fullmatch(r"[\u4e00-\u9fff]{2,18}(?:触觉|感知|技术|科技|机器人|硬件|软件|设备|材料|能源|半导体)企业", name):
        return False
    if re.fullmatch(r"[\u4e00-\u9fff]{2,24}(?:文化创意|消费|家居|生活方式|文创)品牌", name):
        return False
    if "等" in name:
        return False
    if PROJECT_PREDICATE_RE.search(name) and not re.search(r"(公司|企业|项目|团队|实验室|研究院|研究中心|工程中心|课题组)$", name):
        return False
    subject_marker = re.search(
        r"(股份有限公司|有限责任公司|有限公司|公司|企业|项目|团队|实验室|研究院|研究所|研究中心|工程中心|"
        r"课题组|创新群体|创新联合体|中试基地|产业基地|创新平台|技术平台|研发平台|试验平台|装置|系统|产品|计划)$",
        name,
    )
    english_words = re.findall(r"[A-Za-z][A-Za-z0-9-]*", name)
    verified_english_brand = bool(re.fullmatch(r"[A-Z][A-Z0-9.-]*(?:\s+[A-Z][A-Z0-9.-]*){1,3}", name))
    if len(english_words) >= 2 and not verified_english_brand and not subject_marker and not re.search(
        r"(AI|Labs?|Laboratory|Institute|Center|Centre|Technologies|Technology|Robotics|Bio|Systems?|Platform|Project)$",
        name,
        flags=re.IGNORECASE,
    ):
        return False
    if not subject_marker and re.search(r"(的|了|是|以|在|为|将|把|被|对于|关于|正在|让|到|体会|感受|觉得|知道|认识|深刻|意识|理解|了解|学到|得到)", name):
        return False
    # 拒绝明显是完整句子的名称
    if re.search(r"[？！。！]", name):
        return False
    # 拒绝长英文标题（>40字符且纯英文，通常是论文标题）
    if re.fullmatch(r"[A-Za-z0-9\s:,\-()\[\]&;+]+", name) and len(name) > 40:
        return False
    # 拒绝问句
    if re.search(r"如何|为什么|是否|怎么|怎样|什么", name) and not subject_marker:
        return False
    # 拒绝明显的多公司融资综述标题
    if re.search(r"\d+\s*[家个]", name) and re.search(r"(?:企业|公司|融资|上市)", name):
        return False
    # 拒绝以年份/日期开头的泛化描述
    if re.match(r"^(?:19|20)\d{2}[年\s]", name) and not subject_marker:
        return False
    # 拒绝以纯数字、日期、序号开头
    if re.match(r"^(?:\d+[月日年个只家项位次]|第\s*\d+\s*期)", name):
        return False
    return True


def normalize_project_candidate(value: Any) -> str:
    candidate = clean_text(str(value or "")).strip("“”\"'「」『』 ,，。！？；;：:")
    described_brand = re.search(r"(?:研发商|制造商|提供商|品牌|独角兽|公司|企业)([\u4e00-\u9fffA-Za-z0-9·&＋+\-\s]{2,24})$", candidate)
    if described_brand:
        candidate = described_brand.group(1)
    candidate = re.sub(r"^(?:超声脑机接口公司|端侧大模型独角兽|清华系端侧大模型独角兽|消费级智能硬件品牌)", "", candidate)
    candidate = re.sub(r"^(?:推出的|研发的|打造的|研制的)", "", candidate)
    candidate = re.sub(r"^投资(?=[\u4e00-\u9fffA-Za-z0-9])", "", candidate)
    candidate = re.sub(r"(?:目前|近期|已经|已|正式|斩)$", "", candidate)
    if re.fullmatch(r"[A-Za-z0-9·&＋+\-]{2,24}团队", candidate, flags=re.IGNORECASE):
        candidate = re.sub(r"团队$", "项目", candidate)
    if re.search(r"(?:系)?初创$", candidate):
        candidate += "项目"
    if (
        candidate
        and not re.search(r"(项目|团队|实验室|研究院|研究中心|工程中心|课题组|基地|平台|装置|系统|产品|计划)$", candidate)
        and re.search(r"(智能体|外骨骼|机器人|模型|芯片|装置|平台|系统|产品)$", candidate)
    ):
        candidate += "项目"
    return candidate


def extract_primary_news_subject(text: str) -> str:
    """Extract the company/brand that is the subject of the opening news event."""
    opening = re.sub(
        r"^(?:(?:36氪|硬氪)?(?:前线|首发)|独家|首发|喜报|快讯|重磅)\s*[|｜:：]\s*",
        "",
        clean_text(text),
        flags=re.IGNORECASE,
    )[:1800]
    name_chars = r"[\u4e00-\u9fffA-Za-z0-9（）()·&＋+\-]"
    descriptor = rf"(?:{name_chars}{{0,28}}(?:企业|公司|品牌|提供商|供应商|研发商|制造商|独角兽))?"
    event = (
        r"(?=(?:(?:近日|日前|近期|今日)\s*)?(?:\d+\s*个?月(?:内)?\s*)?"
        r"(?:(?:连续|已|正式)\s*)*(?:(?:官宣|宣布)\s*)?(?:完成|获得|获|成立(?:于)?|是一家))"
    )
    patterns = (
        rf"(?:获悉|消息显示|公开信息显示)[，,\s]*(?:(?:近日|日前|近期)[，,\s]*)?{descriptor}\s*[「『“\"]?({name_chars}{{2,40}}?)[」』”\"]?\s*{event}",
        rf"(?:^|[。；;\n])\s*(?:(?:近日|日前|近期|今日)[，,\s]*)?(?:\d{{1,4}}年)?\d{{0,2}}月?\d{{0,2}}日?[，,\s]*{descriptor}\s*[「『“\"]?({name_chars}{{2,40}}?)[」』”\"]?\s*{event}",
        rf"(?:^|[。；;\n])\s*({name_chars}{{2,30}})(?=成立(?:于)?|是一家|专注于|致力于)",
        rf"(?:^|[。；;\n])\s*({name_chars}{{2,30}}?)(?=(?:完成|获得|获).{{0,24}}(?:融资|投资))",
        rf"(?:离职|创业者)?创办\s*[「『“\"]({name_chars}{{2,30}})[」』”\"](?=[，,。；;\s])",
        r"(?:制造商|研发商|公司|企业)\s*([A-Z][A-Za-z0-9&+.-]*(?:\s+[A-Z][A-Za-z0-9&+.-]*){0,3})(?=\s*(?:获得|完成|宣布|获))",
        r"(?:^|[，。；;\n])\s*(?:\d{4}年\d{1,2}月[，,\s]*)?([A-Z][A-Za-z0-9&+.-]*(?:\s+[A-Z][A-Za-z0-9&+.-]*){1,3})(?=\s*成立(?:于)?[，,。\s])",
        r"中文名\s*[“「『\"]?([\u4e00-\u9fffA-Za-z0-9·&＋+\-\s]{2,30}?)[”」』\"]?(?=[，,。；;])",
        r"(?:公司(?:名|叫做?)|品牌(?:名|叫做?))\s*[“「『\"]?([\u4e00-\u9fffA-Za-z0-9·&＋+\-\s]{2,30}?)[”」』\"]?(?=[，,。；;])",
    )
    for pattern in patterns:
        for match in re.finditer(pattern, opening):
            candidate = normalize_project_candidate(match.group(1))
            if is_specific_project_subject_name(candidate):
                return candidate
    return ""


def extract_quoted_financing_subject(title: str) -> str:
    normalized = re.sub(
        r"^(?:(?:36氪|硬氪)?(?:前线|首发)|独家|首发|喜报|快讯|重磅)\s*[|｜:：]\s*",
        "",
        clean_text(title),
        flags=re.IGNORECASE,
    )
    pattern = re.compile(
        r"[「『“\"]([\u4e00-\u9fffA-Za-z0-9·&＋+\-\s]{2,30})[」』”\"]"
        r"(?=[^「『“\"]{0,36}(?:完成|获得|获|融资|投资))",
        flags=re.IGNORECASE,
    )
    for match in pattern.finditer(normalized):
        candidate = normalize_project_candidate(match.group(1))
        if is_specific_project_subject_name(candidate) or re.fullmatch(
            r"[A-Z][A-Z0-9.-]*(?:\s+[A-Z][A-Z0-9.-]*){1,3}",
            candidate,
        ):
            return candidate
    return ""


def extract_verified_source_subject(title: str) -> str:
    normalized = clean_text(title)
    for pattern, subject in VERIFIED_SOURCE_SUBJECT_RULES:
        if pattern.search(normalized):
            return subject
    return ""


def extract_title_project_subject(title: str) -> str:
    normalized = re.sub(
        r"^(?:(?:36氪|硬氪)?(?:前线|首发)|独家|首发|喜报|快讯|重磅)\s*[|｜:：]\s*",
        "",
        clean_text(title),
        flags=re.IGNORECASE,
    )
    marker = (
        r"(?:产业技术中试基地|创新联合体|中试基地|产业基地|创新平台|技术平台|研发平台|试验平台|"
        r"示范基地|转化基地|项目|计划|装置|系统)"
    )
    pattern = rf"(?:^|[\s|｜！!。；;，,\n])\s*([\u4e00-\u9fffA-Za-z0-9（）()·&＋+\-]{{2,60}}?{marker})(?=在|落地|签约|揭牌|启用|发布|完成|获|，|。|$)"
    candidates = [
        clean_text(match.group(1)).strip("“”\"'「」『』 ,，。！？；;：:")
        for match in re.finditer(pattern, normalized)
    ]
    candidates = [candidate for candidate in candidates if is_specific_project_subject_name(candidate)]
    return max(candidates, key=len) if candidates else ""


def extract_financing_title_subject(title: str) -> str:
    normalized = re.sub(
        r"^(?:(?:36氪|硬氪)?(?:前线|首发)|独家|首发|喜报|快讯|重磅)\s*[|｜:：]\s*",
        "",
        clean_text(title),
        flags=re.IGNORECASE,
    )
    response_subject = re.match(
        r"^([A-Za-z][A-Za-z0-9&+.-]{1,30})(?=回应(?:融资|投资)报道)",
        normalized,
        flags=re.IGNORECASE,
    )
    match = response_subject or re.search(
        r"(?:^|[，,；;！!])\s*([\u4e00-\u9fffA-Za-z0-9（）()·&＋+\-\s]{2,30}?)"
        r"(?=(?:(?:目前|近期|已经|已|正式)\s*)?(?:斩获|完成|获得|获).{0,24}(?:融资|投资))",
        normalized,
        flags=re.IGNORECASE,
    )
    if not match:
        return ""
    candidate = normalize_project_candidate(match.group(1))
    return candidate if is_specific_project_subject_name(candidate) else ""


def extract_company_name(item: dict, text: str) -> str:
    """Extract a legal entity only when the source contains explicit evidence.

    ``project_name`` is intentionally not used here: it is often a product,
    laboratory or short brand name.  In particular, this function never adds a
    legal suffix that did not occur in the source.
    """
    for field in ("company_name", "legal_entity", "company_full_name"):
        explicit = clean_text(str(item.get(field, ""))).strip("“”\"' ")
        if is_meaningful_company_name(explicit) and not re.search(
            rf"(?:揭牌|推进会|在|于|由).{{2,}}{LEGAL_COMPANY_SUFFIX_RE}$",
            explicit,
        ) and (
            re.search(rf"{LEGAL_COMPANY_SUFFIX_RE}$", explicit)
            or is_specific_project_subject_name(explicit)
        ):
            return explicit[:128]

    aliases = []
    explicit_project_name = clean_text(str(item.get("project_name", ""))).strip("“”\"'「」『』 ")
    if 2 <= len(explicit_project_name) <= 40 and not re.search(r"[，。！？；;]", explicit_project_name):
        aliases.append(explicit_project_name)
    title = clean_text(item.get("title", ""))
    aliases.extend(
        clean_text(match.group(1))
        for match in re.finditer(r"[“\"「『]([^”\"」』]{2,40})[”\"」』]", title)
    )

    company_group = rf"({LEGAL_COMPANY_CHARS_RE}{{2,80}}{LEGAL_COMPANY_SUFFIX_RE})"
    patterns = (
        # Explicitly labelled fields in articles or imported records.
        rf"(?:公司全称|企业全称|工商主体|公司主体|法律主体|主体名称|企业名称)\s*(?:为|是|[:：])\s*[“\"]?{company_group}",
        # Strong news-writing convention: full legal name followed by an alias.
        rf"(?:^|[，。；;：:\n])\s*[“\"]?{company_group}[”\"]?\s*[（(]\s*(?:以下简称|简称)",
        # Source explicitly introduces the target company.
        rf"(?:获悉|消息显示|公开信息显示|项目主体(?:为|是)|项目公司(?:为|是)|公司主体(?:为|是))[\s，,:：]*[“\"]?{company_group}",
        # A legal entity is the grammatical subject of a financing/company event.
        rf"(?:^|[，。；;：:\n])\s*[“\"]?{company_group}[”\"]?\s*"
        rf"(?=近日|日前|宣布|完成|获得|获|成立于|总部|"
        rf"通过(?:港交所|上市聆讯)|提交(?:上市|IPO)|冲刺(?:港股|上市|IPO)|拟(?:上市|IPO))",
    )
    for pattern in patterns:
        match = re.search(pattern, text)
        if not match:
            continue
        candidate = clean_text(match.group(1)).strip("“”\"' ")
        alias_matched = False
        for alias in aliases:
            position = candidate.find(alias)
            if position >= 0:
                alias_matched = True
                # 行政区划通常就是法定主体的一部分（例如“深圳可立点科技有限公司”），
                # 不能因为简称是“可立点科技”就把“深圳”裁掉。只有简称前明确是
                # 新闻描述性前缀时才裁剪，且绝不凭空补充任何公司后缀。
                prefix = candidate[:position]
                if position > 0 and re.search(r"(?:企业|项目|品牌|合作伙伴|投资标的|被投企业)$", prefix):
                    candidate = candidate[position:]
                break
        stem = re.sub(rf"{LEGAL_COMPANY_SUFFIX_RE}$", "", candidate)
        # “某赛道企业甲公司有限公司”包含描述性前缀。只有来源同时明确给出
        # 项目简称且能在候选中定位时才裁剪，否则宁可留空，不猜测主体边界。
        if not alias_matched and re.search(r"(?:企业|公司)", stem):
            continue
        if is_meaningful_company_name(candidate):
            return candidate[:128]
    return ""


def infer_project_name(item: dict, text: str) -> str:
    title = clean_text(item.get("title", ""))
    explicit_project = normalize_project_candidate(item.get("project_name", ""))
    is_academic = item.get("source_group") == "高校公众号" or bool(re.search(r"(?:学术成果|科研成果|课题组|实验室)", title))
    # 学术来源和普通来源的末级兜底都会使用这些主体模式。
    # 必须在分支外定义，避免普通新闻/arXiv 在走到兜底逻辑时触发
    # UnboundLocalError，进而让整个 /api/candidates 接口返回 500。
    subject_patterns = (
        r"((?:[\u4e00-\u9fff]{2,20}(?:大学|学院|研究所|医院))[\u4e00-\u9fff·]{0,16}(?:教授|研究员|博士)?团队)",
        r"([\u4e00-\u9fffA-Za-z0-9·]{2,30}(?:重点实验室|实验室|研究中心|工程中心|研究院|课题组))",
        r"([\u4e00-\u9fff·]{2,6}(?:教授|研究员|博士)?团队)",
    )

    # 新闻综述/榜单标题 → 不提取单一项目名
    if re.search(r"\d+\s*[家个].{0,30}(?:企业|公司).{0,30}(?:融资|投资|上市)", title):
        return "未命名项目"
    if re.search(r"(?:盘点|榜单|汇总|综述|周报|月报|季报|年报|早报|晚报|\d+点\d*氪|氪星|迟到的狂欢|投资狂潮|终极能源之战)", title):
        return "未命名项目"
    if len(re.findall(r"(?:完成|获得|获).{0,24}(?:融资|投资)", title)) >= 2:
        return "未命名项目"

    verified_source_subject = extract_verified_source_subject(title)
    if verified_source_subject:
        return verified_source_subject
    quoted_financing_subject = extract_quoted_financing_subject(title)
    if quoted_financing_subject:
        return quoted_financing_subject
    financing_title_subject = extract_financing_title_subject(title)
    if financing_title_subject and not re.search(r"(?:团队|项目|产品|平台)$", financing_title_subject):
        return financing_title_subject
    combined_article_text = "\n".join([
        clean_text(item.get("summary", "")),
        clean_text(item.get("article_text", "")),
        clean_text(text),
    ])
    quoted_article_subject = extract_quoted_financing_subject(combined_article_text)
    if quoted_article_subject:
        return quoted_article_subject
    title_project = extract_title_project_subject(title)
    if title_project and (
        not is_specific_project_subject_name(explicit_project)
        or (explicit_project in title_project and len(title_project) > len(explicit_project) + 2)
    ):
        return title_project
    legal_company_subject = extract_company_name(item, combined_article_text)
    if legal_company_subject:
        return legal_company_subject
    article_news_subject = extract_primary_news_subject(combined_article_text)
    if article_news_subject:
        return article_news_subject
    if financing_title_subject:
        return financing_title_subject
    title_news_subject = extract_primary_news_subject(title)
    if title_news_subject:
        return title_news_subject
    if is_specific_project_subject_name(explicit_project):
        return explicit_project
    quoted = re.search(r"[“\"]([^”\"]{2,40})[”\"]", title)
    if quoted and is_specific_project_subject_name(quoted.group(1)):
        return quoted.group(1)

    # 高校 / 学术文章：全文 extract_primary_news_subject 容易把
    # “让我深刻体会到科技成果转化是连接实验室…”这类文章内句子误提取为主体名称。
    # 优先提取实验室/课题组/机构+团队，再回退到公司名称，不再用 news_subject 兜底。
    if is_academic:
        for pattern in subject_patterns:
            candidates = [clean_text(match.group(1)) for match in re.finditer(pattern, title + "\n" + text)]
            candidates = [candidate for candidate in candidates if is_specific_project_subject_name(candidate)]
            if candidates:
                return sorted(set(candidates), key=len, reverse=True)[0]

        for company in re.finditer(r"([\u4e00-\u9fffA-Za-z0-9]{2,30}(?:科技|智能|机器人|医药|生物|材料|能源|半导体|电子|医疗|信息|数据|软件|网络|光电|先导院|研究院))(?:有限公司|公司|完成|获|宣布|近日)?", title + " " + text):
            candidate = clean_text(company.group(1))
            if is_specific_project_subject_name(candidate):
                return candidate

        if is_specific_project_subject_name(title):
            return title[:60]
        return "未命名项目"

    # 非学术 / 投资新闻来源：全文 news_subject 用于最后兜底
    news_subject = extract_primary_news_subject(text)
    if news_subject:
        return news_subject

    for company in re.finditer(r"([\u4e00-\u9fffA-Za-z0-9]{2,30}(?:科技|智能|机器人|医药|生物|材料|能源|半导体|电子|医疗|信息|数据|软件|网络|光电|先导院|研究院))(?:有限公司|公司|完成|获|宣布|近日)?", title + " " + text):
        candidate = clean_text(company.group(1))
        if is_specific_project_subject_name(candidate):
            return candidate

    # 高校成果常以实验室、课题组或“机构 + 负责人 + 团队”为主体。
    for pattern in subject_patterns:
        candidates = [clean_text(match.group(1)) for match in re.finditer(pattern, title + "\n" + text)]
        candidates = [candidate for candidate in candidates if is_specific_project_subject_name(candidate)]
        if candidates:
            return sorted(set(candidates), key=len, reverse=True)[0]

    if is_specific_project_subject_name(title):
        return title[:60]
    return "未命名项目"

def project_profile_text(item: dict) -> str:
    title = clean_text(item.get("title", ""))
    summary = clean_text(item.get("summary", ""))
    return "\n".join([
        title,
        summary,
        clean_text(item.get("article_text", "")),
        clean_text(item.get("comment", "")),
        clean_text(item.get("journal_ref", "")),
        clean_text(item.get("private_market_thesis", "")),
    ])


BUSINESS_REGION_PATTERNS = (
    ("北京", r"北京"),
    ("上海", r"上海"),
    ("天津", r"天津"),
    ("重庆", r"重庆"),
    ("河北", r"河北|石家庄|唐山|保定|廊坊|雄安|秦皇岛|邯郸|沧州"),
    ("山西", r"山西|太原|大同|长治|晋城|晋中|运城|临汾|吕梁"),
    ("内蒙古", r"内蒙古|呼和浩特|包头|鄂尔多斯|赤峰|通辽"),
    ("辽宁", r"辽宁|沈阳|大连|鞍山|抚顺|丹东|锦州|营口"),
    ("吉林", r"吉林省|长春|吉林市|延边|四平|通化"),
    ("黑龙江", r"黑龙江|哈尔滨|齐齐哈尔|大庆|牡丹江|佳木斯"),
    ("江苏", r"江苏|南京|苏州|无锡|常州|南通|扬州|镇江|泰州|盐城|徐州|常熟|昆山"),
    ("浙江", r"浙江|杭州|宁波|温州|嘉兴|湖州|绍兴|金华|舟山|台州|丽水|义乌"),
    ("安徽", r"安徽|合肥|芜湖|蚌埠|马鞍山|安庆|滁州|阜阳|六安"),
    ("福建", r"福建|福州|厦门|泉州|漳州|莆田|宁德"),
    ("江西", r"江西|南昌|九江|赣州|景德镇|上饶"),
    ("山东", r"山东|济南|青岛|烟台|潍坊|济宁|威海|临沂"),
    ("河南", r"河南|郑州|开封|洛阳|新乡|许昌|南阳|商丘"),
    ("湖北", r"湖北|武汉|宜昌|襄阳|荆州|黄冈"),
    ("湖南", r"湖南|长沙|株洲|湘潭|衡阳|岳阳|常德|郴州"),
    ("广东", r"广东|广州|深圳|珠海|汕头|佛山|东莞|中山|惠州|南沙|前海"),
    ("广西", r"广西|南宁|柳州|桂林|北海|钦州"),
    ("海南", r"海南|海口|三亚|儋州"),
    ("四川", r"四川|成都|绵阳|德阳|宜宾|乐山|南充"),
    ("贵州", r"贵州|贵阳|遵义|安顺|毕节"),
    ("云南", r"云南|昆明|曲靖|大理|丽江"),
    ("西藏", r"西藏|拉萨|日喀则|林芝"),
    ("陕西", r"陕西|西安|咸阳|宝鸡|渭南|榆林"),
    ("甘肃", r"甘肃|兰州|天水|酒泉|庆阳"),
    ("青海", r"青海|西宁|海东"),
    ("宁夏", r"宁夏|银川|石嘴山|吴忠"),
    ("新疆", r"新疆|乌鲁木齐|克拉玛依|喀什|伊犁"),
    ("香港", r"香港"),
    ("澳门", r"澳门"),
    ("台湾", r"台湾|台北|新北|台中|台南|高雄|新竹"),
)
INSTITUTION_REGION_PATTERNS = (
    ("北京", r"清华大学|北京大学|北京航空航天大学|北京理工大学|中国人民大学|北京师范大学|中国科学院大学|中关村"),
    ("上海", r"复旦大学|同济大学|华东师范大学|上海交通大学|上海交大|上海科技大学|紫竹高新区|张江"),
    ("浙江", r"浙江大学|浙大|西湖大学|之江实验室|良渚实验室"),
    ("江苏", r"南京大学|东南大学|南京航空航天大学|南京理工大学|苏州大学|江南大学"),
    ("安徽", r"中国科学技术大学|中科大|合肥工业大学"),
    ("湖北", r"武汉大学|华中科技大学|华中农业大学|武汉理工大学"),
    ("湖南", r"中南大学|湖南大学|国防科技大学"),
    ("广东", r"中山大学|华南理工大学|南方科技大学|深圳大学|香港中文大学（深圳）"),
    ("四川", r"四川大学|电子科技大学|西南交通大学"),
    ("陕西", r"西安交通大学|西北工业大学|西安电子科技大学"),
    ("天津", r"天津大学|南开大学"),
    ("重庆", r"重庆大学|西南大学"),
    ("福建", r"厦门大学|福州大学"),
    ("山东", r"山东大学|中国海洋大学"),
    ("辽宁", r"大连理工大学|东北大学"),
    ("吉林", r"吉林大学"),
    ("黑龙江", r"哈尔滨工业大学|哈工大"),
)


def normalize_business_region(value: Any) -> str:
    text = clean_text(str(value or ""))
    if not text or text in {"待确认", "待核验", "待核实", "未披露", "未披露/待核实", "不适用", "无", "-"}:
        return ""
    for region, pattern in BUSINESS_REGION_PATTERNS:
        if re.search(pattern, text):
            return region
    return ""


def infer_business_region(item: dict, text: str, company_name: str = "", lab: str = "") -> dict:
    for field in ("business_region", "region", "reg_location", "registered_address", "location", "headquarters"):
        region = normalize_business_region(item.get(field))
        if region:
            return {"region": region, "region_source": "雷达结构化地区", "region_confidence": "高"}

    for value in (company_name, item.get("project_name", "")):
        candidate = clean_text(str(value or ""))
        for region, pattern in BUSINESS_REGION_PATTERNS:
            match = re.search(pattern, candidate)
            if match and match.start() == 0:
                return {"region": region, "region_source": "主体名称行政区划", "region_confidence": "中"}

    explicit_pattern = re.compile(
        r"(?:注册地|注册地址|工商注册地址|注册于|注册在|总部所在地|总部位于|总部设于|总部设在|"
        r"公司所在地|公司位于|企业所在地|企业位于|公司落户|企业落户|"
        r"项目所在地|项目位于|项目落地于|项目落户|基地所在地|基地位于|基地落地于|坐落于)"
        r"\s*[：:为在]?\s*([^，。；;\n]{2,48})"
    )
    location_text = text[:12000]
    subject_tokens = [
        clean_text(str(value or ""))
        for value in (company_name, item.get("project_name", ""))
        if len(clean_text(str(value or ""))) >= 2
    ]
    subject_tokens += [
        re.sub(r"(?:股份有限公司|有限责任公司|有限公司|公司|企业|项目|团队|实验室|研究院|研究所|研究中心)$", "", value)
        for value in subject_tokens
    ]
    subject_tokens = [value for value in unique_keep_order(subject_tokens, 8) if len(value) >= 2]
    for match in explicit_pattern.finditer(location_text):
        context = location_text[max(0, match.start() - 160):min(len(location_text), match.end() + 80)]
        if subject_tokens and not any(token in context for token in subject_tokens):
            continue
        region = normalize_business_region(match.group(1))
        if region:
            return {"region": region, "region_source": "来源原文明确地点", "region_confidence": "中"}

    academic = "高校" in str(item.get("source_group", "")) or bool(re.search(
        r"大学|学院|研究院|研究所|实验室|课题组|教授团队|科研团队|研究团队",
        " ".join([
            clean_text(str(item.get("project_name", ""))),
            clean_text(lab),
            clean_text(str(item.get("source_name", ""))),
        ]),
    ))
    if academic:
        institution_text = " ".join([
            clean_text(str(item.get("school", ""))),
            clean_text(str(item.get("source_name", ""))),
            clean_text(lab),
            clean_text(str(item.get("project_name", ""))),
        ])
        region = normalize_business_region(institution_text)
        if not region:
            for candidate_region, pattern in INSTITUTION_REGION_PATTERNS:
                if re.search(pattern, institution_text):
                    region = candidate_region
                    break
        if region:
            return {"region": region, "region_source": "所属高校/研究机构", "region_confidence": "中"}

    return {"region": "待确认", "region_source": "", "region_confidence": ""}


def build_project_profile(item: dict) -> dict:
    title = clean_text(item.get("title", ""))
    text = project_profile_text(item)
    source = item.get("source", "")
    link = clean_text(item.get("link", ""))
    dimensions = item.get("score_dimensions", [])
    risk_rows = []
    if item.get("filter_reasons"):
        risk_rows.extend(item.get("filter_reasons", [])[:2])
    for dim in dimensions:
        if int(dim.get("max_score", 0) or 0) >= 10 and int(dim.get("score", 0) or 0) <= 2:
            risk_rows.append(f"{dim.get('label', dim.get('code', '关键维度'))}信息不足")
    if not risk_rows:
        risk_rows = select_sentences(text, ("风险", "不确定", "挑战", "壁垒", "监管", "竞争", "成本", "难点"), 2)
    if not risk_rows:
        risk_rows = ["未披露收入、估值或股权结构时，需通过直接尽调核实。"]

    profile = {
        "project_name": infer_project_name(item, text),
        "company_name": extract_company_name(item, text) or "未披露/待核实",
        "project_round": clean_text(item.get("project_financing_round", "")) or extract_round(text),
        "latest_valuation": "、".join(nearby_amounts(text, ("估值", "估价", "投前", "投后", "市值"))) or "未披露/待核实",
        "financing_amount": "、".join(nearby_amounts(text, ("融资", "投资", "募资", "资金", "金额", "轮"))) or "未披露/待核实",
        "institutions": extract_investors(text),
        "industry": extract_industries(text, item.get("project_industries", [])),
        "core_highlights": "；".join(select_sentences(text, ("首创", "唯一", "核心", "专利", "突破", "量产", "客户", "订单", "临床", "注册证", "顶级", "高温超导", "AI", "机器人", "大模型"), 3)) or clean_text(item.get("project_brief", "")) or "待进一步挖掘",
        "risk_notes": "；".join(unique_keep_order(risk_rows, 4)),
        "team_composition": extract_people(text),
        "contact": extract_contact(text, link),
        "lab": extract_lab(text, item.get("school", "") or item.get("source_name", "") if "高校" in item.get("source_group", "") else ""),
        "source_url": link,
        "data_completeness": "自动抽取，缺失字段需人工核实",
    }
    profile.update(infer_business_region(
        item,
        text,
        company_name=profile["company_name"],
        lab=profile["lab"],
    ))

    if source == "arxiv":
        profile.update({
            "project_name": title,
            "company_name": "不适用",
            "paper_title": title,
            "paper_authors": "、".join(item.get("authors", [])[:8]) or "未披露",
            "paper_first_author": item.get("first_author", "") or "未披露",
            "paper_second_author": item.get("second_author", "") or "未披露",
            "paper_categories": "、".join(item.get("categories", [])) or "未披露",
            "paper_venue": item.get("journal_ref", "") or "arXiv/待发表",
            "paper_comment": item.get("comment", "") or "无",
            "paper_pdf_url": item.get("pdf_url", ""),
            "project_round": "不适用",
            "latest_valuation": "不适用",
            "financing_amount": "不适用",
            "institutions": "论文作者机构待从全文核实",
            "contact": link or "未披露/待核实",
        })
    elif source in {"wechat_api", "wechat_985"} or "高校" in item.get("source_group", ""):
        profile["lab"] = extract_lab(text, item.get("school", "") or item.get("source_name", ""))
        # 来源高校/公众号是项目归属线索，不等于投资机构。单独保存，避免
        # 下游把公众号名称写进 funding_rounds[].investors。
        profile["affiliated_institutions"] = item.get("source_name", "") or item.get("school", "")
    return profile


def attach_project_profile(item: dict) -> dict:
    if not isinstance(item.get("project_profile"), dict):
        item["project_profile"] = build_project_profile(item)
    else:
        profile = item["project_profile"]
        current_project_name = normalize_project_candidate(profile.get("project_name"))
        evidence_text = project_profile_text(item)
        refreshed_project_name = infer_project_name(item, evidence_text)
        compact_evidence = re.sub(r"\s+", "", f"{item.get('title', '')}\n{evidence_text}").lower()
        compact_current = re.sub(r"\s+", "", current_project_name).lower()
        compact_refreshed = re.sub(r"\s+", "", refreshed_project_name).lower()
        current_is_evidenced = bool(compact_current) and compact_current in compact_evidence
        refreshed_is_evidenced = (
            refreshed_project_name == "未命名项目"
            or (bool(compact_refreshed) and compact_refreshed in compact_evidence)
        )
        should_refresh_name = (
            not is_specific_project_subject_name(current_project_name)
            or not current_is_evidenced
            or (
                refreshed_project_name != current_project_name
                and is_specific_project_subject_name(refreshed_project_name)
                and refreshed_is_evidenced
                and (
                    compact_current in compact_refreshed
                    or compact_refreshed in re.sub(r"\s+", "", str(item.get("title", ""))).lower()
                )
            )
        )
        if should_refresh_name:
            profile["project_name"] = refreshed_project_name
        if not is_meaningful_company_name(profile.get("company_name")):
            company_name = "不适用" if item.get("source") == "arxiv" else extract_company_name(item, project_profile_text(item))
            profile["company_name"] = company_name or "未披露/待核实"
    if not normalize_business_region(item["project_profile"].get("region")):
        item["project_profile"].update(infer_business_region(
            item,
            project_profile_text(item),
            company_name=item["project_profile"].get("company_name", ""),
            lab=item["project_profile"].get("lab", ""),
        ))
    return item


def disclosure_status(combined: str, primary_hits: list[str], traction_hits: list[str]) -> dict:
    amount_hits = []
    for match in DISCLOSURE_AMOUNT_RE.finditer(combined):
        context = combined[max(0, match.start() - 24): match.end() + 24]
        if re.search(r"(融资|估值|金额|资金|投资|领投|跟投|轮|募资)", context):
            amount_hits.append(match.group(1).strip())
    investor_hits = phrase_hits(combined, ("投资方", "领投", "跟投", "投资机构", "资本", "基金", "股东"), 5)
    customer_hits = [hit for hit in traction_hits if hit in {"客户", "头部客户", "签约", "合作", "订单", "采购", "中标"}]
    return {
        "round_or_stage": "已披露: " + "、".join(primary_hits[:4]) if primary_hits else "未披露/待验证",
        "amount_or_valuation": "可能披露: " + "、".join(amount_hits[:4]) if amount_hits or contains_phrase(combined, "估值") else "未披露/待验证",
        "investors": "可能披露: " + "、".join(investor_hits[:4]) if investor_hits else "未披露/待验证",
        "ownership_or_cap_table": "通常不公开，待直接尽调",
        "revenue_or_customers": "可能披露: " + "、".join(customer_hits[:4]) if customer_hits else "未披露/待验证",
    }


def next_diligence_actions(score_dimensions: list[dict], disclosure: dict, is_filtered: bool) -> list[str]:
    if is_filtered:
        return ["无需进入项目池；仅在出现明确未上市公司、融资、产品或客户信号后再复核。"]
    actions = [
        "确认法律实体、是否未上市、核心股东和可接触窗口。",
        "核实融资轮次、金额、投资方、估值和资金用途；未披露字段通过直接沟通验证。",
    ]
    dimension_codes = {item["code"] for item in score_dimensions if item.get("score", 0) > 0}
    if "differentiation" in dimension_codes:
        actions.append("核查专利/论文/技术归属、排他许可和工程化成熟度。")
    if "traction" in dimension_codes:
        actions.append("访谈客户或合作方，确认试点、订单、收入和复购可能性。")
    if "team_institution" in dimension_codes:
        actions.append("确认创始团队、教授/实验室关系、过往产业化或创业履历。")
    actions.append("梳理直接竞品、上市/未上市可比公司和潜在退出路径。")
    return actions[:5]


def score_private_market_item(group: str, title: str, summary: str, source_name: str = "", source_type: str = "", link: str = "") -> dict:
    combined = "\n".join([title, summary, source_name, group])
    content_text = "\n".join([title, summary])
    primary_hits = phrase_hits(combined, PRIMARY_MARKET_TERMS)
    private_hits = phrase_hits(combined, PRIVATE_ENTITY_TERMS)
    tech_hits = phrase_hits(combined, TECH_DIFFERENTIATION_TERMS)
    team_hits = phrase_hits(combined, TEAM_INSTITUTION_TERMS)
    school_hits = phrase_hits(combined, UNIVERSITY_985, 6)
    market_hits = phrase_hits(combined, MARKET_URGENCY_TERMS)
    traction_hits = phrase_hits(combined, TRACTION_TERMS)
    timing_hits = phrase_hits(combined, TIMING_TERMS)
    secondary_hits = phrase_hits(combined, SECONDARY_MARKET_TERMS)
    funder_hits = phrase_hits(combined, MAJOR_FUNDERS, 5)
    raw_venue_hits = phrase_hits(combined, TOP_VENUES, 5)
    university_value_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_VALUE_TERMS, 10)
    university_strong_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_STRONG_VALUE_TERMS, 10)
    university_outcome_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_RESEARCH_OUTCOME_TERMS, 10)
    university_noise_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_NOISE_TERMS, 8)

    is_wechat = "公众号" in group or source_type == "gsdata_wechat"
    is_university_source = "高校" in group or bool(school_hits)
    is_university_wechat = is_wechat and is_university_source
    is_institution_source = "机构" in group or "创投新闻" in group or "海外项目" in group or source_type == "wechat_chat"
    venue_hits = university_top_venue_hits(content_text, raw_venue_hits, 5) if is_university_wechat else raw_venue_hits
    has_private_signal = bool(primary_hits or private_hits or traction_hits or tech_hits or university_value_hits or "高校成果" in group)
    secondary_only = bool(secondary_hits) and not has_private_signal
    explicit_financing_flash = source_type == "36kr_financing_flash" and bool(primary_hits)
    public_lead_text = "\n".join([title, summary[:1600]])
    low_value_public_lead = bool(LOW_VALUE_PUBLIC_LEAD_RE.search(public_lead_text))
    pure_academic_public_lead = bool(PURE_ACADEMIC_PUBLIC_LEAD_RE.search(public_lead_text))
    commercial_public_lead = bool(COMMERCIAL_PUBLIC_LEAD_RE.search(public_lead_text))
    real_investment_event = bool(REAL_INVESTMENT_EVENT_RE.search(public_lead_text))
    has_legal_company_subject = bool(re.search(r"[\u4e00-\u9fffA-Za-z0-9（）()·&＋+\-]{2,80}(?:股份有限公司|有限责任公司|有限公司)", title))

    fit_score = 0
    if primary_hits:
        fit_score += min(12, 6 + len(primary_hits) * 2)
    if private_hits:
        fit_score += min(8, 3 + len(private_hits) * 2)
    if "高校成果" in group or "专利" in group:
        fit_score += 5
    if is_university_wechat and university_value_hits:
        fit_score += min(6, 2 + len(university_value_hits))
    if is_institution_source:
        fit_score += 3
    fit_score = min(20, fit_score)

    differentiation_score = 0
    if tech_hits:
        differentiation_score += min(10, 4 + len(tech_hits) * 2)
    if is_university_wechat and university_value_hits:
        differentiation_score += min(6, 2 + len([hit for hit in university_value_hits if hit in UNIVERSITY_VALUE_TERMS]) // 2)
    if venue_hits:
        differentiation_score += 3
    if funder_hits:
        differentiation_score += 3
    differentiation_score = min(15, differentiation_score)

    team_score = 0
    if school_hits:
        team_score += min(8, 4 + len(school_hits))
    if team_hits:
        team_score += min(7, 3 + len(team_hits))
    if is_university_source:
        team_score += 3
    if is_university_wechat and university_value_hits and any(hit in content_text for hit in ("实验室", "重点实验室", "工程中心", "研究中心", "课题组", "教授", "院士", "博士")):
        team_score += 4
    team_score = min(15, team_score)

    market_score = 0
    if market_hits:
        market_score += min(11, 4 + len(market_hits) * 2)
    if any(hit in combined for hit in ("千亿", "百亿", "刚需", "国产替代", "出海")):
        market_score += 4
    if is_university_wechat and any(hit in content_text for hit in ("产业化", "成果转化", "临床", "医疗", "芯片", "半导体", "新能源", "新材料", "机器人", "人工智能", "大模型", "国产替代")):
        market_score += 4
    market_score = min(15, market_score)

    traction_score = 0
    if traction_hits:
        traction_score += min(12, 4 + len(traction_hits) * 2)
    if primary_hits:
        traction_score += 3
    traction_score = min(15, traction_score)

    timing_score = 0
    if timing_hits:
        timing_score += min(7, 3 + len(timing_hits))
    if re.search(r"202[5-9]|20[3-9]\d", combined):
        timing_score += 2
    if primary_hits or traction_hits:
        timing_score += 1
    timing_score = min(10, timing_score)

    source_score = 0
    if link:
        source_score += 2
    if source_type == "gsdata_wechat":
        source_score += 6
    elif source_type == "wechat_chat":
        source_score += 6
    elif is_wechat:
        source_score += 5
    elif group in {"高校成果", "专利", "论文"}:
        source_score += 6
    elif group in {"创投新闻", "海外项目"}:
        source_score += 5
    else:
        source_score += 3
    source_score = min(10, source_score)

    dimensions = [
        dimension_item("private_market_fit", "一级市场可投资性", fit_score, 20, "、".join((primary_hits + private_hits)[:8]) or "未发现明确融资/未上市主体信号", (primary_hits + private_hits)[:8]),
        dimension_item("differentiation", "技术/产品差异化", differentiation_score, 15, "、".join((tech_hits + venue_hits + funder_hits)[:8]) or "未发现强技术资产信号", (tech_hits + venue_hits + funder_hits)[:8]),
        dimension_item("team_institution", "团队/高校/机构质量", team_score, 15, "、".join((school_hits + team_hits)[:8]) or "团队/机构信息不足", (school_hits + team_hits)[:8]),
        dimension_item("market_urgency", "市场空间和需求紧迫性", market_score, 15, "、".join(market_hits[:8]) or "市场需求信息不足", market_hits[:8]),
        dimension_item("traction", "客户/融资/合作验证", traction_score, 15, "、".join((traction_hits + primary_hits)[:8]) or "验证信号不足", (traction_hits + primary_hits)[:8]),
        dimension_item("timing", "时间窗口", timing_score, 10, "、".join(timing_hits[:8]) or "近期催化不明显", timing_hits[:8]),
        dimension_item("source_quality", "来源可信度", source_score, 10, source_name or group or "未知来源", [source_name or group]),
    ]
    raw_score = sum(item["score"] for item in dimensions)
    score = raw_score
    boost_reasons = []
    strong_private_financing = fit_score >= 18 and traction_score >= 12
    strong_explicit_financing = (
        bool(primary_hits)
        and (
            explicit_financing_flash
            or re.search(r"(完成|宣布|获得|获|已完成).{0,12}融资", combined)
            or re.search(r"(天使轮|种子轮|Pre-?A|A轮|B轮|C轮|D轮|首轮|新一轮)", combined, flags=re.IGNORECASE)
        )
    )
    strong_university_commercialization = (
        is_university_source
        and differentiation_score >= 10
        and (
            contains_phrase(combined, "成果转化")
            or contains_phrase(combined, "技术转移")
            or contains_phrase(combined, "产业化")
            or "专利" in tech_hits
            or "发明" in tech_hits
        )
    )
    weak_university_terms = {"论文", "人工智能", "AI", "实验室", "博士", "教授", "大模型", "机器人", "临床", "生物医药", "医疗器械", "项目"}
    strong_university_terms = unique_keep_order(university_strong_hits + [hit for hit in university_value_hits if hit not in weak_university_terms], 10)
    university_weak_research_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_WEAK_RESEARCH_TERMS, 8)
    university_concrete_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_CONCRETE_OUTCOME_TERMS, 10)
    university_team_context_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_TEAM_CONTEXT_TERMS, 8)
    university_domain_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_INVESTABLE_TECH_DOMAIN_TERMS, 8)
    university_hard_noise_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_HARD_NOISE_TERMS, 8)
    university_hard_noise_rescue_hits = phrase_hits(content_text if "高校" in group else combined, UNIVERSITY_HARD_NOISE_RESCUE_TERMS, 8)
    hard_noise_has_real_outcome = bool(primary_hits or venue_hits or university_hard_noise_rescue_hits)
    university_contextual_breakthrough = bool(
        re.search(r"(科研|临床痛点|技术|成果|团队|实验室|课题组|项目).{0,24}(重要突破|重大突破|突破)", content_text)
        or re.search(r"(重要突破|重大突破|突破).{0,24}(科研|临床|技术|成果|团队|实验室|课题组|项目)", content_text)
    )
    concrete_university_outcome = bool(
        primary_hits
        or venue_hits
        or university_concrete_hits
        or university_hard_noise_rescue_hits
        or university_contextual_breakthrough
        or (university_outcome_hits and university_team_context_hits)
        or (university_outcome_hits and university_domain_hits)
        or (
            university_weak_research_hits
            and (
                bool(university_team_context_hits)
                or bool(venue_hits)
                or bool(university_concrete_hits)
            )
        )
    )
    has_university_outcome = concrete_university_outcome
    noise_has_specific_university_outcome = bool(
        primary_hits
        or venue_hits
        or university_concrete_hits
        or university_hard_noise_rescue_hits
        or university_team_context_hits
    )
    noisy_university_content = bool(university_noise_hits) and not noise_has_specific_university_outcome
    hard_noisy_university_content = bool(university_hard_noise_hits) and not hard_noise_has_real_outcome
    generic_university_tech_content = (
        is_university_wechat
        and not primary_hits
        and not concrete_university_outcome
        and bool(tech_hits or market_hits or traction_hits)
    )
    strong_university_tech_value = (
        is_university_wechat
        and bool(university_value_hits)
        and not noisy_university_content
        and (
            (bool(strong_university_terms) and differentiation_score >= 8 and team_score >= 7)
            or (bool(strong_university_terms) and differentiation_score >= 8 and market_score >= 6)
            or contains_phrase(content_text, "成果转化")
            or contains_phrase(content_text, "产业化")
            or contains_phrase(content_text, "技术转移")
            or contains_phrase(content_text, "专利")
            or contains_phrase(content_text, "样机")
            or contains_phrase(content_text, "临床")
            or contains_phrase(content_text, "注册证")
            or contains_phrase(content_text, "中试")
            or bool(venue_hits)
            or bool(university_concrete_hits)
            or bool(university_hard_noise_rescue_hits)
            or university_contextual_breakthrough
            or bool(university_outcome_hits and (university_team_context_hits or university_domain_hits))
            or bool(university_weak_research_hits and concrete_university_outcome)
        )
    )
    if not secondary_only and strong_private_financing:
        score = max(score, 65)
        boost_reasons.append("强一级市场融资/客户验证信号，早期披露不完整时仍进入保留池。")
    if not secondary_only and strong_explicit_financing:
        score = max(score, 62)
        boost_reasons.append("明确一级市场融资事件，融资金额/估值/客户等披露不完整时仍进入保留池。")
    if not secondary_only and strong_university_commercialization:
        score = max(score, 62)
        boost_reasons.append("高校成果转化+强技术资产信号，适合进入观察池。")
    if not secondary_only and strong_university_tech_value:
        score = max(score, UNIVERSITY_TECH_RETAIN_SCORE)
        boost_reasons.append("高校公众号科技成果有潜在高价值，即使未披露融资/公司也进入观察池。")

    filter_reasons = []
    if secondary_only:
        filter_reasons.append("主要是二级市场/上市公司交易信息，未发现明确一级市场机会。")
    if is_university_wechat and low_value_public_lead and not real_investment_event:
        filter_reasons.append("高校来源中的获奖、教学、课程或人物任职资讯，没有明确公司、融资或估值，不构成投资线索。")
    if (
        is_university_wechat
        and pure_academic_public_lead
        and not commercial_public_lead
        and not (has_legal_company_subject or real_investment_event)
    ):
        filter_reasons.append("高校来源中的纯论文或课题组研究，缺少公司主体、投资事实或成果转化信号。")
    if (
        is_university_wechat
        and not has_legal_company_subject
        and not real_investment_event
        and not commercial_public_lead
    ):
        filter_reasons.append("高校来源缺少公司主体、投资事实或成果转化/产业化信号，不构成可执行投资线索。")
    if is_university_wechat and hard_noisy_university_content:
        filter_reasons.append("高校来源中的招生、课程、资源、活动或转载类内容，缺少可投资科技成果线索。")
    if is_university_wechat and noisy_university_content and not primary_hits:
        filter_reasons.append("高校来源中的泛资讯/活动/课程内容，未发现具体科研成果、专利、论文、实验室项目或转化信号。")
    if is_university_wechat and not primary_hits and not concrete_university_outcome and score >= UNIVERSITY_TECH_RETAIN_SCORE:
        filter_reasons.append("高校来源缺少可追踪的具体科研成果、顶刊论文、专利、实验室项目、样机、临床研究或产业化线索。")
    if generic_university_tech_content and score >= PRIVATE_MARKET_RETAIN_SCORE:
        filter_reasons.append("高校来源中的泛科技趋势/工具/活动内容，缺少具体团队成果、专利、顶刊、样机、临床研究或产业化线索。")
    if not has_private_signal and score < 40:
        filter_reasons.append("没有具体未上市公司、成果转化、融资、产品或客户验证信号。")

    if filter_reasons:
        score = min(score, 35)
        decision = "filter"
        decision_label = "过滤"
    elif score >= 80:
        decision = "high_priority"
        decision_label = "高优先级"
    elif score >= PRIVATE_MARKET_RETAIN_SCORE or (strong_university_tech_value and score >= UNIVERSITY_TECH_RETAIN_SCORE):
        decision = "watchlist"
        decision_label = "保留观察"
    elif score >= 40:
        decision = "low_priority"
        decision_label = "低优先级"
    else:
        decision = "filter"
        decision_label = "过滤"
        filter_reasons.append("综合分低于保留阈值。")

    worth_attention = decision in {"high_priority", "watchlist"}
    disclosure = disclosure_status(combined, primary_hits, traction_hits)
    signals = []
    for item in dimensions:
        if item["score"] > 0:
            signals.append({
                "code": item["code"],
                "score": item["score"],
                "detail": f"{item['label']}: {item['detail']}",
            })
    if secondary_hits:
        signals.append({"code": "secondary_market_context", "score": 0, "detail": "二级市场词: " + "、".join(secondary_hits[:8])})
    if strong_university_tech_value:
        signals.append({"code": "university_tech_value", "score": max(0, score - raw_score), "detail": "高校科技前景信号: " + "、".join(university_value_hits[:10])})
    if score > raw_score:
        for reason in boost_reasons:
            signals.append({"code": "private_market_retention_boost", "score": score - raw_score, "detail": reason})
    if filter_reasons:
        signals.append({"code": "filtered_reason", "score": 0, "detail": "；".join(filter_reasons)})

    top_reasons = [item["detail"] for item in dimensions if item["score"] >= max(5, item["max_score"] * 0.45)]
    if worth_attention:
        thesis = "；".join(top_reasons[:3]) or "存在一级市场相关信号，建议进入观察池。"
    else:
        thesis = "；".join(filter_reasons[:2]) or "暂未达到一级市场保留阈值。"

    return {
        "score": min(100, score),
        "worth_attention": worth_attention,
        "signals": signals,
        "decision": decision,
        "decision_label": decision_label,
        "skill_version": SKILL_SCORER_VERSION,
        "score_dimensions": dimensions,
        "filter_reasons": filter_reasons,
        "disclosure_status": disclosure,
        "primary_market_hits": primary_hits,
        "secondary_market_hits": secondary_hits,
        "private_market_thesis": thesis,
        "next_actions": next_diligence_actions(dimensions, disclosure, decision == "filter"),
    }


def score_extra_fields(score: dict) -> dict:
    return {
        "decision": score.get("decision", ""),
        "decision_label": score.get("decision_label", ""),
        "skill_version": score.get("skill_version", SKILL_SCORER_VERSION),
        "score_dimensions": score.get("score_dimensions", []),
        "filter_reasons": score.get("filter_reasons", []),
        "disclosure_status": score.get("disclosure_status", {}),
        "primary_market_hits": score.get("primary_market_hits", []),
        "secondary_market_hits": score.get("secondary_market_hits", []),
        "private_market_thesis": score.get("private_market_thesis", ""),
        "next_actions": score.get("next_actions", []),
    }


def score_item(title: str, summary: str, comment: str, journal_ref: str, authors: list[str], watch_authors: list[str]) -> dict:
    signals = []
    combined = "\n".join([title, summary, comment, journal_ref])
    first_author = authors[0] if authors else ""
    second_author = authors[1] if len(authors) > 1 else ""
    watch = {normalize_name(x) for x in watch_authors if x.strip()}

    hits = [a for a in [first_author, second_author] if normalize_name(a) in watch]
    if hits:
        signals.append({"code": "watch_author_first_second", "score": 35, "detail": "重点作者位于第一/第二作者: " + ", ".join(hits)})

    venues = [venue for venue in TOP_VENUES if contains_phrase(combined, venue)]
    if venues:
        signals.append({"code": "top_venue", "score": 35, "detail": "高水平会议/期刊信号: " + ", ".join(venues[:5])})

    funders = [funder for funder in MAJOR_FUNDERS if contains_phrase(combined, funder)]
    if funders:
        signals.append({"code": "major_funding", "score": 30, "detail": "重大基金/机构支持信号: " + ", ".join(funders[:5])})

    schools = [school for school in UNIVERSITY_985 if contains_phrase(combined, school)]
    if schools:
        signals.append({"code": "affiliation_985", "score": 15, "detail": "985 院校信号: " + ", ".join(schools[:6])})

    if re.search(r"(?i)(github\.com|code is available|source code|dataset is available|huggingface\.co)", combined):
        signals.append({"code": "artifact_available", "score": 8, "detail": "代码、模型或数据集可用"})

    paper_score = min(100, 20 + sum(x["score"] for x in signals))
    private_score = score_private_market_item("论文", title, combined, "arXiv", "arxiv")
    score = max(paper_score, private_score["score"])
    worth_attention = (
        score >= PRIVATE_MARKET_RETAIN_SCORE
        or private_score["worth_attention"]
        or any(x["code"] in {"watch_author_first_second", "top_venue", "major_funding"} for x in signals)
    )
    if worth_attention and private_score["decision"] == "filter":
        private_score["decision"] = "watchlist"
        private_score["decision_label"] = "保留观察"
        private_score["filter_reasons"] = []
    return {
        "score": score,
        "worth_attention": worth_attention,
        "first_author": first_author,
        "second_author": second_author,
        "signals": signals + private_score["signals"],
        "decision": private_score["decision"],
        "decision_label": private_score["decision_label"],
        "skill_version": private_score["skill_version"],
        "score_dimensions": private_score["score_dimensions"],
        "filter_reasons": private_score["filter_reasons"],
        "disclosure_status": private_score["disclosure_status"],
        "primary_market_hits": private_score["primary_market_hits"],
        "secondary_market_hits": private_score["secondary_market_hits"],
        "private_market_thesis": private_score["private_market_thesis"],
        "next_actions": private_score["next_actions"],
    }


def build_query(categories: list[str], keywords: list[str], days: int) -> str:
    parts = []
    if categories:
        parts.append("(" + " OR ".join(f"cat:{x.strip()}" for x in categories if x.strip()) + ")")
    if keywords:
        parts.append("(" + " OR ".join(f'all:"{x.strip().replace(chr(34), "")}"' for x in keywords if x.strip()) + ")")
    end = datetime.now(timezone.utc)
    start = end - timedelta(days=days)
    parts.append(f"submittedDate:[{start:%Y%m%d%H%M} TO {end:%Y%m%d%H%M}]")
    return " AND ".join(parts)


def append_jsonl(path: Path, rows: list[dict]) -> int:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    seen = set()
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                item = json.loads(line)
                key = source_key(item)
                if key:
                    seen.add(key)
            except Exception:
                pass
    written = 0
    with path.open("a", encoding="utf-8") as fh:
        for row in rows:
            key = source_key(row)
            if key and key in seen:
                continue
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")
            if key:
                seen.add(key)
            written += 1
    return written


def upsert_jsonl(path: Path, rows: list[dict]) -> int:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    existing = []
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                item = json.loads(line)
            except Exception:
                continue
            if source_key(item):
                existing.append(item)
    merged = {source_key(item): item for item in existing if source_key(item)}
    written = 0
    for row in rows:
        key = source_key(row)
        if not key:
            continue
        if key not in merged:
            written += 1
        merged[key] = row
    ordered = sorted(
        merged.values(),
        key=lambda item: (item.get("worth_attention", False), item.get("attention_score", 0), item.get("published_at", "")),
        reverse=True,
    )
    path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in ordered) + ("\n" if ordered else ""), encoding="utf-8")
    return written


def replace_chat_candidates_for_groups(path: Path, rows: list[dict], group_keys: set[tuple[str, str]]) -> int:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    existing = []
    if path.exists():
        for line in path.read_text(encoding="utf-8").splitlines():
            try:
                item = json.loads(line)
            except Exception:
                continue
            item_date = clean_text(item.get("date", "")) or clean_text(item.get("published_at", ""))[:10]
            item_group_key = clean_text(item.get("source_key", ""))
            if item.get("source") == "wechat_chat" and (item_date, item_group_key) in group_keys:
                continue
            existing.append(item)
    before_keys = {source_key(item) for item in existing if source_key(item)}
    merged = {source_key(item): item for item in existing if source_key(item)}
    for row in rows:
        key = source_key(row)
        if key:
            merged[key] = row
    ordered = sorted(
        merged.values(),
        key=lambda item: (item.get("worth_attention", False), item.get("attention_score", 0), item.get("published_at", "")),
        reverse=True,
    )
    path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in ordered) + ("\n" if ordered else ""), encoding="utf-8")
    after_keys = {source_key(item) for item in ordered if source_key(item)}
    return len(after_keys - before_keys)


def safe_filename(value: str) -> str:
    value = clean_text(value)
    value = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", value).strip("._")
    return value[:80] or "unknown"


def parse_message_datetime(value: str) -> datetime | None:
    value = clean_text(value)
    if not value:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(value[:19], fmt)
        except Exception:
            pass
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return None


def chat_message_date(row: dict) -> str:
    for field in ("message_time", "send_time", "msg_time", "pushed_at"):
        parsed = parse_message_datetime(row.get(field, ""))
        if parsed:
            return parsed.date().isoformat()
    return today_china().date().isoformat()


def chat_message_time(row: dict) -> str:
    for field in ("message_time", "send_time", "msg_time"):
        value = clean_text(row.get(field, ""))
        if value:
            return value
    return clean_text(row.get("pushed_at", "")) or utc_now_iso()


def normalize_chat_file(value: WechatChatFile | dict[str, Any] | None) -> dict:
    if isinstance(value, WechatChatFile):
        value = value.dict()
    if not isinstance(value, dict):
        return {}
    return {
        "file_serial_no": clean_text(str(value.get("file_serial_no", ""))),
        "file_name": clean_text(str(value.get("file_name", ""))),
        "file_url": clean_text(str(value.get("file_url", ""))),
    }


def decode_raw_msg_content(value: str) -> str:
    value = clean_text(value)
    if not value:
        return ""
    try:
        return base64.b64decode(value).decode("utf-8", errors="ignore").strip()
    except Exception:
        return ""


def normalize_chat_message(payload: WechatChatPushRequest, message: WechatChatMessageIn) -> dict:
    data = message.dict()
    file_data = normalize_chat_file(data.get("file"))
    content = (
        clean_text(data.get("msg_content_decoded", ""))
        or html_to_text(data.get("msg_content", ""))
        or html_to_text(data.get("content", ""))
        or decode_raw_msg_content(data.get("raw_msg_content", ""))
    )
    cite_content = html_to_text(data.get("cite_content", ""))
    msg_key = clean_text(data.get("msg_key", ""))
    group_name = clean_text(data.get("group_name", "")) or "未命名群"
    group_serial_no = clean_text(data.get("group_serial_no", "")) or normalize_name(group_name)
    sender_name = clean_text(data.get("sender_name", "")) or "未知成员"
    message_time = clean_text(data.get("message_time", "")) or clean_text(data.get("send_time", "")) or clean_text(data.get("msg_time", ""))
    source_id = hashlib.md5(
        f"{payload.merchant_no}:{group_serial_no}:{msg_key or sender_name}:{message_time}:{content}".encode("utf-8")
    ).hexdigest()
    row = {
        "source": "wechat_chat_message",
        "source_id": f"chatmsg:{source_id}",
        "merchant_no": clean_text(payload.merchant_no),
        "pushed_at": clean_text(payload.pushed_at),
        "received_at": utc_now_iso(),
        "msg_key": msg_key,
        "group_name": group_name,
        "group_serial_no": group_serial_no,
        "sender_name": sender_name,
        "sender_serial_no": clean_text(data.get("sender_serial_no", "")),
        "cite_content": cite_content,
        "msg_content": content,
        "file": file_data,
        "msg_time": clean_text(data.get("msg_time", "")),
        "send_time": clean_text(data.get("send_time", "")),
        "message_time": message_time,
        "msg_type": data.get("msg_type", ""),
    }
    row["date"] = chat_message_date(row)
    row["sort_time"] = chat_message_time(row)
    return row


def chat_daily_path(date_value: str, group_name: str, group_serial_no: str = "") -> Path:
    group_part = safe_filename(group_serial_no or group_name)
    return WECHAT_CHAT_MESSAGES_DIR / date_value / f"{group_part}.jsonl"


def chat_candidate_group_path(date_value: str, group_name: str, group_serial_no: str = "") -> Path:
    group_part = safe_filename(group_serial_no or group_name)
    return WECHAT_CHAT_CANDIDATES_DIR / date_value / f"{group_part}.jsonl"


def read_chat_daily_messages(date_value: str, group_name: str, group_serial_no: str = "") -> list[dict]:
    path = chat_daily_path(date_value, group_name, group_serial_no)
    rows = []
    if not path.exists():
        return rows
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            rows.append(json.loads(line))
        except Exception:
            continue
    return sorted(rows, key=lambda item: item.get("sort_time", ""))


CHAT_INVESTMENT_HINT_TERMS = (
    "融资", "天使轮", "种子轮", "Pre-A", "A轮", "B轮", "估值", "BP", "商业计划书",
    "路演", "项目", "尽调", "看项目", "推荐项目", "FA", "投资人", "创始人", "创业公司",
    "未上市", "客户", "订单", "量产", "专利", "临床", "注册证", "样机", "芯片",
    "机器人", "人工智能", "大模型", "生物医药", "医疗器械", "新材料", "新能源",
)


def chat_message_signal_text(row: dict) -> str:
    file_data = row.get("file") if isinstance(row.get("file"), dict) else {}
    file_text = " ".join(clean_text(str(file_data.get(key, ""))) for key in ("file_name", "file_url"))
    return "\n".join([
        clean_text(row.get("cite_content", "")),
        clean_text(row.get("msg_content", "")),
        file_text,
    ]).strip()


def chat_context_rows(rows: list[dict], index: int) -> list[dict]:
    start = max(0, index - CHAT_CONTEXT_BEFORE)
    end = min(len(rows), index + CHAT_CONTEXT_AFTER + 1)
    return rows[start:end]


def format_chat_context(rows: list[dict], key_source_id: str = "") -> str:
    lines = []
    for row in rows:
        marker = " *" if row.get("source_id") == key_source_id else ""
        text = chat_message_signal_text(row)
        if not text:
            continue
        lines.append(f"[{chat_message_time(row)}] {row.get('sender_name', '未知成员')}{marker}: {text}")
    return "\n".join(lines)


def chat_candidate_title(row: dict, context_text: str) -> str:
    project_name = infer_project_name({"title": clean_text(row.get("msg_content", ""))}, context_text)
    project_name = project_name if project_name and project_name != "未命名项目" else clean_text(row.get("msg_content", ""))[:32]
    return f"群聊线索｜{project_name or row.get('group_name', '微信群聊')}"


def build_chat_candidate(row: dict, context_rows: list[dict]) -> dict | None:
    signal_text = chat_message_signal_text(row)
    context_text = format_chat_context(context_rows, row.get("source_id", ""))
    scoring_text = "\n".join([signal_text, context_text])
    if not clean_text(scoring_text):
        return None
    row_hint_hits = phrase_hits(signal_text, CHAT_INVESTMENT_HINT_TERMS, 10)
    if not row_hint_hits:
        return None
    hint_hits = unique_keep_order(row_hint_hits + phrase_hits(scoring_text, CHAT_INVESTMENT_HINT_TERMS, 10), 10)
    title = chat_candidate_title(row, scoring_text)
    group_name = row.get("group_name", "微信群聊")
    score = score_private_market_item("微信群聊", title, scoring_text, group_name, "wechat_chat")
    if not score["worth_attention"] and score["score"] >= 45 and hint_hits:
        score["score"] = max(score["score"], PRIVATE_MARKET_RETAIN_SCORE)
        score["worth_attention"] = True
        score["decision"] = "watchlist"
        score["decision_label"] = "保留观察"
        score["filter_reasons"] = []
        score["signals"].append({"code": "wechat_chat_hint", "score": 8, "detail": "群聊投资线索词: " + "、".join(hint_hits[:8])})
        score["private_market_thesis"] = "群聊中出现投资相关线索，需结合上下文人工核实。"
    if not score["worth_attention"]:
        return None
    source_id = row.get("source_id", "").replace("chatmsg:", "wechat_chat:")
    key_time = chat_message_time(row)
    summary = (
        f"{row.get('sender_name', '未知成员')} 在 {key_time} 于「{group_name}」提出关键信息。"
        f"命中信号：{'、'.join(hint_hits[:8])}。"
    )
    item = {
        "source": "wechat_chat",
        "source_id": source_id,
        "fingerprint": hashlib.md5(source_id.encode("utf-8")).hexdigest()[:16],
        "title": title,
        "summary": summary,
        "article_text": "关键信息\n" + signal_text + "\n\n聊天前后文\n" + context_text,
        "article_text_length": len(context_text),
        "source_name": group_name,
        "source_group": "微信群聊",
        "source_key": row.get("group_serial_no", "") or group_name,
        "source_type": "wechat_chat",
        "date": row.get("date", "") or key_time[:10],
        "merchant_no": row.get("merchant_no", ""),
        "group_name": group_name,
        "group_serial_no": row.get("group_serial_no", ""),
        "sender_name": row.get("sender_name", ""),
        "sender_serial_no": row.get("sender_serial_no", ""),
        "key_message_time": key_time,
        "published_at": key_time,
        "updated_at": row.get("received_at", ""),
        "categories": ["微信群聊", group_name, row.get("sender_name", "")],
        "chat_context": [
            {
                "time": chat_message_time(context_row),
                "sender_name": context_row.get("sender_name", ""),
                "content": chat_message_signal_text(context_row),
                "is_key_message": context_row.get("source_id") == row.get("source_id"),
            }
            for context_row in context_rows
            if chat_message_signal_text(context_row)
        ],
        "file": row.get("file", {}),
        "attention_score": score["score"],
        "worth_attention": score["worth_attention"],
        "signals": score["signals"],
        **score_extra_fields(score),
        "collected_at": utc_now_iso(),
    }
    return attach_project_profile(item)


def mine_chat_candidates(date_value: str, group_name: str, group_serial_no: str = "") -> list[dict]:
    rows = read_chat_daily_messages(date_value, group_name, group_serial_no)
    candidates = []
    for index, row in enumerate(rows):
        candidate = build_chat_candidate(row, chat_context_rows(rows, index))
        if candidate:
            candidates.append(candidate)
    deduped = {source_key(item): item for item in candidates if source_key(item)}
    return sorted(deduped.values(), key=lambda item: (item.get("attention_score", 0), item.get("published_at", "")), reverse=True)


def ingest_wechat_chat_push(req: WechatChatPushRequest) -> dict:
    normalized = [normalize_chat_message(req, message) for message in req.messages]
    normalized = [row for row in normalized if row.get("msg_content") or row.get("cite_content") or row.get("file")]
    by_path: dict[Path, list[dict]] = {}
    for row in normalized:
        path = chat_daily_path(row["date"], row["group_name"], row.get("group_serial_no", ""))
        by_path.setdefault(path, []).append(row)
    written_messages = 0
    for path, rows in by_path.items():
        written_messages += append_jsonl(path, rows)
    affected_groups = {
        (row["date"], row["group_name"], row.get("group_serial_no", ""))
        for row in normalized
    }
    candidates = []
    group_candidate_files = []
    for date_value, group_name, group_serial_no in affected_groups:
        group_candidates = mine_chat_candidates(date_value, group_name, group_serial_no)
        group_candidate_path = chat_candidate_group_path(date_value, group_name, group_serial_no)
        group_candidate_path.parent.mkdir(parents=True, exist_ok=True)
        group_candidate_path.write_text(
            "\n".join(json.dumps(row, ensure_ascii=False) for row in group_candidates) + ("\n" if group_candidates else ""),
            encoding="utf-8",
        )
        group_candidate_files.append(str(group_candidate_path))
        candidates.extend(group_candidates)
    group_keys = {(date_value, group_serial_no or normalize_name(group_name)) for date_value, group_name, group_serial_no in affected_groups}
    written_candidates = replace_chat_candidates_for_groups(WECHAT_CHAT_CANDIDATES_FILE, candidates, group_keys)
    return {
        "received": len(req.messages),
        "stored": written_messages,
        "groups": len(affected_groups),
        "candidates": len(candidates),
        "written_candidates": written_candidates,
        "candidate_file": str(WECHAT_CHAT_CANDIDATES_FILE),
        "group_candidate_files": group_candidate_files,
        "message_files": [str(path) for path in sorted(by_path)],
        "items": candidates[:50],
    }


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def default_auto_status() -> dict:
    return {
        "enabled": AUTO_CRAWL_ENABLED,
        "running": False,
        "interval_seconds": AUTO_CRAWL_INTERVAL_SECONDS,
        "groups": AUTO_CRAWL_GROUPS,
        "last_started_at": "",
        "last_finished_at": "",
        "next_run_at": "",
        "last_result": None,
        "last_error": "",
        "consecutive_error_runs": 0,
        "run_count": 0,
    }


def read_auto_status() -> dict:
    if not AUTO_STATUS_FILE.exists():
        return default_auto_status()
    try:
        status = json.loads(AUTO_STATUS_FILE.read_text(encoding="utf-8"))
    except Exception:
        status = default_auto_status()
    merged = default_auto_status()
    merged.update(status)
    return merged


def write_auto_status(status: dict) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    AUTO_STATUS_FILE.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")


def default_wechat_daily_status() -> dict:
    return {
        "enabled": WECHAT_DAILY_ENABLED,
        "running": False,
        "run_hour": WECHAT_DAILY_RUN_HOUR,
        "run_minute": WECHAT_DAILY_RUN_MINUTE,
        "groups": ["高校", "机构"],
        "days": 7,
        "last_started_at": "",
        "last_finished_at": "",
        "last_target_date": "",
        "last_date_end": "",
        "next_run_at": "",
        "last_result": None,
        "last_error": "",
        "consecutive_error_runs": 0,
        "run_count": 0,
    }


def read_wechat_daily_status() -> dict:
    if not WECHAT_DAILY_STATUS_FILE.exists():
        return default_wechat_daily_status()
    try:
        status = json.loads(WECHAT_DAILY_STATUS_FILE.read_text(encoding="utf-8"))
    except Exception:
        status = default_wechat_daily_status()
    merged = default_wechat_daily_status()
    merged.update(status)
    return merged


def write_wechat_daily_status(status: dict) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    WECHAT_DAILY_STATUS_FILE.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")


def today_china() -> datetime:
    return datetime.now(CHINA_TZ)


def yesterday_china_date() -> str:
    return (today_china().date() - timedelta(days=1)).isoformat()


def wechat_default_start_date(days: int = 7) -> str:
    return (today_china().date() - timedelta(days=max(1, days))).isoformat()


def next_wechat_daily_run_at(now: datetime | None = None) -> datetime:
    now = now or today_china()
    candidate = now.replace(hour=WECHAT_DAILY_RUN_HOUR, minute=WECHAT_DAILY_RUN_MINUTE, second=0, microsecond=0)
    if candidate <= now:
        candidate = candidate + timedelta(days=1)
    return candidate


def load_gsdata_credentials() -> tuple[str, str]:
    app_key = os.getenv("GSDATA_APP_KEY", "").strip()
    app_secret = os.getenv("GSDATA_APP_SECRET", "").strip()
    if app_key and app_secret:
        return app_key, app_secret
    if GSDATA_CREDENTIALS_FILE.exists():
        data = json.loads(GSDATA_CREDENTIALS_FILE.read_text(encoding="utf-8"))
        app_key = clean_text(data.get("app_key", ""))
        app_secret = clean_text(data.get("app_secret", ""))
    if not app_key or not app_secret:
        raise RuntimeError("缺少 GSData app_key/app_secret，请设置环境变量或 data/gsdata_credentials.json。")
    return app_key, app_secret


def gsdata_credentials_configured() -> bool:
    try:
        load_gsdata_credentials()
        return True
    except (OSError, ValueError, RuntimeError, json.JSONDecodeError):
        return False


GSDATA_HEALTH_CACHE_SECONDS = 5 * 60
gsdata_health_cache: dict[str, Any] = {"checked_monotonic": 0.0, "result": None}


def probe_gsdata_health(force: bool = False) -> dict[str, Any]:
    now = time.monotonic()
    cached = gsdata_health_cache.get("result")
    if (
        not force
        and isinstance(cached, dict)
        and now - float(gsdata_health_cache.get("checked_monotonic", 0.0)) < GSDATA_HEALTH_CACHE_SECONDS
    ):
        return dict(cached)
    if not gsdata_credentials_configured():
        result = {"ok": False, "status": "not_configured", "error": "GSData 凭据未配置"}
    else:
        accounts = load_wechat_api_accounts()
        if not accounts:
            result = {"ok": False, "status": "accounts_missing", "error": "公众号账号清单为空"}
        else:
            start, end = date_range_for_wechat_api(yesterday_china_date(), 1)
            params = {
                "wx_name": accounts[0]["wx_name"],
                "posttime_start": start,
                "posttime_end": end,
                "order": "desc",
                "sort": "posttime",
                "page": "1",
                "limit": "1",
            }
            token = gsdata_access_token(params, GSDATA_WECHAT_ROUTER)
            try:
                response = httpx.get(
                    GSDATA_API_URL,
                    params=params,
                    headers={"access-token": token},
                    timeout=8,
                )
                response.raise_for_status()
                payload = response.json()
                if payload.get("success"):
                    result = {"ok": True, "status": "authenticated", "error": ""}
                else:
                    result = {
                        "ok": False,
                        "status": "rejected",
                        "error": clean_text(str(payload.get("msg") or payload.get("message") or "GSData 返回失败")),
                    }
            except Exception as exc:
                result = {"ok": False, "status": "unreachable", "error": clean_text(str(exc))[:300]}
    result["checked_at"] = utc_now_iso()
    gsdata_health_cache["checked_monotonic"] = now
    gsdata_health_cache["result"] = dict(result)
    return result


def gsdata_access_token(params: dict[str, str], router: str) -> str:
    app_key, app_secret = load_gsdata_credentials()
    joined = "".join(f"{key}{params[key]}" for key in sorted(params))
    string_a = f"_{joined}_"
    sign = hashlib.md5(f"{app_secret}{string_a}{app_secret}".encode("utf-8")).hexdigest()
    return base64.b64encode(f"{app_key}:{sign}:{router}".encode("utf-8")).decode("ascii")


def load_wechat_api_accounts() -> list[dict]:
    if not WECHAT_ACCOUNTS_XLSX.exists():
        return []
    wb = load_workbook(WECHAT_ACCOUNTS_XLSX, read_only=False, data_only=True)
    accounts = []
    for ws in wb.worksheets:
        headers = [clean_text(str(ws.cell(1, col).value or "")) for col in range(1, ws.max_column + 1)]
        name_col = next((idx + 1 for idx, value in enumerate(headers) if value == "公众号"), 1)
        wx_col = next((idx + 1 for idx, value in enumerate(headers) if value in {"帐号名", "账号名", "微信号"}), 2)
        group = "机构" if "机构" in ws.title else "高校"
        for row in range(2, ws.max_row + 1):
            account_name = clean_text(str(ws.cell(row, name_col).value or ""))
            wx_name = clean_text(str(ws.cell(row, wx_col).value or ""))
            if not account_name or not wx_name:
                continue
            accounts.append({
                "group": group,
                "sheet": ws.title,
                "account_name": account_name,
                "wx_name": wx_name,
            })
    return accounts


def filter_wechat_api_accounts(req: WechatApiRunRequest) -> list[dict]:
    accounts = load_wechat_api_accounts()
    groups = {group for group in req.groups if group.strip()}
    wx_names = {name.casefold() for name in req.wx_names if name.strip()}
    if groups:
        accounts = [account for account in accounts if account["group"] in groups]
    if wx_names:
        accounts = [account for account in accounts if account["wx_name"].casefold() in wx_names]
    if req.max_accounts:
        accounts = accounts[: req.max_accounts]
    return accounts


def date_range_for_wechat_api(date_value: str, days: int = 1) -> tuple[str, str]:
    target = datetime.strptime(date_value, "%Y-%m-%d").date()
    start = datetime.combine(target, datetime.min.time())
    end = start + timedelta(days=max(1, days))
    return start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S")


def normalize_article_line(text: str) -> str:
    text = re.sub(r"[\u200b-\u200f\ufeff]", "", text or "")
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", text)
    text = re.sub(r"\s+([，。！？；：、）》】])", r"\1", text)
    text = re.sub(r"([（《【])\s+", r"\1", text)
    text = re.sub(r"([“‘])\s+", r"\1", text)
    text = re.sub(r"\s+([”’])", r"\1", text)
    return text.strip()


def is_article_noise_line(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    if not compact:
        return True
    if re.fullmatch(r"[0-9０-９]+", compact):
        return True
    if compact in {
        "点击", "蓝字", "关注", "点击蓝字", "点击蓝字关注", "关注我们", "设为星标",
        "星标我们", "分享", "收藏", "点赞", "在看", "阅读原文", "阅读全文",
    }:
        return True
    if re.fullmatch(r"(点击|长按|扫码|识别|回复).{0,18}(关注|蓝字|二维码|原文|查看).{0,18}", compact):
        return True
    if len(compact) <= 40 and re.search(r"(编辑精选|专刊征稿|往期推荐|相关推荐|人物专访|榜样的力量)", compact):
        return True
    if re.fullmatch(r"(编辑|责编|审核|来源|转载)[:：].{0,18}", compact):
        return True
    return False


def unique_article_lines(lines: list[str]) -> list[str]:
    rows = []
    seen = set()
    for line in lines:
        line = normalize_article_line(line)
        if is_article_noise_line(line):
            continue
        key = re.sub(r"\s+", "", line)
        if key in seen:
            continue
        rows.append(line)
        seen.add(key)
    return rows


def article_html_to_text(value: str) -> str:
    if not value:
        return ""
    soup = BeautifulSoup(value, "html.parser")
    for tag in soup(["script", "style", "svg", "noscript", "iframe"]):
        tag.decompose()
    for tag in soup.select("[style*='display:none'], [hidden]"):
        tag.decompose()
    lines = []
    for tag in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote"]):
        line = normalize_article_line(tag.get_text(" ", strip=True))
        if line:
            lines.append(line)
    for row in soup.find_all("tr"):
        cells = [normalize_article_line(cell.get_text(" ", strip=True)) for cell in row.find_all(["th", "td"])]
        cells = [cell for cell in cells if cell]
        if cells:
            lines.append(" | ".join(cells))
    if len(lines) < 3:
        lines = [normalize_article_line(piece) for piece in soup.get_text("\n", strip=True).splitlines()]
    return "\n\n".join(unique_article_lines(lines))[:WECHAT_ARTICLE_MAX_CHARS]


def extract_wechat_article_text_from_html(html: str) -> str:
    soup = BeautifulSoup(html or "", "html.parser")
    content = soup.select_one("#js_content") or soup.select_one(".rich_media_content") or soup.select_one("#img-content")
    return article_html_to_text(str(content)) if content else ""


def extract_wechat_sn(link: str) -> str:
    if not link:
        return ""
    try:
        parsed = urlparse(link)
        value = parse_qs(parsed.query).get("sn", [""])[0]
        if value:
            return clean_text(value)
    except Exception:
        pass
    match = re.search(r"[?&]sn=([^&#]+)", link)
    return clean_text(match.group(1)) if match else ""


def fetch_gsdata_wechat_content(news_local_url: str, client: httpx.Client | None = None) -> tuple[str, str]:
    news_local_url = clean_text(news_local_url)
    if not news_local_url:
        return "", "missing_news_local_url"
    params = {"news_local_url": news_local_url}
    token = gsdata_access_token(params, GSDATA_WECHAT_CONTENT_ROUTER)
    try:
        requester = client or httpx
        response = requester.get(GSDATA_API_URL, params=params, headers={"access-token": token}, timeout=20)
        response.raise_for_status()
        payload = response.json()
    except Exception as exc:
        return "", f"gsdata_content_error: {exc}"
    if not payload.get("success"):
        return "", "gsdata_content_failed: " + json.dumps(payload.get("data", payload), ensure_ascii=False)[:200]
    data = payload.get("data") or {}
    html = data.get("news_content") or data.get("content") or data.get("html") or ""
    text = article_html_to_text(html)
    return text, "gsdata_content_ok" if text else "gsdata_content_empty"


def fetch_wechat_mp_article_text(link: str) -> tuple[str, str]:
    link = clean_text(link)
    if not link or "mp.weixin.qq.com" not in link:
        return "", "missing_mp_link"
    try:
        response = httpx.get(link, headers=WECHAT_ARTICLE_HEADERS, follow_redirects=True, timeout=15)
        response.raise_for_status()
    except Exception as exc:
        return "", f"mp_article_error: {exc}"
    final_url = str(response.url)
    if "wappoc_appmsgcaptcha" in final_url:
        return "", "mp_article_captcha"
    text = extract_wechat_article_text_from_html(response.text)
    return text, "mp_article_ok" if text else "mp_article_empty"


def fetch_wechat_article_text(article: dict, client: httpx.Client | None = None) -> tuple[str, str]:
    text, status = fetch_gsdata_wechat_content(article.get("news_local_url", ""), client)
    if text:
        return text, status
    fallback_text, fallback_status = fetch_wechat_mp_article_text(article.get("news_url", ""))
    if fallback_text:
        return fallback_text, fallback_status
    return "", f"{status}; {fallback_status}"


def score_wechat_api_article(group: str, account_name: str, title: str, summary: str, link: str = "") -> dict:
    return score_private_market_item(f"{group}公众号", title, summary, account_name, "gsdata_wechat", link)


def wechat_api_item(account: dict, article: dict, content_client: httpx.Client | None = None) -> dict:
    title = clean_text(article.get("news_title", ""))
    digest = html_to_text(article.get("news_digest", ""))
    link = clean_text(article.get("news_url", ""))
    article_text, article_fetch_status = fetch_wechat_article_text(article, content_client)
    summary = digest or clean_text(article_text[:300])
    scoring_text = "\n".join(value for value in (digest, article_text) if value)
    source_id = clean_text(article.get("news_uuid", "")) or link
    fingerprint = hashlib.md5((source_id or f"{account['wx_name']}:{title}").encode()).hexdigest()[:16]
    source_group = f"{account['group']}公众号"
    score = score_wechat_api_article(account["group"], account["account_name"], title, scoring_text or summary, link)
    item = {
        "source": "wechat_api",
        "source_id": source_id or fingerprint,
        "fingerprint": fingerprint,
        "title": title,
        "summary": summary,
        "article_text": article_text,
        "article_text_length": len(article_text),
        "article_fetch_status": article_fetch_status,
        "source_name": account["account_name"],
        "source_group": source_group,
        "source_key": account["wx_name"],
        "source_type": "gsdata_wechat",
        "account_name": account["account_name"],
        "wx_name": account["wx_name"],
        "wx_nickname": clean_text(article.get("wx_nickname", "")),
        "news_author": clean_text(article.get("news_author", "")),
        "news_local_url": clean_text(article.get("news_local_url", "")),
        "source_url": clean_text(article.get("source_url", "")),
        "categories": [source_group, account["account_name"]],
        "published_at": clean_text(article.get("news_posttime", "")),
        "updated_at": clean_text(article.get("news_entertime", "")),
        "link": link,
        "cover_url": clean_text(article.get("cover_url", "")),
        "read_count": article.get("news_read_count", ""),
        "like_count": article.get("news_like_count", ""),
        "old_like_num": article.get("news_old_like_num", ""),
        "share_num": article.get("share_num", ""),
        "attention_score": score["score"],
        "worth_attention": score["worth_attention"],
        "signals": score["signals"],
        **score_extra_fields(score),
        "collected_at": utc_now_iso(),
    }
    return attach_project_profile(item)


def fetch_wechat_api_account(account: dict, date_value: str, days: int, limit: int) -> tuple[list[dict], dict | None]:
    start, end = date_range_for_wechat_api(date_value, days)
    rows = []
    seen = set()
    page_size = min(50, max(1, limit))
    max_pages = max(1, min(10, (limit + page_size - 1) // page_size))
    with httpx.Client(timeout=20, follow_redirects=True) as content_client:
        for page in range(1, max_pages + 1):
            params = {
                "wx_name": account["wx_name"],
                "posttime_start": start,
                "posttime_end": end,
                "order": "desc",
                "sort": "posttime",
                "page": str(page),
                "limit": str(page_size),
            }
            token = gsdata_access_token(params, GSDATA_WECHAT_ROUTER)
            try:
                response = None
                for attempt in range(3):
                    response = httpx.get(GSDATA_API_URL, params=params, headers={"access-token": token}, timeout=20)
                    if response.status_code != 429:
                        break
                    time.sleep(1.5 * (attempt + 1))
                response.raise_for_status()
                payload = response.json()
            except Exception as exc:
                return rows, {"account_name": account["account_name"], "wx_name": account["wx_name"], "error": str(exc)}
            if not payload.get("success"):
                return rows, {"account_name": account["account_name"], "wx_name": account["wx_name"], "error": json.dumps(payload, ensure_ascii=False)[:500]}
            data = payload.get("data") or {}
            articles = data.get("newsList") or []
            if not articles:
                break
            for article in articles:
                if not article.get("news_title"):
                    continue
                item = wechat_api_item(account, article, content_client)
                key = source_key(item)
                if key and key not in seen:
                    rows.append(item)
                    seen.add(key)
                if len(rows) >= limit:
                    break
            if len(rows) >= limit or len(articles) < page_size:
                break
    return rows, None


def fetch_wechat_api_batch(req: WechatApiRunRequest) -> dict:
    date_value = req.date or wechat_default_start_date(req.days)
    accounts = filter_wechat_api_accounts(req)
    rows = []
    errors = []
    account_results = []
    worker_count = min(WECHAT_API_MAX_WORKERS, max(1, len(accounts)))
    with ThreadPoolExecutor(max_workers=worker_count) as executor:
        futures = {
            executor.submit(fetch_wechat_api_account, account, date_value, req.days, req.limit_per_account): account
            for account in accounts
        }
        for future in as_completed(futures):
            account = futures[future]
            try:
                account_rows, error = future.result()
            except Exception as exc:
                account_rows = []
                error = {"account_name": account["account_name"], "wx_name": account["wx_name"], "error": str(exc)}
            if error:
                errors.append(error)
            rows.extend(account_rows)
            account_results.append({
                "group": account["group"],
                "account_name": account["account_name"],
                "wx_name": account["wx_name"],
                "fetched": len(account_rows),
                "error": error["error"] if error else "",
            })
    retained_rows = [row for row in rows if row.get("worth_attention")]
    written = append_jsonl(WECHAT_API_FILE, retained_rows)
    return {
        "date": date_value,
        "days": req.days,
        "date_end": (datetime.strptime(date_value, "%Y-%m-%d").date() + timedelta(days=req.days)).isoformat(),
        "accounts": len(accounts),
        "fetched": len(rows),
        "retained": len(retained_rows),
        "filtered": len(rows) - len(retained_rows),
        "written": written,
        "worth_attention": len(retained_rows),
        "errors": errors,
        "account_results": account_results,
        "items": retained_rows[:100],
    }


def read_candidates() -> list[dict[str, Any]]:
    rows = []
    for path in (ARXIV_FILE, WECHAT_FILE, WECHAT_API_FILE, WECHAT_CHAT_CANDIDATES_FILE, INVESTMENT_FILE):
        if not path.exists():
            continue
        for line_no, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
            try:
                item = json.loads(line)
            except Exception:
                continue
            item["_line"] = line_no
            item["_file"] = str(path)
            attach_project_profile(item)
            rows.append(item)
    deduped = {source_key(row): row for row in rows if source_key(row)}
    return sorted(deduped.values(), key=lambda x: (x.get("worth_attention", False), x.get("attention_score", 0), x.get("published_at", "")), reverse=True)


def fetch_arxiv(req: ArxivRunRequest) -> list[dict]:
    params = {
        "search_query": build_query(req.categories, req.keywords, req.days),
        "start": 0,
        "max_results": req.max_results,
        "sortBy": "submittedDate",
        "sortOrder": "descending",
    }
    response = httpx.get(ARXIV_API_URL, params=params, timeout=30, follow_redirects=True)
    response.raise_for_status()
    feed = feedparser.parse(response.text)
    return parse_arxiv_entries(feed.entries, req.watch_authors, req.max_results)


def fetch_arxiv_rss_fallback(req: ArxivRunRequest) -> list[dict]:
    """Fallback collector using arXiv RSS when export API is rate-limited."""
    rows = []
    seen = set()
    for category in req.categories:
        if len(rows) >= req.max_results:
            break
        url = ARXIV_RSS_URL.format(category=category.strip())
        feed = feedparser.parse(url)
        parsed = parse_arxiv_entries(feed.entries, req.watch_authors, req.max_results)
        for item in parsed:
            key = source_key(item)
            if key in seen:
                continue
            seen.add(key)
            rows.append(item)
            if len(rows) >= req.max_results:
                break
    return rows


def _fetch_arxiv_safe(req: ArxivRunRequest) -> list[dict]:
    """Safe arxiv fetcher for auto crawler: tries API first, falls back to RSS."""
    try:
        return fetch_arxiv(req)
    except httpx.HTTPStatusError as exc:
        if exc.response.status_code == 429:
            return fetch_arxiv_rss_fallback(req)
        raise


def parse_arxiv_entries(entries: list[dict], watch_authors: list[str], limit: int) -> list[dict]:
    rows = []
    for entry in entries[:limit]:
        arxiv_id = extract_entry_arxiv_id(entry)
        authors = [a.get("name", "") for a in entry.get("authors", []) if a.get("name")]
        title = clean_text(entry.get("title", ""))
        summary = clean_text(entry.get("summary", ""))
        comment = clean_text(entry.get("arxiv_comment", ""))
        journal_ref = clean_text(entry.get("arxiv_journal_ref", ""))
        score = score_item(title, summary, comment, journal_ref, authors, watch_authors)
        fingerprint = hashlib.md5((arxiv_id or title).encode()).hexdigest()[:16]
        item = {
            "source": "arxiv",
            "source_id": arxiv_id,
            "fingerprint": fingerprint,
            "title": title,
            "summary": summary,
            "authors": authors,
            "first_author": score["first_author"],
            "second_author": score["second_author"],
            "source_group": "论文",
            "categories": [tag.get("term", "") for tag in entry.get("tags", []) if tag.get("term")],
            "published_at": entry_datetime(entry, "published"),
            "updated_at": entry_datetime(entry, "updated"),
            "link": f"https://arxiv.org/abs/{arxiv_id}" if arxiv_id else entry.get("link", ""),
            "pdf_url": f"https://arxiv.org/pdf/{arxiv_id}.pdf" if arxiv_id else "",
            "comment": comment,
            "journal_ref": journal_ref,
            "attention_score": score["score"],
            "worth_attention": score["worth_attention"],
            "signals": score["signals"],
            **score_extra_fields(score),
            "collected_at": datetime.now(timezone.utc).isoformat(),
        }
        rows.append(attach_project_profile(item))
    return rows


def default_wechat_sources() -> list[dict]:
    return [
        {
            "school": item["school"],
            "province": item["province"],
            "accounts": [{"name": f"{item['school']}公众号", "rss_url": ""}],
        }
        for item in UNIVERSITIES_985_SOURCES
    ]


def ensure_wechat_sources_file() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if not WECHAT_SOURCES_FILE.exists():
        save_wechat_sources(default_wechat_sources())


def load_wechat_sources() -> list[dict]:
    ensure_wechat_sources_file()
    try:
        data = json.loads(WECHAT_SOURCES_FILE.read_text(encoding="utf-8"))
    except Exception:
        data = {"sources": default_wechat_sources()}
    sources = data.get("sources", data if isinstance(data, list) else [])
    known = {item["school"]: item for item in default_wechat_sources()}
    for source in sources:
        if source.get("school") in known:
            known[source["school"]].update({
                "province": source.get("province") or known[source["school"]]["province"],
                "accounts": source.get("accounts") or known[source["school"]]["accounts"],
            })
    return list(known.values())


def save_wechat_sources(sources: list[dict | WechatSource]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    cleaned = []
    for source in sources:
        if isinstance(source, BaseModel):
            source = source.model_dump() if hasattr(source, "model_dump") else source.dict()
        accounts = []
        for account in source.get("accounts", []):
            if isinstance(account, BaseModel):
                account = account.model_dump() if hasattr(account, "model_dump") else account.dict()
            accounts.append({
                "name": clean_text(account.get("name", "")),
                "rss_url": clean_text(account.get("rss_url", "")),
            })
        cleaned.append({
            "school": clean_text(source.get("school", "")),
            "province": clean_text(source.get("province", "")),
            "accounts": accounts or [{"name": "", "rss_url": ""}],
        })
    WECHAT_SOURCES_FILE.write_text(json.dumps({"sources": cleaned}, ensure_ascii=False, indent=2), encoding="utf-8")


def html_to_text(value: str) -> str:
    if not value:
        return ""
    if "<" not in value and ">" not in value:
        return clean_text(value)
    return clean_text(BeautifulSoup(value, "html.parser").get_text(" "))


def parse_feed_datetime(entry: dict) -> str:
    for field in ("published", "updated", "created"):
        parsed = entry.get(f"{field}_parsed")
        if parsed:
            return datetime(*parsed[:6], tzinfo=timezone.utc).isoformat()
        if entry.get(field):
            return clean_text(entry.get(field, ""))
    return ""


def score_wechat_article(school: str, title: str, summary: str) -> dict:
    return score_private_market_item("高校公众号", title, summary, school, "wechat_rss")


def fetch_wechat_985(req: WechatRunRequest) -> dict:
    sources = load_wechat_sources()
    rows = []
    errors = []
    feed_count = 0

    with httpx.Client(timeout=25, follow_redirects=True, headers={"User-Agent": "project-discovery-radar/1.0"}) as client:
        for source in sources:
            school = source.get("school", "")
            province = source.get("province", "")
            for account in source.get("accounts", []):
                account_name = clean_text(account.get("name", "")) or school
                rss_url = clean_text(account.get("rss_url", ""))
                if not rss_url:
                    continue
                feed_count += 1
                try:
                    response = client.get(rss_url)
                    response.raise_for_status()
                    feed = feedparser.parse(response.text)
                    if feed.bozo and not feed.entries:
                        raise ValueError(str(feed.bozo_exception))
                except Exception as exc:
                    errors.append({"school": school, "account": account_name, "rss_url": rss_url, "error": str(exc)})
                    continue

                for entry in feed.entries[: req.max_entries_per_feed]:
                    title = clean_text(entry.get("title", ""))
                    summary = html_to_text(entry.get("summary", "") or entry.get("description", ""))
                    link = clean_text(entry.get("link", ""))
                    if not title:
                        continue
                    fingerprint = hashlib.md5((link or f"{school}:{account_name}:{title}").encode()).hexdigest()[:16]
                    score = score_wechat_article(school, title, summary)
                    item = {
                        "source": "wechat_985",
                        "source_id": link or fingerprint,
                        "fingerprint": fingerprint,
                        "title": title,
                        "summary": summary,
                        "school": school,
                        "province": province,
                        "account_name": account_name,
                        "authors": [],
                        "categories": ["985公众号", school],
                        "published_at": parse_feed_datetime(entry),
                        "updated_at": entry_datetime(entry, "updated"),
                        "link": link,
                        "attention_score": score["score"],
                        "worth_attention": score["worth_attention"],
                        "signals": score["signals"],
                        **score_extra_fields(score),
                        "collected_at": datetime.now(timezone.utc).isoformat(),
                    }
                    rows.append(attach_project_profile(item))

    retained_rows = [row for row in rows if row.get("worth_attention")]
    written = append_jsonl(WECHAT_FILE, retained_rows)
    return {
        "feeds": feed_count,
        "fetched": len(rows),
        "retained": len(retained_rows),
        "filtered": len(rows) - len(retained_rows),
        "written": written,
        "worth_attention": len(retained_rows),
        "needs_config": feed_count == 0,
        "message": "还没有配置 985 院校公众号 feed URL。" if feed_count == 0 else "",
        "errors": errors,
        "items": retained_rows,
    }


def score_investment_item(group: str, title: str, summary: str, source_name: str) -> dict:
    return score_private_market_item(group, title, summary, source_name, "investment_source")


def make_investment_item(source: dict, title: str, summary: str, link: str, published_at: str = "") -> dict:
    title = clean_text(title)
    summary = html_to_text(summary)[:1200]
    link = clean_text(link)
    fingerprint = hashlib.md5((link or f"{source['key']}:{title}").encode()).hexdigest()[:16]
    score = score_investment_item(source["group"], title, summary, source["name"])
    item = {
        "source": "investment",
        "source_id": link or fingerprint,
        "fingerprint": fingerprint,
        "title": title,
        "summary": summary,
        "source_name": source["name"],
        "source_group": source["group"],
        "source_key": source["key"],
        "source_type": source["type"],
        "categories": [source["group"], source["name"]],
        "published_at": published_at,
        "updated_at": "",
        "link": link,
        "attention_score": score["score"],
        "worth_attention": score["worth_attention"],
        "signals": score["signals"],
        **score_extra_fields(score),
        "collected_at": datetime.now(timezone.utc).isoformat(),
    }
    return attach_project_profile(item)


def fetch_url_text(client: httpx.Client, url: str) -> str:
    response = client.get(url)
    response.raise_for_status()
    return response.text


def parse_rss_source(source: dict, text: str, limit: int) -> list[dict]:
    feed = feedparser.parse(text)
    rows = []
    for entry in feed.entries[:limit]:
        title = entry.get("title", "")
        summary = entry.get("summary", "") or entry.get("description", "")
        link = entry.get("link", "")
        if not clean_text(title):
            continue
        rows.append(make_investment_item(source, title, summary, link, parse_feed_datetime(entry)))
    return rows


def parse_html_list_source(source: dict, html: str, limit: int) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    seen = set()
    for anchor in soup.select("a[href]"):
        title = clean_text(anchor.get_text(" "))
        href = clean_text(anchor.get("href", ""))
        if len(title) < 6 or title in seen:
            continue
        combined = f"{title} {href}"
        if not any(contains_phrase(combined, keyword) for keyword in (*PROJECT_KEYWORDS, *INVESTMENT_KEYWORDS)):
            continue
        link = urljoin(source["url"], href)
        context = anchor.find_parent(["li", "tr", "div", "article"])
        summary = clean_text(context.get_text(" ")) if context else title
        rows.append(make_investment_item(source, title, summary[:600], link))
        seen.add(title)
        if len(rows) >= limit:
            break
    return rows


def parse_wanfang_source(source: dict, html: str, limit: int) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    seen = set()
    for anchor in soup.select("a[href]"):
        title = clean_text(anchor.get_text(" "))
        href = clean_text(anchor.get("href", ""))
        if len(title) < 8 or title in seen:
            continue
        if not any(contains_phrase(title, keyword) for keyword in ("人工智能", "深度学习", "机器学习", "机器人", "医疗", "材料", "芯片", "算法")):
            continue
        link = urljoin("https://s.wanfangdata.com.cn", href)
        context = anchor.find_parent(["div", "li"])
        summary = clean_text(context.get_text(" ")) if context else title
        rows.append(make_investment_item(source, title, summary[:600], link))
        seen.add(title)
        if len(rows) >= limit:
            break
    return rows


def parse_36kr_financing_flash_source(source: dict, html: str, limit: int) -> list[dict]:
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    seen = set()
    yesterday_label = yesterday_china_date()
    for block in soup.select("div.css-xle9x"):
        title_el = block.select_one("a.title")
        time_el = block.select_one(".item-other .time")
        if not title_el or not time_el:
            continue
        time_text = clean_text(time_el.get_text(" "))
        if time_text not in {"昨天", yesterday_label}:
            continue
        title = clean_text(title_el.get_text(" "))
        href = clean_text(title_el.get("href", ""))
        if not title or title in seen:
            continue
        desc_el = block.select_one(".item-desc span")
        summary = clean_text(desc_el.get_text(" ")) if desc_el else title
        link = urljoin(source["url"], href)
        item = make_investment_item(source, title, summary, link, yesterday_label)
        item["source_type"] = "36kr_financing_flash"
        item["source_time_label"] = time_text
        item["categories"] = [source["group"], source["name"], "昨日融资快报"]
        rows.append(item)
        seen.add(title)
        if len(rows) >= limit:
            break
    return rows


def pitchhub_date_from_publish_time(value: Any) -> str:
    try:
        timestamp = int(value)
    except Exception:
        return ""
    # 36Kr PitchHub publishTime appears to be epoch milliseconds.
    if timestamp > 10_000_000_000:
        timestamp = timestamp / 1000
    return datetime.fromtimestamp(timestamp, CHINA_TZ).date().isoformat()


def pitchhub_item_from_flow(source: dict, raw: dict) -> dict | None:
    material = raw.get("templateMaterial") or {}
    title = clean_text(material.get("widgetTitle", ""))
    summary = clean_text(material.get("widgetContent", ""))
    item_id = raw.get("itemId") or material.get("itemId")
    route = clean_text(raw.get("route", ""))
    published_at = pitchhub_date_from_publish_time(material.get("publishTime"))
    if not title or not item_id:
        return None
    if route.startswith("detail_newsflash"):
        link = f"https://36kr.com/newsflashes/{item_id}"
    elif route.startswith("detail_article"):
        link = f"https://36kr.com/p/{item_id}"
    else:
        link = f"https://36kr.com/newsflashes/{item_id}"
    item = make_investment_item(source, title, summary, link, published_at)
    item["source_type"] = "36kr_financing_flash"
    item["source_time_label"] = published_at
    item["categories"] = [source["group"], source["name"], "昨日融资快报"]
    project = raw.get("projectCard") or {}
    if project:
        item["project_name"] = clean_text(project.get("name", ""))
        item["project_brief"] = clean_text(project.get("briefIntro", ""))
        financing_round = project.get("lastestFinancingRound") or {}
        item["project_financing_round"] = clean_text(financing_round.get("name", ""))
        industries = [clean_text(x.get("name", "")) for x in project.get("tradeList", []) if isinstance(x, dict)]
        item["project_industries"] = [x for x in industries if x]
        item["project_profile"] = build_project_profile(item)
    return item


def extract_pitchhub_page_callback(html: str) -> tuple[str, bool]:
    callback_match = re.search(r'"pageCallback":"([^"]+)"', html)
    has_next_match = re.search(r'"hasNextPage":(\d+)', html)
    return (
        callback_match.group(1) if callback_match else "",
        bool(has_next_match and has_next_match.group(1) == "1"),
    )


def fetch_pitchhub_flow_page(client: httpx.Client, callback: str, page_size: int) -> dict:
    body = {
        "partner_id": "web",
        "timestamp": int(datetime.now(timezone.utc).timestamp() * 1000),
        "partner_version": "1.0.0",
        "param": {
            "pageSize": page_size,
            "pageEvent": 1,
            "pageCallback": callback,
            "siteId": 1,
            "platformId": 2,
        },
    }
    response = client.post(
        PITCHHUB_FLOW_URL,
        json=body,
        headers={
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": "https://pitchhub.36kr.com",
            "Referer": "https://pitchhub.36kr.com/financing-flash",
        },
    )
    response.raise_for_status()
    payload = response.json()
    if payload.get("code") != 0:
        raise ValueError(json.dumps(payload, ensure_ascii=False)[:500])
    return payload.get("data") or {}


def fetch_36kr_financing_flash_pages(client: httpx.Client, source: dict, first_html: str, limit: int) -> list[dict]:
    rows = parse_36kr_financing_flash_source(source, first_html, limit)
    seen = {source_key(row) for row in rows if source_key(row)}
    callback, has_next = extract_pitchhub_page_callback(first_html)
    yesterday = yesterday_china_date()
    max_pages = int(source.get("max_pages", 6) or 6)
    page = 1
    while has_next and callback and len(rows) < limit and page < max_pages:
        data = fetch_pitchhub_flow_page(client, callback, min(20, max(1, limit)))
        page += 1
        has_next = bool(data.get("hasNextPage"))
        callback = clean_text(data.get("pageCallback", ""))
        page_items = data.get("itemList") or []
        saw_older_than_yesterday = False
        for raw in page_items:
            item = pitchhub_item_from_flow(source, raw)
            if not item:
                continue
            published_date = clean_text(item.get("published_at", ""))
            if published_date == yesterday:
                key = source_key(item)
                if key and key not in seen:
                    rows.append(item)
                    seen.add(key)
                    if len(rows) >= limit:
                        break
            elif published_date and published_date < yesterday:
                saw_older_than_yesterday = True
        if saw_older_than_yesterday:
            break
    return rows


def enabled_investment_sources(groups: list[str]) -> list[dict]:
    selected = {group for group in groups if group.strip()}
    sources = []
    for source in SOURCE_DEFAULTS:
        if not source.get("enabled", True):
            continue
        if selected and source["group"] not in selected:
            continue
        sources.append(dict(source))
    return sources


def should_retain_source_candidate(row: dict) -> bool:
    # 论文不使用公司/融资新闻的 worth_attention 门槛。先保留到雷达候选池，
    # 再由下游 AI 论文审查核验完整标题和技术贡献，避免纯论文被误判为无投资事件。
    return bool(row.get("worth_attention")) or row.get("source_group") == "论文"


def fetch_investment_sources(req: InvestmentRunRequest) -> dict:
    rows = []
    errors = []
    source_results = []
    sources = enabled_investment_sources(req.groups)
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; project-discovery-radar/1.0)",
        "Accept": "text/html,application/rss+xml,application/atom+xml,application/xml;q=0.9,*/*;q=0.8",
    }
    with httpx.Client(timeout=12, follow_redirects=True, headers=headers) as client:
        for source in sources:
            source_rows = []
            try:
                url = source["url"]
                source_limit = int(source.get("max_entries_per_run") or req.max_entries_per_source)
                if source["type"] == "wanfang_search":
                    keyword = req.keyword or source.get("keyword", "")
                    url = url.format(keyword=quote_plus(keyword))
                text = fetch_url_text(client, url)
                if source["type"] in {"rss", "arxiv_rss"}:
                    source_rows = parse_rss_source(source, text, source_limit)
                elif source["type"] == "html_list":
                    source_rows = parse_html_list_source(source, text, source_limit)
                elif source["type"] == "wanfang_search":
                    source_rows = parse_wanfang_source(source, text, source_limit)
                elif source["type"] == "36kr_financing_flash":
                    source_rows = fetch_36kr_financing_flash_pages(client, source, text, source_limit)
                else:
                    continue
            except Exception as exc:
                errors.append({
                    "source": source["name"],
                    "group": source["group"],
                    "url": source["url"],
                    "error": str(exc),
                })
            rows.extend(source_rows)
            source_results.append({
                "key": source["key"],
                "name": source["name"],
                "group": source["group"],
                "type": source["type"],
                "fetched": len(source_rows),
            })

    deduped = {}
    for row in rows:
        key = source_key(row)
        if key:
            deduped[key] = row
    rows = sorted(deduped.values(), key=lambda item: (item.get("worth_attention", False), item.get("attention_score", 0), item.get("published_at", "")), reverse=True)
    retained_rows = [row for row in rows if should_retain_source_candidate(row)]
    written = append_jsonl(INVESTMENT_FILE, retained_rows)
    return {
        "sources": len(sources),
        "fetched": len(rows),
        "retained": len(retained_rows),
        "filtered": len(rows) - len(retained_rows),
        "written": written,
        "worth_attention": len([row for row in retained_rows if row.get("worth_attention")]),
        "errors": errors,
        "source_results": source_results,
        "items": retained_rows[:100],
    }


async def run_auto_crawl_once() -> dict:
    global auto_crawler_running
    if auto_crawler_running:
        status = read_auto_status()
        status["running"] = True
        status["last_error"] = "上一轮自动抓取仍在运行，跳过本轮。"
        write_auto_status(status)
        return status

    auto_crawler_running = True
    status = read_auto_status()
    status["enabled"] = True
    status["running"] = True
    status["last_started_at"] = utc_now_iso()
    status["last_error"] = ""
    write_auto_status(status)

    try:
        req = InvestmentRunRequest(groups=AUTO_CRAWL_GROUPS, max_entries_per_source=20, keyword="人工智能")
        result = await asyncio.to_thread(fetch_investment_sources, req)

        # 论文：通过 arxiv API 获取更丰富的元数据（作者/分类/pdf_url），存入 arxiv_candidates.jsonl
        arxiv_result = None
        try:
            arxiv_req = ArxivRunRequest(categories=["cs.AI", "cs.CL", "cs.CV", "cs.LG"], days=1, max_results=50)
            arxiv_rows = await asyncio.to_thread(_fetch_arxiv_safe, arxiv_req)
            arxiv_written = append_jsonl(ARXIV_FILE, arxiv_rows)
            arxiv_result = {
                "fetched": len(arxiv_rows),
                "written": arxiv_written,
                "filtered": len(arxiv_rows) - arxiv_written,
                "worth_attention": len([row for row in arxiv_rows if row.get("worth_attention")]),
            }
        except Exception as exc:
            arxiv_result = {"error": str(exc)}

        status = read_auto_status()
        status["last_result"] = {
            "sources": result.get("sources", 0),
            "fetched": result.get("fetched", 0),
            "retained": result.get("retained", 0),
            "filtered": result.get("filtered", 0),
            "written": result.get("written", 0),
            "worth_attention": result.get("worth_attention", 0),
            "errors": len(result.get("errors", [])),
            "error_samples": result.get("errors", [])[:20],
            "source_results": result.get("source_results", []),
        }
        if arxiv_result:
            status["last_result"]["arxiv"] = arxiv_result
        status["run_count"] = int(status.get("run_count", 0)) + 1
        error_count = len(result.get("errors", []))
        status["consecutive_error_runs"] = (
            int(status.get("consecutive_error_runs", 0)) + 1 if error_count else 0
        )
        status["last_error"] = f"{error_count} 个自动采集源失败" if error_count else ""
    except Exception as exc:
        status = read_auto_status()
        status["last_error"] = str(exc)
        status["consecutive_error_runs"] = int(status.get("consecutive_error_runs", 0)) + 1
    finally:
        auto_crawler_running = False
        status["running"] = False
        status["last_finished_at"] = utc_now_iso()
        next_run = datetime.now(timezone.utc) + timedelta(seconds=AUTO_CRAWL_INTERVAL_SECONDS)
        status["next_run_at"] = next_run.isoformat()
        write_auto_status(status)
    return status


async def auto_crawler_loop() -> None:
    await asyncio.sleep(5)
    while True:
        await run_auto_crawl_once()
        await asyncio.sleep(AUTO_CRAWL_INTERVAL_SECONDS)


async def run_wechat_daily_once(req: WechatApiRunRequest | None = None) -> dict:
    global wechat_daily_running
    if wechat_daily_running:
        status = read_wechat_daily_status()
        status["running"] = True
        status["last_error"] = "上一轮公众号每日抽取仍在运行，跳过本轮。"
        write_wechat_daily_status(status)
        return status

    wechat_daily_running = True
    req = req or WechatApiRunRequest(date=wechat_default_start_date(7), days=7, groups=["高校", "机构"], limit_per_account=100)
    status = read_wechat_daily_status()
    status["enabled"] = True
    status["running"] = True
    status["last_started_at"] = utc_now_iso()
    status["days"] = req.days
    status["last_target_date"] = req.date or wechat_default_start_date(req.days)
    status["last_date_end"] = (datetime.strptime(status["last_target_date"], "%Y-%m-%d").date() + timedelta(days=req.days)).isoformat()
    status["last_error"] = ""
    write_wechat_daily_status(status)
    try:
        result = await asyncio.to_thread(fetch_wechat_api_batch, req)
        status = read_wechat_daily_status()
        status["last_result"] = {
            "date": result.get("date", ""),
            "date_end": result.get("date_end", ""),
            "days": result.get("days", 1),
            "accounts": result.get("accounts", 0),
            "fetched": result.get("fetched", 0),
            "retained": result.get("retained", 0),
            "filtered": result.get("filtered", 0),
            "written": result.get("written", 0),
            "worth_attention": result.get("worth_attention", 0),
            "errors": len(result.get("errors", [])),
            "error_samples": result.get("errors", [])[:20],
        }
        status["run_count"] = int(status.get("run_count", 0)) + 1
        error_count = len(result.get("errors", []))
        status["consecutive_error_runs"] = (
            int(status.get("consecutive_error_runs", 0)) + 1 if error_count else 0
        )
        status["last_error"] = f"{error_count} 个公众号账号采集失败" if error_count else ""
    except Exception as exc:
        status = read_wechat_daily_status()
        status["last_error"] = str(exc)
        status["consecutive_error_runs"] = int(status.get("consecutive_error_runs", 0)) + 1
    finally:
        wechat_daily_running = False
        status["running"] = False
        status["last_finished_at"] = utc_now_iso()
        status["next_run_at"] = next_wechat_daily_run_at().isoformat()
        write_wechat_daily_status(status)
    return status


async def wechat_daily_loop() -> None:
    while True:
        next_run = next_wechat_daily_run_at()
        status = read_wechat_daily_status()
        status["enabled"] = True
        status["running"] = False
        status["run_hour"] = WECHAT_DAILY_RUN_HOUR
        status["run_minute"] = WECHAT_DAILY_RUN_MINUTE
        status["groups"] = ["高校", "机构"]
        status["days"] = 7
        status["next_run_at"] = next_run.isoformat()
        write_wechat_daily_status(status)
        await asyncio.sleep(max(1, (next_run - today_china()).total_seconds()))
        await run_wechat_daily_once(WechatApiRunRequest(date=wechat_default_start_date(7), days=7, groups=["高校", "机构"], limit_per_account=100))


@app.on_event("startup")
async def start_auto_crawler():
    global auto_crawler_task, wechat_daily_task
    status = read_auto_status()
    status["enabled"] = AUTO_CRAWL_ENABLED
    status["running"] = False
    status["interval_seconds"] = AUTO_CRAWL_INTERVAL_SECONDS
    status["groups"] = AUTO_CRAWL_GROUPS
    status["next_run_at"] = (
        (datetime.now(timezone.utc) + timedelta(seconds=5)).isoformat()
        if AUTO_CRAWL_ENABLED else ""
    )
    write_auto_status(status)
    if AUTO_CRAWL_ENABLED:
        auto_crawler_task = asyncio.create_task(auto_crawler_loop())

    daily_enabled = WECHAT_DAILY_ENABLED and gsdata_credentials_configured()
    daily_status = read_wechat_daily_status()
    daily_status["enabled"] = daily_enabled
    daily_status["running"] = False
    daily_status["next_run_at"] = next_wechat_daily_run_at().isoformat() if daily_enabled else ""
    if WECHAT_DAILY_ENABLED and not daily_enabled:
        daily_status["last_error"] = "GSData 凭据未配置，公众号定时采集未启动"
    elif daily_status.get("last_error") == "GSData 凭据未配置，公众号定时采集未启动":
        daily_status["last_error"] = ""
    write_wechat_daily_status(daily_status)
    if daily_enabled:
        wechat_daily_task = asyncio.create_task(wechat_daily_loop())


@app.on_event("shutdown")
async def stop_auto_crawler():
    if auto_crawler_task:
        auto_crawler_task.cancel()
        try:
            await auto_crawler_task
        except asyncio.CancelledError:
            pass
    if wechat_daily_task:
        wechat_daily_task.cancel()
        try:
            await wechat_daily_task
        except asyncio.CancelledError:
            pass


@app.get("/")
async def index():
    return FileResponse(STATIC_DIR / "index.html")


@app.get("/api/health")
async def health(deep: bool = True):
    configured = gsdata_credentials_configured()
    gsdata_health = (
        await asyncio.to_thread(probe_gsdata_health)
        if deep and configured
        else {
            "ok": configured,
            "status": "configured_not_probed" if configured else "not_configured",
            "error": "" if configured else "GSData 凭据未配置",
        }
    )
    return {
        "status": "ok",
        "name": "project-discovery-radar",
        "data_dir": str(DATA_DIR),
        "accounts_file": str(WECHAT_ACCOUNTS_XLSX),
        "gsdata_configured": configured,
        "gsdata_health": gsdata_health,
        "auto_crawl_enabled": AUTO_CRAWL_ENABLED,
        "wechat_daily_enabled": WECHAT_DAILY_ENABLED and configured,
    }


@app.get("/api/auto/status")
async def auto_status():
    status = read_auto_status()
    status["running"] = auto_crawler_running
    return status


@app.post("/api/auto/run-now")
async def auto_run_now():
    return await run_auto_crawl_once()


@app.post("/api/wechat-chat/push")
async def wechat_chat_push(req: WechatChatPushRequest):
    return ingest_wechat_chat_push(req)


@app.get("/api/wechat-chat/messages")
async def wechat_chat_messages(
    date: str = "",
    group_name: str = "",
    group_serial_no: str = "",
    limit: int = Query(default=500, ge=1, le=5000),
):
    date_value = date or today_china().date().isoformat()
    rows = []
    if group_serial_no:
        rows = read_chat_daily_messages(date_value, group_name or group_serial_no, group_serial_no)
    else:
        day_dir = WECHAT_CHAT_MESSAGES_DIR / date_value
        if day_dir.exists():
            for path in sorted(day_dir.glob("*.jsonl")):
                for line in path.read_text(encoding="utf-8").splitlines():
                    try:
                        item = json.loads(line)
                    except Exception:
                        continue
                    if group_name and item.get("group_name") != group_name:
                        continue
                    rows.append(item)
            rows = sorted(rows, key=lambda item: (item.get("group_name", ""), item.get("sort_time", "")))
    return {"date": date_value, "total": len(rows), "items": rows[:limit]}


@app.get("/api/wechat-chat/groups")
async def wechat_chat_groups(date: str = ""):
    date_value = date or today_china().date().isoformat()
    day_dir = WECHAT_CHAT_MESSAGES_DIR / date_value
    groups = []
    if day_dir.exists():
        for path in sorted(day_dir.glob("*.jsonl")):
            rows = []
            for line in path.read_text(encoding="utf-8").splitlines():
                try:
                    rows.append(json.loads(line))
                except Exception:
                    continue
            if not rows:
                continue
            first = rows[0]
            group_name = first.get("group_name", path.stem)
            group_serial_no = first.get("group_serial_no", path.stem)
            candidate_path = chat_candidate_group_path(date_value, group_name, group_serial_no)
            candidate_count = 0
            if candidate_path.exists():
                candidate_count = len([line for line in candidate_path.read_text(encoding="utf-8").splitlines() if line.strip()])
            groups.append({
                "date": date_value,
                "group_name": group_name,
                "group_serial_no": group_serial_no,
                "message_count": len(rows),
                "candidate_count": candidate_count,
                "message_file": str(path),
                "candidate_file": str(candidate_path),
                "first_message_time": rows[0].get("sort_time", ""),
                "last_message_time": rows[-1].get("sort_time", ""),
            })
    return {"date": date_value, "total": len(groups), "groups": groups}


@app.get("/api/summary")
async def summary():
    rows = read_candidates()
    sources = Counter(row.get("source", "unknown") for row in rows)
    return {
        "total": len(rows),
        "worth_attention": len([x for x in rows if x.get("worth_attention")]),
        "avg_score": round(sum(x.get("attention_score", 0) for x in rows) / max(1, len(rows)), 1),
        "sources": dict(sources),
        "data_file": str(DATA_DIR),
    }


@app.get("/api/candidates")
async def candidates(
    q: str = "",
    source: str = "",
    source_key: str = "",
    group: str = "",
    attention_only: bool = False,
    min_score: int = Query(default=0, ge=0, le=100),
    limit: int = Query(default=200, ge=1, le=500),
    sort: Literal["score", "collected"] = "score",
    cursor: str = "",
):
    q_folded = q.casefold().strip()
    rows = []
    for item in read_candidates():
        if source and item.get("source") != source:
            continue
        if source_key and item.get("source_key") != source_key:
            continue
        if group and item.get("source_group") != group:
            continue
        if attention_only and not item.get("worth_attention"):
            continue
        if item.get("attention_score", 0) < min_score:
            continue
        if q_folded:
            haystack = "\n".join([
                item.get("title", ""),
                item.get("summary", ""),
                item.get("article_text", ""),
                " ".join(item.get("authors", [])),
                item.get("school", ""),
                item.get("account_name", ""),
                item.get("wx_name", ""),
                item.get("wx_nickname", ""),
                item.get("source_name", ""),
                item.get("source_group", ""),
                " ".join(signal.get("detail", "") for signal in item.get("signals", [])),
            ]).casefold()
            if q_folded not in haystack:
                continue
        rows.append(item)
    total = len(rows)
    if sort == "collected":
        rows = sorted(rows, key=candidate_cursor_key, reverse=True)
        if cursor:
            try:
                cursor_key = decode_candidate_cursor(cursor)
            except ValueError as exc:
                raise HTTPException(status_code=400, detail="无效的候选记录游标") from exc
            rows = [item for item in rows if candidate_cursor_key(item) < cursor_key]
    page = rows[:limit]
    has_more = len(rows) > limit
    next_cursor = (
        encode_candidate_cursor(candidate_cursor_key(page[-1]))
        if sort == "collected" and has_more and page
        else ""
    )
    return {
        "items": page,
        "total": total,
        "sort": sort,
        "has_more": has_more,
        "next_cursor": next_cursor,
    }


@app.get("/api/wechat/sources")
async def wechat_sources():
    sources = load_wechat_sources()
    configured = sum(
        1
        for source in sources
        for account in source.get("accounts", [])
        if clean_text(account.get("rss_url", ""))
    )
    return {"sources": sources, "total": len(sources), "configured_feeds": configured}


@app.post("/api/wechat/sources")
async def update_wechat_sources(req: WechatSourcesRequest):
    save_wechat_sources(req.sources)
    return await wechat_sources()


@app.post("/api/wechat/run")
async def wechat_run(req: WechatRunRequest):
    try:
        return fetch_wechat_985(req)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"985 公众号抓取失败: {exc}")


@app.get("/api/wechat-api/accounts")
async def wechat_api_accounts():
    accounts = load_wechat_api_accounts()
    groups = Counter(account["group"] for account in accounts)
    return {"accounts": accounts, "total": len(accounts), "groups": dict(groups)}


@app.get("/api/wechat-api/daily-status")
async def wechat_api_daily_status():
    status = read_wechat_daily_status()
    status["running"] = wechat_daily_running
    return status


@app.post("/api/wechat-api/run")
async def wechat_api_run(req: WechatApiRunRequest):
    try:
        return fetch_wechat_api_batch(req)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"公众号 API 抓取失败: {exc}")


@app.post("/api/wechat-api/run-yesterday")
async def wechat_api_run_yesterday(req: WechatApiRunRequest | None = None):
    req = req or WechatApiRunRequest()
    if not req.date:
        req.date = wechat_default_start_date(req.days)
    if not req.groups:
        req.groups = ["高校", "机构"]
    if req.limit_per_account < 100 and not req.wx_names:
        req.limit_per_account = 100
    return await run_wechat_daily_once(req)


@app.get("/api/investment/sources")
async def investment_sources():
    groups = Counter(source["group"] for source in SOURCE_DEFAULTS)
    return {
        "sources": SOURCE_DEFAULTS,
        "groups": dict(groups),
        "enabled": len([source for source in SOURCE_DEFAULTS if source.get("enabled", True)]),
    }


@app.post("/api/investment/run")
async def investment_run(req: InvestmentRunRequest):
    try:
        return fetch_investment_sources(req)
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"多渠道抓取失败: {exc}")


@app.post("/api/arxiv/run")
async def arxiv_run(req: ArxivRunRequest):
    fallback_used = False
    try:
        rows = fetch_arxiv(req)
    except httpx.HTTPStatusError as exc:
        if exc.response.status_code == 429:
            rows = fetch_arxiv_rss_fallback(req)
            fallback_used = True
            if not rows:
                raise HTTPException(status_code=429, detail="arXiv API 临时限流，RSS 兜底也没有返回内容，请过几分钟再试。")
        else:
            raise HTTPException(status_code=502, detail=f"arXiv API 请求失败: HTTP {exc.response.status_code}")
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"arXiv 抓取失败: {exc}")
    retained_rows = [row for row in rows if should_retain_source_candidate(row)]
    written = append_jsonl(ARXIV_FILE, retained_rows)
    return {
        "fetched": len(rows),
        "retained": len(retained_rows),
        "filtered": len(rows) - len(retained_rows),
        "written": written,
        "worth_attention": len(retained_rows),
        "fallback_used": fallback_used,
        "items": retained_rows,
    }
