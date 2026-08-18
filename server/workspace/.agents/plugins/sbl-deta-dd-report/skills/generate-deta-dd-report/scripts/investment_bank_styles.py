#!/usr/bin/env python3
"""Apply and audit the fixed Chinese investment-banking DOCX style system."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph


STYLE_H1 = "投行 - 一级标题"
STYLE_H2 = "投行 - 二级标题"
STYLE_H3 = "投行 - 三级标题"
STYLE_BODY = "投行 - 正文标准"
STYLE_TABLE_HEADER = "投行 - 表格表头"
STYLE_NOTE = "投行 - 注释小字"
STYLE_TOC_TITLE = "投行 - 目录标题"
STYLE_COVER_TITLE = "投行 - 封面标题"
REQUIRED_STYLES = [
    STYLE_H1, STYLE_H2, STYLE_H3, STYLE_BODY, STYLE_TABLE_HEADER,
    STYLE_NOTE, STYLE_TOC_TITLE, STYLE_COVER_TITLE,
]

FONT_HEITI = "黑体"
FONT_HEITI_ASCII = "SimHei"
FONT_SONGTI = "宋体"
FONT_SONGTI_ASCII = "SimSun"
HEADER_GRAY_FILL = "D9D9D9"
LABEL_GRAY_FILL = "F2F2F2"
TABLE_STYLE = "Table Grid"
TABLE_MARGIN_DXA = {"top": "0", "left": "108", "bottom": "0", "right": "108"}
TWO_COLUMN_DATA_HEADERS = {
    ("事项", "核验结果"),
    ("场景", "解决问题"),
    ("事项", "截至报告日情况"),
    ("主要成本项", "金额/合同"),
    ("退出路径", "实现条件"),
}

H1_RE = re.compile(r"^(?:[1-8]、|投资结论及建议$)")
H2_RE = re.compile(r"^[1-8]\.[1-9]\d*\s")
H3_RE = re.compile(r"^[1-8]\.[1-9]\d*\.[1-9]\d*\s*")
NOTE_RE = re.compile(r"^(?:注|备注|说明|测算口径|口径说明)[：:]")
TOC_INSTRUCTION = ' TOC \\o "1-3" \\h \\z \\u '
TOC_PLACEHOLDER = "目录将在 Microsoft Word 中自动更新。"
TOC_ERROR_MARKERS = (
    "错误！未找到目录项。",
    "错误!未找到目录项。",
    "Error! No table of contents entries found.",
)


def _set_font(element, ascii_name: str, east_asia: str, size: float, bold: bool) -> None:
    font = element.font
    font.name = ascii_name
    font.size = Pt(size)
    font.bold = bold
    rpr = element._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (
        ("ascii", ascii_name), ("hAnsi", ascii_name),
        ("eastAsia", east_asia), ("cs", ascii_name),
    ):
        rfonts.set(qn(f"w:{key}"), value)


def _set_run_font(run, ascii_name: str, east_asia: str, size: float, bold: bool | None = None) -> None:
    run.font.name = ascii_name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, value in (
        ("ascii", ascii_name), ("hAnsi", ascii_name),
        ("eastAsia", east_asia), ("cs", ascii_name),
    ):
        rfonts.set(qn(f"w:{key}"), value)


def _set_spacing_lines(style, before_lines: int, after_pt: float) -> None:
    ppr = style._element.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.attrib.pop(qn("w:before"), None)
    spacing.set(qn("w:beforeLines"), str(before_lines))
    spacing.set(qn("w:after"), str(int(after_pt * 20)))


def _set_first_line_chars(style, chars_hundredths: int) -> None:
    ppr = style._element.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    for attr in ("firstLine", "hanging", "hangingChars"):
        ind.attrib.pop(qn(f"w:{attr}"), None)
    ind.set(qn("w:firstLineChars"), str(chars_hundredths))


def _set_outline(style, level: int | None) -> None:
    ppr = style._element.get_or_add_pPr()
    node = ppr.find(qn("w:outlineLvl"))
    if level is None:
        if node is not None:
            ppr.remove(node)
        return
    if node is None:
        node = OxmlElement("w:outlineLvl")
        ppr.append(node)
    node.set(qn("w:val"), str(level))


def _mark_quick_style(style, priority: int) -> None:
    element = style._element
    if element.find(qn("w:qFormat")) is None:
        element.append(OxmlElement("w:qFormat"))
    ui = element.find(qn("w:uiPriority"))
    if ui is None:
        ui = OxmlElement("w:uiPriority")
        element.append(ui)
    ui.set(qn("w:val"), str(priority))


def _get_or_create_style(doc: Document, name: str):
    return doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def define_styles(doc: Document) -> None:
    specs = {
        STYLE_H1: (FONT_HEITI_ASCII, FONT_HEITI, 16, True, 150, 6, 0),
        STYLE_H2: (FONT_HEITI_ASCII, FONT_HEITI, 14, True, 80, 4, 1),
        STYLE_H3: (FONT_HEITI_ASCII, FONT_HEITI, 12, True, 50, 2, 2),
        STYLE_BODY: (FONT_SONGTI_ASCII, FONT_SONGTI, 12, False, 0, 0, None),
        STYLE_TABLE_HEADER: (FONT_HEITI_ASCII, FONT_HEITI, 12, True, 0, 0, None),
        STYLE_NOTE: (FONT_SONGTI_ASCII, FONT_SONGTI, 10.5, False, 0, 0, None),
        STYLE_TOC_TITLE: (FONT_HEITI_ASCII, FONT_HEITI, 26, True, 0, 12, None),
        STYLE_COVER_TITLE: (FONT_HEITI_ASCII, FONT_HEITI, 22, True, 0, 0, None),
    }
    for priority, (name, spec) in enumerate(specs.items(), 1):
        ascii_name, east_asia, size, bold, before_lines, after_pt, outline = spec
        style = _get_or_create_style(doc, name)
        style.base_style = doc.styles["Normal"]
        _set_font(style, ascii_name, east_asia, size, bold)
        style.paragraph_format.keep_with_next = name in {STYLE_H1, STYLE_H2, STYLE_H3}
        style.paragraph_format.keep_together = name in {STYLE_H1, STYLE_H2, STYLE_H3}
        style.paragraph_format.widow_control = True
        # Page breaks are applied to each H1 paragraph below. Keeping the style
        # itself neutral lets a heading immediately following a section break
        # suppress the redundant break cleanly in Microsoft Word.
        style.paragraph_format.page_break_before = False
        style.paragraph_format.space_before = None
        style.paragraph_format.space_after = Pt(after_pt)
        _set_spacing_lines(style, before_lines, after_pt)
        _set_outline(style, outline)
        _mark_quick_style(style, priority)

    body = doc.styles[STYLE_BODY]
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    body.paragraph_format.line_spacing = Pt(20)
    _set_first_line_chars(body, 200)

    table_header = doc.styles[STYLE_TABLE_HEADER]
    table_header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_header.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    table_header.paragraph_format.line_spacing = Pt(20)
    _set_first_line_chars(table_header, 0)

    note = doc.styles[STYLE_NOTE]
    note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    note.paragraph_format.line_spacing = Pt(16)
    _set_first_line_chars(note, 0)

    doc.styles[STYLE_TOC_TITLE].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles[STYLE_COVER_TITLE].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _toc_instruction_nodes(doc: Document):
    complex_nodes = [
        node for node in doc._element.xpath(".//w:instrText")
        if re.search(r"(?:^|\s)TOC(?:\s|$)", node.text or "", re.I)
    ]
    simple_nodes = [
        node for node in doc._element.xpath(".//w:fldSimple")
        if re.search(r"(?:^|\s)TOC(?:\s|$)", node.get(qn("w:instr")) or "", re.I)
    ]
    return complex_nodes, simple_nodes


def _enable_field_updates_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _insert_paragraph_before(paragraph, text: str = ""):
    element = OxmlElement("w:p")
    paragraph._p.addprevious(element)
    inserted = Paragraph(element, paragraph._parent)
    if text:
        inserted.add_run(text)
    return inserted


def _insert_paragraph_after(paragraph):
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def _append_native_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = TOC_INSTRUCTION
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = TOC_PLACEHOLDER
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)


def _remove_static_toc_entries(doc: Document, title) -> int:
    """Remove manual TOC lines before inserting the native field."""
    paragraphs = list(doc.paragraphs)
    title_index = next(
        (idx for idx, paragraph in enumerate(paragraphs) if paragraph._p is title._p),
        None,
    )
    if title_index is None:
        return 0
    start = title_index + 1
    removed = 0
    for paragraph in paragraphs[start:]:
        if _style_name(paragraph) in {STYLE_H1, STYLE_H2, STYLE_H3}:
            break
        text = "".join(paragraph._p.xpath(".//w:t/text()")).strip()
        style = _style_name(paragraph).upper()
        looks_manual = bool(
            style.startswith("TOC")
            or re.search(r"\t\s*[IVXLCDM\d]+\s*$", text, re.I)
        )
        if not text or not looks_manual:
            continue
        ppr = paragraph._p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            continue
        paragraph._p.getparent().remove(paragraph._p)
        removed += 1
    return removed


def ensure_automatic_toc(doc: Document) -> dict:
    """Ensure a native, clickable Word TOC field covering outline levels 1-3."""
    paragraphs = list(doc.paragraphs)
    title = next((p for p in paragraphs if p.text.strip() == "目录"), None)
    if title is None:
        first_heading = next(
            (p for p in paragraphs if _style_name(p) in {STYLE_H1, STYLE_H2, STYLE_H3}),
            None,
        )
        title = (
            _insert_paragraph_before(first_heading, "目录")
            if first_heading is not None
            else doc.add_paragraph("目录")
        )
    _format_paragraph(title, STYLE_TOC_TITLE)

    complex_nodes, simple_nodes = _toc_instruction_nodes(doc)
    if complex_nodes:
        complex_nodes[0].text = TOC_INSTRUCTION
        complex_nodes[0].set(qn("xml:space"), "preserve")
    elif simple_nodes:
        simple_nodes[0].set(qn("w:instr"), TOC_INSTRUCTION.strip())
    else:
        _remove_static_toc_entries(doc, title)
        toc_paragraph = _insert_paragraph_after(title)
        if "TOC 1" in doc.styles:
            toc_paragraph.style = doc.styles["TOC 1"]
        toc_paragraph.paragraph_format.first_line_indent = Pt(0)
        _append_native_toc_field(toc_paragraph)

    _enable_field_updates_on_open(doc)
    complex_nodes, simple_nodes = _toc_instruction_nodes(doc)
    return {
        "title": True,
        "native_field": bool(complex_nodes or simple_nodes),
        "instruction": TOC_INSTRUCTION.strip(),
        "update_on_open": True,
    }


def _toc_snapshot(doc: Document) -> dict:
    paragraphs = list(doc.paragraphs)
    title_index = next(
        (idx for idx, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "目录"),
        None,
    )
    complex_nodes, simple_nodes = _toc_instruction_nodes(doc)
    instructions = [node.text or "" for node in complex_nodes]
    instructions.extend(node.get(qn("w:instr")) or "" for node in simple_nodes)
    settings_update = doc.settings._element.find(qn("w:updateFields"))
    update_value = settings_update.get(qn("w:val")) if settings_update is not None else None

    zone = []
    if title_index is not None:
        for paragraph in paragraphs[title_index + 1:]:
            if _style_name(paragraph) in {STYLE_H1, STYLE_H2, STYLE_H3}:
                break
            zone.append(paragraph)
    zone_texts = [
        "".join(paragraph._p.xpath(".//w:t/text()")).strip()
        for paragraph in zone
    ]
    entry_paragraphs = [
        paragraph for paragraph, text in zip(zone, zone_texts)
        if text
        and text != TOC_PLACEHOLDER
        and (
            _style_name(paragraph).upper().startswith("TOC")
            or bool(paragraph._p.xpath(".//w:hyperlink"))
        )
    ]
    hyperlink_count = sum(len(paragraph._p.xpath(".//w:hyperlink")) for paragraph in zone)
    heading_count = sum(
        1 for paragraph in paragraphs
        if _style_name(paragraph) in {STYLE_H1, STYLE_H2, STYLE_H3}
    )
    error_texts = [
        text for text in zone_texts
        if any(marker.lower() in text.lower() for marker in TOC_ERROR_MARKERS)
    ]
    normalized = " ".join(instructions)
    return {
        "title_present": title_index is not None,
        "native_field_present": bool(instructions),
        "covers_levels_1_to_3": bool(
            re.search(r'\\o\s+"1-3"', normalized, re.I)
            and re.search(r"(?:^|\s)\\u(?:\s|$)", normalized, re.I)
        ),
        "hyperlink_switch": bool(re.search(r"(?:^|\s)\\h(?:\s|$)", normalized, re.I)),
        "update_on_open": update_value not in {None, "0", "false", "off"},
        "heading_count": heading_count,
        "cached_entry_count": len(entry_paragraphs),
        "cached_hyperlink_count": hyperlink_count,
        "error_texts": error_texts,
        "instructions": instructions,
    }


def _style_name(paragraph) -> str:
    return paragraph.style.name if paragraph.style else ""


def _next_nonempty_block_is_table(paragraph) -> bool:
    node = paragraph._p.getnext()
    while node is not None:
        if node.tag == qn("w:tbl"):
            return True
        if node.tag == qn("w:p"):
            texts = node.xpath(".//w:t/text()")
            if "".join(texts).strip():
                return False
        node = node.getnext()
    return False


def _is_caption(paragraph) -> bool:
    centered = paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    bold = any(run.bold for run in paragraph.runs if run.text.strip())
    return centered and bold and _next_nonempty_block_is_table(paragraph)


def _remove_stale_rendered_page_breaks(doc: Document) -> int:
    """Remove layout-cache markers inherited from an older Word pagination.

    They are not author-entered page breaks. Keeping them beside a new
    pageBreakBefore can create an unexplained blank page after repagination.
    """
    removed = 0
    for node in list(doc._element.xpath(".//w:lastRenderedPageBreak")):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
            removed += 1
    return removed


def _remove_empty_spacers_before_h1(doc: Document) -> int:
    """Drop empty spacer paragraphs immediately before paginated H1s.

    When a preceding table already reaches the page foot, invisible spacer
    paragraphs can flow onto a new page; H1 pageBreakBefore then creates a
    second page, leaving the first one visually blank.
    """
    removed = 0
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        style = _style_name(paragraph)
        if not (style in {"Heading 1", STYLE_H1} or H1_RE.match(text)):
            continue
        previous = paragraph._p.getprevious()
        while previous is not None and previous.tag == qn("w:p"):
            texts = "".join(previous.xpath(".//w:t/text()")).strip()
            ppr = previous.find(qn("w:pPr"))
            has_section = ppr is not None and ppr.find(qn("w:sectPr")) is not None
            has_break = bool(previous.xpath(".//w:br"))
            has_object = bool(previous.xpath(".//w:drawing | .//w:pict | .//w:object"))
            if texts or has_section or has_break or has_object:
                break
            candidate = previous
            previous = previous.getprevious()
            candidate.getparent().remove(candidate)
            removed += 1
    return removed


def _format_paragraph(paragraph, style_name: str, *, table_body: bool = False) -> None:
    paragraph.style = style_name
    if style_name == STYLE_H1:
        previous = paragraph._p.getprevious()
        follows_section_break = False
        if previous is not None and previous.tag == qn("w:p"):
            previous_ppr = previous.find(qn("w:pPr"))
            follows_section_break = (
                previous_ppr is not None
                and previous_ppr.find(qn("w:sectPr")) is not None
            )
        # A new-page section break already satisfies the H1 pagination rule.
        # Adding a second pageBreakBefore here creates a visible blank page.
        paragraph.paragraph_format.page_break_before = (
            paragraph.text.strip() != "8、风险提示与对策"
            and not follows_section_break
        )
        paragraph.paragraph_format.first_line_indent = Pt(0)
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 16, True)
    elif style_name == STYLE_H2:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 14, True)
    elif style_name == STYLE_H3:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 12, True)
    elif style_name == STYLE_BODY:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if table_body else WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(20)
        if table_body:
            paragraph.paragraph_format.first_line_indent = Pt(0)
            ppr = paragraph._p.get_or_add_pPr()
            ind = ppr.find(qn("w:ind"))
            if ind is not None:
                ind.set(qn("w:firstLineChars"), "0")
        for run in paragraph.runs:
            _set_run_font(run, FONT_SONGTI_ASCII, FONT_SONGTI, 12, None)
    elif style_name == STYLE_TABLE_HEADER:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(20)
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 12, True)
    elif style_name == STYLE_NOTE:
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.line_spacing = Pt(16)
        for run in paragraph.runs:
            _set_run_font(run, FONT_SONGTI_ASCII, FONT_SONGTI, 10.5, None)
    elif style_name == STYLE_TOC_TITLE:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 26, True)
    elif style_name == STYLE_COVER_TITLE:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 22, True)


def _format_caption(paragraph) -> None:
    """Use one deterministic caption rule for every table title."""
    paragraph.style = STYLE_NOTE
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(16)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        _set_run_font(run, FONT_HEITI_ASCII, FONT_HEITI, 12, True)


def _shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def _clear_cell_shading(cell) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is not None:
        tcpr.remove(shd)


def _set_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for child in list(borders):
        borders.remove(child)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)


def _set_table_cell_margins(table) -> None:
    tblpr = table._tbl.tblPr
    margins = tblpr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tblpr.append(margins)
    for child in list(margins):
        margins.remove(child)
    for edge, width in TABLE_MARGIN_DXA.items():
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), width)
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def _normalize_row_geometry(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    for height in list(trpr.findall(qn("w:trHeight"))):
        trpr.remove(height)
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def _clean_text(value: str) -> str:
    return " ".join(str(value or "").split())


def _is_data_table(table) -> bool:
    if len(table.columns) >= 3:
        return True
    if len(table.columns) != 2 or not table.rows:
        return False
    pair = tuple(_clean_text(cell.text) for cell in table.rows[0].cells)
    return pair in TWO_COLUMN_DATA_HEADERS


def _is_compact_column(table, cell_idx: int) -> bool:
    values = [_clean_text(row.cells[cell_idx].text) for row in table.rows[1:]]
    values = [value for value in values if value]
    if not values:
        return True
    return max(map(len, values)) <= 16 and all("\n" not in value for value in values)


def _ensure_table_style(doc: Document) -> None:
    if TABLE_STYLE not in doc.styles:
        style = doc.styles.add_style(TABLE_STYLE, WD_STYLE_TYPE.TABLE)
        if "Normal Table" in doc.styles:
            style.base_style = doc.styles["Normal Table"]


def _format_tables(doc: Document, counts: dict) -> None:
    _ensure_table_style(doc)
    for table in doc.tables:
        table.style = TABLE_STYLE
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_table_borders(table)
        _set_table_cell_margins(table)
        data_table = _is_data_table(table)
        compact_columns = {
            idx: _is_compact_column(table, idx) for idx in range(len(table.columns))
        } if data_table else {}
        for row_idx, row in enumerate(table.rows):
            _normalize_row_geometry(row)
            for cell_idx, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _clear_cell_shading(cell)
                is_semantic_header = (
                    (data_table and row_idx == 0)
                    or (not data_table and len(table.columns) == 2 and cell_idx == 0)
                )
                if is_semantic_header:
                    _shade_cell(
                        cell,
                        HEADER_GRAY_FILL if data_table else LABEL_GRAY_FILL,
                    )
                for paragraph in cell.paragraphs:
                    target = STYLE_TABLE_HEADER if is_semantic_header else (
                        STYLE_NOTE if NOTE_RE.match(paragraph.text.strip()) else STYLE_BODY
                    )
                    _format_paragraph(paragraph, target, table_body=target == STYLE_BODY)
                    if target == STYLE_BODY:
                        if data_table and (cell_idx == 0 or compact_columns[cell_idx]):
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    counts[target] += 1
            if data_table and row_idx == 0:
                _repeat_header(row)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = trpr.find(qn("w:tblHeader"))
    if header is None:
        trpr.append(OxmlElement("w:tblHeader"))


def apply_styles(doc: Document) -> dict:
    _remove_stale_rendered_page_breaks(doc)
    _remove_empty_spacers_before_h1(doc)
    define_styles(doc)
    paragraphs = list(doc.paragraphs)
    cover_title_indices = [
        idx for idx, p in enumerate(paragraphs[:20])
        if "尽职调查报告" in p.text and p.text.strip()
    ]
    cover_title_index = cover_title_indices[0] if cover_title_indices else None

    counts = {name: 0 for name in REQUIRED_STYLES}
    for idx, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        old_style = _style_name(paragraph)
        if idx == cover_title_index:
            target = STYLE_COVER_TITLE
        elif text == "目录":
            target = STYLE_TOC_TITLE
        elif old_style in {"Heading 3", STYLE_H3} or H3_RE.match(text):
            target = STYLE_H3
        elif old_style in {"Heading 2", STYLE_H2} or H2_RE.match(text):
            target = STYLE_H2
        elif old_style in {"Heading 1", STYLE_H1} or H1_RE.match(text):
            target = STYLE_H1
        elif NOTE_RE.match(text):
            target = STYLE_NOTE
        elif _is_caption(paragraph):
            paragraph._p.getparent().remove(paragraph._p)
            continue
        elif idx < 20 or old_style.startswith("TOC"):
            continue
        else:
            target = STYLE_BODY
        _format_paragraph(paragraph, target)
        counts[target] += 1

    _format_tables(doc, counts)
    ensure_automatic_toc(doc)
    return counts


def audit_styles(doc: Document, *, require_updated_toc: bool = True) -> dict:
    issues = []
    missing = [name for name in REQUIRED_STYLES if name not in doc.styles]
    if missing:
        issues.append("missing styles: " + ", ".join(missing))
    stale_breaks = len(doc._element.xpath(".//w:lastRenderedPageBreak"))
    if stale_breaks and not require_updated_toc:
        issues.append(f"stale rendered page-break markers remain: {stale_breaks}")
    captions = [p.text.strip() for p in doc.paragraphs if p.text.strip() and _is_caption(p)]
    if captions:
        issues.append("standalone table captions are forbidden: " + " | ".join(captions[:20]))
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not (_style_name(paragraph) == STYLE_H1 or H1_RE.match(text)):
            continue
        previous = paragraph._p.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            previous_text = "".join(previous.xpath(".//w:t/text()")).strip()
            ppr = previous.find(qn("w:pPr"))
            has_section = ppr is not None and ppr.find(qn("w:sectPr")) is not None
            has_break = bool(previous.xpath(".//w:br"))
            pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
            previous_style = pstyle.get(qn("w:val")) if pstyle is not None else ""
            probe = previous
            while probe is not None and probe.tag == qn("w:p"):
                probe_text = "".join(probe.xpath(".//w:t/text()")).strip()
                if probe_text:
                    break
                probe = probe.getprevious()
            probe_ppr = probe.find(qn("w:pPr")) if probe is not None and probe.tag == qn("w:p") else None
            probe_style_node = probe_ppr.find(qn("w:pStyle")) if probe_ppr is not None else None
            probe_style = probe_style_node.get(qn("w:val")) if probe_style_node is not None else ""
            follows_toc_result = probe_style.upper().startswith("TOC")
            if (
                not previous_text
                and not has_section
                and not has_break
                and not previous_style.upper().startswith("TOC")
                and not follows_toc_result
            ):
                issues.append(f"empty spacer paragraph remains before H1: {text}")

    property_contract = {
        STYLE_H1: {
            "font": FONT_HEITI, "size": "32", "bold": True,
            "beforeLines": "150", "pageBreakBefore": False, "outline": "0",
        },
        STYLE_H2: {
            "font": FONT_HEITI, "size": "28", "bold": True,
            "beforeLines": "80", "outline": "1",
        },
        STYLE_H3: {
            "font": FONT_HEITI, "size": "24", "bold": True,
            "beforeLines": "50", "outline": "2",
        },
        STYLE_BODY: {
            "font": FONT_SONGTI, "size": "24", "line": "400",
            "lineRule": "exact", "firstLineChars": "200",
        },
        STYLE_TABLE_HEADER: {
            "font": FONT_HEITI, "size": "24", "bold": True,
            "line": "400", "lineRule": "exact", "jc": "center",
        },
        STYLE_NOTE: {
            "font": FONT_SONGTI, "size": "21", "line": "320",
            "lineRule": "exact",
        },
        STYLE_TOC_TITLE: {
            "font": FONT_HEITI, "size": "52", "bold": True, "jc": "center",
        },
        STYLE_COVER_TITLE: {
            "font": FONT_HEITI, "size": "44", "bold": True, "jc": "center",
        },
    }
    for style_name, expected in property_contract.items():
        if style_name not in doc.styles:
            continue
        element = doc.styles[style_name]._element
        rpr = element.find(qn("w:rPr"))
        ppr = element.find(qn("w:pPr"))
        rfonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
        sz = rpr.find(qn("w:sz")) if rpr is not None else None
        spacing = ppr.find(qn("w:spacing")) if ppr is not None else None
        ind = ppr.find(qn("w:ind")) if ppr is not None else None
        jc = ppr.find(qn("w:jc")) if ppr is not None else None
        outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        actual = {
            "font": rfonts.get(qn("w:eastAsia")) if rfonts is not None else None,
            "size": sz.get(qn("w:val")) if sz is not None else None,
            "bold": rpr.find(qn("w:b")) is not None if rpr is not None else False,
            "beforeLines": spacing.get(qn("w:beforeLines")) if spacing is not None else None,
            "line": spacing.get(qn("w:line")) if spacing is not None else None,
            "lineRule": spacing.get(qn("w:lineRule")) if spacing is not None else None,
            "firstLineChars": ind.get(qn("w:firstLineChars")) if ind is not None else None,
            "jc": jc.get(qn("w:val")) if jc is not None else None,
            "outline": outline.get(qn("w:val")) if outline is not None else None,
            "pageBreakBefore": (
                ppr.find(qn("w:pageBreakBefore")) is not None
                and ppr.find(qn("w:pageBreakBefore")).get(qn("w:val")) not in {"0", "false", "off"}
            ) if ppr is not None else False,
        }
        for key, value in expected.items():
            if actual.get(key) != value:
                issues.append(
                    f"{style_name} property {key}: expected {value}, got {actual.get(key)}"
                )

    paragraphs = list(doc.paragraphs)
    cover = [p for p in paragraphs[:20] if "尽职调查报告" in p.text and p.text.strip()]
    toc = [p for p in paragraphs if p.text.strip() == "目录"]
    if not cover or any(_style_name(p) != STYLE_COVER_TITLE for p in cover[:1]):
        issues.append("cover title is not bound to 投行 - 封面标题")
    if not toc:
        issues.append("automatic TOC title is missing")
    elif any(_style_name(p) != STYLE_TOC_TITLE for p in toc[:1]):
        issues.append("TOC title is not bound to 投行 - 目录标题")

    toc_state = _toc_snapshot(doc)
    if not toc_state["native_field_present"]:
        issues.append("native Word TOC field is missing; static/manual directory is forbidden")
    if not toc_state["covers_levels_1_to_3"]:
        issues.append("TOC field must cover heading levels 1-3")
    if not toc_state["hyperlink_switch"]:
        issues.append("TOC field must enable clickable hyperlinks with the \\h switch")
    if not toc_state["update_on_open"] and not require_updated_toc:
        issues.append("document settings must update fields when opened in Microsoft Word")
    if toc_state["error_texts"]:
        issues.append("TOC contains Word field errors: " + " | ".join(toc_state["error_texts"][:5]))
    if require_updated_toc:
        expected = toc_state["heading_count"]
        if toc_state["cached_entry_count"] != expected:
            issues.append(
                "TOC is stale or incomplete: "
                f"{toc_state['cached_entry_count']} cached entries for {expected} headings"
            )
        if toc_state["cached_hyperlink_count"] != expected:
            issues.append(
                "TOC hyperlinks are stale or incomplete: "
                f"{toc_state['cached_hyperlink_count']} links for {expected} headings"
            )

    heading_mismatches = []
    for p in paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if _style_name(p).upper().startswith("TOC"):
            continue
        expected = None
        if H3_RE.match(text):
            expected = STYLE_H3
        elif H2_RE.match(text):
            expected = STYLE_H2
        elif H1_RE.match(text):
            expected = STYLE_H1
        if expected and _style_name(p) != expected:
            heading_mismatches.append(text)
    if heading_mismatches:
        issues.append("heading style mismatches: " + " | ".join(heading_mismatches[:10]))

    for paragraph in paragraphs:
        if _style_name(paragraph) != STYLE_H1:
            continue
        previous = paragraph._p.getprevious()
        if previous is None or previous.tag != qn("w:p"):
            continue
        previous_ppr = previous.find(qn("w:pPr"))
        follows_section_break = (
            previous_ppr is not None
            and previous_ppr.find(qn("w:sectPr")) is not None
        )
        ppr = paragraph._p.pPr
        break_element = ppr.find(qn("w:pageBreakBefore")) if ppr is not None else None
        break_value = break_element.get(qn("w:val")) if break_element is not None else None
        duplicate_break = break_element is not None and break_value not in {"0", "false", "off"}
        if follows_section_break and duplicate_break:
            issues.append(
                f"{STYLE_H1}: ‘{paragraph.text.strip()}’ follows a section break and must not add pageBreakBefore"
            )

    table_header_mismatches = 0
    table_body_mismatches = 0
    table_format_mismatches = []
    for table_index, table in enumerate(doc.tables, 1):
        data_table = _is_data_table(table)
        style_name = table.style.name if table.style else ""
        if style_name != TABLE_STYLE:
            table_format_mismatches.append(f"表{table_index}样式不是{TABLE_STYLE}")
        if table.alignment != WD_TABLE_ALIGNMENT.CENTER:
            table_format_mismatches.append(f"表{table_index}未水平居中")
        tblpr = table._tbl.tblPr
        margins = tblpr.find(qn("w:tblCellMar"))
        actual_margins = {}
        for edge in TABLE_MARGIN_DXA:
            node = margins.find(qn(f"w:{edge}")) if margins is not None else None
            actual_margins[edge] = node.get(qn("w:w")) if node is not None else None
        if actual_margins != TABLE_MARGIN_DXA:
            table_format_mismatches.append(f"表{table_index}单元格边距不统一：{actual_margins}")
        borders = tblpr.find(qn("w:tblBorders"))
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = borders.find(qn(f"w:{edge}")) if borders is not None else None
            actual = (
                node.get(qn("w:val")) if node is not None else None,
                node.get(qn("w:sz")) if node is not None else None,
                node.get(qn("w:color")) if node is not None else None,
            )
            if actual != ("single", "4", "auto"):
                table_format_mismatches.append(f"表{table_index}{edge}边框不是0.5磅黑色单线")
                break
        for row_idx, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr()
            if trpr.find(qn("w:trHeight")) is not None:
                table_format_mismatches.append(f"表{table_index}第{row_idx + 1}行存在固定行高")
            if trpr.find(qn("w:cantSplit")) is None:
                table_format_mismatches.append(f"表{table_index}第{row_idx + 1}行未设置禁止跨页断行")
            for cell_idx, cell in enumerate(row.cells):
                expected_header = (
                    (data_table and row_idx == 0)
                    or (not data_table and len(table.columns) == 2 and cell_idx == 0)
                )
                shd = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
                actual_fill = shd.get(qn("w:fill")) if shd is not None else None
                expected_fill = (
                    HEADER_GRAY_FILL if data_table and row_idx == 0
                    else LABEL_GRAY_FILL if not data_table and len(table.columns) == 2 and cell_idx == 0
                    else None
                )
                if actual_fill != expected_fill:
                    table_format_mismatches.append(
                        f"表{table_index}第{row_idx + 1}行第{cell_idx + 1}列底纹"
                        f"应为{expected_fill or '无'}，实际{actual_fill or '无'}"
                    )
                if cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER:
                    table_format_mismatches.append(
                        f"表{table_index}第{row_idx + 1}行第{cell_idx + 1}列未垂直居中"
                    )
                if expected_header:
                    for p in cell.paragraphs:
                        if p.text.strip() and _style_name(p) != STYLE_TABLE_HEADER:
                            table_header_mismatches += 1
                else:
                    for p in cell.paragraphs:
                        if p.text.strip() and _style_name(p) not in {STYLE_BODY, STYLE_NOTE}:
                            table_body_mismatches += 1
            if data_table and row_idx == 0:
                header = trpr.find(qn("w:tblHeader"))
                if header is None:
                    table_format_mismatches.append(f"表{table_index}表头行未设置跨页重复")
    if table_header_mismatches:
        issues.append(f"table header style mismatches: {table_header_mismatches}")
    if table_body_mismatches:
        issues.append(f"table body style mismatches: {table_body_mismatches}")
    if table_format_mismatches:
        issues.append("table format mismatches: " + " | ".join(table_format_mismatches[:20]))

    return {
        "status": "pass" if not issues else "fail",
        "required_styles": REQUIRED_STYLES,
        "missing_styles": missing,
        "toc": toc_state,
        "require_updated_toc": require_updated_toc,
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    apply_parser = sub.add_parser("apply")
    apply_parser.add_argument("--input", required=True)
    apply_parser.add_argument("--output", required=True)
    audit_parser = sub.add_parser("audit")
    audit_parser.add_argument("--input", required=True)
    audit_parser.add_argument("--json-output")
    audit_parser.add_argument(
        "--allow-unupdated-toc",
        action="store_true",
        help="pre-Word audit: require a native TOC field but allow stale cached entries",
    )
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    doc = Document(str(input_path))
    if args.command == "apply":
        output_path = Path(args.output).expanduser().resolve()
        if output_path == input_path:
            raise SystemExit("output must differ from input")
        counts = apply_styles(doc)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        result = {"status": "styled", "output": str(output_path), "counts": counts}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    result = audit_styles(doc, require_updated_toc=not args.allow_unupdated_toc)
    if args.json_output:
        Path(args.json_output).write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
