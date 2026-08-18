#!/usr/bin/env python3
"""Audit a Deta-style DOCX for semantic repetition and defensive prose density."""

from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


H1_STYLES = {"Heading 1", "投行 - 一级标题"}
H2_STYLES = {"Heading 2", "投行 - 二级标题"}
DEFENSIVE = [
    "需进一步核对", "尚未与账面勾稽", "尚待", "不承保", "不释放", "不计入",
    "无法支持", "不足以支持", "不得外推", "未完成", "仍需验证", "仍需核对",
]
VALUE_BOUNDARY_PATTERNS = [
    r"尚未", r"仍需", r"需进一步", r"不计入", r"不得外推", r"无法支持",
    r"待核", r"前提是", r"以.+为条件",
]
VALUE_MARKERS = ["形成", "构成", "支撑", "带来", "提升", "扩大", "增强", "提供", "有望", "具备", "实现", "降低", "转化"]
VALUE_TOPIC_ORDER = [
    ("industry_and_paid_base", ["行业", "市场", "产业", "卡位", "付费", "合同", "交付", "验收", "回款", "收入底盘"]),
    ("technical_differentiation", ["差异化", "技术路线", "相较", "相比", "R2S2R", "架构", "技术体系"]),
    ("data_feedback_moat", ["数据闭环", "反馈闭环", "数据资产", "壁垒", "复用", "触觉", "数据入口", "回流"]),
    ("commercial_validation", ["客户", "生态", "合作", "验收", "回款", "复购", "跨产品", "场景协同"]),
    ("team_capability", ["团队", "创始人", "科学家", "工程经验", "产业经验", "研发背景", "项目管理"]),
]
TRANSACTION_VALUE_TERMS = ["老股", "分期支付", "治理安排", "入股成本", "资金暴露", "下行保护", "交易安全边际"]
MASTER_RISK_TERMS = ["商业化", "产品化", "复购", "规模收入", "持续收入", "可持续现金流", "估值兑现", "价值兑现"]
ATTRIBUTION_PATTERNS = [
    r"尽调材料(?:显示|未显示|未提供)?",
    r"(?:现有)?材料(?:显示|未显示|未提供|未发现)",
    r"资料(?:显示|未显示|未提供)",
    r"尚未完成勾稽",
    r"需进一步(?:核对|核验|确认)",
    r"需核验",
    r"以[^。；｜\n]{1,40}为准",
]
ATTRIBUTION_FIXED_TEXT = {"部门｜尽调材料人数/规划｜主要职责"}
CROSS_REFERENCE_PATTERNS = [
    r"详见",
    r"参见",
    r"见第[一二三四五六七八九十0-9]+章",
    r"(?:结构|分析|事项|内容|方案)见\s*\d+(?:\.\d+)+",
    r"按\s*\d+(?:\.\d+)+(?:及|、)\d+(?:\.\d+)+所列",
]
PACKED_NUMBERED_PATTERN = re.compile(
    r"(?:^|[。；！？]\s*)([1-9])\s*[.．、]\s*"
)


def clean(value: str) -> str:
    return " ".join(str(value or "").split())


def section_map(doc: _Document) -> tuple[dict[str, str], dict[str, list[str]]]:
    chunks: dict[str, list[str]] = defaultdict(list)
    paragraphs: dict[str, list[str]] = defaultdict(list)
    current_h1 = "封面"
    current_h2 = "封面"
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            text = clean(paragraph.text)
            style = paragraph.style.name if paragraph.style else ""
            if text and style in H1_STYLES:
                current_h1, current_h2 = text, text
                continue
            if text and style in H2_STYLES:
                current_h2 = text
                continue
            if text:
                chunks[current_h2].append(text)
                paragraphs[current_h2].append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            for row in table.rows:
                row_text = "｜".join(clean(cell.text) for cell in row.cells if clean(cell.text))
                if row_text:
                    chunks[current_h2].append(row_text)
    return {key: "\n".join(value) for key, value in chunks.items()}, paragraphs


def match_count(text: str, patterns: list[str]) -> int:
    return sum(bool(re.search(pattern, text)) for pattern in patterns)


def phrase_count(text: str) -> int:
    return sum(text.count(phrase) for phrase in DEFENSIVE)


def find_overview_value(doc: _Document, label: str) -> str:
    for table in doc.tables[:5]:
        for row in table.rows:
            if len(row.cells) >= 2 and clean(row.cells[0].text) == label:
                return clean(row.cells[1].text)
    return ""


def split_value_items(text: str) -> list[str]:
    numbered = re.split(r"(?=(?:^|[；。])\s*[一二三四五六七八九十0-9]+[、.．])", text)
    items = [clean(re.sub(r"^[；。\s]+", "", item)) for item in numbered if clean(item)]
    if len(items) <= 1:
        items = [clean(item) for item in re.split(r"[；\n]", text) if clean(item)]
    return items


def overview_value_items(paragraphs: dict[str, list[str]], doc: _Document) -> tuple[list[str], list[str]]:
    body = paragraphs.get("1.5 投资价值与风险", [])
    numbered = [clean(re.sub(r"^[1-5]\s*[.．、]\s*", "", item))
                for item in body if re.match(r"^[1-5]\s*[.．、]\s*", item)]
    risks = [item for item in body if item.startswith("主要风险")]
    if numbered:
        return numbered, risks
    legacy = find_overview_value(doc, "投资价值")
    return split_value_items(legacy), risks


def value_chain_complete(item: str) -> bool:
    fact_markers = ["已", "形成", "具备", "合同", "客户", "团队", "技术路线", "产品", "交付", "回款"]
    mechanism_markers = ["通过", "依托", "共享", "连接", "回流", "协同", "从而", "使", "可"]
    implication_markers = VALUE_MARKERS + ["收入", "壁垒", "复购", "商业化", "现金流", "客户黏性", "产品化"]
    return (len(item) >= 80 and any(x in item for x in fact_markers)
            and any(x in item for x in mechanism_markers)
            and any(x in item for x in implication_markers))


def chapter_label(section: str) -> str:
    match = re.match(r"^(\d+)(?:[.、])", section)
    if not match:
        return section
    number = match.group(1)
    names = {
        "1": "1、投资概要", "2": "2、公司概况", "3": "3、产品与技术",
        "4": "4、业务情况", "5": "5、行业和市场", "6": "6、未来发展规划",
        "7": "7、投资方案", "8": "8、风险提示与对策",
    }
    return names.get(number, section)


def audit_editorial_quality(doc: _Document) -> dict:
    sections, paragraphs = section_map(doc)
    issues: list[dict] = []

    clusters = [
        ("technology_chain", [r"真实(?:数据)?采集|真实采集", r"合成数据", r"仿真", r"触觉"], 4, 2,
         "完整技术链跨章节重复；仅第3章展开机制，概要和投资章节改写为结论或投资意义"),
        ("financial_core", [r"1[,，]?302\.83", r"48\.87", r"450\s*万元"], 2, 2,
         "历史财务与收入确认数字组跨章节重复；完整数字组只保留在6.2.3"),
        ("transaction_tuple", [r"2[,，]?500\s*万元", r"(?<![0-9,，])500\s*万元", r"7\.50\s*%", r"分期|首笔"], 3, 2,
         "金额拆分、完全稀释持股与分期机制被完整复述；7.2保留完整说明，概要和结论只写各自所需的决定性边界"),
    ]
    cluster_hits = {}
    for code, patterns, threshold, budget, message in clusters:
        raw_hits = sorted(section for section, text in sections.items()
                          if match_count(text, patterns) >= threshold)
        hits = sorted({chapter_label(section) for section in raw_hits})
        cluster_hits[code] = hits
        non_anchor_hits = []
        if code == "technology_chain":
            non_anchor_hits = [section for section in raw_hits
                               if chapter_label(section) != "3、产品与技术"]
            violates = len(non_anchor_hits) > 1
        elif code == "financial_core":
            non_anchor_hits = [section for section in raw_hits if section != "6.2 财务情况"]
            violates = len(non_anchor_hits) > 1
        elif code == "transaction_tuple":
            non_anchor_hits = [section for section in raw_hits
                               if section != "7.2 公司估值与投资方式"]
            violates = len(non_anchor_hits) > 1
        else:
            violates = len(hits) > budget
        if violates:
            issues.append({"priority": "P2", "code": code,
                           "scope": non_anchor_hits or hits, "issue": message})

    core = find_overview_value(doc, "核心判断")
    if phrase_count(core) > 1 or "“" in core or "”" in core:
        issues.append({"priority": "P2", "code": "defensive_core_judgment", "scope": ["1.1 公司情况"],
                       "issue": "核心判断的不确定性限定或负面强调超过预算；压缩为判断加一个控制"})

    value_items, overview_risk_items = overview_value_items(paragraphs, doc)
    defensive_value_items = [item for item in value_items
                             if any(re.search(pattern, item) for pattern in VALUE_BOUNDARY_PATTERNS)]
    incomplete_value_items = [item for item in value_items if not value_chain_complete(item)]
    if defensive_value_items:
        issues.append({"priority": "P2", "code": "defensive_value", "scope": ["1.5 投资价值"],
                       "issue": "投资价值中夹杂风险边界；将验证条件移至主要风险或第8章"})
    if value_items and len(incomplete_value_items) > max(1, len(value_items) // 3):
        issues.append({"priority": "P2", "code": "fragmented_value_logic", "scope": ["1.5 投资价值"],
                       "issue": "多条投资价值未形成事实基础、价值机制和投资意义的完整正向链"})

    value_topic_results = []
    if len(value_items) == 5:
        for item, (code, markers) in zip(value_items, VALUE_TOPIC_ORDER):
            value_topic_results.append({"topic": code, "matched": any(marker in item for marker in markers)})
        if not all(result["matched"] for result in value_topic_results):
            missing = [result["topic"] for result in value_topic_results if not result["matched"]]
            issues.append({"priority": "P2", "code": "deta_value_sequence", "scope": missing,
                           "issue": "1.5未按行业/付费基础、技术差异、数据闭环、商业验证、团队能力的德塔顺序承担分析责任"})
    elif value_items:
        value_topic_results = [{"topic": code, "matched": False} for code, _ in VALUE_TOPIC_ORDER]

    transaction_value_items = [item for item in value_items if any(term in item for term in TRANSACTION_VALUE_TERMS)]
    if transaction_value_items:
        issues.append({"priority": "P2", "code": "transaction_in_company_value", "scope": ["1.5 投资价值"],
                       "issue": "1.5公司价值被交易保护占用；老股折价、分期、治理和下行保护应完整归入7.2"})

    weak_master_risk = (len(overview_risk_items) != 1
                        or not any(term in overview_risk_items[0] for term in MASTER_RISK_TERMS)
                        or "风险集中于" in overview_risk_items[0])
    if weak_master_risk:
        issues.append({"priority": "P2", "code": "overview_master_risk", "scope": ["1.5 主要风险"],
                       "issue": "1.5主要风险应先指出一个商业化/产品化/复购或估值兑现主风险，再说明其他因素的价值传导，不得只列风险清单"})

    highlight_text = sections.get("7.1 投资亮点", "")
    if phrase_count(highlight_text) > 2:
        issues.append({"priority": "P2", "code": "defensive_highlights", "scope": ["7.1 投资亮点"],
                       "issue": "投资亮点中的防御性措辞超过2处；保留证据边界但将完整风险处置移至第8章"})

    conclusion_paragraphs = [p for p in paragraphs.get("投资结论及建议", []) if clean(p)]
    if len(conclusion_paragraphs) != 3:
        issues.append({"priority": "P2", "code": "conclusion_shape", "scope": ["投资结论及建议"],
                       "issue": "结论应严格保持三个正文段落"})
    elif conclusion_paragraphs:
        if phrase_count("\n".join(conclusion_paragraphs[:2])) > 1:
            issues.append({"priority": "P2", "code": "defensive_conclusion", "scope": ["投资结论及建议"],
                           "issue": "结论前两段展开了过多核查措辞；将控制要求集中到第三段"})
        if len(conclusion_paragraphs[2]) > 220 or conclusion_paragraphs[2].count("；") > 4:
            issues.append({"priority": "P2", "code": "long_conditions", "scope": ["投资结论及建议"],
                           "issue": "结论第三段条件清单过长；压缩为一个并列句和一个未完成后果"})

    normalized_groups: dict[str, list[str]] = defaultdict(list)
    for section, items in paragraphs.items():
        for item in items:
            normalized = re.sub(r"[\s，。；、：:（）()“”'\"-]", "", item)
            if len(normalized) >= 50:
                normalized_groups[normalized].append(section)
    exact_duplicates = [sorted(set(scopes)) for scopes in normalized_groups.values() if len(set(scopes)) > 1]
    if exact_duplicates:
        issues.append({"priority": "P2", "code": "exact_cross_section_repeat", "scope": exact_duplicates[:5],
                       "issue": "发现跨章节重复的长解释句；只在信息锚点保留完整表述"})

    attribution_hits: dict[str, list[str]] = defaultdict(list)
    for section, text in sections.items():
        for line in text.splitlines():
            line = clean(line)
            if not line or line in ATTRIBUTION_FIXED_TEXT:
                continue
            matched = [pattern for pattern in ATTRIBUTION_PATTERNS if re.search(pattern, line)]
            if not matched:
                continue
            # The full revenue-recognition finding is allowed once in its
            # finance anchor and once in the risk/control anchor.
            if section == "6.2 财务情况" and "450万元" in line and "需进一步核对" in line:
                continue
            if chapter_label(section) == "8、风险提示与对策":
                continue
            attribution_hits[section].append(line)
    if attribution_hits:
        issues.append({
            "priority": "P2",
            "code": "diffuse_source_attribution",
            "scope": sorted(attribution_hits),
            "issue": "材料归因、缺口或核验措辞散落在经营正文；可确认事实直接陈述，剩余控制集中到第8章和结论",
        })

    # Reader-facing enumerations must be represented as one real Word
    # paragraph per item. A manual line break or a run-on "1...2..." paragraph
    # is not accepted because it collapses under table reflow and editing.
    packed_numbered: list[str] = []
    for paragraph in doc.paragraphs:
        text = clean(paragraph.text)
        if len(PACKED_NUMBERED_PATTERN.findall(text)) >= 2:
            packed_numbered.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = clean(paragraph.text)
                    if len(PACKED_NUMBERED_PATTERN.findall(text)) >= 2:
                        packed_numbered.append(text)
    if packed_numbered:
        issues.append({
            "priority": "P2",
            "code": "packed_numbered_items",
            "scope": packed_numbered[:5],
            "issue": "编号事项挤在同一段落；每项必须使用独立Word段落并逐条换行",
        })

    cross_reference_hits: dict[str, list[str]] = defaultdict(list)
    for section, text in sections.items():
        for line in text.splitlines():
            line = clean(line)
            if line and any(re.search(pattern, line) for pattern in CROSS_REFERENCE_PATTERNS):
                cross_reference_hits[section].append(line)
    if cross_reference_hits:
        issues.append({
            "priority": "P2",
            "code": "cross_reference_shortcut",
            "scope": sorted(cross_reference_hits),
            "issue": "正文用‘详见/参见/见某章节’替代本节结论；应按本节颗粒度直接写出必要事实或判断",
        })

    return {
        "status": "pass" if not issues else "fail",
        "issues": issues,
        "metrics": {
            "semantic_cluster_sections": cluster_hits,
            "core_judgment_defensive_count": phrase_count(core),
            "investment_value_item_count": len(value_items),
            "investment_value_defensive_item_count": len(defensive_value_items),
            "investment_value_incomplete_item_count": len(incomplete_value_items),
            "investment_value_topic_results": value_topic_results,
            "investment_value_transaction_item_count": len(transaction_value_items),
            "investment_value_master_risk_pass": not weak_master_risk,
            "highlight_defensive_count": phrase_count(highlight_text),
            "source_attribution_non_anchor_count": sum(len(v) for v in attribution_hits.values()),
            "source_attribution_non_anchor_sections": sorted(attribution_hits),
            "packed_numbered_paragraph_count": len(packed_numbered),
            "cross_reference_shortcut_count": sum(len(v) for v in cross_reference_hits.values()),
            "cross_reference_shortcut_sections": sorted(cross_reference_hits),
            "conclusion_paragraph_lengths": [len(p) for p in conclusion_paragraphs],
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True)
    parser.add_argument("--json-output")
    args = parser.parse_args()
    result = audit_editorial_quality(Document(str(Path(args.input).expanduser().resolve())))
    if args.json_output:
        Path(args.json_output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
