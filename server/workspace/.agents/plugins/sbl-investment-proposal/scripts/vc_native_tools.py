#!/usr/bin/env python3
"""无网关投资提案 Skill 的确定性本地工具。

本脚本不调用任何大模型。模型推理由调用 Skill 的当前 Codex 会话完成。
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import stat
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any


TEXT_EXT = {"txt", "md", "csv", "json"}
CONVERT_EXT = {"docx", "pdf", "pptx", "xlsx", "doc", "ppt", "xls"}
FACT_CATEGORIES = {
    "entity", "team", "equity_financing", "technology", "product",
    "customers", "financials", "pricing", "deal_terms", "policy",
    "industry", "risk", "plan",
}
OUTLINE = {
    "一、基本情况简介": 1500,
    "二、财务情况": 450,
    "三、交易条件": 800,
    "四、公司业务预测": 300,
    "五、投资亮点与风险控制": 900,
    "六、结论": 90,
}
BANNED_VOICE = {
    "宣称", "据称", "有条件同意", "不宜作出", "待核验", "待确认",
    "尚未明确", "当前材料", "现有材料", "尽调报表", "交流材料披露", "未提供",
}
DEFAULT_ENGINE = Path(__file__).resolve().with_name("vc_proposal_agent.py")


def read_json(path: Path) -> Any:
    with path.open(encoding="utf-8") as fh:
        return json.load(fh)


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + f".{os.getpid()}.tmp")
    with tmp.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    os.replace(tmp, path)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def clean_name(value: str) -> str:
    value = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "-", value).strip("-._")
    return value[:48] or "vc-project"


def command_init_run(args: argparse.Namespace) -> None:
    source = Path(args.input).expanduser().resolve()
    if not source.exists():
        raise SystemExit(f"输入不存在：{source}")
    root = Path(args.output_root).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    fingerprint = sha256_file(source)[:10] if source.is_file() else hashlib.sha256(
        str(source).encode("utf-8")
    ).hexdigest()[:10]
    run_dir = root / f"{clean_name(source.stem or source.name)}-{time.strftime('%Y%m%d-%H%M%S')}-{fingerprint}"
    run_dir.mkdir(parents=False, exist_ok=False)
    write_json(run_dir / "run_manifest.json", {
        "input": str(source), "input_fingerprint": fingerprint,
        "created_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "model_mode": "current-codex-session", "external_llm_gateway": False,
        "status": "initialized",
    })
    print(json.dumps({"run_dir": str(run_dir)}, ensure_ascii=False))


def safe_extract_zip(source: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as archive:
        infos = archive.infolist()
        if len(infos) > 5000:
            raise SystemExit("ZIP 文件项超过 5000，拒绝解压")
        if sum(i.file_size for i in infos) > 2 * 1024**3:
            raise SystemExit("ZIP 解压后大小超过 2GB，拒绝解压")
        base = destination.resolve()
        for info in infos:
            mode = info.external_attr >> 16
            if stat.S_ISLNK(mode):
                raise SystemExit(f"ZIP 含符号链接，拒绝解压：{info.filename}")
            target = (destination / info.filename).resolve()
            if target != base and base not in target.parents:
                raise SystemExit(f"ZIP 路径越界，拒绝解压：{info.filename}")
        archive.extractall(destination)


def split_text(text: str, size: int = 3000, overlap: int = 200) -> list[str]:
    if len(text) <= size:
        return [text]
    parts: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = text.rfind("\n", start + int(size * 0.6), end)
            if boundary > start:
                end = boundary
        parts.append(text[start:end])
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return parts


def extract_text_locally(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        import pymupdf

        with pymupdf.open(path) as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)
            if len(text.strip()) >= 30:
                return text
            tesseract = shutil.which("tesseract")
            if not tesseract:
                return text
            ocr_blocks: list[str] = []
            with tempfile.TemporaryDirectory(prefix="sbl-pdf-ocr-") as temp_dir:
                for index, page in enumerate(pdf, 1):
                    image_path = Path(temp_dir) / f"page-{index:04d}.png"
                    page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(image_path)
                    result = subprocess.run(
                        [tesseract, str(image_path), "stdout", "-l", "chi_sim+eng"],
                        capture_output=True,
                        timeout=180,
                        check=False,
                    )
                    if result.returncode == 0:
                        ocr_blocks.append(result.stdout.decode("utf-8", "replace"))
            return "\n".join(ocr_blocks) or text
    if ext == "docx":
        from docx import Document

        doc = Document(path)
        blocks = [paragraph.text for paragraph in doc.paragraphs]
        blocks.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in doc.tables
            for row in table.rows
        )
        return "\n".join(blocks)
    if ext == "pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text
        )
    if ext == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        blocks: list[str] = []
        for sheet in workbook.worksheets:
            blocks.append(f"[工作表] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    blocks.append("\t".join(values))
        workbook.close()
        return "\n".join(blocks)
    if ext in {"doc", "ppt", "xls"}:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            raise RuntimeError(f"解析 .{ext} 需要 LibreOffice")
        target_ext = {"doc": "docx", "ppt": "pptx", "xls": "xlsx"}[ext]
        with tempfile.TemporaryDirectory(prefix="sbl-convert-") as temp_dir:
            result = subprocess.run(
                [office, "--headless", "--convert-to", target_ext, "--outdir", temp_dir, str(path)],
                capture_output=True,
                timeout=180,
                check=False,
            )
            converted = Path(temp_dir) / f"{path.stem}.{target_ext}"
            if result.returncode or not converted.is_file():
                detail = (result.stderr or result.stdout).decode("utf-8", "replace")[-800:]
                raise RuntimeError(f"LibreOffice 转换失败：{detail}")
            return extract_text_locally(converted)
    raise RuntimeError(f"不支持的材料格式：.{ext}")


def extract_text(path: Path, extractor: str | None) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext in TEXT_EXT:
        return path.read_text(encoding="utf-8", errors="replace")
    if ext not in CONVERT_EXT:
        return ""
    if extractor:
        result = subprocess.run(
            [extractor, str(path)], capture_output=True, timeout=180, check=False
        )
        if result.returncode == 0:
            return result.stdout.decode("utf-8", "replace")
    return extract_text_locally(path)


def command_prepare(args: argparse.Namespace) -> None:
    source = Path(args.input).expanduser().resolve()
    run_dir = Path(args.run_dir).expanduser().resolve()
    if not source.exists() or not run_dir.is_dir():
        raise SystemExit("输入路径或运行目录不存在")
    packets_dir = run_dir / "packets"
    if packets_dir.exists() and any(packets_dir.iterdir()):
        raise SystemExit("packets 目录非空；为避免混用项目，请创建新的运行目录")
    work_root = source
    if source.is_file() and source.suffix.lower() == ".zip":
        work_root = run_dir / "expanded"
        safe_extract_zip(source, work_root)
    elif source.is_file():
        work_root = source.parent
        candidates = [source]
    else:
        candidates = sorted(p for p in work_root.rglob("*") if p.is_file())
    if source.is_file() and source.suffix.lower() == ".zip":
        candidates = sorted(p for p in work_root.rglob("*") if p.is_file())

    extractor = shutil.which("extract-text")
    sources: list[dict[str, Any]] = []
    fragments: list[dict[str, Any]] = []
    failures: list[dict[str, str]] = []
    for path in candidates:
        if (path.name.startswith(".") and path != source) or "__MACOSX" in path.parts:
            continue
        ext = path.suffix.lower().lstrip(".")
        if ext not in TEXT_EXT | CONVERT_EXT:
            continue
        source_id = path.relative_to(work_root).as_posix()
        try:
            text = extract_text(path, extractor)
            if len(text.strip()) < 30:
                raise RuntimeError("提取文本少于 30 字符")
        except Exception as exc:  # noqa: BLE001
            failures.append({"source_id": source_id, "error": str(exc)})
            continue
        source_row = {
            "source_id": source_id, "path": str(path), "extension": ext,
            "sha256": sha256_file(path), "characters": len(text), "status": "parsed",
        }
        sources.append(source_row)
        for index, part in enumerate(split_text(text), 1):
            fragments.append({
                "fragment_id": hashlib.sha256(
                    f"{source_id}\0{index}\0{part}".encode("utf-8")
                ).hexdigest()[:16],
                "source_id": source_id, "part": index, "text": part,
            })
    if not sources:
        raise SystemExit("未解析到可用材料")

    packet_limit = max(6000, args.packet_chars)
    packets: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    current_chars = 0
    for fragment in fragments:
        chars = len(fragment["text"])
        if current and current_chars + chars > packet_limit:
            packets.append(current)
            current, current_chars = [], 0
        current.append(fragment)
        current_chars += chars
    if current:
        packets.append(current)
    packets_dir.mkdir(parents=True, exist_ok=True)
    for index, packet in enumerate(packets, 1):
        write_json(packets_dir / f"packet-{index:04d}.json", {
            "packet_id": f"P{index:04d}", "fragments": packet,
        })
    manifest = {
        "input": str(source), "sources": sources, "failures": failures,
        "source_count": len(sources), "fragment_count": len(fragments),
        "packet_count": len(packets), "extractor": extractor,
    }
    write_json(run_dir / "source_manifest.json", manifest)
    status = "prepared_with_failures" if failures else "prepared"
    run_manifest = read_json(run_dir / "run_manifest.json")
    run_manifest.update({"status": status, "source_manifest": "source_manifest.json"})
    write_json(run_dir / "run_manifest.json", run_manifest)
    print(json.dumps({"status": status, **{k: manifest[k] for k in (
        "source_count", "fragment_count", "packet_count")},
        "failure_count": len(failures)}, ensure_ascii=False))


def validate_fact(item: Any, where: str, source_ids: set[str] | None = None) -> list[str]:
    issues: list[str] = []
    if not isinstance(item, dict):
        return [f"{where} 不是对象"]
    for key in ("category", "fact", "source_id", "source_quote"):
        if not str(item.get(key, "")).strip():
            issues.append(f"{where} 缺少 {key}")
    if item.get("category") not in FACT_CATEGORIES:
        issues.append(f"{where} category 非法：{item.get('category')}")
    if item.get("materiality") not in {1, 2, 3}:
        issues.append(f"{where} materiality 必须为 1-3")
    if source_ids is not None and item.get("source_id") not in source_ids:
        issues.append(f"{where} source_id 不在材料清单：{item.get('source_id')}")
    quote = re.sub(r"\s+", "", str(item.get("source_quote", "")))
    fact = re.sub(r"\s+", "", str(item.get("fact", "")))
    if len(quote) < 4:
        issues.append(f"{where} source_quote 过短")
    if len(fact) > 180:
        issues.append(f"{where} fact 超过 180 字")
    return issues


def command_merge_facts(args: argparse.Namespace) -> None:
    parts_dir = Path(args.parts_dir).expanduser().resolve()
    manifest = read_json(Path(args.manifest).expanduser().resolve())
    source_ids = {s["source_id"] for s in manifest.get("sources", [])}
    files = sorted(parts_dir.glob("packet-*.json"))
    if not files:
        raise SystemExit("facts_parts 中没有 packet-*.json")
    expected_count = int(manifest.get("packet_count", 0))
    if expected_count and len(files) != expected_count:
        raise SystemExit(
            f"事实分片不完整：应有 {expected_count} 个，实际 {len(files)} 个"
        )
    packet_sources: dict[str, str] = {}
    packets_dir = parts_dir.parent / "packets"
    for packet_path in sorted(packets_dir.glob("packet-*.json")):
        packet = read_json(packet_path)
        for fragment in packet.get("fragments", []):
            packet_sources.setdefault(fragment.get("source_id", ""), "")
            packet_sources[fragment.get("source_id", "")] += fragment.get("text", "")
    facts: list[dict[str, Any]] = []
    issues: list[str] = []
    seen: set[tuple[str, str]] = set()
    for path in files:
        data = read_json(path)
        items = data if isinstance(data, list) else data.get("facts", [])
        if not isinstance(items, list):
            issues.append(f"{path.name} 顶层不是数组或 facts 数组")
            continue
        for index, item in enumerate(items, 1):
            issues.extend(validate_fact(item, f"{path.name}[{index}]", source_ids))
            if isinstance(item, dict) and item.get("source_quote"):
                quote = re.sub(r"\s+", "", str(item["source_quote"]))
                source_text = re.sub(
                    r"\s+", "", packet_sources.get(str(item.get("source_id", "")), "")
                )
                if quote not in source_text:
                    issues.append(f"{path.name}[{index}] source_quote 不在对应材料原文中")
            if not isinstance(item, dict) or not item.get("fact"):
                continue
            key = (str(item.get("source_id")), re.sub(r"\s+", "", str(item["fact"])))
            if key in seen:
                continue
            seen.add(key)
            facts.append(dict(item))
    if issues:
        raise SystemExit("事实分片校验失败：\n- " + "\n- ".join(issues[:50]))
    facts.sort(key=lambda x: (-int(x.get("materiality", 1)), x["source_id"], x["fact"]))
    for index, fact in enumerate(facts, 1):
        fact["id"] = f"F{index:04d}"
        fact.setdefault("date_hint", "无")
    write_json(Path(args.output).expanduser().resolve(), facts)
    print(json.dumps({"facts": len(facts), "materiality_3": sum(
        f["materiality"] == 3 for f in facts)}, ensure_ascii=False))


def conflict_key(text: str) -> str:
    key = re.sub(r"[0-9][0-9,.]*", "#", text)
    key = re.sub(r"约|近|超过|逾|余|大于|不足|左右|上下|预计|已|将", "", key)
    return re.sub(r"\s|[，。、；：（）()《》\"']", "", key)[:28]


def command_build_conflicts(args: argparse.Namespace) -> None:
    facts = read_json(Path(args.facts).expanduser().resolve())
    buckets: dict[str, list[dict[str, Any]]] = {}
    for fact in facts:
        text = str(fact.get("fact", ""))
        if re.search(r"https?://|doi\.org|www\.", text, re.I):
            continue
        key = conflict_key(text)
        if len(key) >= 5:
            buckets.setdefault(key, []).append(fact)
    groups = []
    for key, group in buckets.items():
        values = {re.sub(r"\s+", "", str(x.get("fact", ""))) for x in group}
        if len(group) > 1 and len(values) > 1:
            groups.append({
                "topic_key": key, "facts": [x.get("id") for x in group],
                "statements": [{"id": x.get("id"), "fact": x.get("fact"),
                                  "source_id": x.get("source_id"),
                                  "date_hint": x.get("date_hint")} for x in group],
            })
    groups.sort(key=lambda g: g["topic_key"])
    write_json(Path(args.output).expanduser().resolve(), {"groups": groups})
    print(json.dumps({"conflict_groups": len(groups)}, ensure_ascii=False))


def validate_forecast(data: Any, fact_ids: set[str]) -> list[str]:
    if not isinstance(data, dict):
        return ["forecast 顶层不是对象"]
    years, lines = data.get("years", []), data.get("lines", [])
    if not lines:
        return []
    issues: list[str] = []
    totals = [0.0] * len(years)
    for index, line in enumerate(lines, 1):
        for key in ("qty", "price", "cost", "revenue", "margin"):
            if len(line.get(key, [])) != len(years):
                issues.append(f"产品线{index} {key} 长度与年份不一致")
        ids = line.get("source_fact_ids", [])
        if not ids or any(x not in fact_ids for x in ids):
            issues.append(f"产品线{index} 缺少有效事实 ID")
        if line.get("assumption_type") not in {"管理层目标", "项目组测算"}:
            issues.append(f"产品线{index} assumption_type 非法")
        if all(len(line.get(k, [])) == len(years) for k in ("qty", "price", "revenue")):
            scale = float(line.get("unit_scale", 1))
            for pos, (qty, price, revenue) in enumerate(zip(
                    line["qty"], line["price"], line["revenue"])):
                calculated = round(float(qty) * float(price) * scale, 1)
                if abs(calculated - float(revenue)) > 0.5:
                    issues.append(f"产品线{index} {years[pos]} 收入不自洽")
                totals[pos] += float(revenue)
    stated = data.get("total", [])
    if len(stated) != len(years) or any(
            abs(round(totals[i], 1) - float(stated[i])) > 0.5 for i in range(len(years))):
        issues.append("预测合计与产品线之和不一致")
    return issues


def validate_financials(data: Any, fact_ids: set[str]) -> list[str]:
    if not isinstance(data, dict):
        return ["financials 顶层不是对象"]
    periods, tables = data.get("periods", []), data.get("tables", [])
    if not tables:
        return []
    issues: list[str] = []
    if not periods:
        issues.append("财务表缺少期间")
    for ti, table in enumerate(tables, 1):
        for ri, row in enumerate(table.get("rows", []), 1):
            values = row.get("values", [])
            if len(values) != len(periods):
                issues.append(f"财务表{ti}第{ri}行列数不一致")
            ids = row.get("source_fact_ids", [])
            if any(re.search(r"\d", str(v)) for v in values) and (
                    not ids or any(x not in fact_ids for x in ids)):
                issues.append(f"财务表{ti}第{ri}行缺少有效事实 ID")
    return issues


def validate_kind(kind: str, data: Any, facts: list[dict[str, Any]] | None = None) -> list[str]:
    fact_ids = {x.get("id") for x in (facts or [])}
    if kind == "facts":
        if not isinstance(data, list):
            return ["facts 顶层必须为数组"]
        issues = []
        for index, item in enumerate(data, 1):
            issues.extend(validate_fact(item, f"facts[{index}]"))
            if not str(item.get("id", "")).startswith("F"):
                issues.append(f"facts[{index}] 缺少规范 ID")
        return issues
    if kind == "rulings":
        if not isinstance(data, dict):
            return ["rulings 顶层必须为对象"]
        if not isinstance(data.get("rulings"), list) or not isinstance(data.get("blacklist"), list):
            return ["rulings 必须包含 rulings 和 blacklist 数组"]
        return []
    if kind == "plan":
        if not isinstance(data, dict) or not isinstance(data.get("profile"), dict):
            return ["plan 缺少 profile"]
        sections = data.get("sections", {})
        issues = [f"plan 缺少章节：{name}" for name in OUTLINE if name not in sections]
        if sections.get("六、结论", {}).get("must_cover_ids"):
            issues.append("六、结论 must_cover_ids 必须为空")
        for name, section in sections.items():
            for fact_id in section.get("must_cover_ids", []):
                if fact_ids and fact_id not in fact_ids:
                    issues.append(f"{name} 引用了不存在的事实 ID：{fact_id}")
        return issues
    if kind == "financials":
        return validate_financials(data, fact_ids)
    if kind == "forecast":
        return validate_forecast(data, fact_ids)
    if kind == "draft":
        if not isinstance(data, dict):
            return ["draft 顶层必须为对象"]
        return [f"draft 缺少章节：{name}" for name in OUTLINE if name not in data]
    raise SystemExit(f"未知 kind：{kind}")


def command_validate_json(args: argparse.Namespace) -> None:
    data = read_json(Path(args.input).expanduser().resolve())
    facts = read_json(Path(args.facts).expanduser().resolve()) if args.facts else None
    issues = validate_kind(args.kind, data, facts)
    result = {"kind": args.kind, "status": "pass" if not issues else "fail", "issues": issues}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if issues:
        raise SystemExit(2)


def flatten_body(body: Any) -> str:
    if isinstance(body, dict):
        return "\n".join(str(value) for value in body.values())
    return str(body)


def fact_is_covered(fact: str, text: str) -> bool:
    anchors = [anchor for anchor in re.findall(
        r"[0-9][0-9,.]*%?|[A-Za-z][A-Za-z0-9-]{2,}|[\u4e00-\u9fff]{3,8}", fact
    ) if len(anchor) > 1][:8]
    if not anchors:
        return True
    flat = text.replace(",", "")
    hits = sum(anchor.replace(",", "") in flat for anchor in anchors)
    return hits >= max(2, len(anchors) // 3)


def command_audit_draft(args: argparse.Namespace) -> None:
    input_paths = {name: Path(getattr(args, name)).expanduser().resolve() for name in (
        "draft", "facts", "rulings", "plan", "financials", "forecast"
    )}
    draft = read_json(input_paths["draft"])
    facts = read_json(input_paths["facts"])
    rulings = read_json(input_paths["rulings"])
    plan = read_json(input_paths["plan"])
    financials = read_json(input_paths["financials"])
    forecast = read_json(input_paths["forecast"])
    issues: list[dict[str, str]] = []
    fact_ids = {f.get("id") for f in facts}
    for message in validate_financials(financials, fact_ids) + validate_forecast(forecast, fact_ids):
        issues.append({"section": "结构化数据", "issue": message})
    known_numbers = {n.replace(",", "") for n in re.findall(
        r"[0-9][0-9,.]{2,}", json.dumps(
            [facts, rulings, financials, forecast], ensure_ascii=False
        )
    )}
    facts_by_id = {f.get("id"): f for f in facts}
    blacklist = rulings.get("blacklist", [])
    for section, minimum in OUTLINE.items():
        body = draft.get(section)
        if body is None:
            issues.append({"section": section, "issue": "章节缺失"})
            continue
        text = flatten_body(body)
        if len(text.strip()) < minimum:
            issues.append({"section": section, "issue": f"正文 {len(text.strip())} 字低于下限 {minimum}"})
        for word in BANNED_VOICE:
            if word in text:
                issues.append({"section": section, "issue": f"出现禁用词：{word}"})
        for item in blacklist:
            bad = str(item.get("bad", ""))
            if bad and bad in text:
                issues.append({"section": section, "issue": f"使用被否决口径：{bad}"})
        must_ids = plan.get("sections", {}).get(section, {}).get("must_cover_ids", [])
        missing = [fact_id for fact_id in must_ids if fact_id not in fact_ids]
        for fact_id in missing:
            issues.append({"section": section, "issue": f"规划引用不存在事实：{fact_id}"})
        uncovered = [fact_id for fact_id in must_ids if fact_id in facts_by_id and not
                     fact_is_covered(str(facts_by_id[fact_id].get("fact", "")), text)]
        for fact_id in uncovered:
            issues.append({"section": section, "issue": f"未覆盖必写事实：{fact_id}"})
        orphan = sorted({n.replace(",", "") for n in re.findall(
            r"[0-9][0-9,.]{2,}", text
        ) if n.replace(",", "") not in known_numbers and not re.fullmatch(r"20[2-4]\d", n)})
        if orphan:
            issues.append({"section": section, "issue": f"数字无法溯源：{orphan[:8]}"})
    conclusion = flatten_body(draft.get("六、结论", {})).strip()
    if len(conclusion) > 260:
        issues.append({"section": "六、结论", "issue": "超过 260 字"})
    if not conclusion.startswith("综合考虑"):
        issues.append({"section": "六、结论", "issue": "未以“综合考虑”开头"})
    if not conclusion.endswith("建议按主方案实施本次投资。"):
        issues.append({"section": "六、结论", "issue": "未使用固定结尾"})
    if conclusion.count("。") != 1:
        issues.append({"section": "六、结论", "issue": "必须为一个完整句子"})
    report = {
        "status": "pass" if not issues else "fail",
        "standard": "VC_NATIVE_SKILL_V1", "residual_issue_count": len(issues),
        "external_llm_gateway": False,
        "input_sha256": {name: sha256_file(path) for name, path in input_paths.items()},
        "issues": issues,
    }
    write_json(Path(args.output).expanduser().resolve(), report)
    print(json.dumps({"status": report["status"], "issues": len(issues)}, ensure_ascii=False))
    if issues:
        raise SystemExit(2)


def load_engine(path: Path):
    if not path.is_file():
        raise SystemExit(f"原始渲染引擎不存在：{path}")
    spec = importlib.util.spec_from_file_location("vc_proposal_readonly_engine", path)
    if spec is None or spec.loader is None:
        raise SystemExit("无法加载原始渲染引擎")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    if not hasattr(module, "render_docx"):
        raise SystemExit("原始脚本缺少 render_docx()")
    return module


def command_render(args: argparse.Namespace) -> None:
    qc = read_json(Path(args.qc).expanduser().resolve())
    if qc.get("status") != "pass" or qc.get("residual_issue_count") != 0:
        raise SystemExit("质检未通过，拒绝渲染正式 DOCX")
    expected_hashes = qc.get("input_sha256", {})
    for name in ("draft", "forecast", "financials"):
        path = Path(getattr(args, name)).expanduser().resolve()
        if expected_hashes.get(name) != sha256_file(path):
            raise SystemExit(f"{name} 在质检后发生变化，拒绝渲染")
    output = Path(args.output).expanduser().resolve()
    if "'" in str(output) or "\\" in str(output):
        raise SystemExit("输出路径包含原始渲染器不支持的字符")
    output.parent.mkdir(parents=True, exist_ok=True)
    engine_path = Path(args.engine).expanduser().resolve()
    engine = load_engine(engine_path)
    draft = read_json(Path(args.draft).expanduser().resolve())
    forecast = read_json(Path(args.forecast).expanduser().resolve())
    financials = read_json(Path(args.financials).expanduser().resolve())
    meta = read_json(Path(args.meta).expanduser().resolve())
    fixed = getattr(engine, "FIXED_EXECUTION_AUTHORIZATION", "")
    if not fixed or len(meta.get("preamble", [])) < 2 or meta["preamble"][1] != fixed:
        raise SystemExit("meta.preamble[1] 与原始引擎固定授权语不一致")
    payload_text = json.dumps([draft, forecast, financials, meta], ensure_ascii=False)
    if "__OUT__" in payload_text:
        raise SystemExit("正文包含原始渲染器保留标记 __OUT__")
    engine.render_docx(draft, forecast or None, financials or None, meta, str(output))
    if not output.is_file() or output.stat().st_size < 1000 or not zipfile.is_zipfile(output):
        raise SystemExit("DOCX 生成后完整性检查失败")
    print(json.dumps({
        "status": "pass", "output": str(output), "bytes": output.stat().st_size,
        "engine": str(engine_path), "engine_sha256": sha256_file(engine_path),
    }, ensure_ascii=False))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="无网关投资提案 Skill 本地工具")
    sub = parser.add_subparsers(dest="command", required=True)

    init_run = sub.add_parser("init-run", help="创建隔离运行目录")
    init_run.add_argument("--input", required=True)
    init_run.add_argument("--output-root", required=True)
    init_run.set_defaults(func=command_init_run)

    prepare = sub.add_parser("prepare", help="解析材料并生成工作包")
    prepare.add_argument("--input", required=True)
    prepare.add_argument("--run-dir", required=True)
    prepare.add_argument("--packet-chars", type=int, default=24000)
    prepare.set_defaults(func=command_prepare)

    merge = sub.add_parser("merge-facts", help="合并并验证事实分片")
    merge.add_argument("--parts-dir", required=True)
    merge.add_argument("--manifest", required=True)
    merge.add_argument("--output", required=True)
    merge.set_defaults(func=command_merge_facts)

    conflicts = sub.add_parser("build-conflicts", help="生成冲突候选")
    conflicts.add_argument("--facts", required=True)
    conflicts.add_argument("--output", required=True)
    conflicts.set_defaults(func=command_build_conflicts)

    validate = sub.add_parser("validate-json", help="验证阶段 JSON")
    validate.add_argument("--kind", required=True,
                          choices=["facts", "rulings", "plan", "financials", "forecast", "draft"])
    validate.add_argument("--input", required=True)
    validate.add_argument("--facts")
    validate.set_defaults(func=command_validate_json)

    audit = sub.add_parser("audit-draft", help="执行正式稿硬门禁")
    for name in ("draft", "facts", "rulings", "plan", "financials", "forecast", "output"):
        audit.add_argument(f"--{name}", required=True)
    audit.set_defaults(func=command_audit_draft)

    render = sub.add_parser("render", help="只读导入原始脚本并渲染 DOCX")
    render.add_argument("--engine", default=str(DEFAULT_ENGINE))
    for name in ("draft", "forecast", "financials", "meta", "output", "qc"):
        render.add_argument(f"--{name}", required=True)
    render.set_defaults(func=command_render)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
