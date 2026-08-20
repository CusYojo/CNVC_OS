#!/usr/bin/env python3
"""Deterministic intake, audit, Deta formatting, and render support for DD reports."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from investment_bank_styles import apply_styles, audit_styles
from deta_structure_contract import audit_structure
from editorial_quality_audit import audit_editorial_quality


SKILL_DIR = Path(__file__).resolve().parents[1]
REFERENCE = SKILL_DIR / "assets" / "reference.docx"
EXPECTED_SECTIONS = [
    "1、投资概要", "2、公司概况", "3、产品与技术", "4、业务情况",
    "5、行业和市场", "6、未来发展规划", "7、投资方案",
    "8、风险提示与对策", "投资结论及建议",
]
SAMPLE_TERMS = [
    "北京德塔源创", "德塔源创", "德塔智能", "马晓健", "通脑平台",
]
CRITICAL_EXTS = {".docx", ".pdf", ".xlsx", ".xls", ".csv", ".txt", ".md"}
MATERIAL_CATEGORIES = {
    "customer", "financial", "ownership", "technology", "legal",
    "transaction", "forecast",
}
FONT_BODY = "SimSun"
FONT_BODY_EA = "宋体"
FONT_HEADING = "黑体"


def dump_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_json(path: Path, default=None):
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def safe_extract_zip(src: Path, out: Path, max_bytes: int = 8 * 1024**3) -> None:
    out.mkdir(parents=True, exist_ok=True)
    root = out.resolve()
    total = 0
    with zipfile.ZipFile(src) as zf:
        for info in zf.infolist():
            total += info.file_size
            if total > max_bytes:
                raise SystemExit("archive uncompressed size exceeds safety limit")
            target = (out / info.filename).resolve()
            if root != target and root not in target.parents:
                raise SystemExit(f"unsafe archive path: {info.filename}")
        zf.extractall(out)


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    text = "\n\f\n".join((page.extract_text() or "") for page in reader.pages)
    if len(re.sub(r"\s+", "", text)) >= 20:
        return text
    tesseract = shutil.which("tesseract")
    if not tesseract:
        return text
    import pymupdf
    ocr_pages = []
    pdf = pymupdf.open(str(path))
    with tempfile.TemporaryDirectory(prefix="deta-dd-ocr-") as tmp:
        for page_no, page in enumerate(pdf, 1):
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            image = Path(tmp) / f"page-{page_no:04d}.png"
            pix.save(str(image))
            done = subprocess.run(
                [tesseract, str(image), "stdout", "-l", "chi_sim+eng", "--psm", "6"],
                capture_output=True, text=True, timeout=180,
            )
            ocr_pages.append(done.stdout if done.returncode == 0 else "")
    return "\n\f\n".join(ocr_pages)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## SHEET: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                out.append("\t".join(vals))
    return "\n".join(out)


def extract_xls(path: Path) -> str:
    head = path.read_bytes()[:256].lstrip()
    if head.startswith(b"<?xml") or b"<Workbook" in head:
        root = ET.parse(path).getroot()
        ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
        out = []
        for worksheet in root.findall(".//ss:Worksheet", ns):
            name = worksheet.get("{urn:schemas-microsoft-com:office:spreadsheet}Name", "Sheet")
            out.append(f"## SHEET: {name}")
            for row in worksheet.findall(".//ss:Table/ss:Row", ns):
                vals = []
                for cell in row.findall("ss:Cell", ns):
                    data = cell.find("ss:Data", ns)
                    vals.append("" if data is None or data.text is None else data.text)
                if any(vals):
                    out.append("\t".join(vals))
        return "\n".join(out)
    try:
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None, header=None)
    except Exception as exc:
        raise RuntimeError(f"xls parser unavailable: {exc}") from exc
    out = []
    for name, frame in sheets.items():
        out.append(f"## SHEET: {name}")
        out.append(frame.fillna("").astype(str).to_csv(index=False, header=False, sep="\t"))
    return "\n".join(out)


def extract_text(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext in {".txt", ".md", ".csv", ".tsv"}:
        return path.read_text(encoding="utf-8", errors="replace"), "text"
    if ext == ".docx":
        return extract_docx(path), "python-docx"
    if ext == ".pdf":
        return extract_pdf(path), "pypdf"
    if ext == ".xlsx":
        return extract_xlsx(path), "openpyxl"
    if ext == ".xls":
        return extract_xls(path), "pandas"
    return "", "not-applicable"


def command_init(args) -> int:
    source = Path(args.input).expanduser().resolve()
    run = Path(args.run_dir).expanduser().resolve()
    if not source.exists():
        raise SystemExit(f"input not found: {source}")
    run.mkdir(parents=True, exist_ok=True)
    expanded = run / "expanded"
    if source.is_dir():
        source_root = source
    elif source.suffix.lower() == ".zip":
        if expanded.exists() and any(expanded.iterdir()):
            raise SystemExit(f"expanded directory is not empty: {expanded}")
        safe_extract_zip(source, expanded)
        source_root = expanded
    else:
        raise SystemExit("input must be a ZIP archive or directory")

    packets = run / "packets"
    packets.mkdir(exist_ok=True)
    manifest = []
    files = sorted(p for p in source_root.rglob("*") if p.is_file())
    for idx, path in enumerate(files, 1):
        sid = f"S{idx:04d}"
        rel = path.relative_to(source_root).as_posix()
        record = {
            "source_id": sid,
            "path": str(path),
            "relative_path": rel,
            "extension": path.suffix.lower(),
            "size": path.stat().st_size,
            "sha256": sha256(path),
            "parse_status": "not_applicable",
            "parser": "not-applicable",
            "text_chars": 0,
        }
        if path.suffix.lower() in CRITICAL_EXTS:
            try:
                text, parser = extract_text(path)
                txt_path = packets / f"{sid}.txt"
                txt_path.write_text(text, encoding="utf-8")
                record.update(
                    parse_status="parsed" if text.strip() else "empty",
                    parser=parser,
                    text_chars=len(text),
                    text_path=str(txt_path),
                )
            except Exception as exc:
                record.update(parse_status="failed", parse_error=str(exc))
        manifest.append(record)

    state = {
        "schema_version": 2,
        "project": args.project,
        "investor": args.investor,
        "cutoff_date": args.cutoff_date,
        "currency": args.currency,
        "round": args.round,
        "founding_period": None,
        "financial_gap": True,
        "output_type": "上会尽职调查报告",
        "input": str(source),
        "input_sha256": sha256(source) if source.is_file() else None,
        "reference": str(REFERENCE),
        "reference_sha256": sha256(REFERENCE),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "status": "initialized",
    }
    dump_json(run / "workflow_state.json", state)
    dump_json(run / "source_manifest.json", manifest)
    placeholders = {
        "facts.json": [],
        "rulings.json": [],
        "report_plan.json": {"sections": EXPECTED_SECTIONS},
        "report.json": {
            "schema_version": 2,
            "title": args.project,
            "decision": {
                "recommendation": "",
                "core_judgment": "",
                "rationale_fact_ids": [],
                "transaction": {
                    "amount": "",
                    "valuation": "",
                    "fully_diluted_ownership": "",
                    "payment": "",
                },
                "conditions_precedent": [],
                "termination_boundary": "",
            },
            "value_logic": [],
            "risk_register": [],
            "sections": [],
        },
        "revision_log.json": [],
    }
    for name, data in placeholders.items():
        path = run / name
        if not path.exists():
            dump_json(path, data)
    print(json.dumps({
        "status": "initialized", "run_dir": str(run), "files": len(files),
        "parsed": sum(x["parse_status"] == "parsed" for x in manifest),
        "failed": sum(x["parse_status"] == "failed" for x in manifest),
    }, ensure_ascii=False))
    return 0


def report_text(report) -> str:
    return json.dumps(report or {}, ensure_ascii=False)


def numeric_statement(text: str) -> bool:
    return bool(re.search(r"\d", text or ""))


def collect_text_leaves(value, path="$") -> list[tuple[str, str]]:
    leaves = []
    if isinstance(value, dict):
        for key, child in value.items():
            leaves.extend(collect_text_leaves(child, f"{path}.{key}"))
    elif isinstance(value, list):
        for idx, child in enumerate(value):
            leaves.extend(collect_text_leaves(child, f"{path}[{idx}]"))
    elif isinstance(value, str) and value.strip():
        leaves.append((path, value.strip()))
    return leaves


def repeated_passages(report, min_chars: int = 60) -> list[dict]:
    groups = {}
    for path, text in collect_text_leaves(report):
        normalized = re.sub(r"\s+", "", text)
        if len(normalized) < min_chars:
            continue
        groups.setdefault(normalized, []).append(path)
    return [
        {"text": normalized[:100], "paths": paths}
        for normalized, paths in groups.items() if len(set(paths)) > 1
    ]


def validate_decision_contract(report, fact_ids: set[str]) -> list[str]:
    errors = []
    decision = report.get("decision") if isinstance(report, dict) else None
    if not isinstance(decision, dict):
        return ["decision object is missing"]
    if decision.get("recommendation") not in {"invest", "conditional_invest", "do_not_invest"}:
        errors.append("decision.recommendation is invalid")
    core = str(decision.get("core_judgment") or "").strip()
    if not core:
        errors.append("decision.core_judgment is missing")
    elif len(core) > 220 or len(re.findall(r"[。！？!?]", core)) > 2:
        errors.append("decision.core_judgment must remain one or two concise sentences")
    rationale = set(decision.get("rationale_fact_ids") or [])
    if not rationale:
        errors.append("decision.rationale_fact_ids is empty")
    if rationale - fact_ids:
        errors.append("decision contains unknown rationale fact IDs")
    transaction = decision.get("transaction") or {}
    for key in ["amount", "valuation", "fully_diluted_ownership", "payment"]:
        if not str(transaction.get(key) or "").strip():
            errors.append(f"decision.transaction.{key} is missing")
    conditions = decision.get("conditions_precedent") or []
    if decision.get("recommendation") == "conditional_invest" and not conditions:
        errors.append("conditional investment has no conditions precedent")
    for idx, item in enumerate(conditions):
        if not isinstance(item, dict):
            errors.append(f"condition {idx + 1} is malformed")
            continue
        for key in ["id", "priority", "condition", "action_if_unmet"]:
            if not item.get(key):
                errors.append(f"condition {idx + 1} lacks {key}")
        if item.get("priority") not in {"P0", "P1"}:
            errors.append(f"condition {idx + 1} priority must be P0/P1")
        refs = set(item.get("fact_ids") or [])
        if not refs or refs - fact_ids:
            errors.append(f"condition {idx + 1} lacks valid fact IDs")
    if not str(decision.get("termination_boundary") or "").strip():
        errors.append("decision.termination_boundary is missing")

    values = report.get("value_logic") or []
    layers = {item.get("layer") for item in values if isinstance(item, dict)}
    required_layers = {"verified_base", "growth_option", "transaction_protection"}
    if not required_layers.issubset(layers):
        errors.append("value_logic lacks one or more required layers")
    for idx, item in enumerate(values):
        if not isinstance(item, dict):
            errors.append(f"value item {idx + 1} is malformed")
            continue
        for key in ["claim", "mechanism", "investment_implication", "validation_boundary", "status"]:
            if not item.get(key):
                errors.append(f"value item {idx + 1} lacks {key}")
        if item.get("status") not in {"supported", "partially_supported", "not_applicable"}:
            errors.append(f"value item {idx + 1} has invalid status")
        refs = set(item.get("fact_ids") or [])
        if item.get("status") != "not_applicable" and (not refs or refs - fact_ids):
            errors.append(f"value item {idx + 1} lacks valid fact IDs")

    risks = report.get("risk_register") or []
    if not risks:
        errors.append("risk_register is empty")
    for idx, item in enumerate(risks):
        if not isinstance(item, dict):
            errors.append(f"risk {idx + 1} is malformed")
            continue
        for key in ["id", "priority", "category", "description", "value_transmission",
                    "monitoring_trigger", "pre_control", "action_if_unmet"]:
            if not item.get(key):
                errors.append(f"risk {idx + 1} lacks {key}")
        if item.get("priority") not in {"P0", "P1", "P2", "P3"}:
            errors.append(f"risk {idx + 1} has invalid priority")
        refs = set(item.get("fact_ids") or [])
        if item.get("priority") in {"P0", "P1"} and (not refs or refs - fact_ids):
            errors.append(f"risk {idx + 1} lacks valid fact IDs")

    sections = report.get("sections") or []
    conclusion = next(
        (item for item in sections if isinstance(item, dict) and item.get("title") == "投资结论及建议"),
        None,
    )
    paragraphs = conclusion.get("paragraphs") if conclusion else None
    if not isinstance(paragraphs, list) or len(paragraphs) != 3 or any(not str(p).strip() for p in paragraphs):
        errors.append("reader-facing conclusion must contain exactly three nonempty paragraphs")
    section_text = report_text({"sections": sections})
    for heading in ["7.1 投资亮点", "7.2 公司估值与投资方式"]:
        if heading not in section_text:
            errors.append(f"reader-facing report lacks {heading}")
    return errors


def command_audit(args) -> int:
    run = Path(args.run_dir).expanduser().resolve()
    state = load_json(run / "workflow_state.json", {})
    manifest = load_json(run / "source_manifest.json", [])
    facts = load_json(run / "facts.json", [])
    rulings = load_json(run / "rulings.json", [])
    report = load_json(run / "report.json", {})
    revisions = load_json(run / "revision_log.json", [])

    issues = []
    identity_fields = ["project", "cutoff_date", "currency", "round"]
    g0 = all(state.get(k) for k in identity_fields)
    if not g0:
        issues.append({"priority": "P0", "gate": "G0", "issue": "decision identity is incomplete"})

    failed_critical = [x for x in manifest if x.get("extension") in CRITICAL_EXTS and x.get("parse_status") in {"failed", "empty"}]
    g1 = bool(manifest) and not failed_critical
    for item in failed_critical:
        issues.append({"priority": "P1", "gate": "G1", "issue": f"parse failed: {item.get('relative_path')}"})

    fact_ids = set()
    malformed = []
    for fact in facts:
        fid = fact.get("id")
        if fid:
            fact_ids.add(fid)
        required = ["id", "category", "statement", "status", "source_ids"]
        if any(not fact.get(k) for k in required):
            malformed.append(fid or "<missing-id>")
            continue
        if fact.get("category") in MATERIAL_CATEGORIES or numeric_statement(fact.get("statement", "")):
            if any(not fact.get(k) for k in ["subject", "period", "unit", "scope", "confidence"]):
                malformed.append(fid)
    g2 = bool(facts) and not malformed
    if malformed:
        issues.append({"priority": "P1", "gate": "G2", "issue": "malformed material facts: " + ", ".join(malformed[:20])})

    conflicts = Counter(f.get("conflict_group") for f in facts if f.get("conflict_group"))
    required_rulings = {group for group, count in conflicts.items() if count > 1}
    ruled = {r.get("conflict_group") for r in rulings if r.get("decision") and r.get("source_ids")}
    open_conflicts = sorted(required_rulings - ruled)
    g3 = not open_conflicts
    if open_conflicts:
        issues.append({"priority": "P0", "gate": "G3", "issue": "unruled conflicts: " + ", ".join(open_conflicts)})

    customer_facts = [f for f in facts if f.get("category") == "customer"]
    bad_stage = [f.get("id") for f in customer_facts if not f.get("stage")]
    g4 = bool(customer_facts) and not bad_stage
    if bad_stage:
        issues.append({"priority": "P1", "gate": "G4", "issue": "customer facts lack stage: " + ", ".join(bad_stage[:20])})

    sections = report.get("sections") or []
    section_names = [s.get("title") for s in sections if isinstance(s, dict)]
    g5 = section_names == EXPECTED_SECTIONS
    if not g5:
        issues.append({"priority": "P1", "gate": "G5", "issue": "report sections do not match fixed nine-section structure"})

    rtext = report_text(report)
    contamination = [term for term in SAMPLE_TERMS if term in rtext]
    g6 = not contamination
    if contamination:
        issues.append({"priority": "P0", "gate": "G6", "issue": "sample contamination: " + ", ".join(contamination)})

    cited = set(re.findall(r"F-\d+", rtext))
    unknown_cites = sorted(cited - fact_ids)
    confirmed_fin = [f for f in facts if f.get("category") == "financial" and f.get("status") == "confirmed"]
    finance_periods = {f.get("period") for f in confirmed_fin if f.get("period")}
    finance_types = {f.get("statement_type") for f in confirmed_fin if f.get("statement_type")}
    founding_period = state.get("founding_period")
    finance_gap = state.get("financial_gap")
    required_types = {"balance_sheet", "income_statement", "cash_flow_or_bank"}
    g7 = bool(finance_periods) and required_types.issubset(finance_types) and not finance_gap
    if not g7:
        issues.append({"priority": "P0", "gate": "G7", "issue": "financial coverage or cash/bank reconciliation remains incomplete"})
    if unknown_cites:
        issues.append({"priority": "P1", "gate": "G5", "issue": "unknown fact references: " + ", ".join(unknown_cites[:20])})

    contract_errors = validate_decision_contract(report, fact_ids)
    g8 = not contract_errors
    if contract_errors:
        issues.append({"priority": "P1", "gate": "G8", "issue": "decision contract: " + "; ".join(contract_errors[:20])})

    # Internal decision controls intentionally mirror reader-facing facts.
    # Audit repetition only inside the reader-facing section tree.
    duplicates = repeated_passages({"sections": sections})
    g9 = not duplicates
    if duplicates:
        sample_paths = [" / ".join(item["paths"][:3]) for item in duplicates[:10]]
        issues.append({"priority": "P2", "gate": "G9", "issue": "exact repeated passages: " + "; ".join(sample_paths)})

    open_p0 = sum(x.get("priority") == "P0" and x.get("status") != "closed" for x in revisions if isinstance(x, dict))
    open_p1 = sum(x.get("priority") == "P1" and x.get("status") != "closed" for x in revisions if isinstance(x, dict))
    open_p2 = sum(x.get("priority") == "P2" and x.get("status") != "closed" for x in revisions if isinstance(x, dict))
    gates = {"G0_identity": g0, "G1_sources": g1, "G2_facts": g2, "G3_conflicts": g3,
             "G4_customer_stage": g4, "G5_structure_traceability": g5 and not unknown_cites,
             "G6_sample_contamination": g6, "G7_financial_coverage": g7,
             "G8_decision_contract": g8, "G9_reader_discipline": g9}
    score = max(0, 100 - 15 * sum(not x for x in gates.values()) - 5 * len([i for i in issues if i["priority"] == "P1"]))
    hard_pass = all(gates.values()) and open_p0 == 0 and open_p1 == 0 and open_p2 == 0
    status = "pass" if hard_pass and score >= 85 else (
        "blocked" if open_p0 or any(i["priority"] == "P0" for i in issues) else "conditional"
    )
    qc = {
        "status": status, "score": score, "hard_gates": gates,
        "open_revision_P0": open_p0, "open_revision_P1": open_p1,
        "open_revision_P2": open_p2,
        "issues": issues, "founding_period": founding_period,
        "release_ready": status == "pass" and score >= 85,
    }
    dump_json(run / "qc_report.json", qc)
    print(json.dumps(qc, ensure_ascii=False, indent=2))
    return 0 if qc["release_ready"] else 2


def set_run_font(run, family=FONT_BODY, east_asia=FONT_BODY_EA, size=12, bold=None):
    run.font.name = family
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key, val in (("ascii", family), ("hAnsi", family), ("eastAsia", east_asia), ("cs", family)):
        rfonts.set(qn(f"w:{key}"), val)


def set_cell_shading(cell, fill="E7E6E6"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def replace_style(target: Document, source: Document, name: str, target_normal_id: str):
    old = target.styles[name]._element
    new = copy.deepcopy(source.styles[name]._element)
    # Paragraphs reference the target document's existing styleId. Preserve it
    # while copying the reference style's formatting payload.
    old_style_id = old.get(qn("w:styleId"))
    if old_style_id:
        new.set(qn("w:styleId"), old_style_id)
    # Rebind the copied style's internal references to the target package.
    # Reference and target DOCX files use different numeric style IDs.
    for tag in ("basedOn", "next"):
        node = new.find(qn(f"w:{tag}"))
        if node is not None:
            node.set(qn("w:val"), target_normal_id)
    link = new.find(qn("w:link"))
    if link is not None:
        new.remove(link)
    old.getparent().replace(old, new)


def nonempty_paragraphs(doc: Document):
    return [p for p in doc.paragraphs if p.text.strip()]


def copy_paragraph_format(src, dst):
    if src._p.pPr is not None:
        old = dst._p.pPr
        new = copy.deepcopy(src._p.pPr)
        if old is None:
            dst._p.insert(0, new)
        else:
            dst._p.replace(old, new)
    if src.runs and dst.runs and src.runs[0]._r.rPr is not None:
        for run in dst.runs:
            old = run._r.rPr
            new = copy.deepcopy(src.runs[0]._r.rPr)
            if old is None:
                run._r.insert(0, new)
            else:
                run._r.replace(old, new)


def remove_stray_cover_paragraphs(doc: Document):
    for p in list(doc.paragraphs[:8]):
        if p.text.strip() in {"Hu", "H", "U"}:
            p._element.getparent().remove(p._element)


def command_format(args) -> int:
    src = Path(args.input).expanduser().resolve()
    out = Path(args.output).expanduser().resolve()
    if src == out:
        raise SystemExit("output must differ from input")
    doc = Document(str(src))
    ref = Document(str(REFERENCE))
    remove_stray_cover_paragraphs(doc)

    # Match reference cover roles by position among nonempty paragraphs.
    ref_ne = nonempty_paragraphs(ref)
    dst_ne = nonempty_paragraphs(doc)
    if len(ref_ne) >= 4 and len(dst_ne) >= 4:
        for src_p, dst_p in zip(ref_ne[:4], dst_ne[:4]):
            copy_paragraph_format(src_p, dst_p)

    # Apply the named investment-banking style system after content and cover
    # roles have been identified. This is the sole typography authority.
    style_counts = apply_styles(doc)

    # Apply reference geometry by orientation while preserving section breaks.
    ref_portrait = ref.sections[0]
    ref_land_a = next((s for s in ref.sections if s.orientation == WD_ORIENT.LANDSCAPE), None)
    landscape_seen = 0
    for section in doc.sections:
        if section.orientation == WD_ORIENT.LANDSCAPE and ref_land_a:
            landscape_seen += 1
            candidates = [s for s in ref.sections if s.orientation == WD_ORIENT.LANDSCAPE]
            model = candidates[min(landscape_seen - 1, len(candidates) - 1)]
        else:
            model = ref_portrait
        for attr in ["page_width", "page_height", "left_margin", "right_margin", "top_margin",
                     "bottom_margin", "header_distance", "footer_distance"]:
            setattr(section, attr, getattr(model, attr))

    # Use the exact reference declaration on non-cover headers.
    declaration = "申明：本报告为赛智伯乐内部项目文件，禁止外传，报告内容仅代表机构观点"
    for idx, section in enumerate(doc.sections):
        if idx == 0:
            continue
        header = section.header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.text = declaration
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            set_run_font(run, FONT_BODY, FONT_BODY_EA, 9)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    final_doc = Document(str(out))
    # The formatter guarantees a native TOC field and marks it dirty. Cached
    # entries and page numbers become authoritative only after Word refreshes
    # all fields during the native verification step.
    style_audit = audit_styles(final_doc, require_updated_toc=False)
    structure_audit = audit_structure(final_doc)
    editorial_audit = audit_editorial_quality(final_doc)
    passed = (style_audit["status"] == "pass" and
              structure_audit["status"] == "pass" and
              editorial_audit["status"] == "pass")
    print(json.dumps({
        "status": "formatted" if passed else "format_fail",
        "output": str(out), "sha256": sha256(out),
        "style_counts": style_counts, "style_audit": style_audit,
        "structure_audit": structure_audit,
        "editorial_audit": editorial_audit,
    }, ensure_ascii=False))
    return 0 if passed else 2


def bundled_binary(name: str) -> str | None:
    candidates: list[str | Path | None] = [shutil.which(name)]
    if sys.platform == "darwin":
        candidates.extend([
            Path("/opt/homebrew/bin") / name,
            Path("/usr/local/bin") / name,
        ])
        if name in {"soffice", "libreoffice"}:
            candidates.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(candidate)
    return None


def word_export(docx: Path, pdf: Path) -> tuple[bool, str]:
    if sys.platform != "darwin" or not Path("/Applications/Microsoft Word.app").is_dir():
        return False, "Microsoft Word unavailable"
    launched = subprocess.run(
        ["open", "-a", "Microsoft Word", str(docx)],
        capture_output=True, text=True, timeout=30,
    )
    if launched.returncode != 0:
        return False, (launched.stderr or launched.stdout)[-800:]
    script = r'''
on run argv
  set pdfPath to item 1 of argv
  set docName to item 2 of argv
  tell application "Microsoft Word"
    repeat 20 times
      if exists document docName then exit repeat
      delay 0.5
    end repeat
    if not (exists document docName) then error "Word did not open " & docName
    set docRef to document docName
    if (count of tables of contents of docRef) > 0 then
      update (table of contents 1 of docRef)
      update page numbers (table of contents 1 of docRef)
    end if
    save as docRef file name pdfPath file format format PDF
    close docRef saving yes
  end tell
end run
'''
    done = subprocess.run(["osascript", "-", str(pdf), docx.name], input=script,
                          text=True, capture_output=True, timeout=120)
    return done.returncode == 0 and pdf.exists(), (done.stderr or done.stdout)[-800:]


def libreoffice_export(docx: Path, outdir: Path) -> tuple[Path | None, str]:
    soffice = bundled_binary("soffice") or bundled_binary("libreoffice")
    if not soffice:
        return None, "LibreOffice unavailable"
    profile = Path(tempfile.mkdtemp(prefix="deta-dd-lo-"))
    env = os.environ.copy()
    env["HOME"] = str(profile)
    env["TMPDIR"] = "/private/tmp" if Path("/private/tmp").is_dir() else tempfile.gettempdir()
    done = subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{profile}",
                           "--convert-to", "pdf", "--outdir", str(outdir), str(docx)],
                          capture_output=True, text=True, timeout=180, env=env)
    shutil.rmtree(profile, ignore_errors=True)
    pdf = outdir / f"{docx.stem}.pdf"
    return (pdf if done.returncode == 0 and pdf.exists() else None,
            (done.stderr or done.stdout)[-800:])


def command_verify(args) -> int:
    docx = Path(args.docx).expanduser().resolve()
    out = Path(args.output_dir).expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    pdf = out / f"{docx.stem}.pdf"
    engine = "microsoft-word" if args.word_native else "libreoffice"
    ok, reason = word_export(docx, pdf) if args.word_native else (False, "not requested")
    if not ok:
        pdf_obj, fallback_reason = libreoffice_export(docx, out)
        if not pdf_obj:
            raise SystemExit(reason or fallback_reason)
        pdf = pdf_obj
        engine = "libreoffice"
        reason = reason or fallback_reason

    info_bin, ppm_bin = bundled_binary("pdfinfo"), bundled_binary("pdftoppm")
    if not info_bin or not ppm_bin:
        raise SystemExit("pdfinfo/pdftoppm unavailable")
    info = subprocess.run([info_bin, str(pdf)], capture_output=True, text=True, timeout=30)
    match = re.search(r"^Pages:\s+(\d+)", info.stdout, re.M)
    pages = int(match.group(1)) if match else 0
    for old in out.glob("page-*.png"):
        old.unlink()
    raster = subprocess.run([ppm_bin, "-png", "-r", "110", str(pdf), str(out / "page")],
                            capture_output=True, text=True, timeout=300)
    pngs = sorted(out.glob("page-*.png"))
    doc = Document(str(docx))
    style_audit = audit_styles(doc)
    structure_audit = audit_structure(doc)
    editorial_audit = audit_editorial_quality(doc)
    full_text = "\n".join(p.text for p in doc.paragraphs)
    contamination = [term for term in SAMPLE_TERMS if term in full_text]
    issues = []
    if pages <= 0:
        issues.append("invalid PDF page count")
    if raster.returncode or len(pngs) != pages:
        issues.append(f"PNG count mismatch: {len(pngs)}/{pages}")
    if contamination:
        issues.append("sample contamination: " + ", ".join(contamination))
    if any(p.text.strip() in {"Hu", "H", "U"} for p in doc.paragraphs[:8]):
        issues.append("stray cover text")
    if style_audit["status"] != "pass":
        issues.append("named style audit failed: " + "; ".join(style_audit["issues"]))
    if structure_audit["status"] != "pass":
        issues.append("Deta V5 structure audit failed: " + "; ".join(structure_audit["issues"]))
    if editorial_audit["status"] != "pass":
        issues.append("Deta editorial audit failed: " + "; ".join(
            item["issue"] for item in editorial_audit["issues"][:10]
        ))
    result = {
        "status": "rendered_for_review" if not issues else "fail",
        "engine": engine, "page_count": pages, "png_count": len(pngs),
        "pdf": str(pdf), "output_dir": str(out), "issues": issues,
        "requires_all_page_visual_review": True,
        "style_audit": style_audit,
        "structure_audit": structure_audit,
        "editorial_audit": editorial_audit,
        "engine_note": reason if engine != "microsoft-word" else "",
    }
    dump_json(out / "render_audit.json", result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not issues else 2


def build_parser():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    p = sub.add_parser("init")
    p.add_argument("--input", required=True)
    p.add_argument("--run-dir", required=True)
    p.add_argument("--project", required=True)
    p.add_argument("--investor", default="赛智伯乐或其指定主体")
    p.add_argument("--cutoff-date", default=datetime.now().date().isoformat())
    p.add_argument("--currency", default="CNY")
    p.add_argument("--round", default="待锁定")
    p.set_defaults(func=command_init)

    p = sub.add_parser("audit")
    p.add_argument("--run-dir", required=True)
    p.set_defaults(func=command_audit)

    p = sub.add_parser("format")
    p.add_argument("--input", required=True)
    p.add_argument("--output", required=True)
    p.set_defaults(func=command_format)

    p = sub.add_parser("verify")
    p.add_argument("--docx", required=True)
    p.add_argument("--output-dir", required=True)
    p.add_argument("--word-native", action="store_true")
    p.set_defaults(func=command_verify)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
