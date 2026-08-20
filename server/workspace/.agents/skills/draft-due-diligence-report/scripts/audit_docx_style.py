#!/usr/bin/env python3
"""Audit a generated DOCX against the retained Deta V5 visual contract."""

from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from docx_builder_lib import (
    FOOTER_DISTANCE_CM,
    HEADER_DISTANCE_CM,
    MARGIN_BOTTOM_CM,
    MARGIN_LEFT_CM,
    MARGIN_RIGHT_CM,
    MARGIN_TOP_CM,
    PAGE_HEIGHT_CM,
    PAGE_WIDTH_CM,
    TEXT_WIDTH_DXA,
    LANDSCAPE_TEXT_WIDTH_DXA,
    HEADER_NOTICE,
)


STYLE_SPECS = {
    "尽调-正文": ("仿宋_GB2312", 12.0),
    "尽调-一级标题": ("黑体", 15.0),
    "尽调-二级标题": ("黑体", 15.0),
    "尽调-三级标题": ("黑体", 12.0),
    "尽调-四级标题": ("黑体", 12.0),
    "尽调-表正文": ("仿宋_GB2312", 12.0),
    "尽调-封面主体": ("黑体", 22.0),
    "尽调-封面日期": ("仿宋_GB2312", 14.0),
}

FIXED_REPORT_TITLE = "尽职调查报告"
FORBIDDEN_COVER_TITLE_TOKENS = (
    "公开信息预尽职调查报告",
    "初步尽职调查报告",
    "预尽调报告",
    "投前预审报告",
)

ALLOWED_VISIBLE_PARAGRAPH_STYLES = {
    "尽调-正文",
    "尽调-一级标题",
    "尽调-二级标题",
    "尽调-三级标题",
    "尽调-四级标题",
    "尽调-表标题",
    "尽调-图注",
    "尽调-提示",
    "尽调-封面主体",
    "尽调-封面日期",
    "尽调-目录标题",
    "List Bullet",
    "List Number",
    "toc 1",
    "toc 2",
    "toc 3",
    "toc 4",
}


def cm(value) -> float:
    return float(value.cm)


def spacing_is_zero(value) -> bool:
    return value is None or abs(float(value.pt)) < 0.05


def style_east_asia(style) -> str:
    rpr = style.element.rPr
    if rpr is None or rpr.rFonts is None:
        return style.font.name or ""
    return rpr.rFonts.get(qn("w:eastAsia")) or style.font.name or ""


def all_visible_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def audit(path: Path) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    doc = Document(path)
    cover_title_paragraphs = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.style is not None
        and paragraph.style.name == "尽调-封面主体"
        and paragraph.text.strip()
    ]
    cover_title_text = "\n".join(cover_title_paragraphs)
    if FIXED_REPORT_TITLE not in cover_title_paragraphs:
        errors.append(f"cover must contain a title paragraph exactly equal to '{FIXED_REPORT_TITLE}'")
    for token in FORBIDDEN_COVER_TITLE_TOKENS:
        if token in cover_title_text:
            errors.append(f"forbidden cover title variant: {token}")
    if len(doc.sections) < 3:
        errors.append("document must contain cover, TOC, and body sections")
    for index, section in enumerate(doc.sections, 1):
        if section.orientation == WD_ORIENT.LANDSCAPE:
            expected = {
                "page_width": PAGE_HEIGHT_CM,
                "page_height": PAGE_WIDTH_CM,
                "left_margin": 1.80,
                "right_margin": MARGIN_RIGHT_CM,
                "top_margin": MARGIN_TOP_CM,
                "bottom_margin": MARGIN_BOTTOM_CM,
                "header_distance": HEADER_DISTANCE_CM,
                "footer_distance": FOOTER_DISTANCE_CM,
            }
        else:
            expected = {
                "page_width": PAGE_WIDTH_CM,
                "page_height": PAGE_HEIGHT_CM,
                "left_margin": MARGIN_LEFT_CM,
                "right_margin": MARGIN_RIGHT_CM,
                "top_margin": MARGIN_TOP_CM,
                "bottom_margin": MARGIN_BOTTOM_CM,
                "header_distance": HEADER_DISTANCE_CM,
                "footer_distance": FOOTER_DISTANCE_CM,
            }
        for attr, target in expected.items():
            actual = cm(getattr(section, attr))
            if abs(actual - target) > 0.04:
                errors.append(f"section {index} {attr}: {actual:.3f}cm, expected {target:.3f}cm")

    for name, (font, size) in STYLE_SPECS.items():
        try:
            style = doc.styles[name]
        except KeyError:
            errors.append(f"missing style: {name}")
            continue
        if style_east_asia(style) != font:
            errors.append(f"{name}: East Asian font is '{style_east_asia(style)}', expected '{font}'")
        actual_size = style.font.size.pt if style.font.size else None
        if actual_size is None or abs(actual_size - size) > 0.05:
            errors.append(f"{name}: size is {actual_size}, expected {size}pt")

    try:
        body_style = doc.styles["尽调-正文"]
        pf = body_style.paragraph_format
        if pf.line_spacing is None or abs(float(pf.line_spacing) - 1.5) > 0.01:
            errors.append("body style line spacing must be 1.5")
        if pf.first_line_indent is None or abs(cm(pf.first_line_indent) - 0.847) > 0.04:
            errors.append("body style first-line indent must be 0.847cm")
        if pf.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
            errors.append("body style alignment must be justified")
    except KeyError:
        pass

    for name in ("尽调-正文", "尽调-表标题", "尽调-表正文", "尽调-提示"):
        try:
            pf = doc.styles[name].paragraph_format
        except KeyError:
            continue
        if not spacing_is_zero(pf.space_before) or not spacing_is_zero(pf.space_after):
            errors.append(f"{name}: paragraph spacing before/after must be 0pt")

    for name in ("尽调-一级标题", "尽调-二级标题", "尽调-三级标题", "尽调-四级标题"):
        try:
            pf = doc.styles[name].paragraph_format
        except KeyError:
            continue
        if pf.line_spacing is None or abs(float(pf.line_spacing) - 1.25) > 0.01:
            errors.append(f"{name}: line spacing must be 1.25")
        if pf.keep_with_next is not True:
            errors.append(f"{name}: keep-with-next is required")

    try:
        if doc.styles["尽调-一级标题"].paragraph_format.page_break_before is not True:
            errors.append("level-1 heading must start on a new page")
    except KeyError:
        pass

    try:
        pf = doc.styles["尽调-表标题"].paragraph_format
        if pf.keep_with_next is not True:
            errors.append("table-title style must keep with next")
    except KeyError:
        pass

    text = all_visible_text(doc)
    for marker in ("[[", "TODO", "TBD", "待补充", "此处填写"):
        if marker in text:
            errors.append(f"unresolved placeholder in DOCX: {marker}")

    for index, paragraph in enumerate(doc.paragraphs, 1):
        if not paragraph.text.strip():
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name not in ALLOWED_VISIBLE_PARAGRAPH_STYLES:
            errors.append(f"paragraph {index}: visible text uses inconsistent style '{style_name or 'none'}'")

    for index, table in enumerate(doc.tables, 1):
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        layout = tbl_pr.find(qn("w:tblLayout"))
        actual_width = int(tbl_w.get(qn("w:w"), "0")) if tbl_w is not None else 0
        if tbl_w is None or tbl_w.get(qn("w:type")) != "dxa" or actual_width not in {TEXT_WIDTH_DXA, LANDSCAPE_TEXT_WIDTH_DXA}:
            errors.append(
                f"table {index}: width must be {TEXT_WIDTH_DXA} dxa (portrait) "
                f"or {LANDSCAPE_TEXT_WIDTH_DXA} dxa (landscape)"
            )
        if layout is None or layout.get(qn("w:type")) != "fixed":
            errors.append(f"table {index}: fixed layout required")
        is_key_value = bool(
            len(table.columns) == 2
            and table.rows
            and table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")) is not None
            and table.rows[0].cells[0]._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"), "").upper() == "F2F2F2"
        )
        if table.rows and not is_key_value:
            tr_pr = table.rows[0]._tr.trPr
            repeat = tr_pr.find(qn("w:tblHeader")) if tr_pr is not None else None
            if repeat is None:
                errors.append(f"table {index}: header row must repeat")
            for cell in table.rows[0].cells:
                shd = cell._tc.tcPr.find(qn("w:shd"))
                if shd is None or shd.get(qn("w:fill"), "").upper() != "D9D9D9":
                    errors.append(f"table {index}: header fill must be #D9D9D9")
                    break
        if is_key_value:
            for row_index, row in enumerate(table.rows, 1):
                shd = row.cells[0]._tc.tcPr.find(qn("w:shd"))
                if shd is None or shd.get(qn("w:fill"), "").upper() != "F2F2F2":
                    errors.append(f"table {index} row {row_index}: key label fill must be #F2F2F2")
        for row_index, row in enumerate(table.rows, 1):
            tr_pr = row._tr.trPr
            cant_split = tr_pr.find(qn("w:cantSplit")) if tr_pr is not None else None
            if cant_split is None:
                errors.append(f"table {index} row {row_index}: row must not split across pages")
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() and paragraph.style.name != "尽调-表正文":
                        errors.append(
                            f"table {index} row {row_index}: cell paragraph uses "
                            f"'{paragraph.style.name}', expected '尽调-表正文'"
                        )

    if len(doc.sections) >= 3:
        pg = doc.sections[0]._sectPr.find(qn("w:pgNumType"))
        if pg is None or pg.get(qn("w:start")) != "1":
            errors.append("document page numbering must start at 1 on the cover")
        for section_index, section in enumerate(doc.sections[1:], 2):
            pg = section._sectPr.find(qn("w:pgNumType"))
            if pg is not None and pg.get(qn("w:start")):
                errors.append(f"section {section_index}: page numbering must continue, not restart")
        if HEADER_NOTICE not in "\n".join(p.text for p in doc.sections[0].header.paragraphs):
            errors.append("confidentiality header notice is missing")

    with zipfile.ZipFile(path) as archive:
        xml = "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith(".xml")
        )
    if " TOC " not in xml:
        errors.append("TOC field is missing")
    if " PAGE " not in xml:
        errors.append("PAGE field is missing")
    if "updateFields" not in xml:
        warnings.append("Word updateFields setting is missing")
    return errors, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx")
    args = parser.parse_args()
    path = Path(args.docx).resolve()
    try:
        errors, warnings = audit(path)
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        print(f"ERROR: {exc}")
        return 1
    for item in warnings:
        print(f"WARNING: {item}")
    for item in errors:
        print(f"ERROR: {item}")
    print(f"DOCX style audit: {len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
