#!/usr/bin/env python3
"""Forward-test the table-role layout gate with positive and negative DOCX cases."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from typing import Callable
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET


sys.dont_write_bytecode = True

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
Q = lambda name: f"{{{W}}}{name}"

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))


def _get_or_add(parent: ET.Element, name: str, *, prepend: bool = False) -> ET.Element:
    node = parent.find(f"w:{name}", NS)
    if node is None:
        node = ET.Element(Q(name))
        if prepend:
            parent.insert(0, node)
        else:
            parent.append(node)
    return node


def _tables(document: ET.Element) -> list[ET.Element]:
    return document.findall(".//w:tbl", NS)


def _set_fill(cell: ET.Element, fill: str) -> None:
    tc_pr = _get_or_add(cell, "tcPr", prepend=True)
    shade = _get_or_add(tc_pr, "shd")
    shade.set(Q("val"), "clear")
    shade.set(Q("color"), "auto")
    shade.set(Q("fill"), fill)


def _set_bold(cell: ET.Element) -> None:
    for run in cell.findall(".//w:r", NS):
        r_pr = _get_or_add(run, "rPr", prepend=True)
        bold = _get_or_add(r_pr, "b")
        bold.set(Q("val"), "1")
        bold_cs = _get_or_add(r_pr, "bCs")
        bold_cs.set(Q("val"), "1")


def _set_alignment(cell: ET.Element, alignment: str) -> None:
    for paragraph in cell.findall("w:p", NS):
        p_pr = _get_or_add(paragraph, "pPr", prepend=True)
        jc = _get_or_add(p_pr, "jc")
        jc.set(Q("val"), alignment)


def mutate_key_value_row_fill(document: ET.Element) -> None:
    row = _tables(document)[2].findall("w:tr", NS)[0]
    for cell in row.findall("w:tc", NS):
        _set_fill(cell, "D9D9D9")


def mutate_false_header(document: ET.Element) -> None:
    row = _tables(document)[2].findall("w:tr", NS)[0]
    tr_pr = _get_or_add(row, "trPr", prepend=True)
    _get_or_add(tr_pr, "tblHeader")


def mutate_narrative_alignment(document: ET.Element) -> None:
    cell = _tables(document)[2].findall("w:tr", NS)[0].findall("w:tc", NS)[1]
    _set_alignment(cell, "center")


def mutate_narrative_bold(document: ET.Element) -> None:
    cell = _tables(document)[2].findall("w:tr", NS)[0].findall("w:tc", NS)[1]
    _set_bold(cell)


def mutate_all_rows_cant_split(document: ET.Element) -> None:
    for row in _tables(document)[2].findall("w:tr", NS):
        tr_pr = _get_or_add(row, "trPr", prepend=True)
        _get_or_add(tr_pr, "cantSplit")


def mutate_paragraph_format_loss(document: ET.Element) -> None:
    cell = _tables(document)[2].findall("w:tr", NS)[0].findall("w:tc", NS)[1]
    for paragraph in cell.findall("w:p", NS):
        p_pr = paragraph.find("w:pPr", NS)
        if p_pr is not None:
            paragraph.remove(p_pr)
        for run in paragraph.findall("w:r", NS):
            r_pr = run.find("w:rPr", NS)
            if r_pr is not None:
                run.remove(r_pr)


def mutate_paragraph_spacing(document: ET.Element) -> None:
    cell = _tables(document)[2].findall("w:tr", NS)[0].findall("w:tc", NS)[1]
    for paragraph_node in cell.findall("w:p", NS):
        p_pr = _get_or_add(paragraph_node, "pPr", prepend=True)
        spacing = _get_or_add(p_pr, "spacing")
        spacing.set(Q("line"), "240")
        spacing.set(Q("lineRule"), "auto")
        spacing.set(Q("before"), "0")
        spacing.set(Q("after"), "0")


def mutate_financial_header_fill(document: ET.Element) -> None:
    row = _tables(document)[23].findall("w:tr", NS)[0]
    for cell in row.findall("w:tc", NS):
        _set_fill(cell, "D9D9D9")


def mutate_exact_height(document: ET.Element) -> None:
    row = _tables(document)[2].findall("w:tr", NS)[0]
    tr_pr = _get_or_add(row, "trPr", prepend=True)
    height = _get_or_add(tr_pr, "trHeight")
    height.set(Q("val"), "240")
    height.set(Q("hRule"), "exact")


MUTATIONS: list[tuple[str, str, Callable[[ET.Element], None]]] = [
    ("key_value_row_fill", "TABLE_FILL_COLOR_MISMATCH", mutate_key_value_row_fill),
    ("false_header", "TABLE_FALSE_HEADER_ROW", mutate_false_header),
    ("narrative_alignment", "TABLE_NARRATIVE_ALIGNMENT", mutate_narrative_alignment),
    ("narrative_bold", "TABLE_ROLE_STYLE_MISMATCH", mutate_narrative_bold),
    ("all_rows_cant_split", "TABLE_ROW_SPLIT_POLICY", mutate_all_rows_cant_split),
    ("paragraph_format_loss", "TABLE_PARAGRAPH_FORMAT_LOSS", mutate_paragraph_format_loss),
    ("paragraph_spacing", "TABLE_PARAGRAPH_SPACING_MISMATCH", mutate_paragraph_spacing),
    ("financial_header_fill", "TABLE_FILL_COLOR_MISMATCH", mutate_financial_header_fill),
    ("exact_height", "FIXED_ROW_HEIGHT", mutate_exact_height),
]


def make_mutation(source: Path, output: Path, mutation: Callable[[ET.Element], None]) -> None:
    with ZipFile(source) as src:
        document = ET.fromstring(src.read("word/document.xml"))
        mutation(document)
        replacement = ET.tostring(document, encoding="utf-8", xml_declaration=True)
        with ZipFile(output, "w", ZIP_DEFLATED) as dst:
            for item in src.infolist():
                raw = replacement if item.filename == "word/document.xml" else src.read(item.filename)
                dst.writestr(item, raw)


def run_validator(
    validator: Path,
    report: Path,
    target_company: str | None,
    forbidden: list[str],
) -> tuple[int, dict[str, object]]:
    command = [sys.executable, str(validator), str(report), "--strict-sample-schema"]
    if target_company:
        command.extend(["--target-company", target_company])
    for term in forbidden:
        command.extend(["--forbid-term", term])
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    try:
        result = json.loads(completed.stdout)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            f"validator returned non-JSON output (exit={completed.returncode}): {completed.stderr}"
        ) from exc
    return completed.returncode, result


def error_codes(result: dict[str, object]) -> set[str]:
    return {
        str(item.get("code"))
        for item in result.get("errors", [])
        if isinstance(item, dict) and item.get("code")
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("accepted_report", type=Path)
    parser.add_argument("--target-company")
    parser.add_argument("--forbid-term", action="append", default=[])
    args = parser.parse_args()

    validator = SKILL_ROOT / "scripts" / "validate_dd_report.py"
    accepted = args.accepted_report.resolve()

    positive_code, positive = run_validator(
        validator, accepted, args.target_company, args.forbid_term
    )
    if positive_code != 0 or positive.get("status") != "pass":
        raise SystemExit(f"accepted report failed strict validation: {positive!r}")

    results: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="dd-table-layout-regression-") as temporary:
        temp = Path(temporary)
        from docx import Document
        from table_writer import format_all_tables, write_table

        writer_roundtrip = temp / "writer-roundtrip.docx"
        document = Document(accepted)
        table = document.tables[2]
        payload: list[list[str | list[str]]] = []
        for row in table.rows:
            row_values: list[str | list[str]] = []
            for cell in row.cells:
                paragraphs = [paragraph.text for paragraph in cell.paragraphs if paragraph.text]
                row_values.append(paragraphs if len(paragraphs) > 1 else (paragraphs[0] if paragraphs else ""))
            payload.append(row_values)
        write_table(table, payload, table_number=3)
        format_all_tables(document.tables)
        document.save(writer_roundtrip)
        writer_code, writer_result = run_validator(
            validator, writer_roundtrip, args.target_company, args.forbid_term
        )
        if writer_code != 0 or writer_result.get("status") != "pass":
            raise SystemExit(f"table_writer roundtrip failed strict validation: {writer_result!r}")

        for name, expected_code, mutation in MUTATIONS:
            negative = temp / f"{name}.docx"
            make_mutation(writer_roundtrip, negative, mutation)
            return_code, result = run_validator(
                validator, negative, args.target_company, args.forbid_term
            )
            codes = error_codes(result)
            if return_code == 0 or result.get("status") != "fail":
                raise SystemExit(f"mutation {name} was not rejected")
            if expected_code not in codes:
                raise SystemExit(
                    f"mutation {name} missed {expected_code}: actual={sorted(codes)!r}"
                )
            results.append(
                {
                    "mutation": name,
                    "expected_code": expected_code,
                    "error_codes": sorted(codes),
                }
            )

    print(
        json.dumps(
            {
                "status": "pass",
                "accepted_report": str(accepted),
                "positive_status": positive.get("status"),
                "table_writer_roundtrip_status": writer_result.get("status"),
                "negative_cases": results,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
