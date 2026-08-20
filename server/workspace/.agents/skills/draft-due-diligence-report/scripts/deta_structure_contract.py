#!/usr/bin/env python3
"""Audit a DOCX against the fixed Deta V5 chapter and overview-field contract."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.document import Document as _Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


H1_STYLES = {"Heading 1", "投行 - 一级标题"}
H2_STYLES = {"Heading 2", "投行 - 二级标题"}
H3_STYLES = {"Heading 3", "投行 - 三级标题"}

EXPECTED_H1 = [
    "1、投资概要", "2、公司概况", "3、产品与技术", "4、业务情况",
    "5、行业和市场", "6、未来发展规划", "7、投资方案",
    "8、风险提示与对策", "投资结论及建议",
]
EXPECTED_H2 = [
    "1.1 公司情况", "1.2 交易要点", "1.3 行业概况",
    "1.4 商业模式和经营管理", "1.5 投资价值与风险",
    "2.1 公司基本信息", "2.2 历史沿革", "2.3 公司股东情况及实际控制人情况",
    "2.4 核心团队介绍", "2.5 组织架构", "2.6 关联公司及关联交易",
    "2.7 资质、荣誉及法律合规情况", "3.1 产品概念总览", "3.2 核心技术路线",
    "3.3 产品矩阵", "3.4 场景应用", "3.5 核心技术沿革", "3.6 知识产权及数据权属",
    "4.1 商业模式与销售策略", "4.2 客户验证情况", "4.3 供应商、采购与成本情况",
    "5.1 行业趋势与痛点", "5.2 市场分析", "6.1 业务拓展规划", "6.2 财务情况",
    "7.1 投资亮点", "7.2 公司估值与投资方式", "7.3 退出方案",
]
FIXED_H3_HEAD = [
    "2.3.1 公司股东",
    "2.3.2 公司股权穿透及实际控制",
]
FIXED_H3_TAIL = [
    "5.1.1 政策层面",
    "5.1.2 市场层面",
    "5.1.3 用户层面",
    "5.1.4 竞争格局",
    "5.1.5 产业趋势",
    "5.2.1核心下游市场",
    "5.2.2潜在拓展市场",
    "5.2.3综合评估",
    "6.2.1人力成本预测",
    "6.2.2 资产负债表分析",
    "6.2.3 财务预测",
]
TEAM_H3_MIN = 3
TEAM_H3_MAX = 5
CUSTOMER_H3_COUNT = 4
OVERVIEW_FIELDS = [
    ["公司名称", "成立时间", "注册地址", "实际经营地址", "法定代表人/实际控制人", "主营业务"],
    ["投资主体", "投资方式", "投资金额", "投资估值", "项目类型", "收益与责任安排"],
    ["商业路径", "客户与生态", "组织与人员", "经营现状", "管理基础"],
]
OVERVIEW_SECTIONS = ["1.1 公司情况", "1.2 交易要点", "1.4 商业模式和经营管理"]
BASIC_INFO_FIELDS = [
    "公司名称", "统一社会信用代码", "公司类型", "注册资本", "营业期限",
    "经营范围", "治理登记", "子公司/参股公司", "经营资质",
]
LEGAL_FIELDS = [
    "证照资质", "债权债务", "诉讼仲裁", "行政处罚与失信", "劳动社保", "环保消防安全",
]
FIXED_TABLE_HEADERS = {
    "2.2 历史沿革": [["时间", "事项", "注册资本/融资"]],
    "2.3 公司股东情况及实际控制人情况": [
        ["股东", "认缴注册资本 （万元）", "持股比例"],
        ["主体", "上层持有人或合伙人", "穿透要点"],
    ],
    "2.5 组织架构": [["部门/职能", "人数/规划", "主要职责"]],
    "2.6 关联公司及关联交易": [["事项", "核验结果"]],
    "3.1 产品概念总览": [["技术层", "核心模块", "输入/输出"]],
    "3.3 产品矩阵": [["产品/服务", "目标客户", "商业方式", "当前成熟度判断"]],
    "3.4 场景应用": [["场景", "解决问题"]],
    "3.5 核心技术沿革": [["年份", "代表成果", "与公司技术路线的关系"]],
    "3.6 知识产权及数据权属": [["专利名称", "申请号", "申请日", "发明人"], ["软件名称", "登记号", "登记日"]],
    "4.1 商业模式与销售策略": [["阶段", "主要产品", "收入来源"]],
    "4.3 供应商、采购与成本情况": [["类别", "对方及内容", "合同/账面金额", "成本影响"]],
    "5.2 市场分析": [["细分市场", "近期机会", "中期空间"]],
    "6.1 业务拓展规划": [["阶段", "研发目标", "商业目标", "组织目标"]],
    "7.3 退出方案": [["退出路径", "实现条件"]],
    "8、风险提示与对策": [["风险类别", "具体风险描述", "风险控制建议"]],
}


def _clean(value: str) -> str:
    return " ".join(str(value or "").split())


def _next_nonempty_block_is_table(paragraph: Paragraph) -> bool:
    node = paragraph._p.getnext()
    while node is not None:
        if node.tag == qn("w:tbl"):
            return True
        if node.tag == qn("w:p") and "".join(node.xpath(".//w:t/text()")).strip():
            return False
        node = node.getnext()
    return False


def _is_standalone_table_caption(paragraph: Paragraph) -> bool:
    if paragraph.style and paragraph.style.name in H1_STYLES | H2_STYLES | H3_STYLES:
        return False
    centered = paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    bold = any(run.bold for run in paragraph.runs if run.text.strip())
    return centered and bold and _next_nonempty_block_is_table(paragraph)


def _heading_texts(doc: Document, styles: set[str]) -> list[str]:
    return [_clean(p.text) for p in doc.paragraphs if _clean(p.text) and p.style and p.style.name in styles]


def _tables_by_section(doc: Document) -> dict[str, list[Table]]:
    if not isinstance(doc, _Document):
        raise TypeError("expected python-docx Document")
    out: dict[str, list[Table]] = {}
    current_h1 = ""
    current_h2 = ""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, doc)
            text = _clean(paragraph.text)
            style = paragraph.style.name if paragraph.style else ""
            if text and style in H1_STYLES:
                current_h1, current_h2 = text, ""
            elif text and style in H2_STYLES:
                current_h2 = text
        elif child.tag == qn("w:tbl"):
            key = current_h2 or current_h1
            out.setdefault(key, []).append(Table(child, doc))
    return out


def _section_body_paragraphs(doc: Document) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    current_h1 = ""
    current_h2 = ""
    current_h3 = ""
    for child in doc.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, doc)
        text = _clean(paragraph.text)
        style = paragraph.style.name if paragraph.style else ""
        if text and style in H1_STYLES:
            current_h1, current_h2, current_h3 = text, "", ""
        elif text and style in H2_STYLES:
            current_h2, current_h3 = text, ""
        elif text and style in H3_STYLES:
            current_h3 = text
        elif text:
            key = current_h3 or current_h2 or current_h1
            out.setdefault(key, []).append(text)
    return out


def audit_structure(doc: Document) -> dict:
    issues: list[str] = []
    h1 = _heading_texts(doc, H1_STYLES)
    h2 = _heading_texts(doc, H2_STYLES)
    h3 = _heading_texts(doc, H3_STYLES)

    if h1 != EXPECTED_H1:
        issues.append(f"一级标题不一致：期望9个固定标题，实际{len(h1)}个")
    if h2 != EXPECTED_H2:
        issues.append(f"二级标题不一致：期望28个固定标题，实际{len(h2)}个")
    minimum_h3 = len(FIXED_H3_HEAD) + TEAM_H3_MIN + CUSTOMER_H3_COUNT + len(FIXED_H3_TAIL)
    maximum_h3 = len(FIXED_H3_HEAD) + TEAM_H3_MAX + CUSTOMER_H3_COUNT + len(FIXED_H3_TAIL)
    if not minimum_h3 <= len(h3) <= maximum_h3:
        issues.append(
            f"三级标题数量不一致：期望{minimum_h3}-{maximum_h3}个（核心团队3-5人），实际{len(h3)}个"
        )
    else:
        team_count = len(h3) - len(FIXED_H3_HEAD) - CUSTOMER_H3_COUNT - len(FIXED_H3_TAIL)
        actual_head = h3[:len(FIXED_H3_HEAD)]
        team_start = len(FIXED_H3_HEAD)
        team_end = team_start + team_count
        actual_team = h3[team_start:team_end]
        customer_end = team_end + CUSTOMER_H3_COUNT
        actual_customers = h3[team_end:customer_end]
        actual_tail = h3[customer_end:]

        if actual_head != FIXED_H3_HEAD:
            issues.append("2.3三级标题不一致或顺序错误")
        for index, actual in enumerate(actual_team, 1):
            prefix = f"2.4.{index}"
            label = actual[len(prefix):].lstrip(" .、—-") if actual.startswith(prefix) else ""
            if not actual.startswith(prefix) or len(label) < 2:
                issues.append(f"核心团队第{index}项须以“{prefix}”开头并包含真实姓名/角色：实际“{actual}”")
        for index, actual in enumerate(actual_customers, 1):
            prefix = f"4.2.{index}"
            label = actual[len(prefix):].lstrip(" .、—-") if actual.startswith(prefix) else ""
            if not actual.startswith(prefix) or len(label) < 2:
                issues.append(f"客户验证第{index}项须以“{prefix}”开头并包含真实名称：实际“{actual}”")
        if actual_tail != FIXED_H3_TAIL:
            issues.append("5.1、5.2或6.2固定三级标题不一致或顺序错误")

    section_tables = _tables_by_section(doc)
    section_paragraphs = _section_body_paragraphs(doc)
    for section, expected in zip(OVERVIEW_SECTIONS, OVERVIEW_FIELDS):
        tables = section_tables.get(section, [])
        if len(tables) != 1:
            issues.append(f"{section}应且只应保留一张概要键值表，实际{len(tables)}张")
        else:
            table = tables[0]
            if len(table.columns) != 2:
                issues.append(f"{section}概要表列数不一致：期望2列，实际{len(table.columns)}列")
                continue
            actual = [_clean(row.cells[0].text) for row in table.rows]
            if actual != expected:
                issues.append(
                    f"{section}概要表字段不一致：期望“{'｜'.join(expected)}”，实际“{'｜'.join(actual)}”"
                )

    if section_tables.get("1.3 行业概况"):
        issues.append("1.3行业概况应使用正文，不应使用表格包装普通叙事")
    if section_tables.get("1.5 投资价值与风险"):
        issues.append("1.5投资价值与风险应使用一项一段的正文")
    if section_tables.get("7.1 投资亮点"):
        issues.append("7.1投资亮点应使用正文，不应暴露尽调证据列")

    value_paragraphs = section_paragraphs.get("1.5 投资价值与风险", [])
    value_items = [p for p in value_paragraphs if re.match(r"^[1-5]\.\s*", p)]
    risk_items = [p for p in value_paragraphs if p.startswith("主要风险")]
    if len(value_items) != 5 or len(risk_items) != 1:
        issues.append("1.5投资价值与风险须包含5条完整价值链和1段主要风险传导")
    if len("".join(value_paragraphs)) < 650:
        issues.append("1.5投资价值与风险内容深度不足：正文总量不得低于650个字符")

    depth_contract = {
        "5.1.1 政策层面": (3, 500),
        "5.1.2 市场层面": (3, 600),
    }
    for section, (minimum_paragraphs, minimum_chars) in depth_contract.items():
        paragraphs = section_paragraphs.get(section, [])
        if len(paragraphs) < minimum_paragraphs or len("".join(paragraphs)) < minimum_chars:
            issues.append(
                f"{section}内容深度不足：至少{minimum_paragraphs}个实质段落、{minimum_chars}个字符"
            )

    company_tables = section_tables.get("1.1 公司情况", [])
    if company_tables:
        company_values = {
            _clean(row.cells[0].text): _clean(row.cells[1].text)
            for row in company_tables[0].rows if len(row.cells) >= 2
        }
        registered_address = company_values.get("注册地址", "")
        placeholder = re.search(r"未提供|待核实|资料缺失|不详", registered_address)
        detail_marker = re.search(r"街道|镇|乡|路|街|大道|村|号|室|楼|层|幢|栋", registered_address)
        if registered_address and not placeholder and (
            len(registered_address) < 10 or not detail_marker
        ):
            issues.append(
                "注册地址疑似不完整：应按营业执照/工商登记写至街道（镇）、道路、门牌及房号，不得仅写省市区"
            )

    basic_cols = [[_clean(row.cells[0].text) for row in table.rows]
                  for table in section_tables.get("2.1 公司基本信息", [])]
    legal_cols = [[_clean(row.cells[0].text) for row in table.rows]
                  for table in section_tables.get("2.7 资质、荣誉及法律合规情况", [])]
    if BASIC_INFO_FIELDS not in basic_cols:
        issues.append("2.1公司基本信息字段不一致或缺失")
    if LEGAL_FIELDS not in legal_cols:
        issues.append("2.7法律合规字段不一致或缺失")
    for section, expected_headers in FIXED_TABLE_HEADERS.items():
        actual_headers = [[_clean(cell.text) for cell in table.rows[0].cells]
                          for table in section_tables.get(section, [])]
        for expected in expected_headers:
            if expected not in actual_headers:
                issues.append(f"{section}固定表头不一致或缺失：“" + "｜".join(expected) + "”")

    related_tables = section_tables.get("2.6 关联公司及关联交易", [])
    if related_tables:
        related_labels = [_clean(row.cells[0].text) for row in related_tables[0].rows]
        expected_labels = ["事项", "关联企业", "关联往来", "个人往来"]
        if related_labels != expected_labels:
            issues.append(
                "2.6关联公司及关联交易首列须与德塔V5一致：“"
                + "｜".join(expected_labels) + "”"
            )
        grid_widths = [int(col.get(qn("w:w")) or 0) for col in related_tables[0]._tbl.tblGrid.gridCol_lst]
        if len(grid_widths) == 2 and sum(grid_widths) > 0:
            first_ratio = grid_widths[0] / sum(grid_widths)
            if not 0.38 <= first_ratio <= 0.40:
                issues.append(
                    f"2.6关联交易表列宽未对齐德塔V5约39%/61%比例：实际{first_ratio:.1%}/{1-first_ratio:.1%}"
                )

    table_only_sections = (
        "2.1 公司基本信息",
        "2.2 历史沿革",
        "2.6 关联公司及关联交易",
        "2.7 资质、荣誉及法律合规情况",
    )
    for section in table_only_sections:
        if section_paragraphs.get(section):
            issues.append(f"{section}必须为表格自足页，标题与表格之间或表后不得出现正文段落")
        table_count = len(section_tables.get(section, []))
        if table_count != 1:
            issues.append(f"{section}必须且只能保留一张主表，实际{table_count}张")

    shareholder_paragraphs = section_paragraphs.get("2.3.1 公司股东", [])
    if len(shareholder_paragraphs) != 1:
        issues.append("2.3.1公司股东须仅有一个表前股权总括段")
    elif not all(marker in shareholder_paragraphs[0] for marker in ("截至", "注册资本", "%")):
        issues.append("2.3.1股权总括段须写明截至日、注册资本及持股比例")

    if len(section_paragraphs.get("2.3.2 公司股权穿透及实际控制", [])) != 1:
        issues.append("2.3.2须仅保留一个穿透结论段，表后不得追加风险或交割评论")
    if section_paragraphs.get("2.4 核心团队介绍"):
        issues.append("2.4核心团队介绍标题下不得写通用评价，应直接进入3-5名核心人员小节")

    org_tables = section_tables.get("2.5 组织架构", [])
    if org_tables:
        for row in org_tables[0].rows[1:]:
            value = _clean(row.cells[1].text)
            if not re.search(r"\d+\s*(?:人|名)|规划|拟增|待招|现有", value):
                issues.append(
                    f"2.5组织架构人数/规划列不是人数或规划口径：{_clean(row.cells[0].text)}＝{value}"
                )
    if section_paragraphs.get("2.5 组织架构"):
        issues.append("2.5组织架构表前不得出现分析性正文；表后仅允许客观人员结构统计")

    full_text = "\n".join(p.text for p in doc.paragraphs)
    full_text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    if re.search(r"实际经营安排", full_text):
        issues.append("发现禁用字段“实际经营安排”，固定字段必须为“实际经营地址”")
    if "主要资料依据与使用边界" in full_text:
        issues.append("2.7发现禁用模块“主要资料依据与使用边界”，正式报告不得保留资料清单或使用边界表")
    captions = [_clean(p.text) for p in doc.paragraphs if _clean(p.text) and _is_standalone_table_caption(p)]
    if captions:
        issues.append("正式报告不得保留独立表题：“" + "｜".join(captions[:20]) + "”")
    if "公司现有4项未授权发明专利和3项已登记软件著作权" in full_text:
        issues.append("3.6不得保留专利数量、非员工发明人及境外技术归属的表后评价段")

    forbidden_workflow_labels = (
        "核查结论", "尽调证据", "关键验证", "当前证据", "投资处理",
        "控制方案详见", "报告期内", "截至报告日", "公司披露", "管理层披露",
        "尽调材料人数/规划",
    )
    for label in forbidden_workflow_labels:
        if label in full_text:
            issues.append(f"正式报告发现内部工作流或模糊口径标签：{label}")
    if re.search(r"(?:^|[\s｜|：:])P[12](?:[\s｜|：:]|$)", full_text, re.MULTILINE):
        issues.append("正式报告风险表不得暴露P1/P2内部优先级编码")
    if re.search(r"详见\s*(?:第?\d+(?:\.\d+){0,2}|下文|附录)", full_text):
        issues.append("发现“详见”跳转写法；正式报告应在当前位置写明必要结论")
    if re.search(r"(?:尽调)?材料(?:显示|表明|未提供|未显示)", full_text):
        issues.append("发现高频材料归因/缺口话术；正文应直接陈述事实，缺口集中到风险与交割事项")
    if re.search(r"客户访谈(?:显示|表明|反馈|确认)", full_text):
        issues.append("客户访谈归因必须由真实客户访谈底稿支持；无客户访谈时不得使用该表述")

    return {
        "status": "pass" if not issues else "fail",
        "counts": {"h1": len(h1), "h2": len(h2), "h3": len(h3), "tables": len(doc.tables)},
        "issues": issues,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    audit = sub.add_parser("audit")
    audit.add_argument("--input", required=True)
    audit.add_argument("--json-output")
    args = parser.parse_args()

    result = audit_structure(Document(str(Path(args.input).expanduser().resolve())))
    if args.json_output:
        Path(args.json_output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["status"] == "pass" else 2


if __name__ == "__main__":
    raise SystemExit(main())
