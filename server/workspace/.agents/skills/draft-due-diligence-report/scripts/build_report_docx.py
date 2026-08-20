#!/usr/bin/env python3
"""Build a professional Chinese due-diligence DOCX from report JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_BREAK
from docx.shared import Cm

from audit_ic_completeness import audit as audit_ic_completeness

from docx_builder_lib import (
    LANDSCAPE_TEXT_WIDTH_DXA,
    TEXT_WIDTH_CM,
    TEXT_WIDTH_DXA,
    add_image,
    add_key_value_table,
    add_table,
    add_text_paragraph,
    configure_styles,
    find_paragraph,
    insert_stacked_title,
    numbered_heading,
    replace_marker,
    set_core_properties,
    set_run_font,
    set_update_fields,
)


FIXED_REPORT_TITLE = "尽职调查报告"


def current_table_width_dxa(doc: Document) -> int:
    return LANDSCAPE_TEXT_WIDTH_DXA if doc.sections[-1].orientation == WD_ORIENT.LANDSCAPE else TEXT_WIDTH_DXA


def add_section_break(doc: Document, marker, orientation: str) -> None:
    """Insert a new-page section boundary before marker and configure following pages."""
    if orientation not in {"portrait", "landscape"}:
        raise ValueError(f"unsupported section orientation: {orientation}")

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    boundary = doc.paragraphs[-1]
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.85)

    if orientation == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(2.0)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)

    marker._p.addprevious(boundary._p)


def load_json(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as stream:
        return json.load(stream)


def add_block(doc: Document, marker, block: dict, base_dir: Path, counters: list[int]) -> None:
    kind = block.get("type")
    if kind == "heading":
        level = int(block.get("level", 1))
        title = str(block.get("title", "")).strip()
        if block.get("numbered", False):
            title = numbered_heading(title, level, counters)
        p = doc.add_paragraph(style=f"尽调-{['一级标题','二级标题','三级标题','四级标题'][level-1]}")
        r = p.add_run(title)
        font = "黑体"
        size = [15, 15, 12, 12][level - 1]
        set_run_font(r, font, size, True)
        marker._p.addprevious(p._p)
        return

    if kind == "paragraph":
        add_text_paragraph(doc, marker, str(block.get("text", "")))
        return

    if kind == "bullet":
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = None
        r = p.add_run(str(block.get("text", "")))
        set_run_font(r, "仿宋_GB2312", 12)
        marker._p.addprevious(p._p)
        return

    if kind == "numbered_item":
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = None
        r = p.add_run(str(block.get("text", "")))
        set_run_font(r, "仿宋_GB2312", 12)
        marker._p.addprevious(p._p)
        return

    if kind == "callout":
        label = str(block.get("label", "重点结论")).strip()
        text = str(block.get("text", ""))
        add_text_paragraph(doc, marker, text, style="尽调-提示", bold_prefix=f"{label}：")
        return

    if kind == "table":
        title = str(block.get("title", "")).strip()
        if title:
            p = doc.add_paragraph(title, style="尽调-表标题")
            marker._p.addprevious(p._p)
        add_table(
            doc,
            marker,
            [str(x) for x in block.get("headers", [])],
            [[str(v) for v in row] for row in block.get("rows", [])],
            block.get("column_widths"),
            block.get("alignments"),
            current_table_width_dxa(doc),
        )
        return

    if kind == "key_value_table":
        title = str(block.get("title", "")).strip()
        if title:
            p = doc.add_paragraph(title, style="尽调-表标题")
            marker._p.addprevious(p._p)
        add_key_value_table(
            doc,
            marker,
            [[str(v) for v in row] for row in block.get("rows", [])],
            block.get("column_widths"),
            current_table_width_dxa(doc),
        )
        return

    if kind == "image":
        raw = Path(str(block.get("path", "")))
        path = raw if raw.is_absolute() else base_dir / raw
        if not path.exists():
            raise FileNotFoundError(path)
        add_image(
            doc,
            marker,
            path,
            float(block.get("width_cm", TEXT_WIDTH_CM)),
            str(block.get("caption", "")).strip() or None,
        )
        return

    if kind == "page_break":
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        marker._p.addprevious(p._p)
        return

    if kind == "section_break":
        add_section_break(doc, marker, str(block.get("orientation", "")))
        return

    raise ValueError(f"unsupported block type: {kind}")


def build(
    input_path: Path,
    output_path: Path,
    template_path: Path,
    diligence_data_path: Path,
    evidence_path: Path,
) -> None:
    data = load_json(input_path)
    diligence_data = load_json(diligence_data_path)
    evidence = load_json(evidence_path)
    completeness_errors, completeness_warnings = audit_ic_completeness(
        diligence_data,
        report=data,
        evidence=evidence,
    )
    if completeness_errors or completeness_warnings:
        details = "\n".join(
            [f"ERROR: {item}" for item in completeness_errors]
            + [f"WARNING: {item}" for item in completeness_warnings]
        )
        raise ValueError(
            "IC completeness gate failed; DOCX generation is blocked.\n" + details
        )
    meta = data.get("meta", {})
    report_title = str(meta.get("report_title", "")).strip()
    if report_title != FIXED_REPORT_TITLE:
        raise ValueError(
            f"meta.report_title must be exactly '{FIXED_REPORT_TITLE}', got '{report_title}'"
        )
    doc = Document(template_path)
    configure_styles(doc)

    title_marker = None
    for p in list(doc.paragraphs):
        token = p.text.strip()
        if token == "[[PROJECT_NAME]]":
            replace_marker(p, str(meta.get("project_name") or meta.get("legal_entity") or "项目名称"))
        elif token == "[[REPORT_DATE]]":
            replace_marker(p, str(meta.get("report_date", "")))
        elif token == "[[AUTHOR]]":
            replace_marker(p, str(meta.get("author", "")))
        elif token == "[[REPORT_TITLE]]":
            title_marker = p
    if title_marker is None:
        raise ValueError("report title marker not found")
    insert_stacked_title(doc, title_marker, FIXED_REPORT_TITLE)

    marker = find_paragraph(doc, "[[REPORT_BODY]]")
    counters = [0, 0, 0, 0]
    for block in data.get("blocks", []):
        add_block(doc, marker, block, input_path.parent, counters)
    marker._p.getparent().remove(marker._p)

    set_update_fields(doc)
    set_core_properties(doc)
    doc.core_properties.title = str(meta.get("project_name", "")) + FIXED_REPORT_TITLE
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--diligence-data", required=True)
    parser.add_argument("--evidence", required=True)
    parser.add_argument("--template")
    args = parser.parse_args()
    script_dir = Path(__file__).resolve().parent
    template = Path(args.template) if args.template else script_dir.parent / "assets" / "deta-v5-dd-template.docx"
    try:
        build(
            Path(args.input).resolve(),
            Path(args.output).resolve(),
            template.resolve(),
            Path(args.diligence_data).resolve(),
            Path(args.evidence).resolve(),
        )
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(Path(args.output).resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
