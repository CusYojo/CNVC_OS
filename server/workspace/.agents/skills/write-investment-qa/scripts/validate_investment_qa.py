#!/usr/bin/env python3
"""Validate investment Q&A content, natural-writing signals, and optional DOCX structure."""

from __future__ import annotations

import argparse
import re
import statistics
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


REQUIRED_STYLES = {"QA Title", "QA Subtitle", "QA Question", "QA Body", "QA List", "QA Table"}
FIXED_BODY_FONT = "FangSong"
FIXED_FOOTER_FONT = "Helvetica Neue"
UNRESOLVED = (
    "[待补充]",
    "[来源待核验]",
    "[仅有公司单方口径]",
    "[公开信息未检索到]",
    "[多来源口径冲突]",
    "[数据时点较旧]",
    "[数据质量待核验]",
    "[公司预测]",
    "[项目团队判断]",
)
ABSOLUTES = (
    "绝对领先",
    "完全领先",
    "全球领先",
    "国内唯一",
    "行业唯一",
    "不存在风险",
    "不存在实质性风险",
    "下行风险可控",
    "不会被替代",
    "必然成功",
)
FORBIDDEN_HEADINGS = ("执行摘要", "目录", "信息缺口", "来源", "参考资料", "结论")
RISK_WORDS = ("风险", "反方", "替代", "不成立", "击穿", "最坏", "失败")
DECISION_WORDS = (
    "投资结论",
    "进入下一阶段",
    "不进入下一阶段",
    "暂不进入",
    "前置条件",
    "成立条件",
    "不纳入",
    "不具备",
    "放弃",
)
MANAGER_VOICE_FORBIDDEN = (
    "建议",
    "待验证",
    "待核验",
    "待补充",
    "有待观察",
    "有待验证",
    "进一步核验",
    "进一步验证",
    "后续需要",
    "下一步需要",
    "仍需进一步",
    "可以考虑",
    "倾向认为",
    "若情况属实",
)
EVIDENCE_AUDIT_NARRATION = (
    "本报告",
    "未披露",
    "公开证据",
    "证据边界",
    "证据不足",
    "未获证实",
    "未获公开数据证实",
    "不能据此",
    "无法判断",
    "无法获取",
    "信息无法获取",
    "不作定量结论",
    "不做定量结论",
    "不形成判断",
    "不进行判断",
    "难以判断",
    "受限于公开信息",
    "公开信息有限",
    "信息有限",
    "资料有限",
)
FORMULA_LIMITS = {
    "现有材料显示": 2,
    "对投资判断而言": 2,
    "从投资角度看": 2,
}
CAVEAT_WORDS = (
    "尚不能",
    "有待核验",
    "需要核验",
    "信息有限",
    "公开信息有限",
    "资料有限",
)
ECONOMIC_WORDS = (
    "获客",
    "转化",
    "销售周期",
    "客单价",
    "合同额",
    "续费",
    "复购",
    "增购",
    "收入",
    "毛利",
    "利润",
    "成本",
    "实施人天",
    "交付人天",
    "交付效率",
    "人效",
    "回款",
    "应收",
    "现金流",
    "营运资金",
    "资本开支",
    "资本强度",
    "规模效应",
    "经营杠杆",
    "定价权",
    "切换成本",
    "复用",
    "壁垒",
    "市场匹配",
    "PMF",
    "估值",
    "倍数",
    "溢价",
    "折价",
    "基础价值",
    "期权价值",
)
CAUSAL_WORDS = (
    "因此",
    "意味着",
    "取决于",
    "从而",
    "导致",
    "带动",
    "压低",
    "提高",
    "降低",
    "改变",
    "决定",
    "转化为",
    "对应",
    "只有",
    "否则",
    "反过来",
)
INVESTMENT_TENSION_WORDS = (
    "获客",
    "收入",
    "客单价",
    "续费",
    "复购",
    "壁垒",
    "产品化",
    "复制",
    "单位经济",
    "成本",
    "毛利",
    "现金流",
    "回款",
    "估值",
    "资源",
    "分散",
    "成立",
    "不成立",
    "规模",
    "价值",
)
GENERIC_QUESTION_PATTERNS = (
    r"公司(?:的)?基本情况",
    r"公司(?:的)?核心定位",
    r"公司(?:的)?核心技术竞争力",
    r"公司(?:的)?核心优势",
    r"公司(?:的)?商业模式",
    r"公司(?:的)?未来增长",
    r"公司(?:的)?增长路径",
    r"公司(?:的)?研发投入",
    r"公司(?:的)?成本结构",
    r"公司(?:的)?团队情况",
    r"公司(?:的)?财务情况",
    r"公司(?:的)?风险",
)
PROFILE_RANGES = {
    "communication": (6, 9),
    "investment": (8, 12),
    "diligence": (10, 20),
    "ic": (8, 12),
}


def _strip_markdown(value: str) -> str:
    value = re.sub(r"<!--.*?-->", "", value, flags=re.S)
    value = re.sub(r"[`*_>#|\[\]()]", "", value)
    value = re.sub(r"\s+", "", value)
    return value


def _question_sections(text: str) -> list[tuple[str, str]]:
    matches = list(re.finditer(r"^##\s+(?:Q\s*)?(\d+)\s*[：:、.．]\s*(.+)$", text, flags=re.M | re.I))
    sections: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append((match.group(2).strip(), text[match.end():end].strip()))
    return sections


def validate_markdown(
    path: Path,
    allow_placeholders: bool = False,
    profile: str = "communication",
    native_terms: tuple[str, ...] = (),
    allowed_reader_terms: tuple[str, ...] = (),
) -> tuple[list[str], list[str], int]:
    text = path.read_text(encoding="utf-8")
    errors: list[str] = []
    warnings: list[str] = []

    h1 = re.findall(r"^#\s+(.+)$", text, flags=re.M)
    if len(h1) != 1:
        errors.append(f"Expected exactly one H1 project/company title; found {len(h1)}.")

    q_matches = re.findall(r"^##\s+(?:Q\s*)?(\d+)\s*[：:、.．]\s*(.+)$", text, flags=re.M | re.I)
    if not q_matches:
        errors.append("No valid '## Q1：问题' headings found.")
        return errors, warnings, 0
    numbers = [int(n) for n, _ in q_matches]
    sections = _question_sections(text)
    expected = list(range(1, len(numbers) + 1))
    if numbers != expected:
        errors.append(f"Question numbering must be continuous from 1; found {numbers}.")
    minimum, maximum = PROFILE_RANGES[profile]
    if not minimum <= len(numbers) <= maximum:
        warnings.append(
            f"Profile '{profile}' normally uses {minimum}-{maximum} questions; found {len(numbers)}."
        )

    for heading in FORBIDDEN_HEADINGS:
        if re.search(rf"^#+\s*.*{re.escape(heading)}", text, flags=re.M):
            errors.append(f"Default direct Q&A must not contain standalone '{heading}' heading.")

    if re.search(r"^\s*(?:\*\*)?结论[：:](?:\*\*)?", text, flags=re.M):
        errors.append("Standalone '结论：' label is not allowed; integrate the judgment into the answer.")
    if re.search(r"https?://|\[[^\]]+\]\([^)]+\)", text):
        errors.append("Standard reader-facing edition must not contain raw URLs or Markdown links.")
    if re.search(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)", text, flags=re.M):
        errors.append(
            "Fixed lancheng_qa_a4_fixed output does not allow Markdown lists; "
            "rewrite parallel points as ordinary answer paragraphs."
        )
    if re.search(r"^\s*\|.*\|\s*$", text, flags=re.M):
        errors.append(
            "Fixed lancheng_qa_a4_fixed output does not allow Markdown tables; "
            "rewrite the content as ordinary answer paragraphs."
        )
    if re.search(r"\{\{|\}\}|<TODO>|TODO|TBD|XXX", text, flags=re.I):
        errors.append("Placeholder token found.")
    if not allow_placeholders:
        for marker in UNRESOLVED:
            if marker in text:
                errors.append(f"Internal unresolved marker must not appear in the reader-facing report: {marker}")

    for term in MANAGER_VOICE_FORBIDDEN:
        if term in text and term not in allowed_reader_terms:
            errors.append(
                f"Reader-facing report contains unfinished-work or advice language '{term}'; "
                "replace it with a completed analytical treatment."
            )

    for term in EVIDENCE_AUDIT_NARRATION:
        if term in text and term not in allowed_reader_terms:
            errors.append(
                f"Reader-facing report contains evidence-audit narration '{term}'; "
                "state the observable business condition and its economic implication instead."
            )

    for term in ABSOLUTES:
        if term in text:
            warnings.append(f"Promotional or absolute wording requires explicit evidence and boundary: {term}")

    if profile in {"investment", "ic"} and not any(word in text for word in RISK_WORDS):
        warnings.append("No thesis-breaking risk or counterargument language detected.")
    last_question = q_matches[-1][1]
    last_section = sections[-1][1]
    if profile == "ic" and not any(word in last_question + last_section for word in DECISION_WORDS):
        warnings.append("Final question does not appear to close the investment decision or verification conditions.")

    if re.search(r"预计(?:数量|金额|订单).{0,30}(?:正式订单|已签约|订单总额)", text):
        warnings.append("Possible mixing of forecast/intention values with signed-order status; review the status chain.")
    if "客户包括" in text and not any(x in text for x in ("合同", "订单", "交付", "验收", "回款", "Demo", "访谈")):
        warnings.append("Customer names appear without a clear commercial stage or evidence boundary.")

    answer_lengths = [len(_strip_markdown(body)) for _, body in sections]
    if profile == "communication":
        for index, length in enumerate(answer_lengths, start=1):
            if length < 120:
                warnings.append(f"Q{index} answer is only {length} characters; it may be under-explained.")
            elif length > 900:
                warnings.append(f"Q{index} answer is {length} characters; consider tightening the communication profile.")

    if len(answer_lengths) >= 5 and statistics.mean(answer_lengths) > 0:
        variation = statistics.pstdev(answer_lengths) / statistics.mean(answer_lengths)
        if variation < 0.12:
            warnings.append(
                f"Answer lengths are unusually uniform (coefficient of variation {variation:.2f}); vary depth by topic."
            )

    numbered_sections = sum(
        1 for _, body in sections if len(re.findall(r"(?:\*\*)?[（(]\d+[）)]", body)) >= 2
    )
    if len(sections) >= 5 and numbered_sections / len(sections) >= 0.8:
        warnings.append(
            "At least 80% of answers use the same numbered-subpoint structure; vary answer shapes."
        )

    for phrase, limit in FORMULA_LIMITS.items():
        count = text.count(phrase)
        if count > limit:
            warnings.append(f"Formulaic phrase '{phrase}' appears {count} times; recommended maximum is {limit}.")

    for index, (_, body) in enumerate(sections, start=1):
        caveat_count = sum(body.count(word) for word in CAVEAT_WORDS)
        if profile == "communication" and caveat_count > 3:
            warnings.append(
                f"Q{index} contains {caveat_count} audit/caveat phrases; explain the substance before the limitation."
            )

    if profile in {"investment", "ic"}:
        for index, (_, body) in enumerate(sections, start=1):
            compact_body = _strip_markdown(body)
            if not any(word in compact_body for word in ECONOMIC_WORDS):
                warnings.append(
                    f"Q{index} contains no detected business or financial variable; "
                    "connect the company-specific fact to economics or valuation."
                )
            if not any(word in compact_body for word in CAUSAL_WORDS):
                warnings.append(
                    f"Q{index} contains no detected causal bridge; explain why the fact changes the investment view."
                )

    opening_groups: dict[str, list[int]] = {}
    for index, (_, body) in enumerate(sections, start=1):
        opening = _strip_markdown(body)[:14]
        if opening:
            opening_groups.setdefault(opening, []).append(index)
    for opening, indexes in opening_groups.items():
        if len(indexes) >= 2:
            warnings.append(
                f"Answers {indexes} share the same opening '{opening}…'; rewrite for a less formulaic rhythm."
            )

    headings = [heading for heading, _ in sections]
    if profile in {"communication", "investment"}:
        if native_terms:
            specific = sum(1 for heading in headings if any(term in heading for term in native_terms))
            ratio = specific / len(headings) if headings else 0
            if ratio < 0.6:
                warnings.append(
                    f"Only {specific}/{len(headings)} questions contain supplied company-native terms; target at least 60%."
                )
            if profile in {"investment", "ic"}:
                debate_questions = sum(
                    1
                    for heading in headings
                    if any(term in heading for term in native_terms)
                    and any(word in heading for word in INVESTMENT_TENSION_WORDS)
                )
                if debate_questions / len(headings) < 0.5:
                    warnings.append(
                        f"Only {debate_questions}/{len(headings)} questions combine a company-native term "
                        "with an investment tension; target at least 50%."
                    )
        else:
            generic = sum(
                1 for heading in headings
                if any(re.search(pattern, heading) for pattern in GENERIC_QUESTION_PATTERNS)
            )
            if headings and generic / len(headings) >= 0.5:
                warnings.append(
                    "At least half of the questions match generic investment headings; add named products, cases, stages, or tensions."
                )

    return errors, warnings, len(numbers)


def validate_evidence_ledger(path: Path, question_count: int) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    if not path.exists():
        return [f"Evidence ledger not found: {path}"], warnings

    text = path.read_text(encoding="utf-8")
    required = ("事实键", "主张", "来源及定位", "证据等级", "状态", "对应问题")
    table_lines = [line.strip() for line in text.splitlines() if line.strip().startswith("|")]
    header_index = next(
        (index for index, line in enumerate(table_lines) if all(field in line for field in required)),
        None,
    )
    if header_index is None:
        errors.append(f"Evidence ledger must contain columns: {', '.join(required)}.")
        return errors, warnings

    header = [cell.strip() for cell in table_lines[header_index].strip("|").split("|")]
    rows: list[list[str]] = []
    for line in table_lines[header_index + 1:]:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells):
            continue
        if len(cells) == len(header):
            rows.append(cells)
    if not rows:
        errors.append("Evidence ledger contains no claim rows.")
        return errors, warnings

    indices = {name: header.index(name) for name in required}
    for row_number, row in enumerate(rows, start=1):
        claim = row[indices["主张"]]
        source = row[indices["来源及定位"]]
        grade = row[indices["证据等级"]].upper()
        mapping = row[indices["对应问题"]]
        if not claim or claim == "[待补充]":
            errors.append(f"Evidence ledger row {row_number} has no completed claim.")
        if not source or source == "[待补充]":
            errors.append(f"Evidence ledger row {row_number} has no completed source locator or disclosure basis.")
        if not re.fullmatch(r"[ABCD](?:\s*[/、,]\s*[ABCD])*", grade):
            errors.append(f"Evidence ledger row {row_number} has invalid evidence grade '{grade}'.")
        if not re.search(r"Q\s*\d+", mapping, flags=re.I):
            errors.append(f"Evidence ledger row {row_number} has no Q-number mapping.")

    missing_questions = [
        f"Q{number}"
        for number in range(1, question_count + 1)
        if not re.search(rf"Q\s*{number}(?!\d)", text, flags=re.I)
    ]
    if missing_questions:
        errors.append(
            "Evidence ledger does not cover every report question: " + ", ".join(missing_questions)
        )
    if "[待补充]" in text:
        errors.append("Evidence ledger still contains [待补充] placeholders.")

    return errors, warnings


def validate_rendered_pages(path: Path) -> tuple[list[str], list[str], int]:
    errors: list[str] = []
    warnings: list[str] = []
    if not path.exists() or not path.is_dir():
        return [f"Rendered-page directory not found: {path}"], warnings, 0

    numbered: list[tuple[int, Path]] = []
    for item in path.glob("page-*.png"):
        match = re.fullmatch(r"page-(\d+)\.png", item.name)
        if match:
            numbered.append((int(match.group(1)), item))
    numbered.sort()
    if not numbered:
        return [f"No page-<N>.png files found in {path}."], warnings, 0

    numbers = [number for number, _ in numbered]
    if numbers != list(range(1, len(numbers) + 1)):
        errors.append(f"Rendered page numbering is not continuous: {numbers}.")

    for number, item in numbered:
        data = item.read_bytes()
        if len(data) < 500:
            errors.append(f"Rendered page {number} is unexpectedly small ({len(data)} bytes).")
            continue
        if data[:8] != b"\x89PNG\r\n\x1a\n" or len(data) < 24:
            errors.append(f"Rendered page {number} is not a valid PNG file.")
            continue
        width = int.from_bytes(data[16:20], "big")
        height = int.from_bytes(data[20:24], "big")
        if width < 500 or height < 700:
            errors.append(f"Rendered page {number} resolution is too low: {width}x{height}.")

    return errors, warnings, len(numbered)


def mm(value) -> float:
    return value.mm if value is not None else 0.0


def points(value) -> float:
    return value.pt if value is not None else 0.0


def _has_bool_ppr(paragraph, tag: str) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.find(qn(f"w:{tag}")) is not None


def _check_style(
    doc: Document,
    name: str,
    *,
    font: str,
    size: float,
    bold: bool,
    alignment,
    leading: float,
    before: float,
    after: float,
    first_line: float,
) -> list[str]:
    errors: list[str] = []
    style = doc.styles[name]
    fmt = style.paragraph_format
    if style.font.name != font:
        errors.append(f"Style {name} font is '{style.font.name}'; expected '{font}'.")
    if abs(points(style.font.size) - size) > 0.05:
        errors.append(f"Style {name} font size is {points(style.font.size):.2f} pt; expected {size:.2f} pt.")
    if bool(style.font.bold) is not bold:
        errors.append(f"Style {name} bold is {style.font.bold}; expected {bold}.")
    if fmt.alignment != alignment:
        errors.append(f"Style {name} alignment is {fmt.alignment}; expected {alignment}.")
    if fmt.line_spacing_rule != WD_LINE_SPACING.EXACTLY:
        errors.append(f"Style {name} must use exact line spacing.")
    if abs(points(fmt.line_spacing) - leading) > 0.05:
        errors.append(f"Style {name} leading is {points(fmt.line_spacing):.2f} pt; expected {leading:.2f} pt.")
    if abs(points(fmt.space_before) - before) > 0.05:
        errors.append(f"Style {name} space before is {points(fmt.space_before):.2f} pt; expected {before:.2f} pt.")
    if abs(points(fmt.space_after) - after) > 0.05:
        errors.append(f"Style {name} space after is {points(fmt.space_after):.2f} pt; expected {after:.2f} pt.")
    if abs(points(fmt.first_line_indent) - first_line) > 0.05:
        errors.append(
            f"Style {name} first-line indent is {points(fmt.first_line_indent):.2f} pt; "
            f"expected {first_line:.2f} pt."
        )
    if abs(points(fmt.left_indent)) > 0.05 or abs(points(fmt.right_indent)) > 0.05:
        errors.append(f"Style {name} must have zero left and right indents.")
    return errors


def validate_docx(path: Path, question_count: int) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []
    doc = Document(path)
    if not doc.sections:
        errors.append("DOCX has no section.")
        return errors, warnings
    if len(doc.sections) != 1:
        errors.append(f"DOCX must contain exactly one section; found {len(doc.sections)}.")
    section = doc.sections[0]
    targets = {
        "page width": (mm(section.page_width), 210.0),
        "page height": (mm(section.page_height), 297.0),
        "top margin": (mm(section.top_margin), 25.4),
        "bottom margin": (mm(section.bottom_margin), 25.4),
        "left margin": (mm(section.left_margin), 31.75),
        "right margin": (mm(section.right_margin), 31.75),
        "header distance": (mm(section.header_distance), 0.0),
        "footer distance": (mm(section.footer_distance), 17.65),
    }
    for name, (actual, expected) in targets.items():
        if abs(actual - expected) > 0.35:
            errors.append(f"DOCX {name} is {actual:.2f} mm; expected {expected:.2f} mm.")

    style_names = {style.name for style in doc.styles}
    missing = REQUIRED_STYLES - style_names
    if missing:
        errors.append(f"Missing required DOCX styles: {sorted(missing)}")
    else:
        style_specs = (
            ("QA Title", FIXED_BODY_FONT, 16.0, True, WD_ALIGN_PARAGRAPH.CENTER, 22.8, 0.0, 12.0, 0.0),
            ("QA Subtitle", FIXED_BODY_FONT, 16.0, True, WD_ALIGN_PARAGRAPH.CENTER, 22.8, 0.0, 28.8, 0.0),
            ("QA Question", FIXED_BODY_FONT, 10.5, True, WD_ALIGN_PARAGRAPH.LEFT, 22.8, 22.8, 0.0, 0.0),
            ("QA Body", FIXED_BODY_FONT, 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 22.8, 0.0, 0.0, 21.0),
            ("QA List", FIXED_BODY_FONT, 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 22.8, 0.0, 0.0, 21.0),
            ("QA Table", FIXED_BODY_FONT, 10.5, False, WD_ALIGN_PARAGRAPH.LEFT, 22.8, 0.0, 0.0, 0.0),
        )
        for spec in style_specs:
            errors.extend(
                _check_style(
                    doc,
                    spec[0],
                    font=spec[1],
                    size=spec[2],
                    bold=spec[3],
                    alignment=spec[4],
                    leading=spec[5],
                    before=spec[6],
                    after=spec[7],
                    first_line=spec[8],
                )
            )

    main_paragraphs = doc.paragraphs
    q_paragraphs = [p for p in main_paragraphs if p.style and p.style.name == "QA Question"]
    if len(q_paragraphs) != question_count:
        errors.append(f"DOCX contains {len(q_paragraphs)} QA Question paragraphs; expected {question_count}.")
    for p in q_paragraphs:
        if not _has_bool_ppr(p, "keepNext"):
            errors.append(f"Question heading is not keep-with-next: {p.text[:40]}")
        if not _has_bool_ppr(p, "keepLines"):
            errors.append(f"Question heading is not keep-lines-together: {p.text[:40]}")
        if not _has_bool_ppr(p, "widowControl"):
            errors.append(f"Question heading has no widow/orphan control: {p.text[:40]}")
    if q_paragraphs:
        if abs(points(q_paragraphs[0].paragraph_format.space_before)) > 0.05:
            errors.append("The first question must have zero direct space before.")
        for p in q_paragraphs[1:]:
            direct_before = p.paragraph_format.space_before
            if direct_before is not None and abs(points(direct_before) - 22.8) > 0.05:
                errors.append(f"Question has unexpected direct space before: {p.text[:40]}")

    title_paragraphs = [p for p in main_paragraphs if p.style and p.style.name == "QA Title"]
    subtitle_paragraphs = [p for p in main_paragraphs if p.style and p.style.name == "QA Subtitle"]
    if len(title_paragraphs) != 1:
        errors.append(f"DOCX must contain exactly one QA Title paragraph; found {len(title_paragraphs)}.")
    if len(subtitle_paragraphs) != 1 or (subtitle_paragraphs and subtitle_paragraphs[0].text != "Q & A"):
        errors.append("DOCX must contain exactly one 'Q & A' QA Subtitle paragraph.")
    for paragraph, label in (
        *((p, "title") for p in title_paragraphs),
        *((p, "subtitle") for p in subtitle_paragraphs),
    ):
        for tag in ("keepNext", "keepLines", "widowControl"):
            if not _has_bool_ppr(paragraph, tag):
                errors.append(f"DOCX {label} paragraph is missing {tag}.")

    allowed_main_styles = {
        "QA Title",
        "QA Subtitle",
        "QA Question",
        "QA Body",
    }
    for index, paragraph in enumerate(main_paragraphs, start=1):
        if not paragraph.text.strip():
            errors.append(f"DOCX main body contains a manual blank paragraph at position {index}.")
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name not in allowed_main_styles:
            errors.append(f"DOCX paragraph {index} uses unexpected style '{style_name}'.")
        if style_name == "QA Body" and not _has_bool_ppr(paragraph, "widowControl"):
            errors.append(f"QA Body paragraph {index} has no widow/orphan control.")

    for q_paragraph in q_paragraphs:
        q_index = main_paragraphs.index(q_paragraph)
        following = next(
            (p for p in main_paragraphs[q_index + 1:] if p.text.strip()),
            None,
        )
        if following is None or not following.text.startswith("答复："):
            errors.append(f"Question is not followed by a bold answer label: {q_paragraph.text[:40]}")
        elif not following.runs or not following.runs[0].bold:
            errors.append(f"Answer label is not bold after question: {q_paragraph.text[:40]}")

    for paragraph in main_paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        expected_size = 16.0 if style_name in {"QA Title", "QA Subtitle"} else 10.5
        for run in paragraph.runs:
            if not run.text:
                continue
            if run.font.name != FIXED_BODY_FONT:
                errors.append(f"Run font is '{run.font.name}' in text '{run.text[:20]}'; expected '{FIXED_BODY_FONT}'.")
            if abs(points(run.font.size) - expected_size) > 0.05:
                errors.append(
                    f"Run size is {points(run.font.size):.2f} pt in text '{run.text[:20]}'; "
                    f"expected {expected_size:.2f} pt."
                )
            if style_name in {"QA Title", "QA Subtitle", "QA Question"} and not run.bold:
                errors.append(f"Heading run is not bold: {run.text[:20]}")

    if doc.tables:
        errors.append(f"Fixed lancheng_qa_a4_fixed output does not allow tables; found {len(doc.tables)}.")

    footer_xml = section.footer._element.xml
    if "PAGE" not in footer_xml:
        errors.append("Footer does not contain a PAGE field.")
    footer_style = doc.styles["Footer"]
    if footer_style.font.name != FIXED_FOOTER_FONT:
        errors.append(
            f"Footer style font is '{footer_style.font.name}'; expected '{FIXED_FOOTER_FONT}'."
        )
    if abs(points(footer_style.font.size) - 9.0) > 0.05:
        errors.append(
            f"Footer style size is {points(footer_style.font.size):.2f} pt; expected 9.00 pt."
        )
    footer_format = footer_style.paragraph_format
    if footer_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        errors.append("Footer style must be centered.")
    if footer_format.line_spacing_rule != WD_LINE_SPACING.EXACTLY:
        errors.append("Footer style must use exact line spacing.")
    if abs(points(footer_format.line_spacing) - 12.0) > 0.05:
        errors.append(
            f"Footer style leading is {points(footer_format.line_spacing):.2f} pt; expected 12.00 pt."
        )
    footer_paragraphs = [p for p in section.footer.paragraphs if p._p.xml]
    if len(footer_paragraphs) != 1:
        errors.append(f"Footer must contain exactly one paragraph; found {len(footer_paragraphs)}.")
    elif footer_paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.CENTER:
        errors.append("Footer page number must be centered.")
    if FIXED_FOOTER_FONT not in footer_xml:
        errors.append(f"Footer PAGE field is not explicitly mapped to '{FIXED_FOOTER_FONT}'.")
    if any(p.text.strip() for p in section.header.paragraphs):
        errors.append("Header must be empty in lancheng_qa_a4_fixed profile.")

    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        rels = zf.read("word/_rels/document.xml.rels").decode("utf-8", "ignore") if "word/_rels/document.xml.rels" in names else ""
        document_xml = zf.read("word/document.xml").decode("utf-8", "ignore")
        styles_xml = zf.read("word/styles.xml").decode("utf-8", "ignore") if "word/styles.xml" in names else ""
        if "TargetMode=\"External\"" in rels or "TargetMode='External'" in rels:
            errors.append("DOCX contains an external relationship.")
        visible_text = "\n".join(p.text for p in doc.paragraphs)
        visible_text += "\n" + "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        if re.search(r"https?://", visible_text):
            errors.append("DOCX visible text contains a raw URL.")
        if "[待补充]" in document_xml:
            errors.append("DOCX contains unresolved [待补充] placeholder.")
        for prohibited_tag, label in (
            ("<w:tbl", "table"),
            ("<w:drawing", "drawing"),
            ("<w:pict", "picture"),
            ("<w:hyperlink", "hyperlink"),
            ("<w:br", "manual line/page break"),
            ("<w:numPr", "automatic numbering"),
            ("<w:footnoteReference", "footnote"),
            ("<w:endnoteReference", "endnote"),
        ):
            if prohibited_tag in document_xml:
                errors.append(f"Fixed lancheng_qa_a4_fixed output contains a prohibited {label} element.")
        if FIXED_BODY_FONT not in document_xml + styles_xml:
            errors.append(f"Required fixed body font mapping '{FIXED_BODY_FONT}' was not found.")
        for font in ("方正仿宋_GBK", "仿宋_GB2312", "STFangsong", "Songti SC"):
            if font in document_xml + styles_xml:
                errors.append(f"Font substitution or mixed font mapping found: '{font}'.")

    return errors, warnings


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("--docx", type=Path)
    parser.add_argument("--allow-placeholders", action="store_true")
    parser.add_argument("--evidence-ledger", type=Path)
    parser.add_argument("--render-dir", type=Path)
    parser.add_argument(
        "--final-gate",
        action="store_true",
        help="Require Markdown, DOCX, evidence ledger, rendered pages, and zero warnings.",
    )
    parser.add_argument(
        "--profile",
        choices=tuple(PROFILE_RANGES),
        default="communication",
        help="Writing profile used for question-count, risk, decision, and natural-language checks.",
    )
    parser.add_argument(
        "--native-term",
        action="append",
        default=[],
        help="Company-native term expected in question headings; repeat for multiple terms.",
    )
    parser.add_argument(
        "--allow-reader-term",
        action="append",
        default=[],
        help="Allow an otherwise prohibited term only when an unavoidable official proper noun contains it.",
    )
    args = parser.parse_args()

    errors, warnings, q_count = validate_markdown(
        args.markdown.resolve(),
        args.allow_placeholders,
        args.profile,
        tuple(args.native_term),
        tuple(args.allow_reader_term),
    )
    evidence_checked = False
    rendered_page_count = 0
    if args.evidence_ledger:
        ledger_errors, ledger_warnings = validate_evidence_ledger(
            args.evidence_ledger.resolve(), q_count
        )
        errors.extend(ledger_errors)
        warnings.extend(ledger_warnings)
        evidence_checked = True
    if args.docx:
        docx_errors, docx_warnings = validate_docx(args.docx.resolve(), q_count)
        errors.extend(docx_errors)
        warnings.extend(docx_warnings)
    if args.render_dir:
        render_errors, render_warnings, rendered_page_count = validate_rendered_pages(
            args.render_dir.resolve()
        )
        errors.extend(render_errors)
        warnings.extend(render_warnings)

    if args.final_gate:
        if not args.docx:
            errors.append("Final gate requires --docx.")
        if not args.evidence_ledger:
            errors.append("Final gate requires --evidence-ledger.")
        if not args.render_dir:
            errors.append("Final gate requires --render-dir.")
        if args.allow_placeholders:
            errors.append("Final gate cannot be used with --allow-placeholders.")
        if warnings:
            errors.extend(f"Final-gate warning must be resolved: {item}" for item in warnings)
            warnings = []

    for item in errors:
        print(f"ERROR: {item}")
    for item in warnings:
        print(f"WARNING: {item}")
    print(
        f"RESULT: {len(errors)} error(s), {len(warnings)} warning(s), "
        f"{q_count} question(s), profile={args.profile}, "
        f"evidence_audit={'yes' if evidence_checked else 'no'}, rendered_pages={rendered_page_count}"
    )
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
