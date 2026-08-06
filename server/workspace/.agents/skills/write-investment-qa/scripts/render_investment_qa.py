#!/usr/bin/env python3
"""Render Markdown to the single lancheng_qa_a4_fixed investment Q&A format."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


# FangSong is the locked cross-client font for this skill. The retained PDF uses
# a licensed 方正仿宋_GBK subset; allowing per-run font substitution caused output
# to drift between Word, WPS, LibreOffice, Codex and Claude Code.
DEFAULT_FONT = "FangSong"
FOOTER_FONT = "Helvetica Neue"
PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
TOP_MARGIN_MM = 25.4
BOTTOM_MARGIN_MM = 25.4
SIDE_MARGIN_MM = 31.75
FOOTER_DISTANCE_MM = 17.65
BODY_SIZE_PT = 10.5
BODY_LEADING_PT = 22.8
FIRST_LINE_PT = 21.0
TITLE_SIZE_PT = 16.0
TABLE_SIZE_PT = BODY_SIZE_PT
TABLE_LEADING_PT = BODY_LEADING_PT

FONT_FILES = {
    DEFAULT_FONT: ("Fangsong.ttf", "FangSong.ttf"),
    FOOTER_FONT: ("HelveticaNeue.ttc", "HelveticaNeue.ttf", "Helvetica Neue.ttf"),
}
FONT_ROOTS = (
    Path("/System/Library/Fonts"),
    Path("/Library/Fonts"),
    Path("/Applications/Microsoft Word.app/Contents/Resources/DFonts"),
    Path("/Applications/Microsoft Excel.app/Contents/Resources/DFonts"),
    Path("/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts"),
    Path("/usr/share/fonts"),
    Path("/usr/local/share/fonts"),
    Path("C:/Windows/Fonts"),
)


def assert_fixed_fonts_available() -> None:
    missing: list[str] = []
    for font_name, filenames in FONT_FILES.items():
        found = False
        for root in FONT_ROOTS:
            if not root.exists():
                continue
            for filename in filenames:
                if (root / filename).exists() or any(root.rglob(filename)):
                    found = True
                    break
            if found:
                break
        if not found:
            missing.append(font_name)
    if missing:
        raise ValueError(
            "Fixed lancheng_qa_a4_fixed fonts are unavailable: " + ", ".join(missing) + ". "
            "Install the exact fonts before generating the reader-facing DOCX; font substitution is disabled."
        )


def set_run_font(run, font_name: str, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font_name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, font_name: str, size_pt: float, bold: bool = False) -> None:
    style.font.name = font_name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), font_name)


def set_exact_paragraph_format(
    paragraph_format,
    leading_pt: float,
    before_pt: float = 0,
    after_pt: float = 0,
    first_line_pt: float = 0,
) -> None:
    paragraph_format.line_spacing = Pt(leading_pt)
    paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph_format.space_before = Pt(before_pt)
    paragraph_format.space_after = Pt(after_pt)
    paragraph_format.left_indent = Pt(0)
    paragraph_format.right_indent = Pt(0)
    paragraph_format.first_line_indent = Pt(first_line_pt)


def add_bool_ppr(paragraph, tag: str, value: bool = True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    node = ppr.find(qn(f"w:{tag}"))
    if node is None:
        node = OxmlElement(f"w:{tag}")
        ppr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def ensure_style(doc: Document, name: str):
    styles = doc.styles
    return styles[name] if name in styles else styles.add_style(name, 1)


def configure_styles(doc: Document, font_name: str) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, font_name, BODY_SIZE_PT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_exact_paragraph_format(normal.paragraph_format, BODY_LEADING_PT, first_line_pt=FIRST_LINE_PT)

    title = ensure_style(doc, "QA Title")
    set_style_font(title, font_name, TITLE_SIZE_PT, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_exact_paragraph_format(title.paragraph_format, BODY_LEADING_PT, after_pt=12)

    subtitle = ensure_style(doc, "QA Subtitle")
    set_style_font(subtitle, font_name, TITLE_SIZE_PT, True)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_exact_paragraph_format(subtitle.paragraph_format, BODY_LEADING_PT, after_pt=28.8)

    question = ensure_style(doc, "QA Question")
    set_style_font(question, font_name, BODY_SIZE_PT, True)
    question.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_exact_paragraph_format(question.paragraph_format, BODY_LEADING_PT, before_pt=22.8)

    body = ensure_style(doc, "QA Body")
    set_style_font(body, font_name, BODY_SIZE_PT)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_exact_paragraph_format(body.paragraph_format, BODY_LEADING_PT, first_line_pt=FIRST_LINE_PT)

    list_style = ensure_style(doc, "QA List")
    set_style_font(list_style, font_name, BODY_SIZE_PT)
    list_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_exact_paragraph_format(list_style.paragraph_format, BODY_LEADING_PT, first_line_pt=FIRST_LINE_PT)

    table_style = ensure_style(doc, "QA Table")
    set_style_font(table_style, font_name, TABLE_SIZE_PT)
    table_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_exact_paragraph_format(table_style.paragraph_format, TABLE_LEADING_PT)

    footer_style = doc.styles["Footer"]
    set_style_font(footer_style, FOOTER_FONT, 9.0)
    footer_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_exact_paragraph_format(footer_style.paragraph_format, 12.0)
    # LibreOffice refreshes PAGE fields using the linked character style,
    # whereas Word respects the direct result-run formatting. Configure both.
    if "Footer Char" in doc.styles:
        set_style_font(doc.styles["Footer Char"], FOOTER_FONT, 9.0)


def parse_inline(paragraph, text: str, font_name: str, size_pt: float, base_bold: bool = False) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = base_bold
        if part.startswith("**") and part.endswith("**"):
            bold = True
            part = part[2:-2]
        run = paragraph.add_run(part)
        set_run_font(run, font_name, size_pt, bold)


def clean_question(text: str, fallback_number: int) -> tuple[int, str]:
    text = text.strip()
    match = re.match(r"^(?:Q\s*)?(\d+)\s*[：:、.．]\s*(.+)$", text, flags=re.I)
    if match:
        return int(match.group(1)), match.group(2).strip()
    return fallback_number, text


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def split_table_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_markdown(path: Path) -> tuple[str, list[dict]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    title = ""
    questions: list[dict] = []
    current = None
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# ") and not title:
            title = stripped[2:].strip()
            i += 1
            continue
        if stripped.startswith("## "):
            number, question = clean_question(stripped[3:].strip(), len(questions) + 1)
            current = {"number": number, "question": question, "blocks": []}
            questions.append(current)
            i += 1
            continue
        if current is None:
            i += 1
            continue
        if stripped.startswith("|") or re.match(r"^(?:[-*+]\s+|\d+[.)]\s+)", stripped):
            raise ValueError(
                "Fixed lancheng_qa_a4_fixed output does not allow lists or tables; "
                "rewrite parallel points as ordinary answer paragraphs beginning with 第一、第二、第三."
            )
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i].strip()))
                i += 1
            current["blocks"].append(("table", rows))
            continue
        bullet = re.match(r"^[-*+]\s+(.+)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet:
            current["blocks"].append(("bullet", bullet.group(1)))
        elif numbered:
            current["blocks"].append(("number", numbered.group(1)))
        else:
            current["blocks"].append(("paragraph", stripped))
        i += 1
    if not title:
        raise ValueError("Markdown must contain exactly one '# 项目或公司全称' title.")
    if not questions:
        raise ValueError("Markdown must contain at least one '## Q1：问题' section.")
    numbers = [question["number"] for question in questions]
    if numbers != list(range(1, len(numbers) + 1)):
        raise ValueError(
            "Question numbering must be continuous from Q1 in the fixed format; "
            f"received {numbers}."
        )
    return title, questions


def set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    trpr.append(tag)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 80, bottom: int = 80, end: int = 80) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_column_widths(rows: list[list[str]], total_twips: int) -> list[int]:
    ncols = max(len(r) for r in rows)
    weights = []
    for col in range(ncols):
        longest = max((len(r[col]) if col < len(r) else 0) for r in rows)
        weights.append(max(5, min(36, longest)))
    total_weight = sum(weights)
    widths = [max(720, round(total_twips * w / total_weight)) for w in weights]
    scale = total_twips / sum(widths)
    widths = [round(w * scale) for w in widths]
    widths[-1] += total_twips - sum(widths)
    return widths


def apply_table_geometry(table, widths: list[int], total_twips: int) -> None:
    tbl = table._tbl
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(total_twips))
    tblw.set(qn("w:type"), "dxa")

    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tblpr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblpr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[idx]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table(doc: Document, rows: list[list[str]], font_name: str) -> None:
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    widths = table_column_widths(rows, 8307)
    for r_idx, values in enumerate(rows):
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            text = values[c_idx] if c_idx < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["QA Table"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            parse_inline(paragraph, text, font_name, TABLE_SIZE_PT, base_bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "E7E6E6")
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, widths, 8307)


def add_page_number(section, font_name: str = FOOTER_FONT) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.style = "Footer"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_exact_paragraph_format(paragraph.paragraph_format, 12)
    # A simple field preserves result-run typography more consistently across
    # Word, WPS and LibreOffice than a complex begin/separate/end field.
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = paragraph.add_run()
    set_run_font(run, font_name, 9, False)
    display = OxmlElement("w:t")
    display.text = "1"
    run._r.append(display)
    paragraph._p.remove(run._r)
    field.append(run._r)
    paragraph._p.append(field)


def configure_document(doc: Document, font_name: str) -> None:
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(TOP_MARGIN_MM)
    section.bottom_margin = Mm(BOTTOM_MARGIN_MM)
    section.left_margin = Mm(SIDE_MARGIN_MM)
    section.right_margin = Mm(SIDE_MARGIN_MM)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(FOOTER_DISTANCE_MM)
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs:
        p.text = ""
    configure_styles(doc, font_name)
    add_page_number(section)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = ""
    doc.core_properties.subject = ""
    doc.core_properties.keywords = ""


def render(markdown_path: Path, output_path: Path) -> None:
    font_name = DEFAULT_FONT
    assert_fixed_fonts_available()
    title, questions = parse_markdown(markdown_path)
    doc = Document()
    configure_document(doc, font_name)

    p = doc.add_paragraph(style="QA Title")
    parse_inline(p, title, font_name, TITLE_SIZE_PT, True)
    add_bool_ppr(p, "keepNext", True)
    add_bool_ppr(p, "keepLines", True)
    add_bool_ppr(p, "widowControl", True)

    p = doc.add_paragraph(style="QA Subtitle")
    parse_inline(p, "Q & A", font_name, TITLE_SIZE_PT, True)
    add_bool_ppr(p, "keepNext", True)
    add_bool_ppr(p, "keepLines", True)
    add_bool_ppr(p, "widowControl", True)

    for q_idx, question in enumerate(questions):
        p = doc.add_paragraph(style="QA Question")
        if q_idx == 0:
            p.paragraph_format.space_before = Pt(0)
        parse_inline(p, f"{q_idx + 1}、{question['question']}", font_name, BODY_SIZE_PT, True)
        add_bool_ppr(p, "keepNext", True)
        add_bool_ppr(p, "keepLines", True)
        add_bool_ppr(p, "widowControl", True)

        first_answer = True
        for kind, payload in question["blocks"]:
            if kind == "paragraph":
                p = doc.add_paragraph(style="QA Body")
                if first_answer:
                    label = p.add_run("答复：")
                    set_run_font(label, font_name, BODY_SIZE_PT, True)
                    # Accept both the recommended answer-only Markdown and the
                    # common human-authored form beginning with “答复：”.
                    payload = re.sub(
                        r"^(?:\*\*)?答复[：:](?:\*\*)?\s*",
                        "",
                        payload,
                        count=1,
                    )
                parse_inline(p, payload, font_name, BODY_SIZE_PT)
                add_bool_ppr(p, "widowControl", True)
                first_answer = False
            elif kind in ("bullet", "number"):
                if first_answer:
                    p = doc.add_paragraph(style="QA Body")
                    label = p.add_run("答复：")
                    set_run_font(label, font_name, BODY_SIZE_PT, True)
                    first_answer = False
                style = "List Bullet" if kind == "bullet" else "List Number"
                p = doc.add_paragraph(style=style)
                p.paragraph_format.left_indent = Pt(27)
                p.paragraph_format.first_line_indent = Pt(-13.5)
                p.paragraph_format.line_spacing = Pt(18)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                parse_inline(p, payload, font_name, BODY_SIZE_PT)
            elif kind == "table":
                if first_answer:
                    p = doc.add_paragraph(style="QA Body")
                    label = p.add_run("答复：")
                    set_run_font(label, font_name, BODY_SIZE_PT, True)
                    first_answer = False
                add_table(doc, payload, font_name)
        if first_answer:
            p = doc.add_paragraph(style="QA Body")
            label = p.add_run("答复：[待补充]")
            set_run_font(label, font_name, BODY_SIZE_PT, True)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    try:
        render(args.markdown.resolve(), args.output.resolve())
    except ValueError as exc:
        parser.error(str(exc))
    print(f"Created {args.output.resolve()} (font={DEFAULT_FONT}, format=lancheng_qa_a4_fixed)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
