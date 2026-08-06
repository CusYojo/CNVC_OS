#!/usr/bin/env python3
"""Render a direct-Q&A Markdown report as a formal Chinese DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor


Q_RE = re.compile(r"^##\s+Q(\d+)[：:]\s*(.+)$")
H3_RE = re.compile(r"^###\s+(.+)$")
LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
INLINE_RE = re.compile(
    r"(\[([^\]]+)\]\((https?://[^)]+)\)|\*\*(.+?)\*\*|`([^`]+)`)"
)
TABLE_TITLE_RE = re.compile(r"^表\s*\d+[：:]\s*.+$")
NOTE_RE = re.compile(r"^(?:来源|数据来源|注)[：:]\s*.+$")
NUMBERED_RE = re.compile(r"^\s*(\d+)[.、]\s+(.+)$")
BULLET_RE = re.compile(r"^\s*[-*]\s+(.+)$")
CONCLUSION_LABEL_RE = re.compile(
    r"^\s*(?:\*\*)?结论(?:如下)?[：:](?:\*\*)?",
    re.MULTILINE,
)
CJK_RUN_RE = re.compile(r"([\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef\u3000-\u303f]+)")


PAGE_WIDTH_MM = 210
PAGE_HEIGHT_MM = 297
TOP_MM = 27
BOTTOM_MM = 28
LEFT_MM = 31.8
RIGHT_MM = 30
HEADER_MM = 15
FOOTER_MM = 16
CONTENT_WIDTH_DXA = round((PAGE_WIDTH_MM - LEFT_MM - RIGHT_MM) / 25.4 * 1440)
TABLE_INDENT_DXA = 120

# macOS/Word/WPS 兼容名称。项目运行时已经安装 STFangsong；使用 PostScript
# 名称可避免 LibreOffice 将 Songti SC/Heiti SC 的中文字符渲染为缺字方框。
SONGTI = "STFangsong"
HEITI = "STHeiti"
LATIN = "Times New Roman"
ACCENT = "3F5870"
INK = "1F2730"
MUTED = "6D747C"
GRID = "A9B0B7"
ALT_FILL = "F4F6F8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_run_font(
    run,
    *,
    size: float,
    bold: bool = False,
    color: str = INK,
    east_asia: str | None = None,
    force_cjk: bool = False,
) -> None:
    cjk_font = east_asia or (HEITI if bold else SONGTI)
    base_font = cjk_font if force_cjk else LATIN
    run.font.name = base_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), base_font)
    r_fonts.set(qn("w:hAnsi"), base_font)
    r_fonts.set(qn("w:cs"), base_font)
    r_fonts.set(qn("w:eastAsia"), cjk_font)
    if force_cjk:
        r_fonts.set(qn("w:hint"), "eastAsia")
        lang = r_pr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            r_pr.append(lang)
        lang.set(qn("w:eastAsia"), "zh-CN")


def add_text_runs(
    paragraph,
    text: str,
    *,
    size: float,
    bold: bool = False,
    color: str = INK,
    east_asia: str | None = None,
) -> None:
    for part in CJK_RUN_RE.split(text):
        if not part:
            continue
        force_cjk = bool(CJK_RUN_RE.fullmatch(part))
        run = paragraph.add_run(part)
        set_run_font(
            run,
            size=size,
            bold=bold,
            color=color,
            east_asia=east_asia,
            force_cjk=force_cjk,
        )


def set_style_font(style, *, size: float, bold: bool = False, color: str = INK) -> None:
    style.font.name = LATIN
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), LATIN)
    r_fonts.set(qn("w:hAnsi"), LATIN)
    r_fonts.set(qn("w:cs"), LATIN)
    r_fonts.set(qn("w:eastAsia"), HEITI if bold else SONGTI)


def set_exact_spacing(paragraph, points: float) -> None:
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(points)


def add_hyperlink(paragraph, label: str, url: str, *, size: float, bold: bool = False) -> None:
    relation_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)

    for part in CJK_RUN_RE.split(label):
        if not part:
            continue
        cjk = bool(CJK_RUN_RE.fullmatch(part))
        selected_font = (HEITI if bold else SONGTI) if cjk else LATIN
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), selected_font)
        fonts.set(qn("w:hAnsi"), selected_font)
        fonts.set(qn("w:cs"), selected_font)
        fonts.set(qn("w:eastAsia"), HEITI if bold else SONGTI)
        if cjk:
            fonts.set(qn("w:hint"), "eastAsia")
        run_properties.append(fonts)
        if cjk:
            lang = OxmlElement("w:lang")
            lang.set(qn("w:eastAsia"), "zh-CN")
            run_properties.append(lang)

        size_el = OxmlElement("w:sz")
        size_el.set(qn("w:val"), str(round(size * 2)))
        run_properties.append(size_el)
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), str(round(size * 2)))
        run_properties.append(size_cs)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "315B83")
        run_properties.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        run_properties.append(underline)
        if bold:
            run_properties.append(OxmlElement("w:b"))
        run.append(run_properties)

        text = OxmlElement("w:t")
        text.text = part
        run.append(text)
        hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str, *, size: float, default_bold: bool = False) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            add_text_runs(
                paragraph,
                text[cursor : match.start()],
                size=size,
                bold=default_bold,
            )
        if match.group(2) is not None:
            add_hyperlink(
                paragraph,
                match.group(2),
                match.group(3),
                size=size,
                bold=default_bold,
            )
        elif match.group(4) is not None:
            add_text_runs(paragraph, match.group(4), size=size, bold=True)
        else:
            add_text_runs(paragraph, match.group(5), size=size, bold=default_bold)
        cursor = match.end()
    if cursor < len(text):
        add_text_runs(paragraph, text[cursor:], size=size, bold=default_bold)


def set_paragraph_border_bottom(paragraph, *, color: str = "D5D9DD", size: str = "4") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    set_style_font(normal, size=10.5)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.left_indent = Pt(0)
    normal.paragraph_format.right_indent = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(20)

    title = styles.add_style("QA Title", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(title, size=20, bold=True, color="17212B")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.left_indent = Pt(0)
    title.paragraph_format.right_indent = Pt(0)
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    title.paragraph_format.line_spacing = Pt(24)
    title.paragraph_format.keep_with_next = True

    q_heading = styles["Heading 1"]
    set_style_font(q_heading, size=14, bold=True, color="1F3448")
    q_heading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    q_heading.paragraph_format.left_indent = Pt(0)
    q_heading.paragraph_format.right_indent = Pt(0)
    q_heading.paragraph_format.first_line_indent = Pt(0)
    q_heading.paragraph_format.space_before = Pt(12)
    q_heading.paragraph_format.space_after = Pt(6)
    q_heading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    q_heading.paragraph_format.line_spacing = Pt(21)
    q_heading.paragraph_format.keep_with_next = True

    subheading = styles["Heading 2"]
    set_style_font(subheading, size=11, bold=True, color="2F465B")
    subheading.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subheading.paragraph_format.left_indent = Pt(0)
    subheading.paragraph_format.right_indent = Pt(0)
    subheading.paragraph_format.first_line_indent = Pt(0)
    subheading.paragraph_format.space_before = Pt(8)
    subheading.paragraph_format.space_after = Pt(4)
    subheading.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    subheading.paragraph_format.line_spacing = Pt(18)
    subheading.paragraph_format.keep_with_next = True

    caption = styles.add_style("QA Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(caption, size=10.5, bold=True, color="283746")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.left_indent = Pt(0)
    caption.paragraph_format.right_indent = Pt(0)
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    caption.paragraph_format.line_spacing = Pt(16)
    caption.paragraph_format.keep_with_next = True

    note = styles.add_style("QA Note", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(note, size=9, color=MUTED)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.left_indent = Pt(0)
    note.paragraph_format.right_indent = Pt(0)
    note.paragraph_format.first_line_indent = Pt(0)
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(6)
    note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    note.paragraph_format.line_spacing = Pt(14)

    for style_name in ("QA Bullet", "QA Number"):
        style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, size=10.5)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.left_indent = Pt(27)
        style.paragraph_format.right_indent = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(18)


def configure_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(PAGE_WIDTH_MM)
    section.page_height = Mm(PAGE_HEIGHT_MM)
    section.top_margin = Mm(TOP_MM)
    section.bottom_margin = Mm(BOTTOM_MM)
    section.left_margin = Mm(LEFT_MM)
    section.right_margin = Mm(RIGHT_MM)
    section.header_distance = Mm(HEADER_MM)
    section.footer_distance = Mm(FOOTER_MM)
    section.different_first_page_header_footer = True


def configure_settings(document: Document) -> None:
    settings = document.settings.element
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        for item in compat.findall(qn("w:compatSetting")):
            if item.get(qn("w:name")) == "compatibilityMode":
                item.set(qn("w:val"), "15")


def add_field(paragraph, instruction: str, display_text: str = "1") -> None:
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = display_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_header_footer(section, project_label: str) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    set_exact_spacing(paragraph, 12)
    add_text_runs(paragraph, f"{project_label}｜Q&A", size=9, color=MUTED)
    set_paragraph_border_bottom(paragraph)

    first_header = section.first_page_header
    first_paragraph = first_header.paragraphs[0]
    first_paragraph.text = ""
    first_paragraph.paragraph_format.space_before = Pt(0)
    first_paragraph.paragraph_format.space_after = Pt(0)

    for footer in (section.footer, section.first_page_footer):
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_exact_spacing(paragraph, 12)
        add_text_runs(paragraph, "第 ", size=9, color=MUTED)
        add_field(paragraph, "PAGE", "1")
        add_text_runs(paragraph, " 页", size=9, color=MUTED)


def next_numbering_id(numbering, tag: str) -> int:
    values = []
    for element in numbering.findall(qn(tag)):
        attr = "w:abstractNumId" if tag == "w:abstractNum" else "w:numId"
        value = element.get(qn(attr))
        if value is not None:
            values.append(int(value))
    return max(values, default=-1) + 1


def create_numbering(document: Document, *, bullet: bool) -> int:
    numbering = document.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum")
    num_id = next_numbering_id(numbering, "w:num")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if bullet else "%1.")
    level.append(level_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), LATIN)
    fonts.set(qn("w:hAnsi"), LATIN)
    fonts.set(qn("w:eastAsia"), SONGTI)
    r_pr.append(fonts)
    level.append(r_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])


def table_widths(rows: list[list[str]]) -> list[int]:
    columns = max(len(row) for row in rows)
    header = rows[0] if rows else []
    if columns == 4 and "证据强度" in header:
        return [1850, 2550, 750, CONTENT_WIDTH_DXA - 5150]
    if columns == 4 and "主要优势" in header:
        return [1350, 2050, 2200, CONTENT_WIDTH_DXA - 5600]
    if columns == 5 and "可能性" in header and "关键预警信号" in header:
        return [1550, 700, 700, 2700, CONTENT_WIDTH_DXA - 5650]
    scores: list[float] = []
    for column in range(columns):
        values = [row[column] if column < len(row) else "" for row in rows]
        longest = max(
            4.0,
            max(sum(1.0 if ord(char) < 128 else 1.65 for char in value) for value in values),
        )
        scores.append(min(34.0, longest))
    total = sum(scores)
    widths = [round(CONTENT_WIDTH_DXA * score / total) for score in scores]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    columns = max(len(row) for row in rows)
    return [row + [""] * (columns - len(row)) for row in rows]


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)
        borders.append(element)
    tbl_pr.append(borders)


def set_cell_geometry(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.insert(0, tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")

    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        item = margins.find(qn(f"w:{name}"))
        if item is None:
            item = OxmlElement(f"w:{name}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def add_table(document: Document, lines: list[str]) -> None:
    rows = parse_table_rows(lines)
    widths = table_widths(rows)
    table = document.add_table(rows=len(rows), cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    table_width = tbl_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        tbl_pr.append(table_width)
    table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    table_indent.set(qn("w:type"), "dxa")
    tbl_pr.append(table_indent)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    set_table_borders(table)

    for row_index, (row, row_values) in enumerate(zip(table.rows, rows)):
        tr_pr = row._tr.get_or_add_trPr()
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

        for column_index, (cell, value) in enumerate(zip(row.cells, row_values)):
            set_cell_geometry(cell, widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                set_cell_shading(cell, ACCENT)
            elif row_index % 2 == 0:
                set_cell_shading(cell, ALT_FILL)

            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            set_exact_spacing(paragraph, 14)

            short_value = len(re.sub(r"[*_`]", "", value)) <= 8
            if row_index == 0 or short_value:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, value, size=9, default_bold=row_index == 0)
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)

    trailing = document.add_paragraph()
    trailing.paragraph_format.space_before = Pt(0)
    trailing.paragraph_format.space_after = Pt(0)
    set_exact_spacing(trailing, 4)


def build_docx(markdown_path: Path, output_path: Path) -> None:
    text = markdown_path.read_text(encoding="utf-8")
    if CONCLUSION_LABEL_RE.search(text):
        raise ValueError(
            "The report contains a standalone conclusion label. "
            "Remove all '结论：' paragraphs before rendering."
        )
    lines = text.splitlines()
    title_line = next((line for line in lines if line.startswith("# ")), None)
    if title_line is None:
        raise ValueError("The Markdown report must contain one H1 title.")
    title_text = title_line[2:].strip()
    project_label = re.sub(
        r"(?:(?:标准版)?(?:内部版?|仅供内部使用)?\s*)?Q&A\s*报告$",
        "",
        title_text,
    ).strip()
    if not project_label:
        raise ValueError("The report title must include a project name before 'Q&A 报告'.")

    document = Document()
    configure_styles(document)
    configure_settings(document)
    section = document.sections[0]
    configure_section(section)
    configure_header_footer(section, project_label)
    document.core_properties.title = title_text
    document.core_properties.subject = "项目 Q&A 报告"
    document.core_properties.author = ""
    document.core_properties.keywords = ""

    bullet_num_id = create_numbering(document, bullet=True)
    active_number_num_id: int | None = None
    previous_kind = ""
    question_count = 0

    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            previous_kind = "" if previous_kind not in {"number", "bullet"} else previous_kind
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, table_lines)
            previous_kind = "table"
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="QA Title")
            add_text_runs(
                paragraph,
                line[2:].strip(),
                size=20,
                bold=True,
                east_asia=SONGTI,
            )
            previous_kind = "title"
            index += 1
            continue

        q_match = Q_RE.match(line)
        if q_match:
            question_count += 1
            paragraph = document.add_paragraph(style="Heading 1")
            if question_count > 1:
                paragraph.paragraph_format.page_break_before = True
            add_inline(
                paragraph,
                f"Q{q_match.group(1)}：{q_match.group(2).strip()}",
                size=14,
                default_bold=True,
            )
            previous_kind = "question"
            index += 1
            continue

        subheading = H3_RE.match(line)
        if subheading:
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, subheading.group(1).strip(), size=11, default_bold=True)
            previous_kind = "subheading"
            index += 1
            continue

        if TABLE_TITLE_RE.match(line):
            paragraph = document.add_paragraph(style="QA Table Caption")
            add_text_runs(
                paragraph,
                line,
                size=10.5,
                bold=True,
                east_asia=SONGTI,
            )
            previous_kind = "caption"
            index += 1
            continue

        if NOTE_RE.match(line):
            paragraph = document.add_paragraph(style="QA Note")
            add_inline(paragraph, line, size=9)
            previous_kind = "note"
            index += 1
            continue

        numbered = NUMBERED_RE.match(line)
        if numbered:
            if previous_kind != "number":
                active_number_num_id = create_numbering(document, bullet=False)
            paragraph = document.add_paragraph(style="QA Number")
            apply_numbering(paragraph, active_number_num_id or create_numbering(document, bullet=False))
            add_inline(paragraph, numbered.group(2), size=10.5)
            previous_kind = "number"
            index += 1
            continue

        bullet = BULLET_RE.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="QA Bullet")
            apply_numbering(paragraph, bullet_num_id)
            add_inline(paragraph, bullet.group(1), size=10.5)
            previous_kind = "bullet"
            index += 1
            continue

        paragraph = document.add_paragraph(style="Normal")
        add_inline(paragraph, line, size=10.5)
        previous_kind = "body"
        index += 1

    if question_count == 0:
        raise ValueError("No Q headings were found.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Input Markdown report")
    parser.add_argument("output", type=Path, help="Output DOCX path")
    args = parser.parse_args()
    build_docx(args.input.resolve(), args.output.resolve())
    print(f"Created {args.output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
