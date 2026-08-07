#!/usr/bin/env python3
import argparse
from copy import deepcopy
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = SKILL_DIR / "assets" / "primary-layout-authority.docx"
DEFAULT_SECONDARY_REFERENCE = SKILL_DIR / "assets" / "secondary-layout-authority.docx"
EXPECTED_REFERENCE_SHA256 = "0686dc7cd3bc3f098f6d046239c84ae885e1df03bf8719057e8688a14ec90385"
EXPECTED_SECONDARY_SHA256 = "2693a3836cfc25ea3ef848c35678a5b4b583b02b3b7b9ece1182cf52654d93c8"
ROLE_NAMES = {"星实-题目", "星实-一标", "星实-二标", "星实-正文"}


def part(path, name):
    with zipfile.ZipFile(path) as zf:
        return zf.read(name)


def c14n(data):
    return etree.tostring(etree.fromstring(data), method="c14n", exclusive=True)


def xml_hash(path, name):
    return hashlib.sha256(c14n(part(path, name))).hexdigest()


def style_hashes(path):
    root = etree.fromstring(part(path, "word/styles.xml"))
    found = {}
    for style in root.xpath("//w:style", namespaces=NS):
        name = style.find("w:name", namespaces=NS)
        if name is not None and name.get(qn("w:val")) in ROLE_NAMES:
            found[name.get(qn("w:val"))] = hashlib.sha256(
                etree.tostring(style, method="c14n", exclusive=True)
            ).hexdigest()
    return found


def styles_by_id(path):
    root = etree.fromstring(part(path, "word/styles.xml"))
    return {style.get(qn("w:styleId")): style for style in root.xpath("//w:style", namespaces=NS)}


def normalized_style_hash(style):
    clone = deepcopy(style)
    clone.attrib.pop(qn("w:styleId"), None)
    clone.attrib.pop(qn("w:default"), None)
    name = clone.find(qn("w:name"))
    if name is not None:
        clone.remove(name)
    return hashlib.sha256(etree.tostring(clone, method="c14n", exclusive=True)).hexdigest()


def child(parent, name):
    return None if parent is None else parent.find(qn(f"w:{name}"))


def attr(el, name="val"):
    return None if el is None else el.get(qn(f"w:{name}"))


def table_kind(table):
    borders = child(child(table._tbl, "tblPr"), "tblBorders")
    if attr(child(borders, "left")) in {"none", "nil"} and attr(child(borders, "insideV")) in {"none", "nil"}:
        return "three-line"
    return "grid"


def check_margins(table, index, errors):
    margins = child(child(table._tbl, "tblPr"), "tblCellMar")
    expected = {"top": "0", "left": "108", "bottom": "0", "right": "108"}
    if margins is None:
        errors.append(f"TABLE_MARGIN_MISSING:{index}")
        return
    for side, value in expected.items():
        if attr(child(margins, side), "w") != value:
            errors.append(f"TABLE_MARGIN:{index}:{side}:{value}")


def check_grid(table, index, errors):
    borders = child(child(table._tbl, "tblPr"), "tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = child(borders, edge)
        if attr(border) != "single" or attr(border, "sz") != "4":
            errors.append(f"GRID_BORDER:{index}:{edge}")
    for ci, cell in enumerate(table.rows[0].cells):
        fill = attr(child(child(cell._tc, "tcPr"), "shd"), "fill")
        if fill != "C0C0C0":
            errors.append(f"GRID_HEADER_FILL:{index}:{ci}")


def check_three_line(table, index, errors):
    borders = child(child(table._tbl, "tblPr"), "tblBorders")
    for edge in ("left", "right", "insideV"):
        if attr(child(borders, edge)) not in {"none", "nil"}:
            errors.append(f"THREE_LINE_VERTICAL:{index}:{edge}")
    for ci, cell in enumerate(table.rows[0].cells):
        cell_borders = child(child(cell._tc, "tcPr"), "tcBorders")
        if attr(child(cell_borders, "top")) != "single" or attr(child(cell_borders, "top"), "sz") != "12":
            errors.append(f"THREE_LINE_TOP:{index}:{ci}")
        if attr(child(cell_borders, "bottom")) != "single" or attr(child(cell_borders, "bottom"), "sz") != "4":
            errors.append(f"THREE_LINE_HEADER_BOTTOM:{index}:{ci}")
    for ci, cell in enumerate(table.rows[-1].cells):
        cell_borders = child(child(cell._tc, "tcPr"), "tcBorders")
        if attr(child(cell_borders, "bottom")) != "single" or attr(child(cell_borders, "bottom"), "sz") != "12":
            errors.append(f"THREE_LINE_BOTTOM:{index}:{ci}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    parser.add_argument("--primary-reference", default=str(DEFAULT_REFERENCE))
    parser.add_argument("--secondary-reference", default=str(DEFAULT_SECONDARY_REFERENCE))
    parser.add_argument(
        "--require-table-kind",
        action="append",
        choices=("grid", "three-line"),
        default=[],
    )
    parser.add_argument("--report")
    args = parser.parse_args()
    target = Path(args.docx)
    reference = Path(args.primary_reference)
    secondary_reference = Path(args.secondary_reference)
    errors = []

    reference_sha = hashlib.sha256(reference.read_bytes()).hexdigest()
    if reference_sha != EXPECTED_REFERENCE_SHA256:
        errors.append("REFERENCE_SHA256_MISMATCH")
    secondary_sha = hashlib.sha256(secondary_reference.read_bytes()).hexdigest()
    if secondary_sha != EXPECTED_SECONDARY_SHA256:
        errors.append("SECONDARY_REFERENCE_SHA256_MISMATCH")

    target_styles = styles_by_id(target)
    secondary_styles = styles_by_id(secondary_reference)
    secondary_table_style = secondary_styles.get("11")
    if secondary_table_style is None or secondary_table_style.get(qn("w:type")) != "table":
        errors.append("SECONDARY_TABLE_STYLE_MISSING")
        secondary_table_style_hash = None
    else:
        secondary_table_style_hash = normalized_style_hash(secondary_table_style)

    doc = Document(target)
    ref = Document(reference)
    section = doc.sections[0]
    ref_section = ref.sections[0]
    geometry_names = (
        "page_width", "page_height", "top_margin", "right_margin", "bottom_margin",
        "left_margin", "header_distance", "footer_distance",
    )
    geometry = {}
    for name in geometry_names:
        actual, expected = getattr(section, name), getattr(ref_section, name)
        geometry[name] = [int(actual), int(expected)]
        if actual != expected:
            errors.append(f"SECTION_GEOMETRY:{name}")

    expected_styles = style_hashes(reference)
    actual_styles = style_hashes(target)
    for name in sorted(ROLE_NAMES):
        if name not in actual_styles:
            errors.append(f"STYLE_MISSING:{name}")
        elif actual_styles[name] != expected_styles.get(name):
            errors.append(f"STYLE_MISMATCH:{name}")

    recurring_parts = {}
    for name in ("word/header1.xml", "word/footer1.xml"):
        try:
            actual, expected = xml_hash(target, name), xml_hash(reference, name)
            recurring_parts[name] = [actual, expected]
            if actual != expected:
                errors.append(f"RECURRING_PART_MISMATCH:{name}")
        except KeyError:
            errors.append(f"RECURRING_PART_MISSING:{name}")

    counts = {name: sum(1 for p in doc.paragraphs if p.style and p.style.name == name) for name in ROLE_NAMES}
    if counts["星实-题目"] != 1:
        errors.append("TITLE_STYLE_COUNT")
    if counts["星实-一标"] != 6:
        errors.append("H1_STYLE_COUNT")
    if counts["星实-二标"] != 11:
        errors.append("H2_STYLE_COUNT")
    if counts["星实-正文"] < 10:
        errors.append("BODY_STYLE_COUNT")

    nonempty = [p for p in doc.paragraphs if p.text.strip()]
    if nonempty:
        first_run = next((r for r in nonempty[0].runs if r.text), None)
        rpr = None if first_run is None else child(first_run._element, "rPr")
        fonts = child(rpr, "rFonts")
        if nonempty[0].style.name != "星实-题目" or attr(child(rpr, "sz")) != "32" or (None if fonts is None else fonts.get(qn("w:eastAsia"))) != "宋体":
            errors.append("TITLE_DIRECT_FORMAT")
    else:
        errors.append("BODY_EMPTY")

    kinds = []
    for index, table in enumerate(doc.tables):
        kind = table_kind(table)
        kinds.append(kind)
        table_style_id = attr(child(child(table._tbl, "tblPr"), "tblStyle"))
        table_style = target_styles.get(table_style_id)
        if table_style is None:
            errors.append(f"TABLE_STYLE_MISSING:{index}:{table_style_id}")
        elif table_style.get(qn("w:type")) != "table":
            errors.append(f"TABLE_STYLE_WRONG_TYPE:{index}:{table_style_id}")
        elif kind == "three-line" and secondary_table_style_hash and normalized_style_hash(table_style) != secondary_table_style_hash:
            errors.append(f"THREE_LINE_STYLE_MISMATCH:{index}:{table_style_id}")
        layout = attr(child(child(table._tbl, "tblPr"), "tblLayout"), "type")
        if layout != "fixed":
            errors.append(f"TABLE_LAYOUT_NOT_FIXED:{index}")
        check_margins(table, index, errors)
        if kind == "grid":
            check_grid(table, index, errors)
        else:
            check_three_line(table, index, errors)
    for required_kind in args.require_table_kind:
        if required_kind not in kinds:
            errors.append(f"{required_kind.upper().replace('-', '_')}_TABLE_MISSING")

    report = {
        "target": str(target.resolve()),
        "reference_sha256": reference_sha,
        "secondary_reference_sha256": secondary_sha,
        "error_count": len(errors),
        "errors": errors,
        "geometry": geometry,
        "style_counts": counts,
        "table_kinds": kinds,
        "required_table_kinds": args.require_table_kind,
        "recurring_parts": recurring_parts,
        "passed": not errors,
    }
    text = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        Path(args.report).write_text(text, encoding="utf-8")
    print(text)
    raise SystemExit(0 if not errors else 1)


if __name__ == "__main__":
    main()
