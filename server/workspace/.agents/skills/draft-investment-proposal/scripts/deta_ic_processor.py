#!/usr/bin/env python3
"""Deterministic controller and renderer for evidence-driven VC proposals.

The model-facing work (fact extraction, conflict rulings, thesis construction and
writing) is intentionally represented by JSON stage files.  This program owns
the reproducible parts: input locking, source preparation, gate checks, template
contamination checks, DOCX rendering and run manifests.  It never calls an LLM
or an external model gateway.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
DEFAULT_NATIVE_TOOL = ROOT / "vc_native_tools.py"
TEMPLATE_FILENAME = "德塔式精简工商字段投资提案_固定模板V7.docx"
APPROVED_TEMPLATE_SHA256 = "849a6e1ec86c9576f52332ffbf8dc048dea5d929daa438810f701c4e2116da15"
DEFAULT_TEMPLATE = ROOT.parent / "assets" / TEMPLATE_FILENAME


DEFAULT_DOC_RENDERER = Path(os.environ.get("SBL_DOCX_RENDERER", "render_docx.py"))
REQUIRED_STAGE_FILES = (
    "source_ledger.json",
    "facts.json",
    "rulings.json",
    "section_coverage.json",
    "investment_thesis_tree.json",
    "forecast.json",
    "transaction.json",
    "proposal.json",
    "red_team.json",
    "quality_scorecard.json",
    "revision_log.json",
    "contamination_report.json",
)
DEFAULT_BANNED_SAMPLE_TERMS = (
    "德塔智能",
    "北京德塔源创智能科技有限公司",
    "通用人工智能研究院",
    "马晓健",
    "黄恩远",
)
FORMAL_RELEASE_BANNED_TERMS = (
    "事实ID",
    "事实 ID",
    "证据包",
    "冲突裁决",
    "红队",
    "质量评分",
    "内部底稿",
    "工作底稿",
    "项目组测算",
    "专项核查",
    "管理层列示",
    "不同材料",
    "投资团队",
    "投资团队基准",
    "投资团队下行",
    "公司提供",
    "已提供",
    "据管理层",
    "管理层称",
    "材料显示",
    "资料显示",
    "经核查",
    "经复核",
    "销售合同",
    "银行流水",
    "权属清单",
    "知识产权清单",
    "数据资产清单",
    "退出价值敏感性",
    "该路径能否形成持续价值",
    "主要取决于",
    "成功关键",
    "依赖条件",
    "投资意向书要求",
    "本意向书要求",
    "投资方要求",
    "协议要求",
    "本文件要求",
    "要求公司建立",
    "建议公司建立",
    "要求公司完善",
    "建议公司完善",
    "具备稳定边界",
    "有望",
    "MOIC",
    "IRR",
)
CONCLUSION_ONLY_TERMS = (
    "项目组",
    "投资价值",
)
CONCLUSION_DIMENSIONS = (
    "赛道窗口",
    "技术路径及壁垒",
    "商业化验证",
    "团队构成",
    "综合入股估值或价格折让",
    "下行保护",
    "主方案",
)
FORMAL_RELEASE_BANNED_PATTERNS = (
    (
        "引用投资意向书或外部文件提出要求",
        r"(?:投资意向书|本意向书|投资方|协议|本文件).{0,12}(?:要求|规定)",
    ),
    (
        "自行要求公司建立或完善制度",
        r"(?:建议|要求|需|应)公司.{0,12}(?:建立|完善|制定).{0,18}(?:制度|机制|体系|流程)",
    ),
    (
        "成功关键或依赖条件分析",
        r"(?:能否|成功|持续价值).{0,18}(?:取决于|关键在于|依赖于)",
    ),
    (
        "形成或实现结果的依赖分析",
        r"(?:形成|实现|达到|保持).{0,18}(?:取决于|依赖于|关键在于)",
    ),
    (
        "评价性稳定边界结论",
        r"使.{0,24}具备稳定边界",
    ),
)
POLISHING_MODE = "polish_existing_proposal"
POLISHING_SOURCE_FILE = "source_proposal.json"
FORBIDDEN_FORMAL_TABLES = (
    "补充经营与单位经济性指标",
    "现金流量表（摘要）",
    "退出价值敏感性",
)
REQUIRED_BALANCE_ITEMS = (
    "流动资产：", "货币资金", "短期投资", "应收票据", "应收账款", "预付账款",
    "应收股利", "应收利息", "其他应收款", "存货", "其中：原材料", "在产品",
    "库存商品", "周转材料", "其他流动资产", "流动资产合计", "非流动资产：",
    "长期债券投资", "长期股权投资", "固定资产原价", "减：累计折旧", "固定资产账面价值",
    "在建工程", "工程物资", "固定资产清理", "生产性生物资产", "无形资产", "开发支出",
    "长期待摊费用", "其他非流动资产", "非流动资产合计", "资产总计",
    "流动负债：", "短期借款", "应付票据", "应付账款", "预收账款", "应付职工薪酬",
    "应交税费", "应付利息", "应付利润", "其他应付款", "其他流动负债", "流动负债合计",
    "非流动负债：", "长期借款", "长期应付款", "递延收益", "其他非流动负债",
    "非流动负债合计", "负债合计", "所有者权益（或股东权益）：", "实收资本（或股本）",
    "资本公积", "盈余公积", "未分配利润", "所有者权益（或股东权益）合计",
    "负债和所有者权益（或股东权益）总计",
)
REQUIRED_PROFIT_ITEMS = (
    "一、营业收入", "减：营业成本", "税金及附加", "其中：消费税", "营业税",
    "城市维护建设税", "资源税", "土地增值税", "城镇土地使用税、房产税、车船税、印花税",
    "教育费附加、矿产资源补偿费、排污费", "销售费用", "其中：商品维修费",
    "广告费和业务宣传费", "管理费用", "其中：开办费", "业务招待费", "研究费用",
    "财务费用", "其中：利息费用（收入以“-”号填列）", "加：投资收益（损失以“-”号填列）",
    "二、营业利润（亏损以“-”号填列）", "加：营业外收入", "其中：政府补助",
    "减：营业外支出", "其中：坏账损失", "无法收回的长期债券投资损失",
    "无法收回的长期股权投资损失", "自然灾害等不可抗力因素造成的损失", "税收滞纳金",
    "三、利润总额（亏损总额以“-”号填列）", "减：所得税费用",
    "四、净利润（净亏损以“-”号填列）",
)
REQUIRED_FORECAST_COLUMNS = (
    "营业收入（税前）构成", "2026E", "2027E", "2028E", "2029E", "2030E",
)
REQUIRED_FORECAST_TITLE = "营 业 收 ⼊ 五 年 预 测"
REQUIRED_FORECAST_METRICS = (
    "平均单价（元）", "平均成本（元）", "毛利率",
)
OBJECTIVE_FACT_SUBSECTIONS = (
    "（一）公司简介",
    "（二）核心团队",
    "（四）技术路线",
    "（五）运营摘要",
)
UNIFIED_SOCIAL_CREDIT_CODE_PATTERN = re.compile(
    r"(?:统一社会信用代码(?:为|：|:)?\s*)?[0-9A-Z]{18}"
)
OBJECTIVE_FACT_BANNED_PHRASES = (
    "拟于",
    "计划",
    "尚需",
    "仍需",
    "取决于",
    "可形成",
    "宜作为",
    "建议",
    "内部测试",
)
SECTION_COVERAGE_RULES = {
    "（一）公司简介": {"label": "公司简介", "min_dimensions": 4, "min_paragraphs": 3},
    "（二）核心团队": {"label": "核心团队", "min_dimensions": 4, "min_paragraphs": 3},
    "（三）公司股权结构": {"label": "公司股权结构", "min_dimensions": 3, "min_paragraphs": 2},
    "（四）技术路线": {"label": "技术路线", "min_dimensions": 4, "min_paragraphs": 4},
    "（五）运营摘要": {"label": "运营摘要", "min_dimensions": 3, "min_paragraphs": 3},
    "（一）投资亮点": {"label": "投资亮点", "min_dimensions": 4, "min_paragraphs": 4},
}
BODY_FONT = "宋体"
HEADING_FONT = "黑体"
PAGE_WIDTH_CM = 16.2
PAGE_HEIGHT_CM = 21.0
PAGE_MARGIN_VERTICAL_CM = 1.5
PAGE_MARGIN_HORIZONTAL_CM = 1.8
AVAILABLE_TABLE_WIDTH_DXA = 7140
TABLE_HEADER_FONT_SIZE = 9.5
TABLE_BODY_FONT_SIZE = 9.0


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def read_json(path: Path) -> Any:
    with path.open(encoding="utf-8") as stream:
        return json.load(stream)


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def run_checked(command: list[str], *, cwd: Path | None = None) -> subprocess.CompletedProcess:
    result = subprocess.run(command, cwd=cwd, text=True, capture_output=True)
    if result.returncode:
        raise SystemExit(
            f"command failed ({result.returncode}): {' '.join(command)}\n"
            f"{result.stdout[-2000:]}\n{result.stderr[-2000:]}"
        )
    return result


def init_run(args: argparse.Namespace) -> None:
    input_path = Path(args.input).expanduser().resolve()
    if not input_path.is_file():
        raise SystemExit(f"input archive not found: {input_path}")
    native_tool = Path(args.native_tool).expanduser().resolve()
    output_root = Path(args.output_root).expanduser().resolve()
    output_root.mkdir(parents=True, exist_ok=True)
    init_result = run_checked(
        [sys.executable, str(native_tool), "init-run", "--input", str(input_path),
         "--output-root", str(output_root)]
    )
    response = json.loads(init_result.stdout.strip().splitlines()[-1])
    run_dir = Path(response["run_dir"])
    prepare_result = run_checked(
        [sys.executable, str(native_tool), "prepare", "--input", str(input_path),
         "--run-dir", str(run_dir)]
    )
    prepared = json.loads(prepare_result.stdout.strip().splitlines()[-1])
    state = {
        "workflow": "DETA_IC_DENSE_FACT_RELEASE_V9",
        "status": "prepared",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "input": str(input_path),
        "input_sha256": sha256(input_path),
        "reference_pdf": str(Path(args.reference).expanduser().resolve()),
        "reference_sha256": sha256(Path(args.reference).expanduser().resolve()),
        "model_mode": "current-codex-session",
        "external_llm_gateway": False,
        "source_count": prepared.get("source_count"),
        "packet_count": prepared.get("packet_count"),
        "failure_count": prepared.get("failure_count"),
        "next_gate": "G0_ENTITY_BINDING",
    }
    write_json(run_dir / "workflow_state.json", state)
    packaged_workflow = ROOT.parent / "assets" / "WORKFLOW.md"
    if packaged_workflow.exists():
        shutil.copy2(packaged_workflow, run_dir / "WORKFLOW.md")
    (run_dir / "AGENT_WORK_ORDER.md").write_text(
        "# Agent work order\n\n"
        "Use the current Claude model only. First lock the user-confirmed investor full name, "
        "short name, target full name and forbidden old investor names in proposal.meta.entity_binding. "
        "Process packets one by one; write facts_parts, "
        "then facts, rulings, section_coverage, thesis, operating/financial/transaction models, proposal, "
        "red-team review, revisions and the quality scorecard. Do not call an external "
        "model API. Hard gates and schemas are defined in WORKFLOW.md.\n",
        encoding="utf-8",
    )
    print(json.dumps({"status": "prepared", "run_dir": str(run_dir), **prepared}, ensure_ascii=False))


def all_text(value: Any) -> str:
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        return "\n".join(all_text(v) for v in value.values())
    if isinstance(value, list):
        return "\n".join(all_text(v) for v in value)
    return ""


def parse_forecast_number(value: Any) -> float:
    text = str(value).strip().replace(",", "").replace("%", "")
    return float(text)


def proposal_subsections(proposal: dict[str, Any]) -> dict[str, dict[str, Any]]:
    found: dict[str, dict[str, Any]] = {}
    for section in proposal.get("sections", []):
        for subsection in section.get("subsections", []):
            title = str(subsection.get("title", "")).strip()
            if title:
                found[title] = subsection
    return found


def audit_entity_binding(proposal: dict[str, Any], failures: list[str]) -> None:
    """Lock the user-confirmed investor and target identities across the formal report."""
    meta = proposal.get("meta", {})
    binding = meta.get("entity_binding")
    if not isinstance(binding, dict):
        failures.append("proposal.meta.entity_binding缺失，必须先锁定投资主体和目标主体")
        return

    required = (
        "investment_entity_full_name",
        "investment_entity_short_name",
        "target_entity_full_name",
    )
    missing = [key for key in required if not str(binding.get(key, "")).strip()]
    if missing:
        failures.append("entity_binding缺少必填字段：" + "、".join(missing))
        return

    investor_full = str(binding["investment_entity_full_name"]).strip()
    investor_short = str(binding["investment_entity_short_name"]).strip()
    target_full = str(binding["target_entity_full_name"]).strip()
    execution_full = str(binding.get("execution_entity_full_name", investor_full)).strip()
    title_lines = meta.get("title_lines", [])
    title_first = str(title_lines[0]).strip() if title_lines else ""
    title_text = all_text(title_lines)
    intro_text = all_text(meta.get("intro_paragraphs", []))
    formal_payload = {
        "meta": {key: value for key, value in meta.items() if key != "entity_binding"},
        "sections": proposal.get("sections", []),
    }
    proposal_text = all_text(formal_payload)

    if investor_full not in title_first:
        failures.append("报告标题未使用用户确认的投资主体全称：" + investor_full)
    if target_full not in title_text:
        failures.append("报告标题未使用用户确认的目标主体全称：" + target_full)
    if investor_full not in intro_text:
        failures.append("报告导语未使用用户确认的投资主体全称：" + investor_full)
    if f"以下简称“{investor_short}”" not in intro_text:
        failures.append("报告导语未按entity_binding定义投资主体简称：" + investor_short)
    if str(meta.get("signature_entity", "")).strip() != execution_full:
        failures.append("报告落款与已确认的投资/执行主体不一致：" + execution_full)

    forbidden = binding.get("forbidden_investment_entities", [])
    if not isinstance(forbidden, list):
        failures.append("forbidden_investment_entities必须为数组")
    else:
        for entity in forbidden:
            entity_text = str(entity).strip()
            if entity_text and entity_text in proposal_text:
                failures.append("正式提案出现已禁用的投资主体或简称：" + entity_text)


def audit_section_coverage(
    proposal: dict[str, Any], coverage: dict[str, Any], failures: list[str], warnings: list[str]
) -> None:
    if coverage.get("status") != "pass":
        failures.append("section_coverage.status is not pass")
    coverage_sections = coverage.get("sections", {})
    if not isinstance(coverage_sections, dict):
        failures.append("section_coverage.sections must be an object")
        return
    subsections = proposal_subsections(proposal)
    for title, rule in SECTION_COVERAGE_RULES.items():
        label = str(rule["label"])
        entry = coverage_sections.get(title) or coverage_sections.get(label)
        if not isinstance(entry, dict):
            failures.append(f"章节事实覆盖矩阵缺少：{label}")
            continue
        available = entry.get("available_dimensions", [])
        used = entry.get("used_dimensions", [])
        if not isinstance(available, list) or not isinstance(used, list):
            failures.append(f"{label}的available_dimensions和used_dimensions必须为数组")
            continue
        unknown = [dimension for dimension in used if dimension not in available]
        if unknown:
            failures.append(f"{label}使用了未在可用事实中登记的维度：{'、'.join(map(str, unknown))}")
        subsection = subsections.get(title, {})
        actual_paragraphs = subsection.get("paragraphs", [])
        if len(available) >= int(rule["min_dimensions"]):
            if len(used) < int(rule["min_dimensions"]):
                failures.append(
                    f"{label}已有{len(available)}个可用事实维度，但正文仅使用{len(used)}个，内容展开不足"
                )
            if len(actual_paragraphs) < int(rule["min_paragraphs"]):
                failures.append(
                    f"{label}材料充分时正文至少应有{rule['min_paragraphs']}个主题段落，当前为{len(actual_paragraphs)}个"
                )
        else:
            missing_dimensions = entry.get("missing_dimensions", [])
            if entry.get("source_status") != "insufficient" or not missing_dimensions:
                failures.append(
                    f"{label}材料不足时必须标记source_status=insufficient并记录missing_dimensions"
                )
            else:
                warnings.append(f"{label}因材料不足保持精炼：{len(available)}个可用事实维度")
        declared_count = entry.get("paragraph_count")
        if declared_count is not None and int(declared_count) != len(actual_paragraphs):
            failures.append(
                f"{label}覆盖矩阵paragraph_count={declared_count}，与正文{len(actual_paragraphs)}段不一致"
            )


def audit_conclusion(
    proposal: dict[str, Any], coverage: dict[str, Any], failures: list[str]
) -> None:
    conclusion_sections = [
        section for section in proposal.get("sections", [])
        if str(section.get("title", "")).strip() == "六、结论"
    ]
    if len(conclusion_sections) != 1:
        failures.append("正式提案必须且只能包含一个“六、结论”章节")
        return

    conclusion = conclusion_sections[0]
    paragraphs = conclusion.get("paragraphs", [])
    coverage_sections = coverage.get("sections", {})
    entry = None
    if isinstance(coverage_sections, dict):
        entry = coverage_sections.get("六、结论") or coverage_sections.get("结论")
    if not isinstance(entry, dict):
        failures.append("章节事实覆盖矩阵缺少：结论")
        return

    available = entry.get("available_dimensions", [])
    used = entry.get("used_dimensions", [])
    if not isinstance(available, list) or not isinstance(used, list):
        failures.append("结论的available_dimensions和used_dimensions必须为数组")
        return
    unknown = [dimension for dimension in used if dimension not in available]
    if unknown:
        failures.append("结论使用了未在可用事实中登记的维度：" + "、".join(map(str, unknown)))
    nonstandard = [dimension for dimension in used if dimension not in CONCLUSION_DIMENSIONS]
    if nonstandard:
        failures.append("结论使用了未定义的投决维度：" + "、".join(map(str, nonstandard)))
    declared_count = entry.get("paragraph_count")
    if declared_count is not None and int(declared_count) != len(paragraphs):
        failures.append(
            f"结论覆盖矩阵paragraph_count={declared_count}，与正文{len(paragraphs)}段不一致"
        )

    mode = str(entry.get("decision_mode", "")).strip()
    if mode not in {"recommendation", "neutral"}:
        failures.append("结论必须登记decision_mode为recommendation或neutral")
        return
    conclusion_text = all_text(paragraphs).strip()
    if mode == "recommendation":
        if len(paragraphs) != 1:
            failures.append("推荐模式的“六、结论”必须写成一个完整投决段落")
        if len(used) < 4:
            failures.append("推荐模式的结论至少应使用4个已确认投决维度")
        if "主方案" not in used:
            failures.append("推荐模式的结论必须在依据映射中登记“主方案”")
        if not conclusion_text.startswith("综合考虑"):
            failures.append("推荐模式的结论必须以“综合考虑”起笔")
        if "项目组认为" not in conclusion_text:
            failures.append("推荐模式的结论必须使用“项目组认为”作为投决主语")
        required_ending = "建议在充分控制风险的情况下按主方案实施本次投资。"
        if not conclusion_text.endswith(required_ending):
            failures.append(f"推荐模式的结论必须以“{required_ending}”收束")
    else:
        if "项目组认为" in conclusion_text or "投资价值" in conclusion_text:
            failures.append("neutral模式不得写入“项目组认为”或“投资价值”判断")


def audit_revenue_forecast(table: dict[str, Any], failures: list[str]) -> None:
    title = str(table.get("title", ""))
    normalized_title = unicodedata.normalize("NFKC", title).replace(" ", "")
    if normalized_title != "营业收入五年预测":
        return
    if title != REQUIRED_FORECAST_TITLE:
        failures.append(f"公司业务预测表题必须为：{REQUIRED_FORECAST_TITLE}")
    columns = tuple(str(value) for value in table.get("columns", []))
    if columns != REQUIRED_FORECAST_COLUMNS:
        failures.append(
            "营业收入五年预测表头必须为：营业收入（税前）构成｜2026E｜2027E｜2028E｜2029E｜2030E"
        )
    rows = table.get("rows", [])
    if not rows or len(rows) % 5:
        failures.append("营业收入五年预测必须按每项业务五行分组：收入、销售数量、平均单价、平均成本、毛利率")
        return
    for start in range(0, len(rows), 5):
        block = rows[start:start + 5]
        business = str(block[0][0]).strip() if block and block[0] else f"第{start // 5 + 1}项业务"
        if any(len(row) != 6 for row in block):
            failures.append(f"{business}预测行列数不完整")
            continue
        quantity_label = str(block[1][0]).replace(" ", "")
        labels = [str(block[index][0]).replace(" ", "") for index in range(2, 5)]
        if not quantity_label.startswith("销售数量（") or not quantity_label.endswith("）"):
            failures.append(f"{business}缺少带计量单位的销售数量行")
        if tuple(labels) != REQUIRED_FORECAST_METRICS:
            failures.append(f"{business}必须依次列示平均单价（元）、平均成本（元）和毛利率")
        if business in ("销售数量（项目）", "销售数量（万小时）", *REQUIRED_FORECAST_METRICS):
            failures.append(f"第{start // 5 + 1}个预测分组缺少业务名称及业务收入")
        try:
            revenues = [parse_forecast_number(value) for value in block[0][1:]]
            quantities = [parse_forecast_number(value) for value in block[1][1:]]
            prices = [parse_forecast_number(value) for value in block[2][1:]]
            costs = [parse_forecast_number(value) for value in block[3][1:]]
            margins = [parse_forecast_number(value) for value in block[4][1:]]
        except (TypeError, ValueError):
            failures.append(f"{business}预测包含不可重算的数量、金额或比例")
            continue
        if "万小时" in quantity_label:
            revenue_factor = 1.0
        elif "项目" in quantity_label:
            revenue_factor = 0.0001
        else:
            failures.append(f"{business}销售数量单位缺少已定义的收入换算规则")
            continue
        for year_index, year in enumerate(REQUIRED_FORECAST_COLUMNS[1:]):
            expected_revenue = quantities[year_index] * prices[year_index] * revenue_factor
            revenue_tolerance = max(1.0, abs(revenues[year_index]) * 0.001)
            if abs(revenues[year_index] - expected_revenue) > revenue_tolerance:
                failures.append(
                    f"{business}{year}收入无法由销售数量×平均单价重算："
                    f"表内{revenues[year_index]:g}万元，重算{expected_revenue:g}万元"
                )
            if prices[year_index] <= 0:
                failures.append(f"{business}{year}平均单价必须大于0")
                continue
            expected_margin = (prices[year_index] - costs[year_index]) / prices[year_index] * 100
            if abs(margins[year_index] - expected_margin) > 0.6:
                failures.append(
                    f"{business}{year}毛利率无法由平均单价与平均成本复核："
                    f"表内{margins[year_index]:g}%，重算{expected_margin:.1f}%"
                )


def audit_artifacts(artifacts: Path, banned_terms: list[str]) -> dict[str, Any]:
    missing = [name for name in REQUIRED_STAGE_FILES if not (artifacts / name).is_file()]
    failures: list[str] = []
    warnings: list[str] = []
    if missing:
        failures.append("missing stage files: " + ", ".join(missing))
        return {"status": "fail", "hard_failures": failures, "warnings": warnings}

    quality = read_json(artifacts / "quality_scorecard.json")
    qc = read_json(artifacts / "contamination_report.json")
    score = quality.get("decision_grade_score", quality.get("total_score", 0))
    if quality.get("status") != "pass":
        failures.append("quality_scorecard.status is not pass")
    if float(score or 0) < 85:
        failures.append(f"decision-grade score below 85: {score}")
    if quality.get("hard_failures"):
        failures.extend(f"quality hard gate: {x}" for x in quality["hard_failures"])
    if qc.get("status") != "pass" or qc.get("findings"):
        failures.append("stored contamination report is not clean")

    proposal = read_json(artifacts / "proposal.json")
    proposal_text = all_text(proposal)
    audit_entity_binding(proposal, failures)
    non_conclusion_payload = {
        "meta": proposal.get("meta", {}),
        "sections": [
            section for section in proposal.get("sections", [])
            if str(section.get("title", "")).strip() != "六、结论"
        ],
    }
    non_conclusion_text = all_text(non_conclusion_payload)
    editing_mode = str(proposal.get("meta", {}).get("editing_mode", "")).strip()
    if editing_mode == POLISHING_MODE:
        source_path = artifacts / POLISHING_SOURCE_FILE
        if not source_path.is_file():
            failures.append(
                f"原文润色模式缺少只读基线文件：{POLISHING_SOURCE_FILE}"
            )
        else:
            source_proposal = read_json(source_path)
            source_financial = [
                section for section in source_proposal.get("sections", [])
                if "财务" in str(section.get("title", ""))
            ]
            output_financial = [
                section for section in proposal.get("sections", [])
                if "财务" in str(section.get("title", ""))
            ]
            if source_financial != output_financial:
                failures.append(
                    "原文润色模式下财务部分必须保持原文完全不变：标题、文字、表格、表头、科目、金额、期间、单位、顺序或格式存在差异"
                )
    for term in banned_terms:
        if term and term in proposal_text:
            failures.append(f"sample-template contamination found: {term}")
    for term in FORMAL_RELEASE_BANNED_TERMS:
        if term in proposal_text:
            failures.append(f"formal release contains internal/return-model language: {term}")
    for term in CONCLUSION_ONLY_TERMS:
        if term in non_conclusion_text:
            failures.append(f"formal release uses conclusion-only decision language outside 六、结论: {term}")
    for label, pattern in FORMAL_RELEASE_BANNED_PATTERNS:
        if re.search(pattern, proposal_text):
            failures.append(f"formal release violates objective-only language gate: {label}")

    coverage = read_json(artifacts / "section_coverage.json")
    audit_section_coverage(proposal, coverage, failures, warnings)
    audit_conclusion(proposal, coverage, failures)

    required_sections = [
        "基本情况", "财务", "交易", "预测", "投资亮点", "风险控制", "结论"
    ]
    section_titles = "\n".join(s.get("title", "") for s in proposal.get("sections", []))
    for label in required_sections:
        if label not in section_titles and label not in proposal_text:
            failures.append(f"required report module missing: {label}")

    forecast_sections = [
        section for section in proposal.get("sections", [])
        if section.get("title") == "四、公司业务预测"
    ]
    if len(forecast_sections) != 1:
        failures.append("正式提案必须且只能包含一个“四、公司业务预测”章节")
    else:
        forecast_section = forecast_sections[0]
        if forecast_section.get("paragraphs"):
            failures.append("公司业务预测标题后不得保留说明文字，必须直接进入预测表")
        if not forecast_section.get("page_break_before"):
            failures.append("公司业务预测章节应从新页开始，以保持预测表整表呈现")
        forecast_tables = forecast_section.get("tables", [])
        if len(forecast_tables) != 1:
            failures.append("公司业务预测章节必须且只能包含一张预测表")
        elif str(forecast_tables[0].get("title", "")) != REQUIRED_FORECAST_TITLE:
            failures.append(f"公司业务预测表题必须为：{REQUIRED_FORECAST_TITLE}")
        else:
            forecast_table = forecast_tables[0]
            header_size = float(forecast_table.get("header_font_size", TABLE_HEADER_FONT_SIZE))
            body_size = float(forecast_table.get("body_font_size", TABLE_BODY_FONT_SIZE))
            if header_size != TABLE_HEADER_FONT_SIZE or body_size != TABLE_BODY_FONT_SIZE:
                failures.append(
                    "营业收入五年预测表字体字号必须与利润表一致："
                    f"表头黑体{TABLE_HEADER_FONT_SIZE:g}pt、正文宋体{TABLE_BODY_FONT_SIZE:g}pt"
                )

    for section in proposal.get("sections", []):
        for subsection in section.get("subsections", []):
            subsection_title = str(subsection.get("title", ""))
            subsection_text = all_text(subsection.get("paragraphs", []))
            if subsection_title == "（一）公司简介" and (
                "统一社会信用代码" in subsection_text
                or UNIFIED_SOCIAL_CREDIT_CODE_PATTERN.search(subsection_text)
            ):
                failures.append("公司简介不得呈现统一社会信用代码或18位信用代码")
            if subsection_title in OBJECTIVE_FACT_SUBSECTIONS:
                for phrase in OBJECTIVE_FACT_BANNED_PHRASES:
                    if phrase in subsection_text:
                        failures.append(
                            f"{subsection_title}仅可保留确认后的客观事实，出现判断或待核查表述：{phrase}"
                        )
            if subsection_title == "（二）风险控制":
                risk_paragraphs = subsection.get("paragraphs", [])
                if not risk_paragraphs:
                    failures.append("风险控制不得为空")
                if risk_paragraphs and not 3 <= len(risk_paragraphs) <= 5:
                    failures.append("风险控制应选择3至5项与项目事实直接相关的风险")
                for index, paragraph in enumerate(risk_paragraphs, start=1):
                    risk_text = str(paragraph)
                    if "风险—应对措施：" not in risk_text:
                        failures.append(f"风险控制第{index}条必须采用“风险—应对措施：”格式")
                    if len(risk_text) > 120:
                        failures.append(f"风险控制第{index}条过长，应保持简洁")
            if subsection.get("paragraphs_after"):
                failures.append(f"formal release must not place explanatory prose below tables: {subsection.get('title', '')}")
            for table in subsection.get("tables", []):
                title = str(table.get("title", ""))
                columns = table.get("columns", [])
                if title in FORBIDDEN_FORMAL_TABLES:
                    failures.append(f"formal release contains forbidden table: {title}")
                if "资产负债表" in title and columns != ["资产", "期末余额", "负债和所有者权益", "期末余额"]:
                    failures.append("资产负债表表头必须为：资产｜期末余额｜负债和所有者权益｜期末余额")
                if title.startswith("利润表") and columns != ["项目", "本年累计金额", "本期金额"]:
                    failures.append("利润表表头必须为：项目｜本年累计金额｜本期金额")
                row_text = all_text(table.get("rows", []))
                if "资产负债表" in title:
                    missing_items = [item for item in REQUIRED_BALANCE_ITEMS if item not in row_text]
                    if missing_items:
                        failures.append("资产负债表科目不完整：" + "、".join(missing_items))
                if title.startswith("利润表"):
                    normalized = row_text.replace('"-"', '“-”')
                    missing_items = [item for item in REQUIRED_PROFIT_ITEMS if item not in normalized]
                    if missing_items:
                        failures.append("利润表科目不完整：" + "、".join(missing_items))
                audit_revenue_forecast(table, failures)
        if section.get("paragraphs_after"):
            failures.append(f"formal release must not place explanatory prose below tables: {section.get('title', '')}")
        for table in section.get("tables", []):
            title = str(table.get("title", ""))
            if title in FORBIDDEN_FORMAL_TABLES:
                failures.append(f"formal release contains forbidden table: {title}")
            audit_revenue_forecast(table, failures)

    facts = read_json(artifacts / "facts.json")
    if not isinstance(facts, list) or len(facts) < 20:
        failures.append("facts.json has fewer than 20 atomic facts")
    else:
        ids = [f.get("id") or f.get("fact_id") for f in facts]
        if any(not value for value in ids):
            warnings.append("some facts do not carry an id/fact_id")

    transaction = read_json(artifacts / "transaction.json")
    if not any(key in transaction for key in ("post_money", "capitalization", "main_plan")):
        failures.append("transaction model lacks post-money/capitalization data")
    forecast = read_json(artifacts / "forecast.json")
    forecast_text = all_text(forecast)
    if "downside" not in json.dumps(forecast, ensure_ascii=False).lower() and "下行" not in forecast_text:
        failures.append("forecast lacks a downside case")

    return {
        "status": "pass" if not failures else "fail",
        "standard": "DETA_IC_DENSE_FACT_RELEASE_V9",
        "score": score,
        "hard_failures": failures,
        "warnings": warnings,
        "checked_files": {name: sha256(artifacts / name) for name in REQUIRED_STAGE_FILES},
        "checked_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    }


def set_cell_margins(cell, top=55, start=80, bottom=55, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "left" if edge == "start" else "right" if edge == "end" else edge
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name=BODY_FONT, size=10.5, bold=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])
    set_run_font(run, BODY_FONT, 9)


def configure_document(doc: Document) -> None:
    # Force the current Word layout model.  Old FE/compatibility layout can
    # introduce a stray Latin glyph before the first Chinese title in Word.
    settings = doc.settings._element
    compat = settings.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        settings.append(compat)
    for child in list(compat):
        if child.tag == qn("w:useFELayout"):
            compat.remove(child)
        elif child.tag == qn("w:compatSetting") and child.get(qn("w:name")) == "compatibilityMode":
            child.set(qn("w:val"), "15")
    if not any(
        child.tag == qn("w:compatSetting") and child.get(qn("w:name")) == "compatibilityMode"
        for child in compat
    ):
        mode = OxmlElement("w:compatSetting")
        mode.set(qn("w:name"), "compatibilityMode")
        mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        mode.set(qn("w:val"), "15")
        compat.insert(0, mode)
    section = doc.sections[0]
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(PAGE_MARGIN_VERTICAL_CM)
    section.bottom_margin = Cm(PAGE_MARGIN_VERTICAL_CM)
    section.left_margin = Cm(PAGE_MARGIN_HORIZONTAL_CM)
    section.right_margin = Cm(PAGE_MARGIN_HORIZONTAL_CM)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.first_line_indent = Pt(21)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(3.5)
    for style_name, size, before, after in (
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = HEADING_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), HEADING_FONT)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_xml = footer_p._p.xml
    if " PAGE " not in footer_xml and ">PAGE<" not in footer_xml:
        add_page_field(footer_p)


def add_paragraph(doc: Document, text: str, *, bold=False, align=None, size=10.5,
                  indent=True, keep_next=False, page_break=False) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.first_line_indent = Pt(21) if indent else Pt(0)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = keep_next
    p.paragraph_format.page_break_before = page_break
    run = p.add_run(str(text))
    set_run_font(run, HEADING_FONT if bold else BODY_FONT, size, bold)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    tr_pr.append(cant)


def add_table(doc: Document, spec: dict[str, Any]) -> None:
    title = spec.get("title")
    if title:
        add_paragraph(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                      size=10.5, indent=False, keep_next=True)
    if spec.get("note"):
        add_paragraph(doc, spec["note"], align=WD_ALIGN_PARAGRAPH.RIGHT,
                      size=9, indent=False, keep_next=True)
    columns = [str(x) for x in spec.get("columns", [])]
    rows = spec.get("rows", [])
    if not columns:
        return
    raw_widths = [int(value) for value in (spec.get("widths_dxa") or [])]
    if not raw_widths:
        raw_widths = [int(AVAILABLE_TABLE_WIDTH_DXA / len(columns))] * len(columns)
    raw_total = sum(raw_widths)
    if raw_total != AVAILABLE_TABLE_WIDTH_DXA:
        widths = [max(1, round(value * AVAILABLE_TABLE_WIDTH_DXA / raw_total)) for value in raw_widths]
        widths[-1] += AVAILABLE_TABLE_WIDTH_DXA - sum(widths)
    else:
        widths = raw_widths
    table = doc.add_table(rows=1, cols=len(columns))
    table.autofit = False
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total_width = sum(int(x) for x in widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width)))
        grid.append(grid_col)
    for index, width in enumerate(widths):
        table.columns[index].width = Twips(int(width))
    is_revenue_forecast = str(spec.get("title", "")) == REQUIRED_FORECAST_TITLE
    if is_revenue_forecast:
        # The forecast table must share the exact type scale of the income
        # statement.  Ignore stale per-project overrides from earlier releases.
        header_font_size = TABLE_HEADER_FONT_SIZE
        body_font_size = TABLE_BODY_FONT_SIZE
    else:
        header_font_size = float(spec.get("header_font_size", TABLE_HEADER_FONT_SIZE))
        body_font_size = float(
            spec.get("body_font_size", TABLE_BODY_FONT_SIZE if spec.get("dense") else 9.5)
        )
    for index, value in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.width = Twips(int(widths[index]))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=30 if spec.get("dense") else 55,
                         bottom=30 if spec.get("dense") else 55)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(value)
        set_run_font(run, HEADING_FONT, header_font_size, True)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shd)
    set_repeat_table_header(table.rows[0])
    bold_rows = set(spec.get("bold_rows", []))
    bold_cells = {tuple(x) for x in spec.get("bold_cells", [])}
    alignments = spec.get("alignments") or []
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.width = Twips(int(widths[col_index]))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=30 if spec.get("dense") else 55,
                             bottom=30 if spec.get("dense") else 55)
            p = cell.paragraphs[0]
            if col_index < len(alignments):
                alignment = str(alignments[col_index]).lower()
                p.alignment = {
                    "left": WD_ALIGN_PARAGRAPH.LEFT,
                    "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "center": WD_ALIGN_PARAGRAPH.CENTER,
                }.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(str(value))
            set_run_font(run, BODY_FONT, body_font_size,
                         row_index in bold_rows or (row_index, col_index) in bold_cells)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_subsection(doc: Document, subsection: dict[str, Any]) -> None:
    if subsection.get("title"):
        doc.add_paragraph(subsection["title"], style="Heading 2")
    for text in subsection.get("paragraphs", []):
        add_paragraph(doc, text)
    for table in subsection.get("tables", []):
        add_table(doc, table)
    for text in subsection.get("paragraphs_after", []):
        add_paragraph(doc, text)


def validate_approved_template(template_path: Path) -> str:
    template_path = template_path.expanduser().resolve()
    if not template_path.is_file():
        raise SystemExit(f"approved template is missing: {template_path}")
    template_sha256 = sha256(template_path)
    if template_sha256 != APPROVED_TEMPLATE_SHA256:
        raise SystemExit(
            "approved template fingerprint mismatch: "
            f"expected {APPROVED_TEMPLATE_SHA256}, got {template_sha256}"
        )
    try:
        Document(template_path)
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"approved template cannot be opened: {exc}") from exc
    return template_sha256


def clear_template_body(doc: Document) -> None:
    body = doc._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def render_proposal(
    proposal_path: Path,
    output_path: Path,
    template_path: Path,
) -> str:
    proposal = read_json(proposal_path)
    template_path = template_path.expanduser().resolve()
    template_sha256 = validate_approved_template(template_path)
    doc = Document(template_path)
    clear_template_body(doc)
    configure_document(doc)
    meta = proposal.get("meta", {})
    for index, line in enumerate(meta.get("title_lines", [])):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(12 if index == 0 else 0)
        p.paragraph_format.space_after = Pt(6 if index == 0 else 18)
        p.paragraph_format.keep_with_next = index == 0
        run = p.add_run(line)
        set_run_font(run, HEADING_FONT, 17, True)
    if meta.get("salutation"):
        add_paragraph(doc, meta["salutation"], bold=True, indent=False)
    for text in meta.get("intro_paragraphs", []):
        add_paragraph(doc, text)
    for section in proposal.get("sections", []):
        section_heading = doc.add_paragraph(section.get("title", ""), style="Heading 1")
        section_heading.paragraph_format.page_break_before = bool(section.get("page_break_before"))
        for subsection in section.get("subsections", []):
            add_subsection(doc, subsection)
        for text in section.get("paragraphs", []):
            add_paragraph(doc, text)
        for table in section.get("tables", []):
            add_table(doc, table)
    if meta.get("signature_entity") or meta.get("date"):
        signature_values = [meta.get("signature_entity", ""), meta.get("date", "")]
        for index, value in enumerate(signature_values):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = index == 0
            run = p.add_run(str(value))
            set_run_font(run, BODY_FONT, 10.5)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return template_sha256


def command_audit(args: argparse.Namespace) -> None:
    artifacts = Path(args.artifacts).expanduser().resolve()
    banned = list(DEFAULT_BANNED_SAMPLE_TERMS) + list(args.banned_term or [])
    report = audit_artifacts(artifacts, banned)
    output = artifacts / "processor_audit.json"
    write_json(output, report)
    print(json.dumps({"audit": str(output), **report}, ensure_ascii=False))
    if report["status"] != "pass":
        raise SystemExit(2)


def command_render(args: argparse.Namespace) -> None:
    artifacts = Path(args.artifacts).expanduser().resolve()
    banned = list(DEFAULT_BANNED_SAMPLE_TERMS) + list(args.banned_term or [])
    report = audit_artifacts(artifacts, banned)
    write_json(artifacts / "processor_audit.json", report)
    if report["status"] != "pass":
        raise SystemExit("render blocked by processor_audit.json")
    proposal_path = artifacts / "proposal.json"
    output = Path(args.output).expanduser().resolve()
    template = Path(args.template).expanduser().resolve()
    template_sha256 = validate_approved_template(template)
    rendered_template_sha256 = render_proposal(proposal_path, output, template)
    if rendered_template_sha256 != template_sha256:
        raise SystemExit("render blocked: template provenance was lost during rendering")
    manifest = {
        "status": "rendered",
        "workflow": "DETA_IC_DENSE_FACT_RELEASE_V9",
        "proposal": str(proposal_path),
        "proposal_sha256": sha256(proposal_path),
        "docx": str(output),
        "docx_sha256": sha256(output),
        "template": str(template),
        "template_sha256": template_sha256,
        "template_enforced": True,
        "renderer_mode": "clone-approved-docx",
        "audit": str(artifacts / "processor_audit.json"),
        "external_llm_gateway": False,
        "rendered_at": datetime.now().astimezone().isoformat(timespec="seconds"),
    }
    write_json(output.with_suffix(".manifest.json"), manifest)
    print(json.dumps(manifest, ensure_ascii=False))


def command_verify(args: argparse.Namespace) -> None:
    docx = Path(args.docx).expanduser().resolve()
    output_dir = Path(args.output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if args.word_native:
        if sys.platform != "darwin" or not Path("/Applications/Microsoft Word.app").is_dir():
            raise SystemExit("--word-native requires Microsoft Word on macOS")
        pdf = output_dir / f"{docx.stem}.pdf"
        script = r'''
on run argv
  set sourcePath to item 1 of argv
  set exportPath to item 2 of argv
  tell application "Microsoft Word"
    activate
    set docRef to open file name sourcePath
    delay 2
    save as docRef file name exportPath file format format PDF
    close docRef saving no
  end tell
end run
'''
        exported = subprocess.run(
            ["osascript", "-", str(docx), str(pdf)], input=script,
            text=True, capture_output=True, timeout=90,
        )
        if exported.returncode or not pdf.is_file():
            raise SystemExit(f"Word PDF export failed: {(exported.stderr or exported.stdout)[-1000:]}")
        pdftoppm = shutil.which("pdftoppm")
        if not pdftoppm:
            raise SystemExit("pdftoppm is required for Word-native page rendering")
        run_checked([pdftoppm, "-png", "-r", "140", str(pdf), str(output_dir / "page")])
        pages = sorted(output_dir.glob("page-*.png"))
        if not pages:
            raise SystemExit("visual verification failed: no rendered PNG pages")
        print(json.dumps({"status": "rendered_for_review", "engine": "microsoft-word",
                          "page_count": len(pages), "pdf": str(pdf),
                          "output_dir": str(output_dir)}, ensure_ascii=False))
        return
    renderer = Path(args.renderer).expanduser().resolve()
    if renderer.is_file():
        command = [sys.executable, str(renderer), str(docx), "--output_dir", str(output_dir)]
        if args.emit_pdf:
            command.append("--emit_pdf")
        run_checked(command)
    else:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        pdftoppm = shutil.which("pdftoppm")
        if not office or not pdftoppm:
            raise SystemExit(
                "no document renderer found; install LibreOffice and Poppler, "
                "or use --word-native on macOS"
            )
        profile = output_dir / ".libreoffice-profile"
        profile.mkdir(parents=True, exist_ok=True)
        run_checked([
            office,
            "--headless",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx),
        ])
        pdf = output_dir / f"{docx.stem}.pdf"
        if not pdf.is_file():
            raise SystemExit("LibreOffice PDF export failed")
        run_checked([pdftoppm, "-png", "-r", "140", str(pdf), str(output_dir / "page")])
    pages = sorted(output_dir.glob("page-*.png"))
    if not pages:
        raise SystemExit("visual verification failed: no rendered PNG pages")
    print(json.dumps({"status": "rendered_for_review", "page_count": len(pages),
                      "output_dir": str(output_dir)}, ensure_ascii=False))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Deta-style evidence-driven IC processor")
    sub = parser.add_subparsers(dest="command", required=True)
    p_init = sub.add_parser("init", help="lock an input archive and prepare model packets")
    p_init.add_argument("--input", required=True)
    p_init.add_argument("--reference", required=True)
    p_init.add_argument("--output-root", default=str(ROOT / "runs"))
    p_init.add_argument("--native-tool", default=str(DEFAULT_NATIVE_TOOL))
    p_init.set_defaults(func=init_run)
    p_audit = sub.add_parser("audit", help="run hard gates on completed stage artifacts")
    p_audit.add_argument("--artifacts", required=True)
    p_audit.add_argument("--banned-term", action="append")
    p_audit.set_defaults(func=command_audit)
    p_render = sub.add_parser("render", help="audit then render proposal.json as fixed-template DOCX")
    p_render.add_argument("--artifacts", required=True)
    p_render.add_argument("--output", required=True)
    p_render.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    p_render.add_argument("--banned-term", action="append")
    p_render.set_defaults(func=command_render)
    p_verify = sub.add_parser("verify", help="render DOCX to PNG/PDF for visual QA")
    p_verify.add_argument("--docx", required=True)
    p_verify.add_argument("--output-dir", required=True)
    p_verify.add_argument("--renderer", default=str(DEFAULT_DOC_RENDERER))
    p_verify.add_argument("--emit-pdf", action="store_true")
    p_verify.add_argument("--word-native", action="store_true",
                          help="use Microsoft Word with an explicit document handle on macOS")
    p_verify.set_defaults(func=command_verify)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
