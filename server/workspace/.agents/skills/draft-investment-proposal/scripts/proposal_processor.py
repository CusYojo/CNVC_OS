#!/usr/bin/env python3
"""Render Agent-approved proposal content into the Skill-owned DOCX layout."""

from __future__ import annotations

import hashlib
import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips


LAYOUT_AUTHORITY_SHA256 = "0686dc7cd3bc3f098f6d046239c84ae885e1df03bf8719057e8688a14ec90385"
BODY_FONT = "仿宋"
HEADING_FONT = "黑体"
LEVEL2_FONT = "楷体"
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
PAGE_MARGIN_TOP_CM = 2.54
PAGE_MARGIN_RIGHT_CM = 3.17
PAGE_MARGIN_BOTTOM_CM = 2.64
PAGE_MARGIN_LEFT_CM = 3.17
AVAILABLE_TABLE_WIDTH_DXA = 8306
TABLE_HEADER_FONT_SIZE = 9.5
TABLE_BODY_FONT_SIZE = 9.0


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_json(path: Path) -> Any:
    with path.open(encoding="utf-8") as handle:
        return json.load(handle)


def validate_layout_authority(template_path: Path) -> str:
    resolved = template_path.expanduser().resolve()
    if not resolved.is_file():
        raise SystemExit(f"Skill layout authority is missing: {resolved}")
    actual_sha256 = sha256(resolved)
    if actual_sha256 != LAYOUT_AUTHORITY_SHA256:
        raise SystemExit(
            "Skill layout authority fingerprint mismatch: "
            f"expected {LAYOUT_AUTHORITY_SHA256}, got {actual_sha256}"
        )
    try:
        Document(resolved)
    except Exception as exc:  # noqa: BLE001
        raise SystemExit(f"Skill layout authority cannot be opened: {exc}") from exc
    return actual_sha256


def set_run_font(run, name: str = BODY_FONT, size: float = 10.5, bold=None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, BODY_FONT, 9)


def configure_document(document: Document) -> None:
    settings = document.settings._element
    compatibility = settings.find(qn("w:compat"))
    if compatibility is None:
        compatibility = OxmlElement("w:compat")
        settings.append(compatibility)
    for child in list(compatibility):
        if child.tag == qn("w:useFELayout"):
            compatibility.remove(child)
        elif (
            child.tag == qn("w:compatSetting")
            and child.get(qn("w:name")) == "compatibilityMode"
        ):
            child.set(qn("w:val"), "15")
    if not any(
        child.tag == qn("w:compatSetting")
        and child.get(qn("w:name")) == "compatibilityMode"
        for child in compatibility
    ):
        mode = OxmlElement("w:compatSetting")
        mode.set(qn("w:name"), "compatibilityMode")
        mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        mode.set(qn("w:val"), "15")
        compatibility.insert(0, mode)

    section = document.sections[0]
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(PAGE_MARGIN_TOP_CM)
    section.right_margin = Cm(PAGE_MARGIN_RIGHT_CM)
    section.bottom_margin = Cm(PAGE_MARGIN_BOTTOM_CM)
    section.left_margin = Cm(PAGE_MARGIN_LEFT_CM)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.space_after = Pt(0)
    for style_name, font_name in (("Heading 1", HEADING_FONT), ("Heading 2", LEVEL2_FONT)):
        style = document.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(14)
        style.font.bold = style_name == "Heading 1"
        style.font.color.rgb = RGBColor(0, 0, 0)
        properties = style._element.get_or_add_rPr()
        fonts = properties.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), font_name)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if " PAGE " not in footer._p.xml and ">PAGE<" not in footer._p.xml:
        add_page_field(footer)


def add_paragraph(
    document: Document,
    text: str,
    *,
    bold: bool = False,
    align=None,
    size: float = 12,
    indent: bool = True,
    keep_next: bool = False,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = Pt(24)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Pt(24) if indent else Pt(0)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = keep_next
    run = paragraph.add_run(str(text))
    set_run_font(run, HEADING_FONT if bold else BODY_FONT, size, bold)


def set_cell_margins(cell, *, top=55, start=80, bottom=55, end=80) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def normalized_widths(spec: dict[str, Any], column_count: int) -> list[int]:
    widths = [int(value) for value in (spec.get("widths_dxa") or [])]
    if len(widths) != column_count or sum(widths) <= 0:
        base = int(AVAILABLE_TABLE_WIDTH_DXA / column_count)
        widths = [base] * column_count
        widths[-1] += AVAILABLE_TABLE_WIDTH_DXA - sum(widths)
        return widths
    total = sum(widths)
    if total != AVAILABLE_TABLE_WIDTH_DXA:
        widths = [max(1, round(value * AVAILABLE_TABLE_WIDTH_DXA / total)) for value in widths]
        widths[-1] += AVAILABLE_TABLE_WIDTH_DXA - sum(widths)
    return widths


def add_table(document: Document, spec: dict[str, Any]) -> None:
    title = str(spec.get("title", ""))
    if title:
        add_paragraph(
            document,
            title,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            size=10.5,
            indent=False,
            keep_next=True,
        )
    note = str(spec.get("note", ""))
    if note:
        add_paragraph(
            document,
            note,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
            size=9,
            indent=False,
            keep_next=True,
        )
    columns = [str(value) for value in spec.get("columns", [])]
    if not columns:
        return
    widths = normalized_widths(spec, len(columns))
    table = document.add_table(rows=1, cols=len(columns))
    table.autofit = False
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for index, width in enumerate(widths):
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)
        table.columns[index].width = Twips(width)

    dense = bool(spec.get("dense"))
    header_size = float(spec.get("header_font_size", TABLE_HEADER_FONT_SIZE))
    body_size = float(spec.get("body_font_size", TABLE_BODY_FONT_SIZE if dense else 9.5))
    for index, value in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.width = Twips(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=30 if dense else 55, bottom=30 if dense else 55)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_run_font(run, HEADING_FONT, header_size, True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shading)
    set_repeat_table_header(table.rows[0])

    alignments = [str(value).lower() for value in (spec.get("alignments") or [])]
    bold_rows = {int(value) for value in (spec.get("bold_rows") or [])}
    bold_cells = {tuple(value) for value in (spec.get("bold_cells") or [])}
    for row_index, values in enumerate(spec.get("rows", [])):
        row = table.add_row()
        set_cant_split(row)
        for column_index, value in enumerate(values[:len(columns)]):
            cell = row.cells[column_index]
            cell.width = Twips(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=30 if dense else 55, bottom=30 if dense else 55)
            paragraph = cell.paragraphs[0]
            alignment = alignments[column_index] if column_index < len(alignments) else ""
            paragraph.alignment = {
                "left": WD_ALIGN_PARAGRAPH.LEFT,
                "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "center": WD_ALIGN_PARAGRAPH.CENTER,
            }.get(alignment, WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(
                run,
                BODY_FONT,
                body_size,
                row_index in bold_rows or (row_index, column_index) in bold_cells,
            )


def add_subsection(document: Document, subsection: dict[str, Any]) -> None:
    title = str(subsection.get("title", ""))
    if title:
        document.add_paragraph(title, style="Heading 2")
    for text in subsection.get("paragraphs", []):
        add_paragraph(document, str(text))
    for table in subsection.get("tables", []):
        add_table(document, table)
    for text in subsection.get("paragraphs_after", []):
        add_paragraph(document, str(text))


def clear_template_body(document: Document) -> None:
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def render_proposal(
    proposal_path: Path,
    output_path: Path,
    template_path: Path,
) -> str:
    proposal = read_json(proposal_path)
    if not isinstance(proposal, dict) or not isinstance(proposal.get("meta"), dict):
        raise ValueError("proposal must contain a meta object")
    sections = proposal.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ValueError("proposal.sections must be a non-empty array")
    if any(not isinstance(section, dict) or not section.get("title") for section in sections):
        raise ValueError("each proposal section must contain a title")
    resolved_template = template_path.expanduser().resolve()
    template_sha256 = validate_layout_authority(resolved_template)
    document = Document(resolved_template)
    clear_template_body(document)
    configure_document(document)

    metadata = proposal.get("meta", {})
    for index, line in enumerate(metadata.get("title_lines", [])):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_before = Pt(12 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(6 if index == 0 else 18)
        paragraph.paragraph_format.keep_with_next = index == 0
        run = paragraph.add_run(str(line))
        set_run_font(run, HEADING_FONT, 16, True)
    if metadata.get("salutation"):
        add_paragraph(document, str(metadata["salutation"]), bold=True, indent=False)
    for text in metadata.get("intro_paragraphs", []):
        add_paragraph(document, str(text))

    for section in proposal.get("sections", []):
        heading = document.add_paragraph(str(section.get("title", "")), style="Heading 1")
        heading.paragraph_format.page_break_before = bool(section.get("page_break_before"))
        for subsection in section.get("subsections", []):
            add_subsection(document, subsection)
        for text in section.get("paragraphs", []):
            add_paragraph(document, str(text))
        for table in section.get("tables", []):
            add_table(document, table)

    signature_values = [metadata.get("signature_entity", ""), metadata.get("date", "")]
    if any(signature_values):
        for index, value in enumerate(signature_values):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = index == 0
            run = paragraph.add_run(str(value))
            set_run_font(run, BODY_FONT, 10.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return template_sha256


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--proposal-path", required=True, type=Path)
    parser.add_argument("--output-path", required=True, type=Path)
    parser.add_argument("--template-path", required=True, type=Path)
    args = parser.parse_args()
    if args.output_path.suffix.lower() != ".docx":
        parser.error("--output-path must end in .docx")
    if args.output_path.resolve() in (args.proposal_path.resolve(), args.template_path.resolve()):
        parser.error("output must not overwrite input or layout authority")
    try:
        render_proposal(args.proposal_path, args.output_path, args.template_path)
    except (ValueError, OSError, KeyError, TypeError) as exc:
        parser.exit(1, f"proposal rendering failed: {type(exc).__name__}: {exc}\n")
    print(json.dumps({"status": "rendered", "bytes": args.output_path.stat().st_size}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
