#!/usr/bin/env python3
"""Validate the JiaLiang Brain Science case style in a DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


SKILL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = SKILL_DIR / "assets" / "primary-layout-authority.docx"
EXPECTED_REFERENCE_SHA256 = "0686dc7cd3bc3f098f6d046239c84ae885e1df03bf8719057e8688a14ec90385"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def child(parent, name):
    return None if parent is None else parent.find(qn(f"w:{name}"))


def attr(element, name="val"):
    return None if element is None else element.get(qn(f"w:{name}"))


def ppr(paragraph, name):
    return child(child(paragraph._p, "pPr"), name)


def first_text_run(paragraph):
    return next((run for run in paragraph.runs if run.text.strip()), None)


def run_prop(run, name):
    return None if run is None else child(child(run._element, "rPr"), name)


def check_tables(doc, errors):
    metrics = []
    for table_index, table in enumerate(doc.tables, 1):
        tbl_pr = child(table._tbl, "tblPr")
        borders = child(tbl_pr, "tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = child(borders, edge)
            if attr(border) != "single" or attr(border, "sz") != "4":
                errors.append(f"TABLE:{table_index}:BORDER:{edge}:{attr(border)}:{attr(border, 'sz')}")
        header_fills = []
        for cell_index, cell in enumerate(table.rows[0].cells, 1):
            fill = attr(child(child(cell._tc, "tcPr"), "shd"), "fill")
            header_fills.append(fill)
            if fill != "C0C0C0":
                errors.append(f"TABLE:{table_index}:HEADER_FILL:{cell_index}:{fill}")
        if child(child(table.rows[0]._tr, "trPr"), "tblHeader") is None:
            errors.append(f"TABLE:{table_index}:REPEATING_HEADER")
        metrics.append({"table": table_index, "rows": len(table.rows), "columns": len(table.columns), "header_fills": header_fills})
    return metrics


def validate(target: Path, reference: Path) -> dict[str, object]:
    errors: list[str] = []
    reference_sha = hashlib.sha256(reference.read_bytes()).hexdigest()
    if reference_sha != EXPECTED_REFERENCE_SHA256:
        errors.append("REFERENCE_SHA256_MISMATCH")

    with zipfile.ZipFile(target) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    refs = root.xpath("//w:sectPr/w:headerReference | //w:sectPr/w:footerReference", namespaces=NS)
    ref_kinds = sorted(etree.QName(node).localname for node in refs)
    if ref_kinds != ["footerReference", "headerReference"]:
        errors.append(f"HEADER_FOOTER_REFERENCES:{ref_kinds}")

    doc = Document(target)
    if len(doc.sections) != 1:
        errors.append(f"SECTION_COUNT:{len(doc.sections)}")
    section = doc.sections[0]
    expected_geometry = {
        "page_width": 11906 * 635,
        "page_height": 16838 * 635,
        "top_margin": 1440 * 635,
        "right_margin": 1797 * 635,
        "bottom_margin": 1497 * 635,
        "left_margin": 1797 * 635,
        "header_distance": 851 * 635,
        "footer_distance": 992 * 635,
    }
    geometry = {}
    for name, expected in expected_geometry.items():
        actual = int(getattr(section, name))
        geometry[name] = actual
        if actual != expected:
            errors.append(f"SECTION_GEOMETRY:{name}:{actual}:EXPECTED:{expected}")

    paragraphs = [paragraph for paragraph in doc.paragraphs if paragraph.text.strip()]
    if not paragraphs:
        return {"passed": False, "errors": ["DOCUMENT_EMPTY"]}

    title = paragraphs[0]
    title_run = first_text_run(title)
    if title.style is None or title.style.name != "星实-题目":
        errors.append("TITLE_STYLE")
    if attr(run_prop(title_run, "sz")) != "32":
        errors.append(f"TITLE_SIZE:{attr(run_prop(title_run, 'sz'))}")
    if run_prop(title_run, "b") is None:
        errors.append("TITLE_NOT_BOLD")
    title_spacing = ppr(title, "spacing")
    if attr(title_spacing, "before") != "0" or attr(title_spacing, "after") != "120":
        errors.append(f"TITLE_SPACING:{attr(title_spacing, 'before')}:{attr(title_spacing, 'after')}")

    counts = {"salutation": 0, "body": 0, "h1": 0, "h2": 0}
    for index, paragraph in enumerate(paragraphs[1:], 2):
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        indent = ppr(paragraph, "ind")
        run = first_text_run(paragraph)
        if "投资决策委员会成员" in text:
            counts["salutation"] += 1
            if style != "星实-正文":
                errors.append(f"SALUTATION:{index}:STYLE:{style}")
            if attr(indent, "left") not in {None, "0"} or attr(indent, "firstLine") != "0" or attr(indent, "firstLineChars") != "0":
                errors.append(f"SALUTATION:{index}:INDENT:{attr(indent, 'left')}:{attr(indent, 'firstLine')}:{attr(indent, 'firstLineChars')}")
            if attr(run_prop(run, "sz")) != "24" or run_prop(run, "b") is None:
                errors.append(f"SALUTATION:{index}:RUN")
        elif style == "星实-正文":
            counts["body"] += 1
            if attr(indent, "firstLine") != "480":
                errors.append(f"BODY:{index}:INDENT:{attr(indent, 'firstLine')}")
            for body_run in paragraph.runs:
                if body_run.text.strip() and attr(run_prop(body_run, "sz")) != "24":
                    errors.append(f"BODY:{index}:SIZE:{attr(run_prop(body_run, 'sz'))}")
        elif style == "星实-一标":
            counts["h1"] += 1
            spacing = ppr(paragraph, "spacing")
            if attr(spacing, "before") != "240" or attr(spacing, "after") != "0":
                errors.append(f"H1:{index}:SPACING:{attr(spacing, 'before')}:{attr(spacing, 'after')}")
        elif style == "星实-二标":
            counts["h2"] += 1

    if counts["salutation"] != 1:
        errors.append(f"SALUTATION_COUNT:{counts['salutation']}")
    if not 3 <= counts["h1"] <= 6:
        errors.append(f"H1_COUNT:{counts['h1']}")
    if counts["h2"] == 0:
        errors.append("H2_COUNT:0")

    for role, paragraph in zip(("SIGNOFF", "DATE"), paragraphs[-2:]):
        indent = ppr(paragraph, "ind")
        run = first_text_run(paragraph)
        if attr(ppr(paragraph, "jc")) != "right":
            errors.append(f"{role}:ALIGNMENT:{attr(ppr(paragraph, 'jc'))}")
        if attr(indent, "left") != "0" or attr(indent, "firstLine") != "0":
            errors.append(f"{role}:INDENT:{attr(indent, 'left')}:{attr(indent, 'firstLine')}")
        if attr(run_prop(run, "sz")) != "24" or attr(run_prop(run, "szCs")) != "24":
            errors.append(f"{role}:SIZE:{attr(run_prop(run, 'sz'))}:{attr(run_prop(run, 'szCs'))}")

    tables = check_tables(doc, errors)
    return {
        "target": str(target.resolve()),
        "reference": str(reference.resolve()),
        "reference_sha256": reference_sha,
        "passed": not errors,
        "error_count": len(errors),
        "errors": errors,
        "geometry": geometry,
        "paragraph_counts": counts,
        "tables": tables,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()
    report = validate(args.docx.resolve(), args.reference.resolve())
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    print(rendered)
    if args.report:
        args.report.write_text(rendered, encoding="utf-8")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
