#!/usr/bin/env python3
"""Shared deterministic DOCX helpers for the retained Deta V5 visual system."""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_LEFT_CM = 2.20
MARGIN_RIGHT_CM = 2.00
MARGIN_TOP_CM = 2.00
MARGIN_BOTTOM_CM = 1.80
HEADER_DISTANCE_CM = 0.75
FOOTER_DISTANCE_CM = 0.85
TEXT_WIDTH_CM = 16.80
TEXT_WIDTH_DXA = 9524
LANDSCAPE_TEXT_WIDTH_CM = 25.90
LANDSCAPE_TEXT_WIDTH_DXA = 14683
HEADER_NOTICE = "申明：本报告为内部项目文件，禁止外传，报告内容仅代表尽调机构观点"


def _set_font(element, name: str) -> None:
    rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_run_font(run, name: str, size_pt: float, bold: bool | None = None) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_font(run._element, name)


def set_style_font(style, name: str, size_pt: float, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    _set_font(style.element, name)


def set_outline_level(style, level: int) -> None:
    ppr = style.element.get_or_add_pPr()
    existing = ppr.find(qn("w:outlineLvl"))
    if existing is None:
        existing = OxmlElement("w:outlineLvl")
        ppr.append(existing)
    existing.set(qn("w:val"), str(level))


def get_or_add_style(doc: Document, name: str, base: str = "Normal"):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        try:
            style.base_style = doc.styles[base]
        except KeyError:
            pass
    return style


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, "仿宋_GB2312", 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    cover = get_or_add_style(doc, "尽调-封面主体")
    set_style_font(cover, "黑体", 22, True)
    cover.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(0)
    cover.paragraph_format.space_after = Pt(0)

    cover_meta = get_or_add_style(doc, "尽调-封面日期")
    set_style_font(cover_meta, "仿宋_GB2312", 14, False)
    cover_meta.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_meta.paragraph_format.space_before = Pt(0)
    cover_meta.paragraph_format.space_after = Pt(0)

    toc_title = get_or_add_style(doc, "尽调-目录标题")
    set_style_font(toc_title, "黑体", 18, True)
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.line_spacing = 1.5
    toc_title.paragraph_format.keep_with_next = True

    body = get_or_add_style(doc, "尽调-正文")
    set_style_font(body, "仿宋_GB2312", 12)
    pf = body.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.first_line_indent = Cm(0.847)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True

    heading_specs = [
        ("尽调-一级标题", "黑体", 15, 0, 12, 6, True),
        ("尽调-二级标题", "黑体", 15, 1, 8, 4, False),
        ("尽调-三级标题", "黑体", 12, 2, 6, 2, False),
        ("尽调-四级标题", "黑体", 12, 3, 4, 2, False),
    ]
    for name, font, size, outline, before, after, page_break in heading_specs:
        style = get_or_add_style(doc, name)
        set_style_font(style, font, size, True)
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.25
        pf.left_indent = Cm(0)
        pf.first_line_indent = Cm(0)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.keep_together = True
        pf.page_break_before = page_break
        set_outline_level(style, outline)

    table_title = get_or_add_style(doc, "尽调-表标题")
    set_style_font(table_title, "仿宋_GB2312", 12, True)
    table_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_title.paragraph_format.keep_with_next = True
    table_title.paragraph_format.space_before = Pt(0)
    table_title.paragraph_format.space_after = Pt(0)

    table_body = get_or_add_style(doc, "尽调-表正文")
    set_style_font(table_body, "仿宋_GB2312", 12)
    table_body.paragraph_format.space_before = Pt(0)
    table_body.paragraph_format.space_after = Pt(0)
    table_body.paragraph_format.line_spacing = 1.25

    caption = get_or_add_style(doc, "尽调-图注")
    set_style_font(caption, "仿宋_GB2312", 10.5)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(0)

    callout = get_or_add_style(doc, "尽调-提示")
    set_style_font(callout, "仿宋_GB2312", 12)
    callout.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    callout.paragraph_format.line_spacing = 1.5
    callout.paragraph_format.first_line_indent = Cm(0.847)
    callout.paragraph_format.left_indent = Cm(0)
    callout.paragraph_format.right_indent = Cm(0)
    callout.paragraph_format.space_before = Pt(0)
    callout.paragraph_format.space_after = Pt(0)

    for list_name in ("List Bullet", "List Number"):
        try:
            style = doc.styles[list_name]
        except KeyError:
            continue
        set_style_font(style, "仿宋_GB2312", 12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)


def configure_section(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.left_margin = Cm(MARGIN_LEFT_CM)
    section.right_margin = Cm(MARGIN_RIGHT_CM)
    section.top_margin = Cm(MARGIN_TOP_CM)
    section.bottom_margin = Cm(MARGIN_BOTTOM_CM)
    section.header_distance = Cm(HEADER_DISTANCE_CM)
    section.footer_distance = Cm(FOOTER_DISTANCE_CM)


def set_paragraph_bottom_border(paragraph, size_eighths: int = 6, color: str = "000000") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def add_header_notice(section) -> None:
    section.different_first_page_header_footer = True
    for header in (section.header, section.first_page_header):
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.text = ""
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(HEADER_NOTICE)
        set_run_font(run, "仿宋_GB2312", 9)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, "仿宋_GB2312", 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE \\* MERGEFORMAT "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)


def set_page_number_start(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect_pr.append(pg)
    pg.set(qn("w:start"), str(start))


def clear_footer(section) -> None:
    for p in section.footer.paragraphs:
        p.text = ""


def configure_footer_page_number(section) -> None:
    section.different_first_page_header_footer = True
    for footer in (section.footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = ""
        add_page_field(p)


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在打开文档后自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def set_core_properties(doc: Document) -> None:
    cp = doc.core_properties
    cp.title = "专业投资尽职调查报告"
    cp.subject = "Investment Due Diligence"
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.keywords = ""


def create_sanitized_template(output_path: str | Path) -> Path:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_header_notice(doc.sections[0])
    configure_footer_page_number(doc.sections[0])
    set_page_number_start(doc.sections[0], 1)

    for _ in range(6):
        doc.add_paragraph(style="尽调-正文")
    doc.add_paragraph("[[PROJECT_NAME]]", style="尽调-封面主体")
    for _ in range(1):
        doc.add_paragraph(style="尽调-正文")
    doc.add_paragraph("[[REPORT_TITLE]]", style="尽调-封面主体")
    for _ in range(8):
        doc.add_paragraph(style="尽调-正文")
    doc.add_paragraph("[[AUTHOR]]", style="尽调-封面日期")
    doc.add_paragraph("[[REPORT_DATE]]", style="尽调-封面日期")

    toc_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(toc_section)
    inherited_pg = toc_section._sectPr.find(qn("w:pgNumType"))
    if inherited_pg is not None:
        toc_section._sectPr.remove(inherited_pg)
    add_header_notice(toc_section)
    configure_footer_page_number(toc_section)
    doc.add_paragraph("目录", style="尽调-目录标题")
    toc_p = doc.add_paragraph(style="尽调-正文")
    toc_p.paragraph_format.first_line_indent = Cm(0)
    add_toc_field(toc_p)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    inherited_pg = body_section._sectPr.find(qn("w:pgNumType"))
    if inherited_pg is not None:
        body_section._sectPr.remove(inherited_pg)
    add_header_notice(body_section)
    configure_footer_page_number(body_section)
    doc.add_paragraph("[[REPORT_BODY]]", style="尽调-正文")

    set_update_fields(doc)
    set_core_properties(doc)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def replace_marker(paragraph, text: str) -> None:
    paragraph.text = ""
    run = paragraph.add_run(text)
    # The paragraph style carries the final font; direct setting avoids renderer drift.
    if paragraph.style.name == "尽调-封面主体":
        set_run_font(run, "黑体", 22, True)
    elif paragraph.style.name == "尽调-封面日期":
        set_run_font(run, "仿宋_GB2312", 14, False)


def insert_stacked_title(doc: Document, marker, title: str) -> None:
    p = doc.add_paragraph(title, style="尽调-封面主体")
    marker._p.addprevious(p._p)
    marker._p.getparent().remove(marker._p)


def set_cell_margins(cell, top: int = 0, start: int = 108, bottom: int = 0, end: int = 108) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, size_eighths: int = 4, color: str = "000000") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size_eighths))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row) -> None:
    """Keep a logical table row together across Word page boundaries."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_table_geometry(table, ratios: Sequence[float], total_width_dxa: int = TEXT_WIDTH_DXA) -> list[int]:
    if not ratios or any(x <= 0 for x in ratios):
        raise ValueError("column_widths must contain positive ratios")
    total = sum(ratios)
    widths = [round(total_width_dxa * x / total) for x in ratios]
    widths[-1] += total_width_dxa - sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(total_width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
    return widths


def write_cell(cell, text: str, bold: bool = False, alignment: str = "left") -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.style = "尽调-表正文"
    p.paragraph_format.first_line_indent = Cm(0)
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
    run = p.add_run(str(text))
    set_run_font(run, "仿宋_GB2312", 12, bold)
    set_cell_margins(cell, top=72, start=108, bottom=72, end=108)


def add_table(
    doc: Document,
    marker,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    ratios: Sequence[float] | None = None,
    alignments: Sequence[str] | None = None,
    total_width_dxa: int = TEXT_WIDTH_DXA,
):
    if not headers:
        raise ValueError("table headers cannot be empty")
    cols = len(headers)
    ratios = list(ratios or [1 / cols] * cols)
    if len(ratios) != cols:
        raise ValueError("column_widths length must match headers")
    alignments = list(alignments or ["left"] * cols)
    if len(alignments) != cols:
        raise ValueError("alignments length must match headers")
    table = doc.add_table(rows=1, cols=cols)
    set_table_geometry(table, ratios, total_width_dxa)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for idx, value in enumerate(headers):
        set_cell_shading(header.cells[idx], "D9D9D9")
        write_cell(header.cells[idx], value, bold=True, alignment=alignments[idx])
        # Keep the header row with the first data row.  Word otherwise may
        # leave a repeated table header alone at the bottom of a page.
        for paragraph in header.cells[idx].paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    for row_data in rows:
        if len(row_data) != cols:
            raise ValueError("table row length must match headers")
        row = table.add_row()
        set_row_cant_split(row)
        for idx, value in enumerate(row_data):
            write_cell(row.cells[idx], str(value), bold=False, alignment=alignments[idx])
    set_table_geometry(table, ratios, total_width_dxa)
    marker._p.addprevious(table._tbl)
    return table


def add_key_value_table(
    doc: Document,
    marker,
    rows: Sequence[Sequence[str]],
    ratios: Sequence[float] | None = None,
    total_width_dxa: int = TEXT_WIDTH_DXA,
):
    """Add the Deta-style two-column grey-label summary table."""
    if not rows:
        raise ValueError("key-value table rows cannot be empty")
    ratios = list(ratios or [0.29, 0.71])
    table = doc.add_table(rows=0, cols=2)
    for row_data in rows:
        if len(row_data) != 2:
            raise ValueError("key-value table rows must contain label and value")
        row = table.add_row()
        set_row_cant_split(row)
        set_cell_shading(row.cells[0], "F2F2F2")
        write_cell(row.cells[0], str(row_data[0]), bold=True, alignment="center")
        write_cell(row.cells[1], str(row_data[1]), bold=False, alignment="center")
    set_table_geometry(table, ratios, total_width_dxa)
    set_table_borders(table)
    marker._p.addprevious(table._tbl)
    return table


def add_image(doc: Document, marker, path: Path, width_cm: float, caption: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(min(width_cm, TEXT_WIDTH_CM)))
    marker._p.addprevious(p._p)
    if caption:
        c = doc.add_paragraph(caption, style="尽调-图注")
        marker._p.addprevious(c._p)


def add_text_paragraph(doc: Document, marker, text: str, style: str = "尽调-正文", bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        set_run_font(r, "仿宋_GB2312", 12, True)
    r = p.add_run(text)
    set_run_font(r, "仿宋_GB2312", 12)
    marker._p.addprevious(p._p)
    return p


def find_paragraph(doc: Document, exact_text: str):
    for p in doc.paragraphs:
        if p.text.strip() == exact_text:
            return p
    raise ValueError(f"marker not found: {exact_text}")


CHINESE_NUMERALS = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]


def chinese_number(n: int) -> str:
    if n <= 10:
        return CHINESE_NUMERALS[n]
    if n < 20:
        return "十" + CHINESE_NUMERALS[n - 10]
    tens, ones = divmod(n, 10)
    return CHINESE_NUMERALS[tens] + "十" + (CHINESE_NUMERALS[ones] if ones else "")


def numbered_heading(title: str, level: int, counters: list[int]) -> str:
    if level < 1 or level > 4:
        raise ValueError("heading level must be 1-4")
    counters[level - 1] += 1
    for idx in range(level, 4):
        counters[idx] = 0
    if level == 1:
        return f"{counters[0]}、{title}"
    if level == 2:
        return f"{counters[0]}.{counters[1]} {title}"
    if level == 3:
        return f"{counters[0]}.{counters[1]}.{counters[2]} {title}"
    return f"{counters[0]}.{counters[1]}.{counters[2]}.{counters[3]} {title}"
