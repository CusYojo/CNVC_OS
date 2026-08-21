#!/usr/bin/env python3
"""Run structural checks on a PPTX. Visual QA is still mandatory."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


PLACEHOLDER = re.compile(r"(?:TODO|TBD|lorem\s+ipsum|待填写|占位符|XX公司|请输入)", re.I)


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape) -> str:
    if getattr(shape, "has_text_frame", False):
        return shape.text_frame.text or ""
    if getattr(shape, "has_table", False):
        return "\n".join(cell.text for row in shape.table.rows for cell in row.cells)
    return ""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--expected-slides", type=int)
    parser.add_argument("--min-body-pt", type=float, default=12.0)
    parser.add_argument("--ignore-bottom-fraction", type=float, default=0.10)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    path = args.pptx.resolve()
    errors: list[dict] = []
    warnings: list[dict] = []
    font_counts: Counter[str] = Counter()
    size_counts: Counter[float] = Counter()

    if not path.exists():
        raise SystemExit(f"PPTX not found: {path}")

    try:
        with zipfile.ZipFile(path) as package:
            bad = package.testzip()
            if bad:
                errors.append({"type": "zip_corruption", "part": bad})
    except zipfile.BadZipFile:
        raise SystemExit(f"Not a valid PPTX ZIP: {path}")

    prs = Presentation(path)
    slide_count = len(prs.slides)
    if args.expected_slides is not None and slide_count != args.expected_slides:
        errors.append(
            {"type": "slide_count", "expected": args.expected_slides, "actual": slide_count}
        )

    width, height = prs.slide_width, prs.slide_height
    tolerance = int(0.02 * 914400)

    for slide_no, slide in enumerate(prs.slides, 1):
        for index, shape in enumerate(iter_shapes(slide.shapes), 1):
            name = getattr(shape, "name", f"shape-{index}")
            left, top = getattr(shape, "left", 0), getattr(shape, "top", 0)
            right = left + getattr(shape, "width", 0)
            bottom = top + getattr(shape, "height", 0)
            if left < -tolerance or top < -tolerance or right > width + tolerance or bottom > height + tolerance:
                warnings.append(
                    {
                        "type": "out_of_bounds",
                        "slide": slide_no,
                        "shape": name,
                        "box_in": [round(v / 914400, 3) for v in (left, top, right, bottom)],
                    }
                )

            text = shape_text(shape)
            if text and PLACEHOLDER.search(text):
                errors.append({"type": "placeholder", "slide": slide_no, "shape": name, "text": text[:160]})

            if getattr(shape, "has_text_frame", False):
                in_footer = top >= height * (1 - args.ignore_bottom_fraction)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            font_counts[run.font.name] += 1
                        if run.font.size:
                            pt = round(run.font.size.pt, 2)
                            size_counts[pt] += 1
                            if not in_footer and 11.0 <= pt < args.min_body_pt:
                                warnings.append(
                                    {
                                        "type": "small_body_font",
                                        "slide": slide_no,
                                        "shape": name,
                                        "size_pt": pt,
                                        "text": run.text[:80],
                                    }
                                )

    report = {
        "pptx": str(path),
        "slides": slide_count,
        "slide_size_in": [round(width / 914400, 3), round(height / 914400, 3)],
        "errors": errors,
        "warnings": warnings,
        "fonts": dict(font_counts.most_common()),
        "font_sizes_pt": {str(k): v for k, v in sorted(size_counts.items())},
        "note": "This audit cannot detect visual collisions, weak hierarchy, image relevance, or rendering differences.",
    }

    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered, encoding="utf-8")
    print(rendered)
    raise SystemExit(2 if errors else 0)


if __name__ == "__main__":
    main()
