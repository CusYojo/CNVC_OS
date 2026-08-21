#!/usr/bin/env python3
"""Check that required facts/metrics from a baseline document remain in the deck."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from pptx import Presentation


def normalize(value: str) -> str:
    return re.sub(r"[\s,，。；;：:（）()]+", "", value).lower()


def extract_text(path: Path) -> str:
    prs = Presentation(path)
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(parts)


def parse_required(path: Path) -> list[tuple[str, list[str]]]:
    items: list[tuple[str, list[str]]] = []
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        if "\t" in line:
            label, probes = line.split("\t", 1)
        else:
            label, probes = line, line
        items.append((label.strip(), [p.strip() for p in probes.split("||") if p.strip()]))
    return items


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--required", required=True, type=Path, help="One item per line; optional TAB and || aliases")
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    deck = normalize(extract_text(args.pptx.resolve()))
    required = parse_required(args.required.resolve())
    present, missing = [], []
    for label, probes in required:
        matched = next((probe for probe in probes if normalize(probe) in deck), None)
        (present if matched else missing).append({"label": label, "matched": matched, "probes": probes})

    report = {
        "pptx": str(args.pptx.resolve()),
        "required_count": len(required),
        "present_count": len(present),
        "missing_count": len(missing),
        "missing": missing,
        "present": present,
    }
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(rendered, encoding="utf-8")
    print(rendered)
    raise SystemExit(3 if missing else 0)


if __name__ == "__main__":
    main()
