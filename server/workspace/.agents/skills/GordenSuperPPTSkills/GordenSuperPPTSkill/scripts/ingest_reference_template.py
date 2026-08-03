#!/usr/bin/env python3
"""Ingest a user-provided PPT reference into a task-scoped analysis bundle.

The script performs deterministic intake only: source preservation, rendering,
text extraction, contact-sheet generation, palette sampling, page-role hints,
and manifests. Semantic/visual conclusions must still be reviewed by the agent.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from PIL import Image, ImageDraw, ImageOps
except ImportError as exc:  # pragma: no cover - environment failure
    raise SystemExit("Pillow is required: python3 -m pip install pillow") from exc


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
REUSE_LEVELS = ("strict-template", "structure-and-style", "style-only", "hybrid-reference")
SCOPES = ("task-only", "persistent-candidate")

ROLE_PATTERNS: list[tuple[str, tuple[str, ...]]] = [
    ("summary", ("项目摘要", "公司摘要", "executive summary", "概览")),
    ("investment-highlights", ("投资亮点", "核心亮点", "investment highlights")),
    ("competition", ("竞争", "竞品", "对标", "competition", "competitive")),
    ("problem", ("痛点", "问题", "供需", "需求", "problem")),
    ("company", ("公司能力", "公司介绍", "发展历程", "股权结构", "company")),
    ("moat", ("壁垒", "护城河", "竞争优势", "moat")),
    ("ecosystem", ("战略绑定", "产业生态", "产业园", "研究院", "ecosystem")),
    ("market", ("行业", "市场", "赛道", "趋势", "政策", "market")),
    ("product", ("产品", "解决方案", "业务架构", "product", "solution")),
    ("technology", ("技术", "架构", "模型", "算法", "technology", "architecture")),
    ("team", ("团队", "创始人", "科学家", "高管", "ceo", "cto", "team")),
    ("business-model", ("商业模式", "商业与投资", "收入结构", "商业化", "business model")),
    ("financials", ("财务", "盈利预测", "收入预测", "营收", "financial", "forecast")),
    ("investment-plan", ("投资方案", "融资方案", "交易方案", "investment plan", "fundraising")),
    ("validation", ("落地", "验证", "案例", "客户访谈", "场景", "合作", "poc", "validation", "case")),
    ("risk", ("风险", "risk")),
    ("exit", ("退出", "回报", "exit", "return")),
    ("closing", ("感谢聆听", "谢谢", "thank you", "thanks")),
]

GENERIC_TERMS = {
    "公司", "项目", "投资", "报告", "行业", "市场", "产品", "技术", "团队", "情况", "核心",
    "业务", "发展", "能力", "智能", "方案", "分析", "数据", "客户", "合作", "中国", "未来",
    "company", "project", "investment", "report", "market", "product", "technology", "team",
}


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Render and index a user-provided PDF/PPTX/image template for GordenSuperPPTSkill."
    )
    parser.add_argument("input", help="PDF, PPTX, image, or directory containing ordered page images.")
    parser.add_argument("--out-dir", required=True, help="New or empty output directory for the intake bundle.")
    parser.add_argument("--template-name", help="Human-readable template name. Defaults to the input stem.")
    parser.add_argument("--scope", choices=SCOPES, default="task-only")
    parser.add_argument("--reuse-level", choices=REUSE_LEVELS, default="structure-and-style")
    parser.add_argument("--dpi", type=int, default=120, help="PDF/PPTX rendering DPI. Default: 120.")
    parser.add_argument("--contact-cols", type=int, default=4, help="Contact-sheet columns. Default: 4.")
    args = parser.parse_args(argv)
    if args.dpi < 72 or args.dpi > 300:
        parser.error("--dpi must be between 72 and 300")
    if args.contact_cols < 1 or args.contact_cols > 8:
        parser.error("--contact-cols must be between 1 and 8")
    return args


def write_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def natural_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", path.name)]


def ensure_output_dir(path: Path) -> None:
    if path.exists() and any(path.iterdir()):
        raise RuntimeError(f"Output directory must be new or empty: {path}")
    path.mkdir(parents=True, exist_ok=True)


def classify_input(path: Path) -> str:
    if path.is_dir():
        return "image-directory"
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".pptx":
        return "pptx"
    if suffix in IMAGE_SUFFIXES:
        return "image"
    raise ValueError("Supported inputs: PDF, PPTX, PNG/JPEG/WebP/TIFF/BMP, or an image directory.")


def find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    return None


def convert_pptx_to_pdf(source: Path, source_dir: Path) -> tuple[Path, str]:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError(
            "PPTX rendering requires LibreOffice/soffice. Export the deck to PDF or install LibreOffice."
        )
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(source_dir), str(source)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    produced = source_dir / f"{source.stem}.pdf"
    if not produced.exists():
        matches = sorted(source_dir.glob("*.pdf"), key=natural_key)
        if len(matches) != 1:
            raise RuntimeError("LibreOffice did not produce a uniquely identifiable PDF.")
        produced = matches[0]
    stable = source_dir / "rendered-source.pdf"
    if produced != stable:
        produced.replace(stable)
    return stable, f"soffice:{soffice}"


def render_pdf(pdf_path: Path, rendered_dir: Path, dpi: int) -> tuple[list[Path], str]:
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("PDF rendering requires pdftoppm (Poppler).")
    prefix = rendered_dir / "raw-page"
    subprocess.run(
        [pdftoppm, "-png", "-r", str(dpi), str(pdf_path), str(prefix)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    raw_pages = sorted(rendered_dir.glob("raw-page-*.png"), key=natural_key)
    if not raw_pages:
        raise RuntimeError("pdftoppm produced no page images.")
    pages: list[Path] = []
    for index, raw in enumerate(raw_pages, start=1):
        stable = rendered_dir / f"page-{index:03d}.png"
        raw.replace(stable)
        pages.append(stable)
    return pages, f"pdftoppm:{pdftoppm}"


def ingest_images(source: Path, kind: str, rendered_dir: Path) -> tuple[list[Path], str]:
    candidates = [source] if kind == "image" else [p for p in source.iterdir() if p.suffix.lower() in IMAGE_SUFFIXES]
    candidates = sorted(candidates, key=natural_key)
    if not candidates:
        raise RuntimeError("No supported page images found.")
    pages: list[Path] = []
    for index, candidate in enumerate(candidates, start=1):
        with Image.open(candidate) as opened:
            image = ImageOps.exif_transpose(opened).convert("RGB")
            stable = rendered_dir / f"page-{index:03d}.png"
            image.save(stable, format="PNG", optimize=True)
            pages.append(stable)
    return pages, "pillow"


def extract_pdf_text(pdf_path: Path) -> list[str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    reader = PdfReader(str(pdf_path))
    return [(page.extract_text() or "").strip() for page in reader.pages]


def extract_pptx_text(pptx_path: Path) -> list[str]:
    try:
        from pptx import Presentation
    except ImportError:
        return []
    presentation = Presentation(str(pptx_path))
    output: list[str] = []
    for slide in presentation.slides:
        chunks: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text and text.strip():
                chunks.append(text.strip())
        output.append("\n".join(chunks))
    return output


def normalize_page_texts(texts: list[str], page_count: int) -> list[str]:
    normalized = list(texts[:page_count])
    normalized.extend([""] * (page_count - len(normalized)))
    return normalized


def first_title(text: str) -> str:
    for line in text.splitlines():
        cleaned = re.sub(r"\s+", " ", line).strip(" |•·\t")
        if cleaned:
            return cleaned[:120]
    return ""


def infer_role(page_number: int, page_count: int, title: str, text: str) -> str:
    if page_number == 1:
        return "cover"
    strong_roles = {"summary", "investment-highlights", "investment-plan", "exit"}
    full_normalized = re.sub(r"\s+", "", f"{title}\n{text[:1400]}").lower()
    title_normalized = re.sub(r"\s+", "", title).lower()
    if "客户访谈" in title_normalized:
        return "validation"
    if "风险对策" in title_normalized or title_normalized.startswith("风险控制"):
        return "risk"
    if re.search(r"20\d{2}.*收入", title_normalized):
        return "financials"
    for role, patterns in ROLE_PATTERNS:
        if role in strong_roles and any(
            re.sub(r"\s+", "", pattern).lower() in full_normalized for pattern in patterns
        ):
            return role
    heading_lines = [line.strip() for line in text.splitlines() if line.strip()][:3]
    heading = "\n".join(heading_lines) or title
    for candidate in (title, heading):
        normalized = re.sub(r"\s+", "", candidate).lower()
        for role, patterns in ROLE_PATTERNS:
            if any(re.sub(r"\s+", "", pattern).lower() in normalized for pattern in patterns):
                return role
    if sum(keyword in full_normalized for keyword in ("创始人", "ceo", "cto", "核心团队")) >= 2:
        return "team"
    if "客户访谈" in full_normalized:
        return "validation"
    if page_number == page_count and len(text.strip()) < 160:
        return "closing"
    return "content"


def evenly_spaced_indices(count: int, target: int) -> list[int]:
    if count <= target:
        return list(range(count))
    if target <= 1:
        return [0]
    return sorted({round(i * (count - 1) / (target - 1)) for i in range(target)})


def palette_candidates(page_paths: list[Path]) -> list[dict[str, Any]]:
    totals: Counter[tuple[int, int, int]] = Counter()
    total_pixels = 0
    for index in evenly_spaced_indices(len(page_paths), min(12, len(page_paths))):
        with Image.open(page_paths[index]) as opened:
            image = opened.convert("RGB")
            image.thumbnail((180, 110))
            quantized = image.quantize(colors=10, method=Image.Quantize.MEDIANCUT).convert("RGB")
            colors = quantized.getcolors(maxcolors=quantized.width * quantized.height) or []
            for count, color in colors:
                rounded = tuple(int(round(channel / 16) * 16) for channel in color)
                rounded = tuple(min(255, value) for value in rounded)
                totals[rounded] += count
                total_pixels += count
    results: list[dict[str, Any]] = []
    used: set[tuple[int, int, int]] = set()
    for color, count in totals.most_common(6):
        results.append({
            "hex": "#{:02X}{:02X}{:02X}".format(*color),
            "share": round(count / max(1, total_pixels), 4),
            "kind": "dominant",
        })
        used.add(color)
    chromatic = [
        (color, count) for color, count in totals.most_common()
        if max(color) - min(color) >= 32 and color not in used and min(color) < 240
    ]
    for color, count in chromatic[:6]:
        results.append({
            "hex": "#{:02X}{:02X}{:02X}".format(*color),
            "share": round(count / max(1, total_pixels), 4),
            "kind": "accent-candidate",
        })
    return results


def ratio_label(width: int, height: int) -> str:
    ratio = width / height
    standards = [(16 / 9, "16:9"), (4 / 3, "4:3"), (3 / 2, "3:2"), (16 / 10, "16:10")]
    best_ratio, best_label = min(standards, key=lambda item: abs(item[0] - ratio))
    return best_label if abs(best_ratio - ratio) <= 0.04 else f"{ratio:.3f}:1"


def create_contact_sheet(page_paths: list[Path], output: Path, columns: int) -> None:
    cell_width, image_height, label_height, gap = 420, 236, 26, 12
    rows = math.ceil(len(page_paths) / columns)
    canvas = Image.new(
        "RGB",
        (gap + columns * (cell_width + gap), gap + rows * (image_height + label_height + gap)),
        "#F2F3F5",
    )
    draw = ImageDraw.Draw(canvas)
    for index, path in enumerate(page_paths):
        row, column = divmod(index, columns)
        x = gap + column * (cell_width + gap)
        y = gap + row * (image_height + label_height + gap)
        with Image.open(path) as opened:
            image = ImageOps.contain(opened.convert("RGB"), (cell_width, image_height))
            px = x + (cell_width - image.width) // 2
            py = y + (image_height - image.height) // 2
            canvas.paste(image, (px, py))
        draw.text((x + 6, y + image_height + 4), f"Page {index + 1}", fill="#202124")
    canvas.save(output, format="PNG", optimize=True)


def repeated_terms(texts: Iterable[str], limit: int = 30) -> list[dict[str, Any]]:
    counter: Counter[str] = Counter()
    for text in texts:
        terms = re.findall(r"[\u4e00-\u9fff]{2,12}|[A-Za-z][A-Za-z0-9&.+_-]{2,}", text)
        for term in terms:
            normalized = term.lower().strip()
            if normalized in GENERIC_TERMS or normalized.isdigit():
                continue
            counter[normalized] += 1
    return [{"term": term, "count": count} for term, count in counter.most_common(limit)]


def relative(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    source = Path(args.input).expanduser().resolve()
    out_dir = Path(args.out_dir).expanduser().resolve()
    if not source.exists():
        print(f"Input does not exist: {source}", file=sys.stderr)
        return 2

    try:
        kind = classify_input(source)
        ensure_output_dir(out_dir)
        source_dir = out_dir / "source"
        rendered_dir = out_dir / "rendered"
        source_dir.mkdir()
        rendered_dir.mkdir()

        if source.is_file():
            preserved_source = source_dir / source.name
            shutil.copy2(source, preserved_source)
            source_hash = sha256_file(source)
        else:
            preserved_source = source_dir / "pages"
            preserved_source.mkdir()
            source_files = [
                path for path in sorted(source.iterdir(), key=natural_key)
                if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
            ]
            for path in source_files:
                shutil.copy2(path, preserved_source / path.name)
            hashes = [sha256_file(path) for path in source_files]
            source_hash = hashlib.sha256("".join(hashes).encode("ascii")).hexdigest()

        render_source: Path | None = None
        if kind == "pdf":
            render_source = preserved_source
            page_paths, renderer = render_pdf(preserved_source, rendered_dir, args.dpi)
            texts = extract_pdf_text(preserved_source)
        elif kind == "pptx":
            render_source, converter = convert_pptx_to_pdf(preserved_source, source_dir)
            page_paths, renderer = render_pdf(render_source, rendered_dir, args.dpi)
            renderer = f"{converter};{renderer}"
            texts = extract_pptx_text(preserved_source) or extract_pdf_text(render_source)
        else:
            page_paths, renderer = ingest_images(preserved_source, kind, rendered_dir)
            texts = []

        texts = normalize_page_texts(texts, len(page_paths))
        pages: list[dict[str, Any]] = []
        for index, (path, text) in enumerate(zip(page_paths, texts), start=1):
            with Image.open(path) as image:
                width, height = image.size
            title = first_title(text)
            pages.append({
                "page": index,
                "image": relative(path, out_dir),
                "title_hint": title,
                "role_hint": infer_role(index, len(page_paths), title, text),
                "width": width,
                "height": height,
                "text_chars": len(text),
            })

        role_first_pages: dict[str, int] = {}
        for page in pages:
            role_first_pages.setdefault(page["role_hint"], page["page"])
        selected_pages = sorted(set(role_first_pages.values()))

        first_width, first_height = pages[0]["width"], pages[0]["height"]
        average_chars = round(sum(len(text) for text in texts) / max(1, len(texts)), 1)
        density = "low" if average_chars < 220 else "medium" if average_chars < 650 else "high"
        template_name = args.template_name or source.stem

        contact_sheet = out_dir / "contact-sheet.png"
        create_contact_sheet(page_paths, contact_sheet, args.contact_cols)

        write_json(out_dir / "extracted-text.json", {
            "pages": [{"page": index, "text": text} for index, text in enumerate(texts, start=1)]
        })
        write_json(out_dir / "page-index.json", {"pages": pages})
        write_json(out_dir / "page-archetypes.json", {
            "analysis_status": "role_hints_require_agent_review",
            "selected_pages": selected_pages,
            "roles": role_first_pages,
            "instruction": "Review the contact sheet and replace role hints before slide mapping.",
        })
        write_json(out_dir / "sample-fingerprint.json", {
            "cover_text": texts[0][:1200] if texts else "",
            "frequent_terms": repeated_terms(texts),
            "instruction": "Treat these as possible sample residue; verify entities visually before generation.",
        })
        write_json(out_dir / "template-profile.json", {
            "schema_version": 1,
            "template_name": template_name,
            "analysis_status": "draft_requires_agent_visual_review",
            "scope": args.scope,
            "reuse_level": args.reuse_level,
            "canvas": {
                "width_px": first_width,
                "height_px": first_height,
                "aspect_ratio": ratio_label(first_width, first_height),
            },
            "content_dna": {
                "detected_roles": role_first_pages,
                "average_text_chars_per_page": average_chars,
                "density_hint": density,
                "storyline": [],
            },
            "visual_dna": {
                "palette_candidates": palette_candidates(page_paths),
                "fonts": [],
                "grid": "requires_visual_review",
                "title_system": "requires_visual_review",
                "card_and_chart_language": "requires_visual_review",
                "imagery_style": "requires_visual_review",
            },
            "reuse_policy": {
                "borrow": ["story structure", "page roles", "layout skeleton", "visual rhythm", "color proportions"],
                "replace": ["all text", "logos", "people", "product images", "dates", "numbers", "claims"],
            },
        })

        generated_files = [
            "contact-sheet.png", "extracted-text.json", "page-index.json", "page-archetypes.json",
            "sample-fingerprint.json", "template-profile.json",
        ]
        manifest = {
            "schema_version": 1,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "template_name": template_name,
            "scope": args.scope,
            "reuse_level": args.reuse_level,
            "source": {
                "original_path": str(source),
                "preserved_copy": relative(preserved_source, out_dir) if preserved_source else None,
                "kind": kind,
                "sha256": source_hash,
            },
            "render_source": relative(render_source, out_dir) if render_source and out_dir in render_source.parents else str(render_source) if render_source else None,
            "renderer": renderer,
            "page_count": len(page_paths),
            "generated_files": generated_files,
            "promotion_status": "not-promoted" if args.scope == "task-only" else "awaiting-explicit-approval",
        }
        write_json(out_dir / "template-intake-manifest.json", manifest)

        print(json.dumps({
            "template_name": template_name,
            "out_dir": str(out_dir),
            "page_count": len(page_paths),
            "scope": args.scope,
            "reuse_level": args.reuse_level,
            "contact_sheet": str(contact_sheet),
            "profile": str(out_dir / "template-profile.json"),
            "manifest": str(out_dir / "template-intake-manifest.json"),
        }, ensure_ascii=False, indent=2))
        return 0
    except (OSError, RuntimeError, ValueError, subprocess.CalledProcessError) as exc:
        print(f"Template intake failed: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
