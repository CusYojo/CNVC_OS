#!/usr/bin/env python3
"""Prepare evidence, build a template-faithful DOCX, and verify the result."""

from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any, Iterable
from xml.etree import ElementTree as ET


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
CP_NS = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
DC_NS = "http://purl.org/dc/elements/1.1/"
DCTERMS_NS = "http://purl.org/dc/terms/"
XSI_NS = "http://www.w3.org/2001/XMLSchema-instance"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
NS = {"w": W_NS, "w14": W14_NS}
W = f"{{{W_NS}}}"
W14 = f"{{{W14_NS}}}"
CP = f"{{{CP_NS}}}"
DC = f"{{{DC_NS}}}"
DCTERMS = f"{{{DCTERMS_NS}}}"
XSI = f"{{{XSI_NS}}}"
REL = f"{{{REL_NS}}}"
CT = f"{{{CT_NS}}}"
SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "reference.docx"
PUBLIC_VALIDATOR = SKILL_DIR / "scripts" / "validate_public_verification.py"
REQUIRED_SECTIONS = ["公司情况介绍", "投资理由", "投资计划", "投资情形分析"]
TEMPLATE_RESIDUE = [
    "德塔智能", "北京德塔源创", "马晓健", "刘航欣", "黄思远", "朱松纯",
    "24亿元", "2.7亿元", "26.7亿元", "0.56%", "2026年7月1日",
    "2026年   7   月   1   日",
]
REASON_DEFENSIVE_PHRASES = ["但", "仍需", "取决于", "适宜设置为", "交割前应", "需进一步"]
REASON_TITLE_MIN_CHARS = 15
REASON_TITLE_MAX_CHARS = 34
REASON_OVERBROAD_TITLES = {
    "技术方向与产业需求匹配",
    "产品与技术形成递进组合",
    "核心团队能力与产品路线对应",
    "商业化已进入客户交付验证阶段",
    "产业落地与股权融资基础已经形成",
}
DEFENSIVE_PHRASES = ["仍需", "交割前", "以最终", "以交割前", "不构成无条件", "不能据此作绝对结论"]
ITEM7_CHECKLIST_MARKERS = [
    "纳入交割前核验", "包括", "分别取得专项", "外汇登记", "股权代持",
    "持股平台", "知识产权", "数据合规", "财务规范", "政策返还", "审批",
]
COMPLIANCE_ASSERTION_MARKERS = [
    "符合", "不涉及", "未涉及", "不会导致", "未发现", "无其他", "未从事",
    "不构成", "属于", "未超过", "满足",
]
COMPLIANCE_INDECISIVE_MARKERS = [
    "尚待确认", "待确认", "暂不能", "不能测算", "不能计算", "尚不能",
    "无法核对", "无法判断", "待定", "需专项核查", "无法形成",
]
ANALYSIS_PROCESS_PHRASES = [
    "尚未确定", "待明确", "待补充", "仍需", "尚需", "应核对",
    "核查后，方可确定", "核查后方可确定", "不能测算", "不能计算",
    "无法核对", "无法判断", "不宜作出结论", "不宜作出符合性结论",
    "最终仍应核实", "资料未提供", "资料尚未提供",
]
ANALYSIS_PROCESS_PATTERNS = [
    ("应结合……判断", r"应结合[^。；]{0,80}(?:判断|核对|确认|测算|计算)"),
    ("应按……核对/测算", r"应按[^。；]{0,80}(?:核对|判断|测算|计算|列示)"),
    ("核查后方可形成判断", r"核查后[，, ]*方可(?:确定|判断|形成)"),
    ("不能仅凭……形成结论", r"不能仅凭[^。；]{0,80}(?:作出|形成)[^。；]{0,16}结论"),
    ("未确定前拒绝结论", r"未确定前[^。；]{0,80}(?:不宜|不能|无法)[^。；]{0,24}(?:结论|判断)"),
]
ANALYSIS_ROLE_RULES = [
    ("投资限制事项", ("投资限制", "限制事项", "禁止事项")),
    ("返投义务影响", ("返投",)),
    ("关联交易", ("关联交易", "关联关系")),
    ("投资方向", ("投资方向", "投资领域", "产业方向")),
    ("投资配置", ("投资配置", "配置要求", "投资层级")),
    ("投资集中度", ("投资集中度", "集中度", "单项目上限")),
    ("其他法律监管事项", ("违法违规", "法律法规", "监管规定", "法律监管")),
]
CONCLUSION_FORBIDDEN_MARKERS = [
    "暂无法", "无法形成", "不能形成", "尚不能", "无法判断", "待确认", "待定",
]
CONCLUSION_REQUIRED_MARKER = "原则上符合"
PLAN_RESERVATION_MAX_CHARS = 72
COMPANY_PROFILE_MIN_CHARS = 150
TEAM_MIN_PARAGRAPHS = 4
TEAM_MIN_TOTAL_CHARS = 560
TEAM_MIN_MEMBER_CHARS = 90
PRODUCT_MIN_PARAGRAPHS = 2
PRODUCT_MIN_TOTAL_CHARS = 360
REASON_REQUIRED_ITEMS = 5
REASON_MIN_ITEM_CHARS = 80
REASON_MIN_TOTAL_CHARS = 500
PLAN_MIN_TOTAL_CHARS = 330
ANALYSIS_MIN_ITEM_CHARS = 70
ANALYSIS_MIN_TOTAL_CHARS = 600
READINESS_REQUIRED_COMPONENTS = [
    "fund_agreement",
    "transaction_terms",
    "return_investment",
    "concentration",
    "related_party",
]
FUND_CLAUSE_KEYS = [
    "investment_scope",
    "investment_restrictions",
    "return_investment",
    "concentration",
    "configuration",
]
TRANSACTION_TERM_KEYS = [
    "investment_amount",
    "currency",
    "transaction_form",
    "pre_money_valuation",
    "post_money_valuation",
    "post_investment_ownership",
    "fully_diluted_basis",
]
RETURN_INVESTMENT_KEYS = [
    "as_of_date",
    "denominator",
    "multiplier",
    "completed_amount",
    "proposed_eligible_credit",
    "post_investment_headroom",
]
CONCENTRATION_KEYS = [
    "as_of_date",
    "denominator",
    "limit_ratio",
    "existing_aggregated_exposure",
    "proposed_amount",
    "post_investment_ratio",
    "headroom",
]
COMPANY_PROFILE_OWNERSHIP_MARKERS = [
    "持股", "股权比例", "表决权", "资本表", "控股股东", "实际控制人",
]
COMPANY_PROFILE_REGISTRY_MARKERS = [
    "统一社会信用代码", "营业执照号", "营业执照号码", "注册号", "组织机构代码", "证照编号",
]
COMPANY_PROFILE_ROUTINE_REGISTRY_MARKERS = [
    "法定代表人", "注册资本", "实缴资本", "认缴资本",
]
COMPANY_PROFILE_FINANCIAL_DD_PATTERNS = [
    re.compile(
        r"(?:收入确认|开票|发票|交付|验收)[^。；\n]{0,40}"
        r"(?:差异|时点|跨期|调整|不一致|异常)"
    ),
    re.compile(
        r"(?:差异|跨期|调整|不一致|异常)[^。；\n]{0,40}"
        r"(?:收入确认|开票|发票|交付|验收)"
    ),
    re.compile(
        r"(?:财务|审计|税务|回款|应收)[^。；\n]{0,30}"
        r"(?:问题|风险|异常|调整|保留|差异)"
    ),
]
COMPANY_PROFILE_EXACT_FINANCIAL_METRIC_PATTERN = re.compile(
    r"(?:营业收入|营收|净利润|利润总额|毛利率|毛利|应收账款|"
    r"经营活动现金流量净额|经营性现金流|现金流)[^。；\n]{0,24}?"
    r"[+-]?(?:[0-9０-９]{1,3}(?:[,，][0-9０-９]{3})+|[0-9０-９]+)"
    r"(?:[.．][0-9０-９]+)?\s*(?:%|％|元|万元|亿元)"
)
COMPANY_PROFILE_FINANCIAL_DISCLOSURE_BASES = {"audited", "special_audit"}
DECISION_LAYER_META_MARKERS = [
    "公司资料记载", "公司材料记载", "相关材料记载", "材料显示", "资料显示",
    "财务尽调资料显示", "法律尽调所附履历显示", "访谈材料记载", "履历记载",
    "法律尽调载明", "法律尽调核查到", "现阶段按拟任状态表述", "当前按拟任状态表述",
    "现阶段按拟任及双聘安排表述", "当前按拟任及双聘安排表述",
]
VISIBLE_ATTACHMENT_TITLE_PATTERN = re.compile(r"《(?P<title>[^》\n]{2,80})》")
VISIBLE_SOURCE_PROCESS_PATTERNS = [
    re.compile(
        r"(?:根据|依据|参照|按照)(?:现有|相关|所附|上述)?"
        r"(?:资料|材料|文件|报告|台账|资本表|投资意向书)[^。；\n]{0,20}"
    ),
    re.compile(
        r"(?:资料|材料|文件|报告|台账|资本表|投资意向书)[^。；\n]{0,16}"
        r"(?:显示|记载|载明|所列|披露|表明|约定)"
    ),
    re.compile(r"(?:V|v)\s*\d+(?:\.\d+)?(?:版|版本)?"),
    re.compile(r"(?:/Users/|[A-Za-z]:\\)[^\s，。；]{2,120}"),
    re.compile(r"[^\s，。；]{1,80}\.(?:docx|pdf|xlsx|xls|zip|rar)"),
]
ALLOWED_FORMAL_LAW_TITLE_PATTERNS = [
    re.compile(r"^中华人民共和国.+法$"),
    re.compile(r"^[^》]{1,30}(?:法律|条例|行政法规|司法解释|监管规定)$"),
]
VISIBLE_SOURCE_PROSE_SECTIONS = {"公司情况介绍", "投资理由", "投资计划"}
TEAM_DIMENSION_MARKERS = {
    "role": ["创始", "董事", "监事", "总经理", "负责人", "总监", "科学家", "教授", "管理人员", "技术骨干", "拟任"],
    "education": ["本科", "硕士", "博士", "大学", "学院", "工程师"],
    "experience": ["曾任", "曾在", "任职", "工作经历", "主导", "参与", "负责", "历任"],
    "expertise_or_responsibility": [
        "研究方向", "专业方向", "主要负责", "现负责", "负责公司", "承担公司",
        "人工智能", "自动驾驶", "机器人", "算法", "感知", "数据", "数字孪生",
        "触觉", "项目管理", "商业化", "产品",
    ],
}
TEAM_EVALUATIVE_BENEFIT_PATTERNS = [
    ("使其能够为项目或公司发挥作用", re.compile(
        r"使其能够[^。；\n]{0,50}(?:统筹|支撑|支持|连接|推动|促进|补强|助力|服务)"
    )),
    ("推断可支持或支撑公司", re.compile(
        r"(?:可|能够|将)(?:直接)?(?:支持|支撑|连接|补强|助力|推动|促进)(?:公司|本项目|项目)"
    )),
    ("推断可为公司提供帮助", re.compile(
        r"(?:可|能够|将)?为公司[^。；\n]{0,50}(?:提供|带来)"
        r"(?:支持|助力|价值|资源|赋能|技术指导|学术指导|专业意见|合作资源)"
    )),
    ("推断可为平台产品或业务提供帮助", re.compile(
        r"(?:可|能够|将)(?:直接)?为[^。；\n]{0,70}(?:提供|带来)"
        r"[^。；\n]{0,18}(?:支持|助力|价值|资源|赋能|指导)"
    )),
    ("推断有助于公司或项目", re.compile(
        r"(?:有助于|助力)(?:公司|本项目|项目)"
    )),
    ("推断与公司路线相匹配", re.compile(
        r"与公司[^。；\n]{0,50}(?:产品|业务|技术|研发|发展)?(?:路线|方向|需求)"
        r"[^。；\n]{0,24}(?:对应|匹配|契合)关系"
    )),
]
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".xls", ".doc"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | OFFICE_EXTENSIONS | {".pdf"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_text_flexible(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def decode_zip_name(info: zipfile.ZipInfo) -> str:
    name = info.filename
    if info.flag_bits & 0x800:
        return name
    try:
        raw = name.encode("cp437")
    except UnicodeEncodeError:
        return name
    candidates = [name]
    for enc in ("gb18030", "utf-8", "big5"):
        try:
            candidates.append(raw.decode(enc))
        except UnicodeDecodeError:
            pass

    def score(value: str) -> tuple[int, int, int]:
        cjk = sum("\u3400" <= ch <= "\u9fff" for ch in value)
        mojibake = sum(ch in "��������" for ch in value)
        controls = sum(ord(ch) < 32 and ch not in "\t\n\r" for ch in value)
        return (cjk - 4 * mojibake - 8 * controls, -mojibake, -controls)

    return max(candidates, key=score)


def safe_extract_zip(archive: Path, target: Path) -> list[str]:
    warnings: list[str] = []
    target.mkdir(parents=True, exist_ok=True)
    root = target.resolve()
    with zipfile.ZipFile(archive) as zf:
        for info in zf.infolist():
            decoded = decode_zip_name(info).replace("\\", "/")
            parts = [p for p in PurePosixPath(decoded).parts if p not in ("", ".")]
            if not parts or any(p == ".." for p in parts) or PurePosixPath(decoded).is_absolute():
                warnings.append(f"blocked unsafe path: {decoded}")
                continue
            out = target.joinpath(*parts)
            try:
                out.resolve().relative_to(root)
            except ValueError:
                warnings.append(f"blocked escaping path: {decoded}")
                continue
            unix_mode = (info.external_attr >> 16) & 0o170000
            if unix_mode == 0o120000:
                warnings.append(f"skipped symlink: {decoded}")
                continue
            if info.is_dir() or decoded.endswith("/"):
                out.mkdir(parents=True, exist_ok=True)
                continue
            out.parent.mkdir(parents=True, exist_ok=True)
            with zf.open(info) as src, out.open("wb") as dst:
                shutil.copyfileobj(src, dst)
    return warnings


def xml_paragraph_texts(xml_bytes: bytes) -> list[str]:
    root = ET.fromstring(xml_bytes)
    texts: list[str] = []
    for p in root.findall(".//w:p", NS):
        pieces: list[str] = []
        for node in p.iter():
            if node.tag == W + "t" and node.text:
                pieces.append(node.text)
            elif node.tag == W + "tab":
                pieces.append("\t")
            elif node.tag in (W + "br", W + "cr"):
                pieces.append("\n")
        text = "".join(pieces).strip()
        if text:
            texts.append(text)
    return texts


def extract_docx(path: Path) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as zf:
        names = ["word/document.xml"] + sorted(
            n for n in zf.namelist()
            if re.fullmatch(r"word/(header|footer)\d+\.xml", n)
        )
        for name in names:
            if name in zf.namelist():
                parts.extend(xml_paragraph_texts(zf.read(name)))
    return "\n".join(parts)


def run_text_command(command: list[str], timeout: int = 180) -> str:
    result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                            timeout=timeout, check=False)
    if result.returncode != 0:
        raise RuntimeError(result.stderr.decode("utf-8", errors="replace").strip())
    for enc in ("utf-8", "gb18030"):
        try:
            return result.stdout.decode(enc)
        except UnicodeDecodeError:
            pass
    return result.stdout.decode("utf-8", errors="replace")


def extract_pdf(path: Path) -> str:
    tool = shutil.which("pdftotext")
    if not tool:
        raise RuntimeError("pdftotext unavailable")
    return run_text_command([tool, "-layout", str(path), "-"])


def extract_doc(path: Path) -> str:
    tool = shutil.which("textutil") or shutil.which("antiword")
    if not tool:
        raise RuntimeError("textutil/antiword unavailable")
    if Path(tool).name == "textutil":
        return run_text_command([tool, "-convert", "txt", "-stdout", str(path)])
    return run_text_command([tool, str(path)])


def extract_workbook(path: Path) -> str:
    suffix = path.suffix.lower()
    rows: list[str] = []
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                rows.append(f"## Sheet: {sheet.title}")
                count = 0
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if v is None else str(v).strip() for v in row]
                    if any(values):
                        rows.append("\t".join(values).rstrip())
                        count += len(values)
                    if count > 50000:
                        rows.append("[truncated after 50,000 cells]")
                        break
        finally:
            book.close()
    else:
        try:
            import xlrd
        except ImportError:
            xlrd = None
        if xlrd is not None:
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                for sheet in book.sheets():
                    rows.append(f"## Sheet: {sheet.name}")
                    for r in range(min(sheet.nrows, 5000)):
                        values = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                        if any(values):
                            rows.append("\t".join(values).rstrip())
            finally:
                book.release_resources()
        else:
            soffice = shutil.which("soffice")
            mac_soffice = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
            if not soffice and mac_soffice.exists():
                soffice = str(mac_soffice)
            if not soffice:
                raise RuntimeError("xlrd and LibreOffice are both unavailable for .xls extraction")
            with tempfile.TemporaryDirectory(prefix="compliance-xls-") as tmp:
                tmpdir = Path(tmp)
                profile = tmpdir / "lo-profile"
                command = [soffice, "--headless", f"-env:UserInstallation={profile.as_uri()}",
                           "--convert-to", "xlsx", "--outdir", str(tmpdir), str(path)]
                result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                        timeout=180, check=False)
                converted = tmpdir / f"{path.stem}.xlsx"
                if result.returncode != 0 or not converted.exists():
                    detail = result.stderr.decode("utf-8", errors="replace").strip()
                    raise RuntimeError(f"LibreOffice .xls conversion failed: {detail}")
                return extract_workbook(converted)
    return "\n".join(rows)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return read_text_flexible(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix in (".xlsx", ".xls"):
        return extract_workbook(path)
    if suffix == ".doc":
        return extract_doc(path)
    raise RuntimeError(f"unsupported extension: {suffix}")


def classify_source(path: Path) -> tuple[str, str, int]:
    value = str(path).lower()
    rules = [
        (("增资协议", "股东协议", "股权转让协议", "投资协议"), "transaction", "A", 95),
        (("营业执照", "工商", "公司章程"), "corporate/legal", "A", 92),
        (("审计报告", "专审", "财务尽职"), "financial", "B", 90),
        (("法律尽调", "尽调报告"), "legal", "B", 88),
        (("股东名册", "captable"), "ownership", "B", 86),
        (("客户", "销售", "采购", "合同"), "commercial", "C", 80),
        (("利润表", "资产负债表", "现金流量表", "预算"), "financial", "C", 78),
        (("访谈", "会议纪要", "交流纪要"), "interview", "D", 70),
        (("产品", "技术", "商业计划", "战略"), "business/technology", "E", 65),
        (("员工", "花名册", "竞业", "保密", "财务制度"), "operations", "C", 60),
    ]
    for needles, category, grade, rank in rules:
        if any(n.lower() in value for n in needles):
            return category, grade, rank
    return "other", "E", 30


def make_snippet(text: str, limit: int = 1800) -> str:
    clean = re.sub(r"[ \t]+", " ", text)
    clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
    keywords = ["投资", "估值", "股权", "返投", "关联", "收入", "团队", "技术", "客户"]
    positions = [clean.find(k) for k in keywords if clean.find(k) >= 0]
    start = max(0, min(positions) - 200) if positions else 0
    return clean[start:start + limit]


def prepare_command(args: argparse.Namespace) -> int:
    source = Path(args.input).expanduser().resolve()
    workdir = Path(args.workdir).expanduser().resolve()
    workdir.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []
    if source.is_file() and source.suffix.lower() == ".zip":
        expanded = workdir / "expanded"
        if expanded.exists() and any(expanded.iterdir()):
            raise SystemExit(f"Refusing to reuse non-empty extraction directory: {expanded}")
        warnings.extend(safe_extract_zip(source, expanded))
        root = expanded
        source_hash = sha256(source)
    elif source.is_dir():
        root = source
        source_hash = None
    else:
        raise SystemExit("INPUT must be a ZIP archive or directory")

    packets = workdir / "packets"
    packets.mkdir(exist_ok=True)
    manifest: list[dict[str, Any]] = []
    for index, path in enumerate(sorted(p for p in root.rglob("*") if p.is_file()), 1):
        rel = path.relative_to(root).as_posix()
        category, grade, rank = classify_source(path)
        entry: dict[str, Any] = {
            "source_id": f"S{index:04d}", "path": rel, "size": path.stat().st_size,
            "sha256": sha256(path), "extension": path.suffix.lower(),
            "category": category, "evidence_grade": grade, "rank": rank,
            "text_length": 0, "packet": None, "status": "inventoried",
        }
        if path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                text = extract_text(path)
                text = text.replace("\x00", "").strip()
                entry["text_length"] = len(text)
                if text:
                    packet_path = packets / f"{entry['source_id']}.txt"
                    packet_path.write_text(
                        f"SOURCE_ID: {entry['source_id']}\nPATH: {rel}\nCATEGORY: {category}\n"
                        f"EVIDENCE_GRADE: {grade}\n\n{text}\n", encoding="utf-8")
                    entry["packet"] = packet_path.relative_to(workdir).as_posix()
                    entry["status"] = "extracted"
                else:
                    entry["status"] = "empty_or_scan"
                    warnings.append(f"no extractable text: {rel}")
            except Exception as exc:
                entry["status"] = "extract_failed"
                entry["error"] = str(exc)
                warnings.append(f"extract failed: {rel}: {exc}")
        elif path.suffix.lower() in (".rar", ".7z", ".lnk"):
            entry["status"] = "nested_or_shortcut_not_expanded"
            warnings.append(f"manual review required: {rel}")
        manifest.append(entry)

    manifest.sort(key=lambda item: (-item["rank"], item["path"]))
    write_json(workdir / "source_manifest.json", {
        "created_at": dt.datetime.now().astimezone().isoformat(),
        "input": str(source), "input_sha256": source_hash,
        "root": str(root), "warnings": warnings, "sources": manifest,
    })

    pack_lines = ["# Evidence pack", "", f"Input: `{source}`", "", "## Extraction warnings", ""]
    pack_lines.extend(f"- {w}" for w in warnings[:80])
    if not warnings:
        pack_lines.append("- None")
    pack_lines.extend(["", "## Highest-priority sources", ""])
    for item in manifest[:24]:
        pack_lines.append(f"### {item['source_id']} — {item['path']}")
        pack_lines.append(f"Category: {item['category']} | Grade: {item['evidence_grade']} | Status: {item['status']}")
        if item.get("packet"):
            text = (workdir / item["packet"]).read_text(encoding="utf-8", errors="replace")
            pack_lines.extend(["", make_snippet(text), ""])
        else:
            pack_lines.extend(["", "[No extractable packet]", ""])
    (workdir / "evidence_pack.md").write_text("\n".join(pack_lines), encoding="utf-8")

    prompt = f"""# Agent drafting task

Create a Chinese investment-compliance note from the evidence in `{workdir}`.
Read `source_manifest.json` and relevant packet files, build a conflict-aware fact ledger, then author `content.json` using the skill's content schema.
Use the retained standard only for structure and fund facts explicitly identified as such; never carry over the template project facts.
Required sections: {'、'.join(REQUIRED_SECTIONS)}. The final compliance analysis must contain all seven standard checks and a conditional conclusion.
Before drafting, check for missing fund clauses, transaction terms, return-investment/concentration data, related-party materials, and material company evidence. If decisive inputs are missing, ask the user for them in a blocking final response, set `delivery_readiness.status: awaiting_user_input`, record `supplement_request.outcome: awaiting_response`, and stop the turn before authoring `content.json` or building a DOCX. Resume only after the user supplies the inputs or explicitly instructs you to continue with the available materials. For the latter, use `status: proceed_with_available_materials` and record `continuation_authorization` with `authorized: true`, `basis: explicit_user_instruction`, and the user's instruction; keep the gaps in `open_issues` and generate the same standard DOCX without an internal-preview label. User silence is not authorization. Never invent a clause, amount, ratio, calculation, or completed review. Use `status: blocked` only for an unresolved material conflict, known prohibited/non-compliant condition, or contradictory supplied data.
Keep 公司简介 business-first: after the verified full legal-name opening, cover formation/incubation background, business positioning, products and technology, target customers, operating stage, organization and broad commercial status. Do not put the legal representative, registered/paid-in/subscribed capital, licence identifiers or other routine registry fields in visible 公司简介; retain them in the audit layer and use them elsewhere only when they directly affect a compliance issue. Move revenue-recognition, invoicing/acceptance timing, cut-off, audit-adjustment, collection, tax, and other financial due-diligence issues to notes/open_issues or closing conditions. Do not put exact revenue, profit, margin, receivables, or cash-flow figures in 公司简介 unless the user explicitly requested that disclosure and the figures are supported by audited or special-audit evidence with no material conflict; record that exception in `company_profile_financial_disclosure`. Stage-level wording such as “已形成初步商业化收入” is permitted when supported.
Record the target's verified full legal name in `target_company.legal_name` and begin the first 公司简介 paragraph with that exact name. After the opening identification, use 公司/标的公司 for the target. In 公司情况介绍、投资理由 and 投资计划, state facts directly. Do not expose attachment filenames, dates/versions, paths, or acquisition-process wording; keep them in source_ids, notes and the fact ledger. Use 我方 for the internal investment side, 本基金 for a determined fund, 指定基金主体 for an undetermined vehicle, and 管理人/基金管理人 only when the legal role matters. The full issuer name in `closing.company` is allowed only in the closing signature and must not appear in any visible section.
For every 核心团队 member paragraph, populate `role_title` and `person_name`, then begin the visible text with their exact concatenation in the order `role_title + person_name`. Write the company role/title first and the person's name second, followed by education or professional training, representative employers/projects/research results, professional direction, and a documented current company responsibility when available. Do not use a name-first lead. Do not append an inferred statement about how the person can help, support, connect, strengthen or match the company/project; put any collective team-investment thesis once in 投资理由 instead. Objective wording such as `主要负责公司机器人系统、控制及数采硬件平台` is allowed; evaluative wording such as `其能力可支持公司规模化交付` or `其背景与公司产品路线具有直接对应关系` is not. Examples: `公司联合创始人、CTO刘航欣，……` and `公司拟任首席科学家Abdulmotaleb El Saddik，……`. Do not reproduce promotional hyperbole merely because it appears in a reference sample.
"""
    (workdir / "agent_prompt.md").write_text(prompt, encoding="utf-8")
    print(json.dumps({"workdir": str(workdir), "sources": len(manifest),
                      "warnings": len(warnings), "manifest": str(workdir / "source_manifest.json")},
                     ensure_ascii=False))
    return 0


def paragraph_text(p: ET.Element) -> str:
    return "".join(t.text or "" for t in p.findall(".//w:t", NS))


def find_paragraph(paragraphs: Iterable[ET.Element], exact: str | None = None,
                   startswith: str | None = None) -> ET.Element:
    for p in paragraphs:
        text = paragraph_text(p).strip()
        if exact is not None and text == exact:
            return p
        if startswith is not None and text.startswith(startswith):
            return p
    raise ValueError(f"template exemplar not found: {exact or startswith}")


def clone_rpr(p: ET.Element) -> ET.Element | None:
    first = p.find("w:r", NS)
    if first is None:
        return None
    rpr = first.find("w:rPr", NS)
    return copy.deepcopy(rpr) if rpr is not None else None


def set_bold(rpr: ET.Element | None, bold: bool) -> ET.Element:
    if rpr is None:
        rpr = ET.Element(W + "rPr")
    for name in ("b", "bCs"):
        node = rpr.find(f"w:{name}", NS)
        if bold and node is None:
            ET.SubElement(rpr, W + name)
        elif not bold and node is not None:
            rpr.remove(node)
    return rpr


def add_text_run(p: ET.Element, text: str, rpr: ET.Element | None = None,
                 bold: bool | None = None) -> None:
    r = ET.SubElement(p, W + "r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    if bold is not None:
        current = r.find("w:rPr", NS)
        if current is not None:
            r.remove(current)
        r.insert(0, set_bold(copy.deepcopy(rpr), bold))
    t = ET.SubElement(r, W + "t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text


def split_numbered_lead(text: str) -> tuple[str, str]:
    """Split a Deta numbered item into its bold conclusion lead and reasoning."""
    positions = [position for position in (text.find("："), text.find(":"), text.find("。"))
                 if position >= 0]
    if not positions:
        return text, ""
    boundary = min(positions) + 1
    return text[:boundary], text[boundary:]


def format_closing_date(text: str) -> str:
    """Match the authoritative Deta sample's spaced Chinese-date treatment."""
    match = re.fullmatch(r"\s*(\d{4})年\s*(\d{1,2})月\s*(\d{1,2})日\s*", text)
    if not match:
        return text
    year, month, day = match.groups()
    return f"{year}年   {int(month)}   月   {int(day)}   日"


def metadata_timestamp(content: dict[str, Any]) -> str:
    """Return a stable W3CDTF timestamp tied to the document's stated date."""
    closing_date = str(content.get("closing", {}).get("date", ""))
    match = re.fullmatch(
        r"\s*(\d{4})年\s*(\d{1,2})月\s*(\d{1,2})日\s*", closing_date
    )
    if match:
        year, month, day = (int(value) for value in match.groups())
        return f"{year:04d}-{month:02d}-{day:02d}T00:00:00Z"
    as_of_date = str(content.get("public_verification", {}).get("as_of_date", ""))
    return f"{as_of_date}T00:00:00Z"


def safe_core_properties(content: dict[str, Any]) -> bytes:
    """Create project-specific core properties without template provenance."""
    for prefix, namespace in (
        ("cp", CP_NS), ("dc", DC_NS), ("dcterms", DCTERMS_NS),
        ("xsi", XSI_NS),
    ):
        ET.register_namespace(prefix, namespace)
    root = ET.Element(CP + "coreProperties")
    timestamp = metadata_timestamp(content)
    company = str(content["closing"]["company"])
    title = str(content["title"])
    created = ET.SubElement(root, DCTERMS + "created")
    created.set(XSI + "type", "dcterms:W3CDTF")
    created.text = timestamp
    ET.SubElement(root, DC + "creator").text = company
    ET.SubElement(root, CP + "lastModifiedBy").text = company
    modified = ET.SubElement(root, DCTERMS + "modified")
    modified.set(XSI + "type", "dcterms:W3CDTF")
    modified.text = timestamp
    ET.SubElement(root, DC + "subject").text = title
    ET.SubElement(root, DC + "title").text = title
    ET.SubElement(root, CP + "revision").text = "1"
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def strip_package_metadata_relationships(xml_bytes: bytes) -> bytes:
    """Remove optional extended/custom-property relationships from a DOCX package."""
    ET.register_namespace("", REL_NS)
    root = ET.fromstring(xml_bytes)
    for relationship in list(root):
        relation_type = relationship.get("Type", "")
        target = relationship.get("Target", "")
        if relation_type.endswith(("/extended-properties", "/custom-properties")) or target in {
            "docProps/app.xml", "docProps/custom.xml",
        }:
            root.remove(relationship)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def strip_package_metadata_content_types(xml_bytes: bytes) -> bytes:
    """Remove content-type declarations for deleted metadata parts."""
    ET.register_namespace("", CT_NS)
    root = ET.fromstring(xml_bytes)
    for override in list(root):
        if override.get("PartName") in {"/docProps/app.xml", "/docProps/custom.xml"}:
            root.remove(override)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def metadata_contract_findings(
    docx: Path, content: dict[str, Any]
) -> tuple[list[str], dict[str, Any]]:
    """Validate that generated package metadata is current and privacy-clean."""
    errors: list[str] = []
    metrics: dict[str, Any] = {}
    with zipfile.ZipFile(docx) as archive:
        names = set(archive.namelist())
        forbidden_parts = sorted(names & {"docProps/app.xml", "docProps/custom.xml"})
        if forbidden_parts:
            errors.append(f"forbidden template metadata parts remain: {forbidden_parts}")
        metrics["forbidden_parts_present"] = forbidden_parts

        if "docProps/core.xml" not in names:
            errors.append("missing project-specific core properties")
            metrics["core_properties_status"] = "missing"
        else:
            core = ET.fromstring(archive.read("docProps/core.xml"))
            expected = {
                DC + "title": str(content["title"]),
                DC + "subject": str(content["title"]),
                DC + "creator": str(content["closing"]["company"]),
                CP + "lastModifiedBy": str(content["closing"]["company"]),
                DCTERMS + "created": metadata_timestamp(content),
                DCTERMS + "modified": metadata_timestamp(content),
                CP + "revision": "1",
            }
            actual = {child.tag: child.text or "" for child in list(core)}
            unexpected = sorted(set(actual) - set(expected))
            mismatches = {
                key: {"expected": value, "actual": actual.get(key)}
                for key, value in expected.items()
                if actual.get(key) != value
            }
            if unexpected:
                errors.append(f"unexpected core-property fields remain: {unexpected}")
            if mismatches:
                errors.append(f"core properties do not match current document: {mismatches}")
            metrics["core_properties_status"] = (
                "pass" if not unexpected and not mismatches else "fail"
            )

        relationships = ET.fromstring(archive.read("_rels/.rels"))
        stale_relationships = [
            {
                "type": relationship.get("Type"),
                "target": relationship.get("Target"),
            }
            for relationship in list(relationships)
            if relationship.get("Type", "").endswith(
                ("/extended-properties", "/custom-properties")
            )
            or relationship.get("Target", "") in {
                "docProps/app.xml", "docProps/custom.xml",
            }
        ]
        if stale_relationships:
            errors.append(f"stale metadata relationships remain: {stale_relationships}")
        metrics["stale_metadata_relationships"] = stale_relationships

        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        stale_content_types = [
            node.get("PartName") for node in list(content_types)
            if node.get("PartName") in {"/docProps/app.xml", "/docProps/custom.xml"}
        ]
        if stale_content_types:
            errors.append(f"stale metadata content types remain: {stale_content_types}")
        metrics["stale_metadata_content_types"] = stale_content_types
    metrics["status"] = "pass" if not errors else "fail"
    return errors, metrics


def clone_with_text(exemplar: ET.Element, text: str,
                    label: str | None = None,
                    bold_lead: bool = False) -> ET.Element:
    p = copy.deepcopy(exemplar)
    p.attrib.pop(W14 + "paraId", None)
    p.attrib.pop(W14 + "textId", None)
    rpr = clone_rpr(p)
    for child in list(p):
        if child.tag != W + "pPr":
            p.remove(child)
    if label and bold_lead:
        lead, reasoning = split_numbered_lead(text)
        add_text_run(p, label + lead, rpr, bold=True)
        if reasoning:
            add_text_run(p, reasoning, rpr, bold=False)
    elif label:
        add_text_run(p, label, rpr, bold=True)
        add_text_run(p, text, rpr, bold=False)
    else:
        add_text_run(p, text, rpr)
    return p


def clone_blank_paragraph(exemplar: ET.Element) -> ET.Element:
    p = copy.deepcopy(exemplar)
    p.attrib.pop(W14 + "paraId", None)
    p.attrib.pop(W14 + "textId", None)
    for child in list(p):
        if child.tag != W + "pPr":
            p.remove(child)
    return p


def analysis_process_language_hits(text: str) -> list[str]:
    hits = [phrase for phrase in ANALYSIS_PROCESS_PHRASES if phrase in text]
    hits.extend(
        label for label, pattern in ANALYSIS_PROCESS_PATTERNS
        if re.search(pattern, text)
    )
    return sorted(set(hits))


def decision_semantic_findings(content: dict[str, Any]) -> dict[str, Any]:
    sections = {
        section.get("heading"): section
        for section in content.get("sections", [])
        if isinstance(section, dict)
    }
    analysis_blocks = sections.get("投资情形分析", {}).get("blocks", [])
    numbered = [block for block in analysis_blocks if block.get("type") == "numbered"]
    item_results: list[dict[str, Any]] = []
    for index, block in enumerate(numbered[:7], start=1):
        text = block.get("text", "").strip()
        lead = re.split(r"[：:。；;]", text, maxsplit=1)[0]
        assertions = [marker for marker in COMPLIANCE_ASSERTION_MARKERS if marker in lead]
        indecisive = [marker for marker in COMPLIANCE_INDECISIVE_MARKERS if marker in lead]
        process_language = analysis_process_language_hits(text)
        role_name, role_markers = ANALYSIS_ROLE_RULES[index - 1]
        role_hits = [marker for marker in role_markers if marker in text]
        item_results.append({
            "item": index,
            "lead": lead,
            "assertions": assertions,
            "indecisive": indecisive,
            "process_language": process_language,
            "required_role": role_name,
            "role_markers": role_hits,
            "role_pass": bool(role_hits),
            "pass": bool(assertions) and not indecisive and not process_language,
        })

    conclusions = [block for block in analysis_blocks if block.get("type") == "conclusion"]
    conclusion_text = conclusions[0].get("text", "").strip() if len(conclusions) == 1 else ""
    conclusion_forbidden = [
        marker for marker in CONCLUSION_FORBIDDEN_MARKERS if marker in conclusion_text
    ]
    has_condition = "前提下" in conclusion_text or "条件下" in conclusion_text
    conclusion_pass = (
        len(conclusions) == 1
        and CONCLUSION_REQUIRED_MARKER in conclusion_text
        and has_condition
        and not conclusion_forbidden
    )
    return {
        "items": item_results,
        "invalid_items": [result for result in item_results if not result["pass"]],
        "invalid_process_items": [
            result for result in item_results if result["process_language"]
        ],
        "invalid_roles": [result for result in item_results if not result["role_pass"]],
        "conclusion_text": conclusion_text,
        "conclusion_forbidden": conclusion_forbidden,
        "conclusion_has_condition": has_condition,
        "conclusion_pass": conclusion_pass,
    }


def compact_char_count(text: str) -> int:
    return len(re.sub(r"\s+", "", text or ""))


def reason_title_specificity_findings(reasons: list[dict[str, Any]]) -> dict[str, Any]:
    """Require each investment-reason lead to state a specific investment thesis."""
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, block in enumerate(reasons, start=1):
        text = str(block.get("text", "")).strip()
        lead = re.split(r"[：:。]", text, maxsplit=1)[0].strip()
        chars = compact_char_count(lead)
        problems: list[str] = []
        if chars < REASON_TITLE_MIN_CHARS:
            problems.append(f"shorter than {REASON_TITLE_MIN_CHARS} compact characters")
        if chars > REASON_TITLE_MAX_CHARS:
            problems.append(f"longer than {REASON_TITLE_MAX_CHARS} compact characters")
        if lead in REASON_OVERBROAD_TITLES:
            problems.append("over-broad category label instead of a project-specific thesis")
        if lead in seen:
            problems.append("duplicate summary title")
        seen.add(lead)
        rows.append({
            "item": index,
            "title": lead,
            "chars": chars,
            "problems": problems,
            "pass": not problems,
        })
    return {
        "status": "pass" if all(row["pass"] for row in rows) else "fail",
        "items": rows,
        "invalid_items": [row for row in rows if not row["pass"]],
    }


def company_subsection_blocks(section: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    active: str | None = None
    for block in section.get("blocks", []):
        if block.get("type") == "subheading":
            active = block.get("text", "").strip()
            groups.setdefault(active, [])
        elif active:
            groups.setdefault(active, []).append(block)
    return groups


def content_completeness_findings(content: dict[str, Any]) -> dict[str, Any]:
    """Check section-role placement, team coverage, and sample-calibrated richness."""
    sections = {
        section.get("heading"): section
        for section in content.get("sections", [])
        if isinstance(section, dict)
    }
    errors: list[str] = []
    warnings: list[str] = []
    company = sections.get("公司情况介绍", {})
    company_groups = company_subsection_blocks(company)
    expected_subsections = ["公司简介", "核心团队", "产品及技术"]
    actual_subsections = [
        block.get("text", "").strip()
        for block in company.get("blocks", [])
        if block.get("type") == "subheading"
    ]
    if actual_subsections != expected_subsections:
        errors.append(
            f"公司情况介绍 subheadings must be exactly {expected_subsections}; got {actual_subsections}"
        )

    profile_blocks = [
        block for block in company_groups.get("公司简介", [])
        if block.get("type") == "paragraph"
    ]
    target_company = content.get("target_company")
    target_legal_name = (
        target_company.get("legal_name", "").strip()
        if isinstance(target_company, dict) else ""
    )
    profile_opening = profile_blocks[0].get("text", "").strip() if profile_blocks else ""
    if not target_legal_name:
        errors.append("target_company.legal_name is required")
        profile_legal_name_status = "missing"
    elif not profile_opening.startswith(target_legal_name):
        errors.append(
            "the first 公司简介 paragraph must begin with target_company.legal_name; "
            f"expected {target_legal_name!r}, got {profile_opening[:80]!r}"
        )
        profile_legal_name_status = "fail"
    else:
        profile_legal_name_status = "pass"
    profile_text = "".join(block.get("text", "") for block in profile_blocks)
    profile_chars = compact_char_count(profile_text)
    ownership_hits = [
        marker for marker in COMPANY_PROFILE_OWNERSHIP_MARKERS if marker in profile_text
    ]
    if re.search(r"(?:股东|股权|表决权|资本表)[^。；]{0,30}[0-9０-９]+(?:\.[0-9０-９]+)?[%％]", profile_text):
        ownership_hits.append("股权百分比")
    ownership_hits = sorted(set(ownership_hits))
    registry_hits = [
        marker for marker in COMPANY_PROFILE_REGISTRY_MARKERS if marker in profile_text
    ]
    registry_hits.extend(
        re.findall(r"(?<![0-9A-Z])[0-9A-Z]{18}(?![0-9A-Z])", profile_text)
    )
    registry_hits = sorted(set(registry_hits))
    routine_registry_hits = sorted(set(
        marker for marker in COMPANY_PROFILE_ROUTINE_REGISTRY_MARKERS
        if marker in profile_text
    ))
    financial_dd_hits = sorted(set(
        match.group(0)
        for pattern in COMPANY_PROFILE_FINANCIAL_DD_PATTERNS
        for match in pattern.finditer(profile_text)
    ))
    financial_metric_hits = sorted(set(
        match.group(0)
        for match in COMPANY_PROFILE_EXACT_FINANCIAL_METRIC_PATTERN.finditer(profile_text)
    ))
    disclosure = content.get("company_profile_financial_disclosure")
    disclosure_present = disclosure is not None
    disclosure_sources = disclosure.get("source_ids", []) if isinstance(disclosure, dict) else []
    financial_disclosure_allowed = (
        isinstance(disclosure, dict)
        and disclosure.get("allowed") is True
        and disclosure.get("user_requested") is True
        and disclosure.get("basis") in COMPANY_PROFILE_FINANCIAL_DISCLOSURE_BASES
        and disclosure.get("no_material_conflict") is True
        and isinstance(disclosure_sources, list)
        and bool(disclosure_sources)
        and all(isinstance(source_id, str) and source_id.strip() for source_id in disclosure_sources)
    )
    if not disclosure_present:
        financial_disclosure_status = "not_requested"
    elif financial_disclosure_allowed:
        financial_disclosure_status = "allowed"
    else:
        financial_disclosure_status = "invalid"
    if profile_chars < COMPANY_PROFILE_MIN_CHARS:
        warnings.append(
            f"公司简介 is too thin: {profile_chars} chars; minimum {COMPANY_PROFILE_MIN_CHARS}"
        )
    if ownership_hits:
        errors.append(
            "公司简介 contains ownership/cap-table information that belongs in 投资计划 or "
            f"投资情形分析: {ownership_hits}"
        )
    if registry_hits:
        errors.append(
            "公司简介 contains licence/registry identifiers that belong in the audit layer, "
            f"not visible narrative prose: {registry_hits}"
        )
    if routine_registry_hits:
        errors.append(
            "公司简介 must be business-first and must not display the legal representative, "
            "registered/paid-in/subscribed capital or other routine registry fields; keep them "
            f"in the audit layer or a directly relevant compliance analysis: {routine_registry_hits}"
        )
    if financial_dd_hits:
        errors.append(
            "公司简介 contains financial due-diligence findings that belong in notes/open_issues "
            f"or closing conditions, not profile prose: {financial_dd_hits}"
        )
    if disclosure_present and not financial_disclosure_allowed:
        errors.append(
            "company_profile_financial_disclosure is invalid; exact profile financial metrics "
            "require allowed=true, user_requested=true, basis=audited|special_audit, "
            "no_material_conflict=true, and non-empty source_ids"
        )
    if financial_metric_hits and not financial_disclosure_allowed:
        errors.append(
            "公司简介 contains exact financial metrics without an explicit, user-requested, "
            "audited and conflict-free disclosure basis: "
            f"{financial_metric_hits}"
        )

    decision_meta_hits: list[dict[str, Any]] = []
    source_narration_hits: list[dict[str, Any]] = []
    issuer_name_hits: list[dict[str, Any]] = []
    closing = content.get("closing", {})
    issuer_name = closing.get("company", "").strip() if isinstance(closing, dict) else ""
    for section in content.get("sections", []):
        heading = section.get("heading", "")
        subsection = ""
        for block_index, block in enumerate(section.get("blocks", []), start=1):
            if block.get("type") == "subheading":
                subsection = block.get("text", "").strip()
                continue
            if block.get("type") not in {"paragraph", "numbered", "conclusion"}:
                continue
            text = block.get("text", "")
            if issuer_name and issuer_name in text:
                issuer_name_hits.append({
                    "section": heading,
                    "subsection": subsection,
                    "block": block_index,
                    "issuer_name": issuer_name,
                    "text": text[:120],
                })
            if heading not in VISIBLE_SOURCE_PROSE_SECTIONS:
                continue
            markers = sorted(set(
                marker for marker in DECISION_LAYER_META_MARKERS if marker in text
            ))
            if markers:
                decision_meta_hits.append({
                    "section": heading,
                    "subsection": subsection,
                    "block": block_index,
                    "markers": markers,
                    "text": text[:120],
                })
            source_markers = visible_source_narration_markers(text)
            if source_markers:
                source_narration_hits.append({
                    "section": heading,
                    "subsection": subsection,
                    "block": block_index,
                    "markers": source_markers,
                    "text": text[:160],
                })
    if decision_meta_hits:
        errors.append(
            "visible decision-layer prose contains evidence-acquisition meta-language; state "
            "the supported fact directly and keep provenance in source_ids/notes/open_issues: "
            f"{decision_meta_hits}"
        )
    if source_narration_hits:
        errors.append(
            "visible company/reason/plan prose contains attachment filenames, internal versions, "
            "paths, or source-process narration; state the fact directly and keep provenance in "
            f"source_ids/notes: {source_narration_hits}"
        )
    if issuer_name_hits:
        errors.append(
            "the full issuer/manager company name from closing.company is allowed only in the "
            "closing signature; use 我方、本基金、指定基金主体 or 管理人 in visible sections: "
            f"{issuer_name_hits}"
        )

    team_blocks = [
        block for block in company_groups.get("核心团队", [])
        if block.get("type") == "paragraph"
    ]
    team_chars = [compact_char_count(block.get("text", "")) for block in team_blocks]
    team_dimension_rows: list[dict[str, Any]] = []
    team_title_before_name_issues: list[dict[str, Any]] = []
    team_evaluative_benefit_hits: list[dict[str, Any]] = []
    for index, block in enumerate(team_blocks, start=1):
        text = block.get("text", "").strip()
        person_name = block.get("person_name", "").strip()
        role_title = block.get("role_title", "").strip()
        expected_prefix = f"{role_title}{person_name}" if role_title and person_name else ""
        title_before_name = bool(expected_prefix) and text.startswith(expected_prefix)
        dimensions = {
            dimension: any(marker in text for marker in markers)
            for dimension, markers in TEAM_DIMENSION_MARKERS.items()
        }
        row = {
            "member": index,
            "person_name": person_name,
            "role_title": role_title,
            "expected_prefix": expected_prefix,
            "title_before_name": title_before_name,
            "chars": compact_char_count(text),
            "dimensions": dimensions,
            "dimension_count": sum(dimensions.values()),
        }
        team_dimension_rows.append(row)
        if not title_before_name:
            team_title_before_name_issues.append(row)
        for label, pattern in TEAM_EVALUATIVE_BENEFIT_PATTERNS:
            for match in pattern.finditer(text):
                team_evaluative_benefit_hits.append({
                    "member": index,
                    "person_name": person_name,
                    "pattern": label,
                    "text": match.group(0),
                })
    if team_title_before_name_issues:
        errors.append(
            "each 核心团队 paragraph must provide role_title/person_name and begin with "
            "the exact role_title + person_name sequence (title before person name): "
            f"{team_title_before_name_issues}"
        )
    if team_evaluative_benefit_hits:
        errors.append(
            "核心团队 biographies must remain objective and must not append inferred project-"
            "benefit or company-fit statements; state documented expertise/current duties and "
            "move the collective team investment thesis to 投资理由: "
            f"{team_evaluative_benefit_hits}"
        )
    if len(team_blocks) < TEAM_MIN_PARAGRAPHS:
        warnings.append(
            f"核心团队 requires at least {TEAM_MIN_PARAGRAPHS} member paragraphs; got {len(team_blocks)}"
        )
    if sum(team_chars) < TEAM_MIN_TOTAL_CHARS:
        warnings.append(
            f"核心团队 is too thin: {sum(team_chars)} chars; minimum {TEAM_MIN_TOTAL_CHARS}"
        )
    thin_members = [row for row in team_dimension_rows if row["chars"] < TEAM_MIN_MEMBER_CHARS]
    if thin_members:
        warnings.append(
            f"核心团队 member paragraphs below {TEAM_MIN_MEMBER_CHARS} chars: {thin_members}"
        )
    incomplete_members = [row for row in team_dimension_rows if row["dimension_count"] < 3]
    if incomplete_members:
        warnings.append(
            "team coverage requires at least three of role/education/experience/"
            "expertise_or_responsibility per member: "
            f"{incomplete_members}"
        )

    product_blocks = [
        block for block in company_groups.get("产品及技术", [])
        if block.get("type") == "paragraph"
    ]
    product_chars = sum(compact_char_count(block.get("text", "")) for block in product_blocks)
    if len(product_blocks) < PRODUCT_MIN_PARAGRAPHS:
        warnings.append(
            f"产品及技术 requires at least {PRODUCT_MIN_PARAGRAPHS} paragraphs; got {len(product_blocks)}"
        )
    if product_chars < PRODUCT_MIN_TOTAL_CHARS:
        warnings.append(
            f"产品及技术 is too thin: {product_chars} chars; minimum {PRODUCT_MIN_TOTAL_CHARS}"
        )

    reasons = [
        block for block in sections.get("投资理由", {}).get("blocks", [])
        if block.get("type") == "numbered"
    ]
    reason_chars = [compact_char_count(block.get("text", "")) for block in reasons]
    if len(reasons) != REASON_REQUIRED_ITEMS:
        errors.append(
            f"投资理由 requires exactly {REASON_REQUIRED_ITEMS} numbered items; got {len(reasons)}"
        )
    thin_reasons = [index for index, chars in enumerate(reason_chars, start=1) if chars < REASON_MIN_ITEM_CHARS]
    if thin_reasons:
        warnings.append(
            f"投资理由 items below {REASON_MIN_ITEM_CHARS} chars: {thin_reasons}"
        )
    if sum(reason_chars) < REASON_MIN_TOTAL_CHARS:
        warnings.append(
            f"投资理由 is too thin: {sum(reason_chars)} chars; minimum {REASON_MIN_TOTAL_CHARS}"
        )
    reason_titles = reason_title_specificity_findings(reasons)
    if reason_titles["invalid_items"]:
        errors.append(
            "investment-reason summary titles must state a project-specific investment thesis "
            "instead of a broad category label: "
            f"{reason_titles['invalid_items']}"
        )

    plan_blocks = [
        block for block in sections.get("投资计划", {}).get("blocks", [])
        if block.get("type") == "paragraph"
    ]
    plan_chars = sum(compact_char_count(block.get("text", "")) for block in plan_blocks)
    if plan_chars < PLAN_MIN_TOTAL_CHARS:
        warnings.append(f"投资计划 is too thin: {plan_chars} chars; minimum {PLAN_MIN_TOTAL_CHARS}")

    analysis_blocks = [
        block for block in sections.get("投资情形分析", {}).get("blocks", [])
        if block.get("type") == "numbered"
    ]
    analysis_process_hits: list[dict[str, Any]] = []
    for index, block in enumerate(analysis_blocks[:7], start=1):
        text = block.get("text", "")
        markers = analysis_process_language_hits(text)
        if markers:
            analysis_process_hits.append({
                "item": index,
                "markers": markers,
                "text": text[:160],
            })
    if analysis_process_hits:
        errors.append(
            "投资情形分析 visible decision layer contains unfinished audit, missing-input, "
            "calculation-request, or refusal-to-conclude language; move it to "
            f"delivery_readiness/notes/open_issues: {analysis_process_hits}"
        )
    analysis_chars = [compact_char_count(block.get("text", "")) for block in analysis_blocks[:7]]
    thin_analysis = [
        index for index, chars in enumerate(analysis_chars, start=1)
        if chars < ANALYSIS_MIN_ITEM_CHARS
    ]
    if thin_analysis:
        warnings.append(
            f"投资情形分析 items below {ANALYSIS_MIN_ITEM_CHARS} chars: {thin_analysis}"
        )
    if sum(analysis_chars) < ANALYSIS_MIN_TOTAL_CHARS:
        warnings.append(
            f"投资情形分析 is too thin: {sum(analysis_chars)} chars; minimum {ANALYSIS_MIN_TOTAL_CHARS}"
        )

    metrics = {
        "company_profile_chars": profile_chars,
        "company_profile_legal_name_status": profile_legal_name_status,
        "target_company_legal_name": target_legal_name,
        "company_profile_ownership_hits": ownership_hits,
        "company_profile_registry_hits": registry_hits,
        "company_profile_routine_registry_hits": routine_registry_hits,
        "company_profile_financial_dd_hits": financial_dd_hits,
        "company_profile_financial_metric_hits": financial_metric_hits,
        "company_profile_financial_disclosure_status": financial_disclosure_status,
        "decision_layer_meta_language_hits": decision_meta_hits,
        "visible_source_narration_hits": source_narration_hits,
        "visible_issuer_name_hits": issuer_name_hits,
        "team_member_paragraphs": len(team_blocks),
        "team_chars": sum(team_chars),
        "team_title_before_name_status": "pass" if not team_title_before_name_issues else "fail",
        "team_title_before_name_issues": team_title_before_name_issues,
        "team_objective_prose_status": "pass" if not team_evaluative_benefit_hits else "fail",
        "team_evaluative_benefit_hits": team_evaluative_benefit_hits,
        "team_coverage": team_dimension_rows,
        "product_paragraphs": len(product_blocks),
        "product_chars": product_chars,
        "reason_items": len(reasons),
        "reason_chars": sum(reason_chars),
        "reason_average_chars": round(sum(reason_chars) / len(reason_chars), 1) if reason_chars else 0,
        "reason_title_specificity_status": reason_titles["status"],
        "reason_title_specificity": reason_titles["items"],
        "plan_chars": plan_chars,
        "analysis_items": len(analysis_blocks),
        "analysis_chars": sum(analysis_chars),
        "analysis_average_chars": round(sum(analysis_chars) / len(analysis_chars), 1) if analysis_chars else 0,
        "analysis_process_language_status": "pass" if not analysis_process_hits else "fail",
        "analysis_process_language_hits": analysis_process_hits,
    }
    metrics["evidence_richness_status"] = "pass" if not warnings else "limited"
    return {
        "status": "pass" if not errors else "fail",
        "errors": errors,
        "warnings": warnings,
        "metrics": metrics,
    }


def _nonempty_source_ids(component: dict[str, Any]) -> bool:
    source_ids = component.get("source_ids")
    return isinstance(source_ids, list) and bool(source_ids) and all(
        isinstance(value, str) and value.strip() for value in source_ids
    )


def _missing_keys(value: dict[str, Any], keys: list[str]) -> list[str]:
    return [key for key in keys if value.get(key) in (None, "", [])]


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _valid_iso_date(value: Any) -> bool:
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return False
    try:
        dt.date.fromisoformat(value)
    except ValueError:
        return False
    return True


def visible_source_narration_markers(text: str) -> list[str]:
    """Return attachment/source-process tokens that must stay out of decision prose."""
    markers: list[str] = []
    for match in VISIBLE_ATTACHMENT_TITLE_PATTERN.finditer(text):
        title = match.group("title").strip()
        if not any(pattern.fullmatch(title) for pattern in ALLOWED_FORMAL_LAW_TITLE_PATTERNS):
            markers.append(match.group(0))
    for pattern in VISIBLE_SOURCE_PROCESS_PATTERNS:
        markers.extend(match.group(0) for match in pattern.finditer(text))
    return sorted(set(markers))


def delivery_readiness_findings(content: dict[str, Any]) -> dict[str, Any]:
    """Require a user decision before an incomplete evidence packet can proceed."""
    readiness = content.get("delivery_readiness")
    errors: list[str] = []
    warnings: list[str] = []
    metrics: dict[str, Any] = {
        "status": readiness.get("status") if isinstance(readiness, dict) else "missing",
        "components": {},
    }
    if not isinstance(readiness, dict):
        return {
            "status": "fail",
            "errors": ["delivery_readiness is required to record the drafting basis and limitations"],
            "warnings": warnings,
            "metrics": metrics,
        }

    declared_status = readiness.get("status")
    allowed_statuses = {
        "ready", "awaiting_user_input", "proceed_with_available_materials", "blocked",
    }
    if declared_status not in allowed_statuses:
        errors.append(
            f"delivery_readiness.status must be one of {sorted(allowed_statuses)}; "
            f"got {declared_status!r}"
        )
    if declared_status == "blocked":
        errors.append("delivery_readiness explicitly records a hard-stop condition")
    if declared_status == "awaiting_user_input":
        errors.append(
            "delivery_readiness is awaiting user input; ask the user in a blocking final "
            "response and stop before authoring or build"
        )
    if not _valid_iso_date(readiness.get("as_of_date")):
        errors.append("delivery_readiness.as_of_date must be YYYY-MM-DD")

    blocking_issues = readiness.get("blocking_issues", [])
    if not isinstance(blocking_issues, list):
        errors.append("delivery_readiness.blocking_issues must be an array")
        blocking_issues = []
    elif blocking_issues:
        errors.append(
            "delivery_readiness contains unresolved material conflicts or known non-compliance: "
            f"{blocking_issues}"
        )

    missing_inputs = readiness.get("missing_decisive_inputs", [])
    if not isinstance(missing_inputs, list):
        errors.append("delivery_readiness.missing_decisive_inputs must be an array")
        missing_inputs = []
    elif missing_inputs:
        warnings.append(f"supplemental inputs were not provided: {missing_inputs}")

    supplement = readiness.get("supplement_request")
    if not isinstance(supplement, dict):
        errors.append("delivery_readiness.supplement_request is required")
    else:
        requested = supplement.get("requested")
        outcome = supplement.get("outcome")
        requested_items = supplement.get("requested_items")
        allowed_outcomes = {
            "not_needed", "provided", "awaiting_response", "not_provided", "declined",
        }
        if not isinstance(requested, bool):
            errors.append("delivery_readiness.supplement_request.requested must be boolean")
        if outcome not in allowed_outcomes:
            errors.append(
                "delivery_readiness.supplement_request.outcome must be one of "
                f"{sorted(allowed_outcomes)}"
            )
        if not isinstance(requested_items, list):
            errors.append("delivery_readiness.supplement_request.requested_items must be an array")
        if declared_status == "awaiting_user_input":
            if requested is not True or outcome != "awaiting_response":
                errors.append(
                    "awaiting_user_input requires a requested supplement and "
                    "outcome=awaiting_response"
                )
            if not isinstance(requested_items, list) or not requested_items:
                errors.append("awaiting_user_input requires non-empty requested_items")
        if declared_status == "proceed_with_available_materials":
            if requested is not True or outcome not in {"not_provided", "declined", "provided"}:
                errors.append(
                    "proceed_with_available_materials requires a recorded supplement request and outcome"
                )

    authorization = readiness.get("continuation_authorization")
    if not isinstance(authorization, dict):
        errors.append("delivery_readiness.continuation_authorization is required")
    else:
        authorized = authorization.get("authorized")
        basis = authorization.get("basis")
        instruction = authorization.get("instruction")
        if not isinstance(authorized, bool):
            errors.append("continuation_authorization.authorized must be boolean")
        if basis not in {"not_required", "awaiting_response", "explicit_user_instruction"}:
            errors.append("continuation_authorization.basis is invalid")
        if declared_status == "proceed_with_available_materials":
            if authorized is not True or basis != "explicit_user_instruction":
                errors.append(
                    "proceed_with_available_materials requires explicit user authorization"
                )
            if not isinstance(instruction, str) or not instruction.strip():
                errors.append(
                    "explicit continuation authorization requires a non-empty user instruction"
                )
        elif declared_status == "awaiting_user_input":
            if authorized is not False or basis != "awaiting_response":
                errors.append(
                    "awaiting_user_input must record authorized=false and basis=awaiting_response"
                )
        elif declared_status in {"ready", "blocked"}:
            if authorized is not False or basis != "not_required":
                errors.append(
                    f"{declared_status} must record authorized=false and basis=not_required"
                )

    if declared_status in {"awaiting_user_input", "proceed_with_available_materials"}:
        if not missing_inputs:
            errors.append(f"{declared_status} requires non-empty missing_decisive_inputs")

    strict = declared_status == "ready"

    def add_gap(message: str) -> None:
        (errors if strict else warnings).append(message)

    for component_name in READINESS_REQUIRED_COMPONENTS:
        component = readiness.get(component_name)
        component_gaps: list[str] = []
        component_errors: list[str] = []
        if not isinstance(component, dict):
            component_gaps.append("component is missing")
            component = {}
        expected_status = (
            "calculated" if component_name in {"return_investment", "concentration"}
            else "verified"
        )
        component_status = component.get("status")
        if component_status not in {"verified", "calculated", "pending", "not_applicable"}:
            component_gaps.append(f"status is incomplete or invalid: {component_status!r}")
        elif component_status != expected_status:
            component_gaps.append(f"status is {component_status!r}; expected {expected_status!r}")
        if not _nonempty_source_ids(component):
            component_gaps.append("source_ids do not contain a decisive source binding")

        if component_name == "fund_agreement":
            clauses = component.get("applicable_clauses")
            if not isinstance(clauses, dict):
                component_gaps.append("applicable_clauses is absent")
            else:
                missing = _missing_keys(clauses, FUND_CLAUSE_KEYS)
                if missing:
                    component_gaps.append(f"missing applicable fund clauses: {missing}")
        elif component_name == "transaction_terms":
            terms = component.get("terms")
            if not isinstance(terms, dict):
                component_gaps.append("terms are absent")
            else:
                missing = _missing_keys(terms, TRANSACTION_TERM_KEYS)
                if missing:
                    component_gaps.append(f"missing definitive transaction terms: {missing}")
                for key in ("investment_amount", "pre_money_valuation", "post_money_valuation"):
                    value = terms.get(key)
                    if value not in (None, "") and (not _is_number(value) or value <= 0):
                        component_gaps.append(f"{key} is not a positive number")
                ownership = terms.get("post_investment_ownership")
                if ownership not in (None, "") and (
                    not _is_number(ownership) or not 0 < ownership <= 1
                ):
                    component_gaps.append("post_investment_ownership is not a ratio in (0, 1]")
                pre_money = terms.get("pre_money_valuation")
                post_money = terms.get("post_money_valuation")
                if _is_number(pre_money) and _is_number(post_money) and post_money < pre_money:
                    component_errors.append(
                        "post_money_valuation must not be below pre_money_valuation"
                    )
        elif component_name == "return_investment":
            calculation = component.get("calculation")
            if not isinstance(calculation, dict):
                component_gaps.append("calculation is absent")
            else:
                missing = _missing_keys(calculation, RETURN_INVESTMENT_KEYS)
                if missing:
                    component_gaps.append(f"missing return-investment calculation fields: {missing}")
                date_value = calculation.get("as_of_date")
                if date_value not in (None, "") and not _valid_iso_date(date_value):
                    component_gaps.append("as_of_date is not YYYY-MM-DD")
                for key in ("denominator", "multiplier"):
                    value = calculation.get(key)
                    if value not in (None, "") and (not _is_number(value) or value <= 0):
                        component_gaps.append(f"{key} is not a positive number")
                for key in ("completed_amount", "proposed_eligible_credit", "post_investment_headroom"):
                    value = calculation.get(key)
                    if value not in (None, "") and (not _is_number(value) or value < 0):
                        component_gaps.append(f"{key} is not a non-negative number")
        elif component_name == "concentration":
            calculation = component.get("calculation")
            if not isinstance(calculation, dict):
                component_gaps.append("calculation is absent")
            else:
                missing = _missing_keys(calculation, CONCENTRATION_KEYS)
                if missing:
                    component_gaps.append(f"missing concentration calculation fields: {missing}")
                date_value = calculation.get("as_of_date")
                if date_value not in (None, "") and not _valid_iso_date(date_value):
                    component_gaps.append("as_of_date is not YYYY-MM-DD")
                denominator = calculation.get("denominator")
                limit_ratio = calculation.get("limit_ratio")
                post_ratio = calculation.get("post_investment_ratio")
                if denominator not in (None, "") and (
                    not _is_number(denominator) or denominator <= 0
                ):
                    component_gaps.append("denominator is not a positive number")
                if limit_ratio not in (None, "") and (
                    not _is_number(limit_ratio) or not 0 < limit_ratio <= 1
                ):
                    component_gaps.append("limit_ratio is not a ratio in (0, 1]")
                for key in (
                    "existing_aggregated_exposure", "proposed_amount",
                    "post_investment_ratio", "headroom",
                ):
                    value = calculation.get(key)
                    if value not in (None, "") and (not _is_number(value) or value < 0):
                        component_gaps.append(f"{key} is not a non-negative number")
                existing = calculation.get("existing_aggregated_exposure")
                proposed = calculation.get("proposed_amount")
                headroom = calculation.get("headroom")
                if all(_is_number(value) for value in (denominator, existing, proposed, post_ratio)) and denominator > 0:
                    expected_ratio = (existing + proposed) / denominator
                    if abs(post_ratio - expected_ratio) > 1e-6:
                        component_errors.append(
                            "post_investment_ratio does not reconcile to "
                            "(existing_aggregated_exposure + proposed_amount) / denominator"
                        )
                if all(_is_number(value) for value in (denominator, limit_ratio, existing, proposed, headroom)):
                    expected_headroom = denominator * limit_ratio - existing - proposed
                    if abs(headroom - expected_headroom) > 0.01:
                        component_errors.append(
                            "headroom does not reconcile to denominator * limit_ratio - "
                            "existing_aggregated_exposure - proposed_amount"
                        )
        elif component_name == "related_party":
            perimeter = component.get("perimeter")
            if not isinstance(perimeter, list) or not perimeter:
                component_gaps.append("perimeter does not list the reviewed parties")

        metrics["components"][component_name] = {
            "status": component.get("status"),
            "source_count": len(component.get("source_ids", []))
            if isinstance(component.get("source_ids"), list) else 0,
            "gaps": component_gaps,
            "errors": component_errors,
        }
        for gap in component_gaps:
            add_gap(f"delivery_readiness.{component_name}: {gap}")
        errors.extend(
            f"delivery_readiness.{component_name}: {error}" for error in component_errors
        )

    transaction = readiness.get("transaction_terms", {})
    concentration = readiness.get("concentration", {})
    transaction_amount = transaction.get("terms", {}).get("investment_amount") \
        if isinstance(transaction, dict) and isinstance(transaction.get("terms"), dict) else None
    proposed_amount = concentration.get("calculation", {}).get("proposed_amount") \
        if isinstance(concentration, dict) and isinstance(concentration.get("calculation"), dict) else None
    if _is_number(transaction_amount) and _is_number(proposed_amount):
        if abs(transaction_amount - proposed_amount) > 0.01:
            errors.append(
                "delivery_readiness transaction investment_amount does not reconcile to "
                "concentration proposed_amount"
            )

    sections = {
        section.get("heading"): section
        for section in content.get("sections", [])
        if isinstance(section, dict)
    }
    analysis_blocks = [
        block for block in sections.get("投资情形分析", {}).get("blocks", [])
        if block.get("type") == "numbered"
    ][:7]
    pending_items = [
        index for index, block in enumerate(analysis_blocks, start=1)
        if block.get("status") == "pending"
    ]
    metrics["pending_compliance_items"] = pending_items
    if pending_items:
        message = f"compliance items remain evidence-limited: {pending_items}"
        (errors if strict else warnings).append(message)

    metrics["blocking_error_count"] = len(errors)
    metrics["warning_count"] = len(warnings)
    return {
        "status": "pass" if not errors else "fail",
        "errors": errors,
        "warnings": warnings,
        "metrics": metrics,
    }


def validate_content(content: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    sections = content.get("sections")
    if not isinstance(sections, list):
        return ["sections must be an array"]
    headings = [s.get("heading") for s in sections if isinstance(s, dict)]
    if headings != REQUIRED_SECTIONS:
        errors.append(f"section order must be exactly {REQUIRED_SECTIONS}; got {headings}")
    closing = content.get("closing", {})
    if not closing.get("company") or not closing.get("date"):
        errors.append("closing.company and closing.date are required")
    public = content.get("public_verification")
    if not isinstance(public, dict):
        errors.append("public_verification binding is required")
    else:
        mode = public.get("mode")
        if mode not in {"online", "online_limited", "offline_user_requested"}:
            errors.append("public_verification.mode is invalid")
        if not isinstance(public.get("record"), str) or not public.get("record"):
            errors.append("public_verification.record is required")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(public.get("as_of_date", ""))):
            errors.append("public_verification.as_of_date must be YYYY-MM-DD")
        if not re.fullmatch(r"[0-9a-f]{64}", str(public.get("sha256", ""))):
            errors.append("public_verification.sha256 must be a lowercase SHA-256")
        if public.get("validation_status") != "pass":
            errors.append("public_verification.validation_status must be pass")
        coverage = public.get("coverage_status")
        if coverage not in {"complete", "limited", "offline"}:
            errors.append("public_verification.coverage_status is invalid")
        if mode == "offline_user_requested" and coverage != "offline":
            errors.append("offline public verification must use coverage_status=offline")
        if mode != "offline_user_requested" and coverage == "offline":
            errors.append("online public verification cannot use coverage_status=offline")
        if public.get("decision_impact") not in {
            "no_material_public_conflict_found", "conditions_added", "not_verified",
        }:
            errors.append("public_verification.decision_impact is invalid")
    analysis = next((s for s in sections if s.get("heading") == "投资情形分析"), None)
    if analysis:
        item_count = sum(b.get("type") == "numbered" for b in analysis.get("blocks", []))
        if item_count < 7:
            errors.append(f"投资情形分析 requires at least 7 numbered checks; got {item_count}")
        if item_count >= 7:
            semantic = decision_semantic_findings(content)
            if semantic["invalid_items"]:
                errors.append(
                    "seven compliance checks must begin with an explicit affirmative judgment "
                    "and must not use indecisive leads or audit-process language: "
                    f"{semantic['invalid_items']}"
                )
            if semantic["invalid_process_items"]:
                errors.append(
                    "investment-scenario visible prose must not narrate unfinished review, "
                    "missing inputs, calculation requests, or refusal to conclude; move those "
                    f"tasks to the audit layer: {semantic['invalid_process_items']}"
                )
            if semantic["invalid_roles"]:
                errors.append(
                    "投资情形分析 must follow the standard role map in order: investment "
                    "restrictions, return-investment impact, related parties, investment "
                    "direction, configuration, concentration, and residual law/regulation; "
                    f"invalid items: {semantic['invalid_roles']}"
                )
            if not semantic["conclusion_pass"]:
                errors.append(
                    "conclusion must state a conditional 原则上符合 judgment and must not say "
                    f"that no conclusion can be formed: text={semantic['conclusion_text']!r}; "
                    f"forbidden={semantic['conclusion_forbidden']}"
                )
    completeness = content_completeness_findings(content)
    errors.extend(completeness["errors"])
    readiness = delivery_readiness_findings(content)
    errors.extend(readiness["errors"])
    return errors


def validate_public_binding(
    content: dict[str, Any], content_path: Path
) -> tuple[list[str], dict[str, Any]]:
    public = content.get("public_verification")
    if not isinstance(public, dict):
        return [], {}
    record_value = public.get("record")
    if not isinstance(record_value, str) or not record_value:
        return [], {}
    record_path = Path(record_value).expanduser()
    if not record_path.is_absolute():
        record_path = (content_path.parent / record_path).resolve()
    metrics: dict[str, Any] = {"record": str(record_path)}
    if not record_path.is_file():
        return [f"public verification record not found: {record_path}"], metrics
    actual_sha = sha256(record_path)
    metrics["sha256"] = actual_sha
    if public.get("sha256") != actual_sha:
        return [
            "public verification record SHA-256 mismatch: "
            f"content={public.get('sha256')!r}, actual={actual_sha}"
        ], metrics
    try:
        record = json.loads(record_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return [f"public verification record is unreadable: {exc}"], metrics
    spec = importlib.util.spec_from_file_location(
        "public_verification_validator", PUBLIC_VALIDATOR
    )
    if spec is None or spec.loader is None:
        return [f"cannot load public verification validator: {PUBLIC_VALIDATOR}"], metrics
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    record_errors = module.validate_record(record) if isinstance(record, dict) else [
        "public verification record root must be an object"
    ]
    errors = [f"public verification: {error}" for error in record_errors]
    if isinstance(record, dict):
        summary = record.get("summary", {})
        scope = record.get("scope", {})
        comparisons = {
            "as_of_date": record.get("as_of_date"),
            "coverage_status": scope.get("coverage_status"),
            "decision_impact": summary.get("decision_impact"),
            "validation_status": summary.get("validation_status"),
        }
        metrics.update(comparisons)
        for key, record_value in comparisons.items():
            if public.get(key) != record_value:
                errors.append(
                    f"public verification binding mismatch for {key}: "
                    f"content={public.get(key)!r}, record={record_value!r}"
                )
    metrics["binding_status"] = "pass" if not errors else "fail"
    return errors, metrics


def build_command(args: argparse.Namespace) -> int:
    content_path = Path(args.content).expanduser().resolve()
    output = Path(args.output).expanduser().resolve()
    template = Path(args.template).expanduser().resolve() if args.template else DEFAULT_TEMPLATE
    content = json.loads(content_path.read_text(encoding="utf-8"))
    errors = validate_content(content)
    build_warnings = (
        content_completeness_findings(content).get("warnings", [])
        + delivery_readiness_findings(content).get("warnings", [])
    )
    binding_errors, _ = validate_public_binding(content, content_path)
    errors.extend(binding_errors)
    if errors:
        raise SystemExit("Invalid content:\n- " + "\n- ".join(errors))
    if not template.exists():
        raise SystemExit(f"Template not found: {template}")

    with zipfile.ZipFile(template) as zin:
        document_xml = zin.read("word/document.xml")
        root = ET.fromstring(document_xml)
        body = root.find("w:body", NS)
        if body is None:
            raise SystemExit("Template document body not found")
        paragraphs = body.findall("w:p", NS)
        ex_title = find_paragraph(paragraphs, exact="关于德塔智能项目投资合规性的说明")
        ex_main = find_paragraph(paragraphs, exact="公司情况介绍")
        ex_sub = find_paragraph(paragraphs, exact="公司简介")
        ex_body = find_paragraph(paragraphs, startswith="北京德塔源创智能科技有限公司")
        ex_numbered = find_paragraph(paragraphs, startswith="2、")
        ex_conclusion = find_paragraph(paragraphs, startswith="综上，")
        ex_closing_company = find_paragraph(paragraphs, exact="浙江赛智伯乐股权投资管理有限公司")
        ex_closing_date = find_paragraph(paragraphs, startswith="2026年")
        ex_blank = next((p for p in paragraphs if not paragraph_text(p).strip()), None)
        if ex_blank is None:
            raise SystemExit("Template transition paragraph not found")
        sectpr = body.find("w:sectPr", NS)
        if sectpr is None:
            raise SystemExit("Template section properties not found")
        for child in list(body):
            if child is not sectpr:
                body.remove(child)

        body.insert(len(body) - 1, clone_with_text(ex_title, content["title"]))
        for section in content["sections"]:
            if section["heading"] == "投资情形分析":
                body.insert(len(body) - 1, clone_blank_paragraph(ex_blank))
            body.insert(len(body) - 1, clone_with_text(ex_main, section["heading"]))
            for block in section.get("blocks", []):
                kind = block.get("type", "paragraph")
                if kind == "subheading":
                    p = clone_with_text(ex_sub, block["text"])
                elif kind == "numbered":
                    label = block.get("label") or ""
                    if label and not re.search(r"[、.)）]\s*$", label):
                        label += "、"
                    p = clone_with_text(
                        ex_numbered, block["text"], label=label, bold_lead=True
                    )
                elif kind == "conclusion":
                    p = clone_with_text(ex_conclusion, block["text"])
                else:
                    p = clone_with_text(ex_body, block["text"])
                body.insert(len(body) - 1, p)
        body.insert(len(body) - 1, clone_blank_paragraph(ex_blank))
        body.insert(len(body) - 1, clone_with_text(ex_closing_company, content["closing"]["company"]))
        body.insert(
            len(body) - 1,
            clone_with_text(ex_closing_date, format_closing_date(content["closing"]["date"])),
        )

        ET.register_namespace("w", W_NS)
        ET.register_namespace("w14", W14_NS)
        new_document_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        new_core_xml = safe_core_properties(content)
        output.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output, "w") as zout:
            for info in zin.infolist():
                if info.filename in {"docProps/app.xml", "docProps/custom.xml"}:
                    continue
                if info.filename == "word/document.xml":
                    data = new_document_xml
                elif info.filename == "docProps/core.xml":
                    data = new_core_xml
                elif info.filename == "_rels/.rels":
                    data = strip_package_metadata_relationships(zin.read(info.filename))
                elif info.filename == "[Content_Types].xml":
                    data = strip_package_metadata_content_types(zin.read(info.filename))
                else:
                    data = zin.read(info.filename)
                zout.writestr(info, data)
    print(json.dumps({
        "output": str(output),
        "sha256": sha256(output),
        "template": str(template),
        "template_sha256": sha256(template),
        "warnings": build_warnings,
    }, ensure_ascii=False))
    return 0


def package_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as zf:
        return {name: hashlib.sha256(zf.read(name)).hexdigest() for name in zf.namelist()}


def direct_run_properties(docx: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(docx) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    rows: list[dict[str, Any]] = []
    for p_index, p in enumerate(root.findall(".//w:body/w:p", NS)):
        p_text = paragraph_text(p).strip()
        if not p_text:
            continue
        for r in p.findall("w:r", NS):
            text = "".join(t.text or "" for t in r.findall(".//w:t", NS)).strip()
            if not text:
                continue
            rpr = r.find("w:rPr", NS)
            fonts = rpr.find("w:rFonts", NS) if rpr is not None else None
            size = rpr.find("w:sz", NS) if rpr is not None else None
            underline = rpr.find("w:u", NS) if rpr is not None else None
            rows.append({
                "paragraph": p_index, "paragraph_text": p_text, "text": text,
                "east_asia_font": fonts.get(W + "eastAsia") if fonts is not None else None,
                "ascii_font": fonts.get(W + "ascii") if fonts is not None else None,
                "hansi_font": fonts.get(W + "hAnsi") if fonts is not None else None,
                "half_points": int(size.get(W + "val")) if size is not None and size.get(W + "val") else None,
                "bold": rpr is not None and rpr.find("w:b", NS) is not None,
                "italic": rpr is not None and rpr.find("w:i", NS) is not None,
                "underlined": underline is not None and underline.get(W + "val", "single") not in ("0", "false", "none"),
            })
    return rows


def direct_paragraph_properties(docx: Path) -> list[dict[str, Any]]:
    with zipfile.ZipFile(docx) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    rows: list[dict[str, Any]] = []
    for p_index, p in enumerate(root.findall(".//w:body/w:p", NS)):
        text = paragraph_text(p).strip()
        if not text:
            rows.append({"paragraph": p_index, "text": "", "blank": True})
            continue
        ppr = p.find("w:pPr", NS)
        spacing = ppr.find("w:spacing", NS) if ppr is not None else None
        ind = ppr.find("w:ind", NS) if ppr is not None else None
        jc = ppr.find("w:jc", NS) if ppr is not None else None
        rows.append({
            "paragraph": p_index,
            "text": text,
            "blank": False,
            "align": jc.get(W + "val") if jc is not None else None,
            "before": spacing.get(W + "before") if spacing is not None else None,
            "after": spacing.get(W + "after") if spacing is not None else None,
            "line": spacing.get(W + "line") if spacing is not None else None,
            "line_rule": spacing.get(W + "lineRule") if spacing is not None else None,
            "first_line": ind.get(W + "firstLine") if ind is not None else None,
            "left": ind.get(W + "left") if ind is not None else None,
            "keep_next": ppr is not None and ppr.find("w:keepNext", NS) is not None,
        })
    return rows


def verify_command(args: argparse.Namespace) -> int:
    content_path = Path(args.content).expanduser().resolve()
    docx = Path(args.docx).expanduser().resolve()
    out = Path(args.out).expanduser().resolve()
    template = Path(args.template).expanduser().resolve() if args.template else DEFAULT_TEMPLATE
    content = json.loads(content_path.read_text(encoding="utf-8"))
    errors = validate_content(content)
    warnings: list[str] = []
    metrics: dict[str, Any] = {}
    completeness = content_completeness_findings(content)
    metrics["content_completeness_status"] = completeness["status"]
    metrics["content_richness"] = completeness["metrics"]
    warnings.extend(completeness.get("warnings", []))
    readiness = delivery_readiness_findings(content)
    metrics["delivery_readiness_status"] = readiness["status"]
    metrics["delivery_readiness"] = readiness["metrics"]
    warnings.extend(readiness.get("warnings", []))
    binding_errors, public_metrics = validate_public_binding(content, content_path)
    errors.extend(binding_errors)
    metrics["public_verification"] = public_metrics
    if not docx.is_file():
        errors.append(f"DOCX not found: {docx}")
        qa = {
            "pass": False,
            "created_at": dt.datetime.now().astimezone().isoformat(),
            "errors": errors,
            "warnings": warnings,
            "metrics": metrics,
            "visual_qa": "not_run_missing_docx",
        }
        write_json(out, qa)
        print(json.dumps({"pass": False, "errors": len(errors), "warnings": len(warnings),
                          "qa": str(out)}, ensure_ascii=False))
        return 2
    metadata_errors, metadata_metrics = metadata_contract_findings(docx, content)
    errors.extend(metadata_errors)
    metrics["document_metadata"] = metadata_metrics

    with zipfile.ZipFile(docx) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
        visible = "\n".join(paragraph_text(p) for p in root.findall(".//w:body/w:p", NS))
        sect = root.find(".//w:sectPr", NS)
        if sect is None:
            errors.append("missing section properties")
        else:
            pgsz = sect.find("w:pgSz", NS)
            expected_size = {"w": "11906", "h": "16838"}
            actual_size = {k: pgsz.get(W + k) if pgsz is not None else None for k in expected_size}
            metrics["page_size_dxa"] = actual_size
            if actual_size != expected_size:
                errors.append(f"page size differs from A4 portrait: {actual_size}")
            pgmar = sect.find("w:pgMar", NS)
            expected = {"top": "1440", "bottom": "1440", "left": "1800", "right": "1800"}
            actual = {k: pgmar.get(W + k) if pgmar is not None else None for k in expected}
            metrics["margins_dxa"] = actual
            if actual != expected:
                errors.append(f"page margins differ from authoritative 1.0/1.25 inch sample: {actual}")

    for section in REQUIRED_SECTIONS:
        if section not in visible:
            errors.append(f"missing visible section: {section}")
    for residue in TEMPLATE_RESIDUE:
        if residue in visible:
            errors.append(f"template residue detected: {residue}")

    sections_by_heading = {s.get("heading"): s for s in content.get("sections", [])}
    reason_blocks = sections_by_heading.get("投资理由", {}).get("blocks", [])
    reason_hits: list[dict[str, Any]] = []
    for index, block in enumerate(reason_blocks, start=1):
        text = block.get("text", "")
        hits = [phrase for phrase in REASON_DEFENSIVE_PHRASES if phrase in text]
        if hits:
            reason_hits.append({"item": index, "phrases": hits})
    if reason_hits:
        errors.append(f"investment reasons contain defensive tails: {reason_hits}")

    defensive_counts = {phrase: visible.count(phrase) for phrase in DEFENSIVE_PHRASES}
    repeated_defensive = {phrase: count for phrase, count in defensive_counts.items() if count > 1}
    if repeated_defensive:
        warnings.append(f"repeated defensive phrases detected: {repeated_defensive}")

    plan_blocks = sections_by_heading.get("投资计划", {}).get("blocks", [])
    plan_reservation_text = plan_blocks[-1].get("text", "").strip() if plan_blocks else ""
    if not plan_reservation_text:
        errors.append("investment plan requires a final reservation sentence")
    else:
        if len(plan_reservation_text) > PLAN_RESERVATION_MAX_CHARS:
            warnings.append(
                f"investment-plan reservation exceeds {PLAN_RESERVATION_MAX_CHARS} characters: "
                f"{len(plan_reservation_text)}"
            )
        if "正式交易文件为准" not in plan_reservation_text:
            warnings.append("investment-plan reservation should end with definitive documents")

    analysis_blocks = sections_by_heading.get("投资情形分析", {}).get("blocks", [])
    numbered_analysis = [b for b in analysis_blocks if b.get("type") == "numbered"]
    semantic = decision_semantic_findings(content)
    item7_text = numbered_analysis[6].get("text", "") if len(numbered_analysis) >= 7 else ""
    item7_hits = [marker for marker in ITEM7_CHECKLIST_MARKERS if marker in item7_text]
    if item7_hits:
        errors.append(f"item 7 contains a closing checklist; move it to open_issues: {item7_hits}")
    conclusion_blocks = [b for b in analysis_blocks if b.get("type") == "conclusion"]
    if len(conclusion_blocks) != 1:
        errors.append(f"requires exactly one conclusion block; got {len(conclusion_blocks)}")
    elif conclusion_blocks:
        conclusion_text = conclusion_blocks[0].get("text", "").strip()
        sentence_count = len(re.findall(r"[。！？]", conclusion_text))
        if sentence_count != 1:
            errors.append(f"conclusion must be one sentence; got {sentence_count}")
        if len(conclusion_text) > 120:
            warnings.append(f"conclusion exceeds 120 characters: {len(conclusion_text)}")

    unsourced: list[str] = []
    pending = 0
    for section in content.get("sections", []):
        for block in section.get("blocks", []):
            if block.get("type") != "subheading" and not block.get("source_ids"):
                unsourced.append(block.get("label", "") + block.get("text", "")[:50])
            if block.get("status") == "pending":
                pending += 1
    if unsourced:
        warnings.append(f"{len(unsourced)} factual blocks lack source_ids")
    if pending:
        warnings.append(f"{pending} blocks are pending verification")
    if not content.get("open_issues"):
        warnings.append("open_issues is empty; confirm this is intentional")

    runs = direct_run_properties(docx)
    font_errors = []
    decoration_errors = []
    bold_role_errors = []
    subheading_texts = {
        block.get("text", "")
        for section in content.get("sections", [])
        for block in section.get("blocks", [])
        if block.get("type") == "subheading"
    }
    for row in runs:
        is_title = row["paragraph"] == 0
        expected_font = "黑体" if is_title else "宋体"
        expected_size = 28 if is_title else 24
        if (
            row["east_asia_font"] != expected_font
            or row["ascii_font"] != expected_font
            or row["hansi_font"] != expected_font
            or row["half_points"] != expected_size
        ):
            font_errors.append(row)
        if row["italic"] or row["underlined"]:
            decoration_errors.append(row)
        paragraph_text_value = row["paragraph_text"]
        if is_title and row["bold"]:
            bold_role_errors.append(row)
        elif paragraph_text_value in REQUIRED_SECTIONS and not row["bold"]:
            bold_role_errors.append(row)
        elif paragraph_text_value in subheading_texts and not row["bold"]:
            bold_role_errors.append(row)
        elif paragraph_text_value[:1].isdigit() and "、" in paragraph_text_value[:3]:
            label_match = re.match(r"^\d+[、.)）]\s*", paragraph_text_value)
            remainder = paragraph_text_value[label_match.end():] if label_match else paragraph_text_value
            lead, _ = split_numbered_lead(remainder)
            expected_bold_text = (label_match.group(0) if label_match else "") + lead
            if row["bold"] != (row["text"] == expected_bold_text):
                bold_role_errors.append(row)
    if font_errors:
        errors.append(f"{len(font_errors)} runs differ from required font/size/western-font contract")
    if decoration_errors:
        errors.append(f"{len(decoration_errors)} runs use forbidden italic or underline formatting")
    if bold_role_errors:
        errors.append(f"{len(bold_role_errors)} runs differ from required bold roles")
    metrics["visible_runs"] = len(runs)
    metrics["font_size_exceptions"] = len(font_errors)
    metrics["decoration_exceptions"] = len(decoration_errors)
    metrics["bold_role_exceptions"] = len(bold_role_errors)

    paragraphs = direct_paragraph_properties(docx)
    paragraph_errors: list[dict[str, Any]] = []
    blank_paragraphs = [row for row in paragraphs if row["blank"]]
    company_text = content.get("closing", {}).get("company", "")
    date_text = format_closing_date(content.get("closing", {}).get("date", ""))
    conclusion_text = semantic.get("conclusion_text", "")
    for row in paragraphs:
        if row["blank"]:
            continue
        text_value = row["text"]
        if row["paragraph"] == 0:
            expected_p = {"align": "center", "before": None, "after": None, "first_line": None, "left": None, "keep_next": False}
        elif text_value in REQUIRED_SECTIONS:
            expected_p = {"align": None, "before": "240", "after": None, "first_line": None, "left": None, "keep_next": True}
        elif text_value in subheading_texts:
            expected_p = {"align": None, "before": None, "after": None, "first_line": "482", "left": "0", "keep_next": True}
        elif text_value == company_text:
            expected_p = {"align": "right", "before": "240", "after": None, "first_line": "420", "left": "300", "keep_next": True}
        elif text_value == date_text:
            expected_p = {"align": "right", "before": "240", "after": None, "first_line": "420", "left": "300", "keep_next": False}
        elif text_value[:1].isdigit() and "、" in text_value[:3]:
            expected_p = {"align": None, "before": "240", "after": None, "first_line": "0", "left": None, "keep_next": False}
        elif text_value == conclusion_text:
            expected_p = {"align": None, "before": "240", "after": None, "first_line": "420", "left": None, "keep_next": False}
        else:
            expected_p = {"align": None, "before": None, "after": None, "first_line": "480", "left": None, "keep_next": False}
        actual_p = {key: row[key] for key in expected_p}
        if actual_p != expected_p or row["line"] != "360" or row["line_rule"] != "auto":
            paragraph_errors.append({
                "paragraph": row["paragraph"],
                "text": text_value[:60],
                "expected": expected_p | {"line": "360", "line_rule": "auto"},
                "actual": actual_p | {"line": row["line"], "line_rule": row["line_rule"]},
            })
    rows_by_index = {row["paragraph"]: row for row in paragraphs}
    blank_followers = [
        rows_by_index.get(row["paragraph"] + 1, {}).get("text")
        for row in blank_paragraphs
    ]
    if len(blank_paragraphs) != 2 or blank_followers != ["投资情形分析", company_text]:
        errors.append(
            "blank transition paragraphs differ from authoritative sample: "
            f"count={len(blank_paragraphs)}, followers={blank_followers}"
        )
    if paragraph_errors:
        errors.append(f"{len(paragraph_errors)} paragraphs differ from alignment/spacing/indent contract")
    metrics["blank_paragraph_count"] = len(blank_paragraphs)
    metrics["paragraph_format_exceptions"] = len(paragraph_errors)

    final_hashes = package_hashes(docx)
    template_hashes = package_hashes(template)
    allowed_metadata_changes = {
        "[Content_Types].xml", "_rels/.rels", "docProps/core.xml",
        "docProps/app.xml", "docProps/custom.xml",
    }
    changed_parts = sorted(
        name for name, digest in template_hashes.items()
        if name != "word/document.xml"
        and name not in allowed_metadata_changes
        and final_hashes.get(name) != digest
    )
    missing_parts = sorted(
        name for name in template_hashes
        if name not in final_hashes and name not in allowed_metadata_changes
    )
    if changed_parts:
        errors.append(f"preserve-only package parts changed: {changed_parts}")
    if missing_parts:
        errors.append(f"template package parts missing: {missing_parts}")
    metrics["preserve_only_parts_changed"] = changed_parts
    metrics["template_parts_missing"] = missing_parts
    metrics["docx_sha256"] = sha256(docx)
    metrics["template_sha256"] = sha256(template)
    metrics["open_issue_count"] = len(content.get("open_issues", []))
    metrics["pending_block_count"] = pending
    metrics["reason_defensive_hits"] = reason_hits
    metrics["defensive_phrase_counts"] = defensive_counts
    metrics["conclusion_character_count"] = len(conclusion_blocks[0].get("text", "")) if conclusion_blocks else 0
    metrics["conclusion_semantic_status"] = "pass" if semantic["conclusion_pass"] else "fail"
    metrics["compliance_lead_status"] = "pass" if not semantic["invalid_items"] else "fail"
    metrics["compliance_invalid_leads"] = semantic["invalid_items"]
    metrics["compliance_role_map_status"] = "pass" if not semantic["invalid_roles"] else "fail"
    metrics["compliance_invalid_roles"] = semantic["invalid_roles"]
    metrics["analysis_process_language_status"] = (
        "pass" if not semantic["invalid_process_items"] else "fail"
    )
    metrics["analysis_process_language_hits"] = semantic["invalid_process_items"]
    metrics["conclusion_forbidden_hits"] = semantic["conclusion_forbidden"]
    metrics["plan_reservation_character_count"] = len(plan_reservation_text)
    metrics["item7_checklist_hits"] = item7_hits
    qa = {
        "pass": not errors,
        "created_at": dt.datetime.now().astimezone().isoformat(),
        "errors": errors, "warnings": warnings, "metrics": metrics,
        "visual_qa": "required_separately_by_documents_skill",
    }
    write_json(out, qa)
    print(json.dumps({"pass": qa["pass"], "errors": len(errors),
                      "warnings": len(warnings), "qa": str(out)}, ensure_ascii=False))
    return 0 if qa["pass"] else 2


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prepare = sub.add_parser("prepare", help="extract and index a ZIP/directory")
    prepare.add_argument("input")
    prepare.add_argument("--workdir", required=True)
    prepare.set_defaults(func=prepare_command)

    build = sub.add_parser("build", help="build DOCX from content JSON")
    build.add_argument("--content", required=True)
    build.add_argument("--output", required=True)
    build.add_argument("--template")
    build.set_defaults(func=build_command)

    verify = sub.add_parser("verify", help="structurally verify content and DOCX")
    verify.add_argument("--content", required=True)
    verify.add_argument("--docx", required=True)
    verify.add_argument("--out", required=True)
    verify.add_argument("--template")
    verify.set_defaults(func=verify_command)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
