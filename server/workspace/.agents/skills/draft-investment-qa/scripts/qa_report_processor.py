#!/usr/bin/env python3
"""Deterministic controller for evidence-backed investment-committee Q&A reports.

The current Claude Code model owns fact extraction, conflict rulings, question design,
analysis and writing. This program owns reproducible input preparation, schema and
quality gates, DOCX rendering, and package verification. It never calls an LLM or
an external model gateway and has no dependency on vc_docs_agent.py.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import zipfile
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:  # pragma: no cover - exercised by dependency preflight
    raise SystemExit("缺少 python-docx；请先运行 python3 -m pip install -r requirements.txt") from exc


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{WORD_NS}}}"

MIN_QUESTIONS = 6
MAX_QUESTIONS = 9
MIN_ANSWER_CHARS = 450
MAX_ANSWER_CHARS = 1000
MIN_NUMBERS_PER_ANSWER = 4
# Songti SC is available in the target macOS/Word environment and survives
# LibreOffice's DOCX-to-PDF path with complete Simplified Chinese glyphs.
FONT_NAME = "Songti SC"
TITLE_PT = 17.5
BODY_PT = 10.5
# STKaiti is the Office/LibreOffice-compatible family alias for macOS Kaiti.
# It preserves the Deta reference's 楷体 look and avoids missing-glyph boxes in
# headless DOCX-to-PDF rendering, while Word resolves it to the same Kaiti asset.
DETA_CJK_FONT_NAME = "STKaiti"
DETA_LATIN_FONT_NAME = "Times New Roman"
# PDF extraction reports 15.95/11.05 pt after WPS rendering; Word stores these
# as the corresponding half-point sizes 16/11 pt in the DOCX package.
DETA_TITLE_PT = 16.0
DETA_BODY_PT = 11.0
DETA_LINE_PT = 14.4
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_TOP_BOTTOM_CM = 1.55
MARGIN_LEFT_RIGHT_CM = 3.05
DETA_MARGIN_TOP_BOTTOM_CM = 2.54
DETA_MARGIN_LEFT_RIGHT_CM = 3.175

PLAN_FIELDS = (
    "id",
    "question",
    "angle",
    "question_archetype",
    "answer_structure",
    "evidence_focus",
    "queries",
    "must_cover_fact_ids",
    "mechanism_chain",
    "competing_explanation",
    "discriminator",
    "economic_bridge",
    "valuation_treatment",
    "boundary",
)

QUESTION_ARCHETYPES = {
    "positioning", "market_need", "mechanism", "customer_case",
    "transferability", "product_role", "revenue_quality", "valuation", "organization",
}

EXTERNAL_SCORE_DIMENSIONS = {
    "核心争议选择与问题锋利度": 10,
    "数据颗粒度与证据状态": 20,
    "赛道定位、产业分工与对标": 10,
    "公司专属机制与商业模式": 15,
    "客户收入效率与现金流传导": 20,
    "团队执行与商业化能力": 10,
    "结论清晰与自然表达": 15,
}

INTERNAL_SCORE_DIMENSIONS = {
    "核心争议选择与问题锋利度": 10,
    "数据颗粒度与证据状态": 15,
    "赛道定位、产业分工与对标": 10,
    "公司专属机制与商业模式": 15,
    "客户收入毛利现金流传导": 15,
    "治理、IP 与组织完整度": 10,
    "风险处理与结论闭合": 10,
    "估值、付款与交易安排": 10,
    "自然清晰的专业表达": 5,
}

# Backward-compatible alias for legacy/internal callers.
SCORE_DIMENSIONS = INTERNAL_SCORE_DIMENSIONS

LEGACY_SCORE_DIMENSIONS = {
    "核心争议选择与问题锋利度": 15,
    "数据颗粒度与证据状态": 15,
    "公司专属经营机制": 20,
    "客户收入毛利现金流传导": 15,
    "竞争性解释与证伪指标": 10,
    "估值付款与资本配置处理": 15,
    "自然叙事与非清单化表达": 10,
}

MANDATORY_DUAL_USE_MODULES = {
    "valuation_benchmark",
    "business_flywheel",
    "financial_reconciliation",
    "governance_ip",
    "commercialization_organization",
}

BANNED_ANSWER_PHRASES = (
    "尚未确定",
    "不宜作出",
    "无法判断",
    "有待观察后再议",
    "当前材料不足以支持",
    "现有材料不足以支持",
    "当前证据尚不足以",
    "现有证据尚不足以",
    "不能确认",
    "无法确认",
    "尚不能确认",
    "目前只能确认",
    "承认边界：",
)

INTERNAL_TRACE_PHRASES = (
    "事实ID",
    "事实 ID",
    "证据包",
    "冲突裁决",
    "红队",
    "质量评分",
    "内部底稿",
    "工作底稿",
    "source_id",
    "fact_id",
)

DUAL_USE_BANNED_PHRASES = (
    "反向解释",
    "另一种解释",
    "估值泡沫",
    "存在隐患",
    "不存在被替代风险",
    "长期不会被替代",
    "无法复制的核心壁垒",
    "无控制权旁落风险",
    "不存在知识产权侵权隐患",
    "全方位资源赋能",
    "阶段约束",
    "价值释放条件",
    "治理资产",
    "资本上限",
    "成长锚",
    "对称资本处理",
    "平台飞轮",
    "建议公司",
    "公司应当",
    "公司应该",
    "公司需要",
    "公司还需要",
    "后续需要",
    "后续建议",
    "接下来重点观察",
    "接下来重点看",
    "建议进一步",
    "建议补充",
    "建议设置",
    "建议进行",
    "投后需完成",
    "尚待验证",
    "仍有待证明",
    "提交审议",
    "纳入投后",
    "投后管理",
    "纳入交割安排",
    "拟列为交割条件",
)

EXTERNAL_QA_BANNED_PHRASES = (
    "经项目组核查", "项目组判断", "项目组认为", "项目组支持",
    "未计入本次定价", "不计入价值", "未纳入估值",
    "分期支付", "款项暂缓", "专项预留", "原股东补偿", "相应款项不释放",
    "据公司披露", "据管理层", "管理层表示", "管理层称", "公司称",
    "这说明", "共同证明", "这体现", "这表明", "由此可见", "据此",
    "客户入口", "高维数据积累", "价值释放", "成长锚", "治理资产", "平台飞轮", "资本上限", "生态闭环",
    "不是以远期技术设想作为主要定价依据", "不包含意向金额",
    "采购范围随着任务往下走自然变宽",
    "这个顺序比单纯的采购金额更能解释复购",
    "对4.5亿元最有分量的支撑",
)

POLISHED_META_PATTERNS = (
    r"这个顺序比.{0,24}更能解释",
    r"对.{0,16}最有分量的支撑",
    r"采购范围随着.{0,18}自然变宽",
)

TECHNICAL_CHECKLIST_PATTERN = re.compile(
    r"(?:第一，)?数据处理能力.{0,60}(?:第二，)?(?:场景与仿真能力|仿真能力).{0,60}"
    r"(?:第三，)?项目交付能力"
)

EXTERNAL_DIRECT_ANSWER_PATTERN = re.compile(
    r"^(?:\d|合理|能|可以|是|已经验证|已得到验证|具备|清晰|较好|稳定|可靠|不存在|不会|当前|现有客户|"
    r"大衍|公司的|触觉手套|合成数据|真实异构数据|现有团队|团队|"
    r"追觅|行之途|客户|客户数量|现有收入|本轮)"
)

EXTERNAL_FRAGMENT_OPENING_PATTERN = re.compile(r"^(?:合理|能|可以|是|具备|清晰|较好)[。！!]" )

EXTERNAL_INVESTMENT_CLOSURE_PATTERN = re.compile(
    r"估值.{0,16}支撑|商业化能力|收入和回款质量|支撑本次投资判断|"
    r"稳定的?收入来源|业务增长来源|持续增长基础|交付.{0,8}能力|"
    r"技术能力.{0,16}(?:商业收入|收入)|投资价值|持续采购|继续采购|复购|"
    r"可重复销售|单独销售|跨行业.{0,12}(?:复制|复用)|客户.{0,12}(?:付费|采购)|"
    r"产品化方向|现金流业务|商业(?:化)?结果|客户任务|估值.{0,16}(?:解释|合理)|采购和回款|"
    r"当前收入|提供参照|按月采购|第二次订单|重复付费|按约结算|"
    r"下一张采购单|一项产品|收入和现金.{0,8}落地|可卖的.{0,8}(?:变多|增加)|"
    r"贡献主要收入|就是答案"
)

VISIBLE_LABEL_PATTERNS = (
    r"^(?:阶段约束|正向价值|产业逻辑|商业模式|价值释放条件|投资建议|风险提示|资本处理|判断依据)[：:]",
    r"^(?:第一层|第二层|第三层)[：:]",
)

GENERIC_CLOSING_PATTERN = re.compile(
    r"(?:已经|已)(?:形成|验证).{0,20}(?:基础|能力|价值)|构成.{0,16}(?:基础|投资价值)|"
    r"具备.{0,16}(?:能力|基础|质量|价值)"
)

CAPITAL_ACTION_PATTERN = re.compile(
    r"加速付款|后续付款|延迟付款|停止付款|不支付|重定价|专项预留|投资款抵扣|老股款|增资款"
)

ADVISORY_TASK_PATTERNS = (
    r"建议以",
    r"建议对",
    r"建议把",
    r"建议在",
    r"未来应",
    r"只有.{0,50}(?:才能证明|才可以证明|才说明)",
    r"(?:是否能够|能否).{0,28}仍需观察",
    r"尚未|尚无|仍需|仍在.{0,12}(?:形成|验证|观察|跟踪)",
    r"待完成|待验证|待跟踪",
    r"接下来|下一步",
)

OPENING_CONCLUSION_PATTERN = re.compile(
    r"经项目组核查|项目组(?:判断|认为|认定)|本次投资|本次定价|本次交易|现有证据已经"
)

CLOSING_CONCLUSION_PATTERN = re.compile(
    r"项目组(?:据此)?(?:维持|支持|认为|判断|提交)|不影响本次投资(?:判断|建议)|"
    r"不构成本次投资的否决因素|本次投资建议|本次(?:估值|定价|交易)(?:仅|未|不|采用|按)"
)

REPORT_STAGES = {
    "final_recommendation",
    "team_recommendation",
    "internal_approved",
    "terms_agreed",
    "signed_or_closing",
}

NEGATIVE_OPENING_PATTERNS = (
    r"^(?:结论是)?尚(?:未|不足)",
    r"^(?:结论是)?只能",
    r"^(?:结论是)?(?:无法|不能|不宜)",
    r"^(?:结论是)?.{0,12}(?:尚不足以支撑|只支撑部分|属于期权价值)",
)

SENSITIVE_CLAIM_PATTERNS = {
    "joint_development": r"联合研发",
    "ip_ownership": r"(?:知识产权|IP).{0,24}(?:100%归属|归属公司|归公司|商用权归)",
    "control_stability": r"(?:实控人|控制权).{0,16}稳定|不谋求控制权",
    "equity_incentive": r"股权激励|员工持股|四年分期兑现|离职回购",
    "noncompete": r"竞业限制|竞业协议|全员竞业",
    "non_substitutability": r"不可替代|不被替代|不存在被替代",
    "industrial_enablement": r"产业(?:股东|合作方).{0,24}(?:赋能|协同)|全方位资源赋能",
}

STAGE_CEILING_BENCHMARK_PATTERNS = (
    r"Physical\s+Intelligence",
    r"Skild\s+AI",
    r"海外.{0,12}(?:数十亿|百亿).{0,8}(?:美元|估值)",
)

DIRECT_BENCHMARK_JUSTIFICATION_PATTERNS = (
    r"(?:因此|由此|说明|证明).{0,28}(?:4(?:\.5)?亿|当前估值|本轮估值).{0,16}(?:合理|性价比|安全边际)",
    r"(?:4(?:\.5)?亿|当前估值|本轮估值).{0,18}(?:处于|属于).{0,12}(?:合理区间|性价比)",
)

BENCHMARK_DIFFERENCE_TOKENS = (
    "轮次", "阶段", "规模", "产品广度", "收入阶段", "不能直接套用", "赛道上限",
)

INFO_QUESTION_PATTERNS = (
    r"公司是做什么的",
    r"行业前景如何",
    r"市场空间(?:有多大|如何)",
    r"请介绍(?:一下)?公司",
    r"公司的核心团队(?:是谁|如何)",
)


def read_json(path: Path) -> Any:
    with path.open(encoding="utf-8") as stream:
        return json.load(stream)


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_extract_zip(archive: Path, destination: Path) -> None:
    destination = destination.resolve()
    with zipfile.ZipFile(archive) as bundle:
        for member in bundle.infolist():
            if member.is_dir() or member.filename.startswith("__MACOSX/"):
                continue
            target = (destination / member.filename).resolve()
            if destination not in target.parents:
                raise SystemExit(f"压缩包含不安全路径：{member.filename}")
            target.parent.mkdir(parents=True, exist_ok=True)
            with bundle.open(member) as source, target.open("wb") as output:
                shutil.copyfileobj(source, output)


def extract_docx(path: Path) -> str:
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""
    workbook = load_workbook(path, read_only=True, data_only=False)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"[工作表] {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(" | ".join(values))
    return "\n".join(lines)


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return ""
    presentation = Presentation(path)
    lines: list[str] = []
    for number, slide in enumerate(presentation.slides, 1):
        lines.append(f"[幻灯片 {number}]")
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text.strip():
                lines.append(text.strip())
    return "\n".join(lines)


def extract_csv(path: Path) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open(encoding=encoding, newline="") as stream:
                return "\n".join(" | ".join(row) for row in csv.reader(stream))
        except UnicodeDecodeError:
            continue
    return ""


def extract_text(path: Path) -> tuple[str, str | None]:
    suffix = path.suffix.lower()
    try:
        if suffix in {".txt", ".md", ".json", ".yaml", ".yml"}:
            for encoding in ("utf-8-sig", "gb18030"):
                try:
                    return path.read_text(encoding=encoding), None
                except UnicodeDecodeError:
                    continue
        if suffix in {".csv", ".tsv"}:
            return extract_csv(path), None
        if suffix == ".docx":
            return extract_docx(path), None
        if suffix == ".pdf":
            text = extract_pdf(path)
            return text, None if text.strip() else "PDF 无可提取文本，需 OCR"
        if suffix in {".xlsx", ".xlsm"}:
            text = extract_xlsx(path)
            return text, None if text.strip() else "Excel 未提取到可用单元格"
        if suffix == ".pptx":
            text = extract_pptx(path)
            return text, None if text.strip() else "PPTX 未提取到可用文本"
        return "", f"暂不支持直接解析 {suffix or '无扩展名'}"
    except Exception as exc:  # noqa: BLE001 - extraction failures belong in manifest
        return "", f"解析失败：{type(exc).__name__}: {exc}"


def chunk_text(text: str, size: int = 6000) -> Iterable[str]:
    compact = text.replace("\x00", "").strip()
    for start in range(0, len(compact), size):
        yield compact[start:start + size]


def prepare_input(input_path: Path, workdir: Path) -> None:
    if not input_path.exists():
        raise SystemExit(f"输入不存在：{input_path}")
    workdir.mkdir(parents=True, exist_ok=True)
    expanded = workdir / "expanded"
    packets_dir = workdir / "packets"
    if expanded.exists() or packets_dir.exists():
        raise SystemExit("工作目录已包含 expanded/ 或 packets/；请使用新的空目录")
    expanded.mkdir()
    packets_dir.mkdir()

    if input_path.is_dir():
        source_root = input_path
    elif input_path.suffix.lower() == ".zip":
        safe_extract_zip(input_path, expanded)
        source_root = expanded
    else:
        target = expanded / input_path.name
        shutil.copy2(input_path, target)
        source_root = expanded

    files = sorted(
        path for path in source_root.rglob("*")
        if path.is_file() and not path.name.startswith(".")
    )
    manifest: list[dict[str, Any]] = []
    evidence_sections: list[str] = []
    packet_number = 0
    for source_number, path in enumerate(files, 1):
        source_id = f"S{source_number:04d}"
        relative = str(path.relative_to(source_root))
        text, warning = extract_text(path)
        packet_ids: list[str] = []
        for part_number, chunk in enumerate(chunk_text(text), 1):
            packet_number += 1
            packet_id = f"P{packet_number:05d}"
            packet_path = packets_dir / f"{packet_id}.txt"
            packet_path.write_text(
                f"source_id: {source_id}\npath: {relative}\npart: {part_number}\n\n{chunk}",
                encoding="utf-8",
            )
            packet_ids.append(packet_id)
        manifest.append({
            "source_id": source_id,
            "path": relative,
            "extension": path.suffix.lower(),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "characters": len(text),
            "packet_ids": packet_ids,
            "warning": warning,
        })
        if text.strip():
            evidence_sections.append(f"## {source_id}｜{relative}\n\n{text}")

    write_json(workdir / "source_manifest.json", {
        "input": str(input_path.resolve()),
        "source_count": len(manifest),
        "packet_count": packet_number,
        "sources": manifest,
    })
    (workdir / "evidence_pack.md").write_text(
        "# 投资 Q&A 证据包\n\n" + "\n\n".join(evidence_sections), encoding="utf-8"
    )
    write_json(workdir / "facts.json", [])
    write_json(workdir / "rulings.json", {"rulings": [], "blacklist": []})
    write_json(workdir / "qa_plan.json", {
        "profile": {
            "narrative_mode": "dual_use",
            "sector_map": [],
            "project_position": "",
            "project_position_fact_ids": [],
            "coverage": {
                module: [] for module in sorted(MANDATORY_DUAL_USE_MODULES)
            },
        },
        "questions": [],
    })
    write_json(workdir / "qa_content.json", {
        "meta": {
            "company": "",
            "title": "",
            "investor": "",
            "date": "",
            "narrative_mode": "dual_use",
            "audience_mode": "external_decision_qa",
            "report_stage": "final_recommendation",
            "investment_stance": "support",
        },
        "overview": {
            "market_question": "",
            "players": [],
            "project_position": "",
            "used_fact_ids": [],
        },
        "items": [],
    })
    write_json(workdir / "quality_scorecard.json", {
        "dimensions": {
            name: {"score": 0, "max": maximum, "reason": ""}
            for name, maximum in EXTERNAL_SCORE_DIMENSIONS.items()
        },
        "total": 0,
        "hard_failures": [],
        "revision_routes": [],
    })
    write_json(workdir / "revision_log.json", [])
    print(f"prepared {len(manifest)} sources and {packet_number} packets in {workdir}")


def normalized_numbers(text: str) -> set[str]:
    return {
        match.replace(",", "")
        for match in re.findall(r"(?<![A-Za-z])\d[\d,.]*", text)
        if len(match.replace(",", "")) >= 3
    }


def answer_text(item: dict[str, Any]) -> str:
    paragraphs = item.get("answer_paragraphs") or []
    return "\n".join(str(paragraph).strip() for paragraph in paragraphs if str(paragraph).strip())


def fact_status_rank(status: str) -> int:
    return {
        "pending": 0,
        "management_stated": 1,
        "calculated": 2,
        "verified": 3,
    }.get(status, -1)


def narrative_mode(plan: Any, content: Any) -> str:
    plan_profile = plan.get("profile", {}) if isinstance(plan, dict) else {}
    meta = content.get("meta", {}) if isinstance(content, dict) else {}
    return str(
        meta.get("narrative_mode")
        or plan_profile.get("narrative_mode")
        or ("reference_faithful" if meta.get("format_profile") == "deta_qa_pdf" else "legacy")
    )


def load_artifacts(root: Path) -> dict[str, Any]:
    required = (
        "facts.json",
        "rulings.json",
        "qa_plan.json",
        "qa_content.json",
        "quality_scorecard.json",
        "revision_log.json",
    )
    missing = [name for name in required if not (root / name).is_file()]
    if missing:
        raise SystemExit("缺少阶段文件：" + "、".join(missing))
    return {name[:-5]: read_json(root / name) for name in required}


def add_issue(issues: list[dict[str, str]], gate: str, level: str, message: str) -> None:
    issues.append({"gate": gate, "level": level, "message": message})


def audit_artifacts(root: Path) -> dict[str, Any]:
    data = load_artifacts(root)
    facts = data["facts"]
    rulings = data["rulings"]
    plan = data["qa_plan"]
    content = data["qa_content"]
    scorecard = data["quality_scorecard"]
    revision_log = data["revision_log"]
    issues: list[dict[str, str]] = []

    if not isinstance(facts, list) or not facts:
        add_issue(issues, "evidence", "error", "facts.json 为空或不是数组")
        facts = []
    fact_map: dict[str, dict[str, Any]] = {}
    for index, fact in enumerate(facts, 1):
        if not isinstance(fact, dict) or not fact.get("id") or not fact.get("fact"):
            add_issue(issues, "evidence", "error", f"事实第 {index} 条缺少 id 或 fact")
            continue
        fact_id = str(fact["id"])
        if fact_id in fact_map:
            add_issue(issues, "evidence", "error", f"事实 ID 重复：{fact_id}")
        fact_map[fact_id] = fact
        if not fact.get("source_ids"):
            add_issue(issues, "evidence", "error", f"事实 {fact_id} 缺少 source_ids")
        if fact.get("status") not in {"verified", "management_stated", "calculated", "pending"}:
            add_issue(issues, "evidence", "error", f"事实 {fact_id} 的 status 无效")

    if not isinstance(rulings, dict) or not {"rulings", "blacklist"}.issubset(rulings):
        add_issue(issues, "rulings", "error", "rulings.json 缺少 rulings 或 blacklist")
        rulings = {"rulings": [], "blacklist": []}

    mode = narrative_mode(plan, content)
    dual_use = mode == "dual_use"
    decision_narrative = mode in {"dual_use", "reference_faithful"}
    if mode not in {"dual_use", "internal_ic", "reference_faithful", "legacy"}:
        add_issue(issues, "plan", "error", f"未知 narrative_mode：{mode}")

    meta = content.get("meta", {}) if isinstance(content, dict) else {}
    if mode == "reference_faithful" and meta.get("format_profile") != "deta_qa_pdf":
        add_issue(
            issues,
            "format",
            "error",
            "reference_faithful 必须显式设置 format_profile=deta_qa_pdf",
        )
    audience_mode = str(meta.get("audience_mode") or "internal_decision_qa")
    report_stage = str(meta.get("report_stage") or "")
    investment_stance = str(meta.get("investment_stance") or "")
    if decision_narrative:
        if audience_mode not in {"external_decision_qa", "internal_decision_qa"}:
            add_issue(issues, "stance", "error", "qa_content.meta 缺少有效 audience_mode")
        if report_stage not in REPORT_STAGES:
            add_issue(issues, "stance", "error", "qa_content.meta 缺少有效 report_stage")
        if investment_stance not in {"support", "oppose"}:
            add_issue(issues, "stance", "error", "qa_content.meta 缺少明确 investment_stance")

    profile = plan.get("profile", {}) if isinstance(plan, dict) else {}
    if decision_narrative:
        sector_map = profile.get("sector_map") or []
        expected_categories = {
            "pure_world_model", "embodiment_oem_internal", "outsourced_data_service"
        }
        actual_categories = {
            str(entry.get("category")) for entry in sector_map if isinstance(entry, dict)
        }
        if actual_categories != expected_categories:
            add_issue(issues, "plan", "error", f"{mode} 赛道三分法不完整")
        for entry in sector_map if isinstance(sector_map, list) else []:
            if not isinstance(entry, dict):
                continue
            fact_ids = entry.get("fact_ids") or []
            if not fact_ids:
                add_issue(issues, "plan", "error", f"赛道类别 {entry.get('category')} 缺少事实支持")
            for fact_id in fact_ids:
                if str(fact_id) not in fact_map:
                    add_issue(issues, "plan", "error", f"赛道定位引用不存在的事实 {fact_id}")
        position_ids = profile.get("project_position_fact_ids") or []
        if not profile.get("project_position") or len(position_ids) < 2:
            add_issue(issues, "plan", "error", f"{mode} 缺少项目定位或至少 2 条事实支持")
        for fact_id in position_ids:
            if str(fact_id) not in fact_map:
                add_issue(issues, "plan", "error", f"项目定位引用不存在的事实 {fact_id}")

    questions = plan.get("questions") if isinstance(plan, dict) else None
    if not isinstance(questions, list):
        add_issue(issues, "plan", "error", "qa_plan.json 缺少 questions 数组")
        questions = []
    if not MIN_QUESTIONS <= len(questions) <= MAX_QUESTIONS:
        add_issue(issues, "plan", "error", f"质疑数量为 {len(questions)}，必须为 6—9 个")

    plan_map: dict[str, dict[str, Any]] = {}
    covered_modules: set[str] = set()
    question_archetypes: set[str] = set()
    numbered_question_count = 0
    numeric_question_count = 0
    evidence_focus_signatures: list[tuple[str, ...]] = []
    for index, question in enumerate(questions, 1):
        if not isinstance(question, dict):
            add_issue(issues, "plan", "error", f"Q{index} 不是对象")
            continue
        missing = [field for field in PLAN_FIELDS if not question.get(field)]
        if missing:
            add_issue(issues, "plan", "error", f"Q{index} 缺少：{'、'.join(missing)}")
        qid = str(question.get("id", f"Q{index}"))
        if qid in plan_map:
            add_issue(issues, "plan", "error", f"问题 ID 重复：{qid}")
        plan_map[qid] = question
        module = str(question.get("module") or "")
        if decision_narrative:
            if not module:
                add_issue(issues, "plan", "error", f"{qid} 缺少 module")
            covered_modules.add(module)
        qtext = str(question.get("question", ""))
        archetype = str(question.get("question_archetype") or "")
        if decision_narrative and archetype not in QUESTION_ARCHETYPES:
            add_issue(issues, "plan", "error", f"{qid} 缺少有效 question_archetype")
        if archetype:
            question_archetypes.add(archetype)
        if question.get("answer_structure") == "numbered":
            numbered_question_count += 1
        evidence_focus = question.get("evidence_focus") or []
        if not isinstance(evidence_focus, list) or not 1 <= len(evidence_focus) <= 3:
            add_issue(issues, "plan", "error", f"{qid} evidence_focus 须为 1—3 项")
        else:
            evidence_focus_signatures.append(tuple(sorted(str(value).strip() for value in evidence_focus)))
        if re.search(r"\d", qtext):
            numeric_question_count += 1
        if audience_mode != "external_decision_qa" and not re.search(r"\d", qtext):
            add_issue(issues, "plan", "error", f"{qid} 未引用任何具体数字或期间")
        if not re.search(r"[？?]$", qtext.strip()):
            add_issue(issues, "plan", "warning", f"{qid} 未以问号收束")
        if any(re.search(pattern, qtext) for pattern in INFO_QUESTION_PATTERNS):
            add_issue(issues, "plan", "error", f"{qid} 是信息型提问而非尖锐质疑")
        if decision_narrative and question.get("answer_structure") not in {"numbered", "natural"}:
            add_issue(issues, "plan", "error", f"{qid} 缺少有效 answer_structure")
        must_ids = question.get("must_cover_fact_ids") or []
        minimum_must_ids = 1 if audience_mode == "external_decision_qa" else 3
        if not isinstance(must_ids, list) or not minimum_must_ids <= len(must_ids) <= 10:
            add_issue(issues, "plan", "error", f"{qid} 必用事实须为 {minimum_must_ids}—10 条")
        for fact_id in must_ids if isinstance(must_ids, list) else []:
            if str(fact_id) not in fact_map:
                add_issue(issues, "plan", "error", f"{qid} 引用了不存在的事实 {fact_id}")
        why_layer = str(question.get("why_layer") or "").strip()
        if decision_narrative and why_layer and len(why_layer) < 10:
            add_issue(issues, "plan", "error", f"{qid} why_layer 过于笼统")

    minimum_archetypes = 5 if len(questions) >= 8 else 3
    if decision_narrative and len(question_archetypes) < minimum_archetypes:
        add_issue(issues, "plan", "error", f"问题原型少于 {minimum_archetypes} 种，问法容易模板化")
    if audience_mode == "external_decision_qa" and questions:
        if numeric_question_count > len(questions) // 2:
            add_issue(issues, "plan", "error", "超过一半的对外问题依赖数字锚点，问题设计过于材料化")
        if numbered_question_count > len(questions) // 2:
            add_issue(issues, "plan", "error", "超过一半的对外回答规划为分点结构，报告模板化")
        for signature in set(evidence_focus_signatures):
            if signature and evidence_focus_signatures.count(signature) > len(questions) // 2:
                add_issue(issues, "plan", "error", f"超过一半问题重复使用同一证据组合：{'、'.join(signature)}")

    if decision_narrative:
        required_modules = (
            MANDATORY_DUAL_USE_MODULES - {"governance_ip"}
            if audience_mode == "external_decision_qa"
            else MANDATORY_DUAL_USE_MODULES
        )
        missing_modules = sorted(required_modules - covered_modules)
        if missing_modules:
            add_issue(issues, "plan", "error", f"{mode} 缺少必要模块：" + "、".join(missing_modules))
        coverage = profile.get("coverage") or {}
        for module in sorted(required_modules):
            mapped = coverage.get(module) if isinstance(coverage, dict) else None
            if not isinstance(mapped, list) or not mapped:
                add_issue(issues, "plan", "error", f"coverage 未映射模块 {module}")
            for qid in mapped if isinstance(mapped, list) else []:
                if str(qid) not in plan_map or str(plan_map[str(qid)].get("module")) != module:
                    add_issue(issues, "plan", "error", f"coverage 映射错误：{module} -> {qid}")
        commercialization_ids = coverage.get("commercialization_organization") if isinstance(coverage, dict) else []
        if isinstance(commercialization_ids, list) and len(commercialization_ids) != 1:
            add_issue(issues, "plan", "error", "商业化与组织扩张必须且只能由 1 个问题承载")

    items = content.get("items") if isinstance(content, dict) else None
    if not isinstance(items, list):
        add_issue(issues, "content", "error", "qa_content.json 缺少 items 数组")
        items = []
    if len(items) != len(questions):
        add_issue(issues, "content", "error", f"正文 {len(items)} 题与规划 {len(questions)} 题不一致")

    if dual_use:
        overview = content.get("overview") if isinstance(content, dict) else None
        if not isinstance(overview, dict):
            add_issue(issues, "content", "error", "dual_use 缺少 overview")
        else:
            if not overview.get("market_question") or not overview.get("project_position"):
                add_issue(issues, "content", "error", "overview 缺少市场争议或项目定位")
            players = overview.get("players") or []
            if not isinstance(players, list) or len(players) != 3:
                add_issue(issues, "content", "error", "overview.players 必须恰好三类")
            overview_ids = overview.get("used_fact_ids") or []
            if len(overview_ids) < 2:
                add_issue(issues, "content", "error", "overview 至少需要 2 条事实支持")
            for fact_id in overview_ids:
                if str(fact_id) not in fact_map:
                    add_issue(issues, "content", "error", f"overview 引用不存在的事实 {fact_id}")

    source_numbers = normalized_numbers(json.dumps(facts, ensure_ascii=False))
    source_numbers |= normalized_numbers(json.dumps(questions, ensure_ascii=False))
    seen_answers: list[tuple[str, str]] = []
    generic_closing_count = 0
    fact_usage: dict[str, list[str]] = {}
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            add_issue(issues, "content", "error", f"正文第 {index} 题不是对象")
            continue
        qid = str(item.get("id", f"Q{index}"))
        planned = plan_map.get(qid)
        if not planned:
            add_issue(issues, "content", "error", f"正文 {qid} 不在规划中")
            continue
        if str(item.get("question", "")).strip() != str(planned.get("question", "")).strip():
            add_issue(issues, "content", "error", f"{qid} 正文问题与规划问题不一致")
        paragraphs = item.get("answer_paragraphs") or []
        paragraph_minimum = 1 if audience_mode == "external_decision_qa" else 2
        paragraph_maximum = 4
        if not isinstance(paragraphs, list) or not paragraph_minimum <= len(paragraphs) <= paragraph_maximum:
            add_issue(
                issues, "content", "error",
                f"{qid} 回答须为 {paragraph_minimum}—{paragraph_maximum} 个自然段",
            )
        roles = item.get("paragraph_roles") or []
        allowed_roles = {
            "opening_position", "evidence_and_reasoning", "supporting_case",
            "business_implication", "open_issues", "transaction_treatment",
        }
        if not isinstance(roles, list) or len(roles) != len(paragraphs):
            add_issue(issues, "content", "error", f"{qid} 段落角色须与正文段数一致")
        elif any(role not in allowed_roles for role in roles):
            add_issue(issues, "content", "error", f"{qid} 包含无效段落角色")
        elif roles and roles[0] != "opening_position":
            add_issue(issues, "content", "error", f"{qid} 首段角色必须为 opening_position")
        text = answer_text(item)
        length = len(re.sub(r"\s+", "", text))
        answer_minimum = 1 if audience_mode == "external_decision_qa" else MIN_ANSWER_CHARS
        answer_maximum = 800 if audience_mode == "external_decision_qa" else MAX_ANSWER_CHARS
        if length < answer_minimum or length > answer_maximum:
            add_issue(issues, "content", "error", f"{qid} 回答 {length} 字，不在 {answer_minimum}—{answer_maximum} 字范围")
        planned_module = str(planned.get("module") or "")
        minimum_numbers = (
            2 if audience_mode == "external_decision_qa"
            and planned_module in {"valuation_benchmark", "financial_reconciliation"}
            else (MIN_NUMBERS_PER_ANSWER if audience_mode != "external_decision_qa" else 0)
        )
        if len(re.findall(r"\d[\d,.]*", text)) < minimum_numbers:
            add_issue(issues, "content", "error", f"{qid} 少于 {minimum_numbers} 个量化事实")
        for phrase in (*BANNED_ANSWER_PHRASES, *INTERNAL_TRACE_PHRASES):
            if phrase in text:
                add_issue(issues, "content", "error", f"{qid} 出现禁用表述“{phrase}”")
        for paragraph in paragraphs if isinstance(paragraphs, list) else []:
            visible = str(paragraph).strip()
            if any(re.search(pattern, visible) for pattern in VISIBLE_LABEL_PATTERNS):
                add_issue(issues, "narrative", "error", f"{qid} 使用了“小标题+冒号”的模板化段首")
        if decision_narrative:
            opening = str(paragraphs[0]).strip() if paragraphs else ""
            closing = str(paragraphs[-1]).strip() if paragraphs else ""
            if audience_mode != "external_decision_qa" and not OPENING_CONCLUSION_PATTERN.search(opening[:90]):
                add_issue(issues, "stance", "error", f"{qid} 首段未以项目组结论回答问题")
            if audience_mode != "external_decision_qa" and not CLOSING_CONCLUSION_PATTERN.search(closing[-120:]):
                add_issue(issues, "stance", "error", f"{qid} 未以明确投资判断收口")
            if audience_mode == "external_decision_qa" and EXTERNAL_FRAGMENT_OPENING_PATTERN.search(opening[:12]):
                add_issue(issues, "narrative", "error", f"{qid} 使用单字模板开场，须改为完整判断句")
            if audience_mode == "external_decision_qa" and GENERIC_CLOSING_PATTERN.search(closing[-90:]):
                generic_closing_count += 1
            answer_structure = str(planned.get("answer_structure") or "")
            numbered_hits = sum(bool(re.search(rf"{token}[，、：:]", text)) for token in ("第一", "第二", "第三"))
            if audience_mode == "external_decision_qa" and answer_structure == "numbered" and numbered_hits < 2:
                add_issue(issues, "narrative", "error", f"{qid} 规划为分点回答但缺少第一/第二框架")
            if audience_mode == "external_decision_qa" and answer_structure == "natural" and numbered_hits >= 2:
                add_issue(issues, "narrative", "error", f"{qid} 规划为自然回答却机械使用分点框架")
            if audience_mode == "external_decision_qa" and TECHNICAL_CHECKLIST_PATTERN.search(text):
                add_issue(issues, "narrative", "error", f"{qid} 以技术能力分类清单替代客户价值论证")
            if audience_mode == "external_decision_qa" and any(
                re.search(pattern, text) for pattern in POLISHED_META_PATTERNS
            ):
                add_issue(issues, "narrative", "error", f"{qid} 使用替听众概括证据分量的工整元话语")
            for pattern in ADVISORY_TASK_PATTERNS:
                if re.search(pattern, text):
                    add_issue(issues, "stance", "error", f"{qid} 出现顾问式或继续尽调式表达")
            if report_stage in {"final_recommendation", "team_recommendation"} and re.search(
                r"双方已(?:约定|确认)|交易文件已(?:明确|约定)|协议已(?:明确|约定)|已写入交易文件",
                text,
            ):
                add_issue(issues, "stance", "error", f"{qid} 将项目组口径误写为双方已约定")
            if report_stage == "final_recommendation" and audience_mode != "external_decision_qa" and not re.search(
                r"项目组(?:据此)?支持(?:按.{0,30})?本次投资|项目组维持本次投资建议",
                closing[-100:],
            ):
                add_issue(issues, "stance", "error", f"{qid} 最终结论稿未明确落脚于支持本次投资")
        if any(re.search(pattern, text, re.I) for pattern in STAGE_CEILING_BENCHMARK_PATTERNS):
            if any(re.search(pattern, text) for pattern in DIRECT_BENCHMARK_JUSTIFICATION_PATTERNS):
                add_issue(issues, "benchmark", "error", f"{qid} 用赛道上限直接证明项目当前估值合理")
            required_difference_tokens = 1 if audience_mode == "external_decision_qa" else 2
            if sum(token in text for token in BENCHMARK_DIFFERENCE_TOKENS) < required_difference_tokens:
                add_issue(issues, "benchmark", "error", f"{qid} 海外/成熟平台对标未充分披露阶段差异")
        if decision_narrative:
            for phrase in DUAL_USE_BANNED_PHRASES:
                if phrase in text:
                    add_issue(issues, "narrative", "error", f"{qid} 出现双用版禁用表述“{phrase}”")
            if audience_mode == "external_decision_qa":
                for phrase in EXTERNAL_QA_BANNED_PHRASES:
                    if phrase in text:
                        add_issue(issues, "narrative", "error", f"{qid} 对外答疑出现内部口径“{phrase}”")
            if any(re.search(pattern, opening) for pattern in NEGATIVE_OPENING_PATTERNS):
                add_issue(issues, "narrative", "error", f"{qid} 以防御性否定开头")
        missing_links = []
        planned_module = str(planned.get("module") or "")
        if planned_module in {"valuation_benchmark", "financial_reconciliation"}:
            required_links = [
                ("客户行为", r"采购|合同转化|复购|提价|交付效率|使用频率|客户ROI|投资回收期"),
                ("收入/现金流", r"收入|毛利|成本|现金流|回款|应收|利润"),
            ]
        elif planned_module == "commercialization_organization":
            required_links = [("商业结果", r"交付|采购|复购|收入|回款|销售")]
        else:
            required_links = [("客户或商业结果", r"客户任务|采购|复购|使用|交付|销售|收入|成本|回款")]
        if audience_mode != "external_decision_qa" and planned_module in {
            "valuation_benchmark", "financial_reconciliation", "governance_ip", "policy_transaction",
        }:
            required_links.append(
                ("资本配置", r"估值|定价|付款|支付|投资额度|基准价值|成长价值|期权价值|折价|里程碑|终止")
            )
        for label, pattern in required_links:
            if not re.search(pattern, text):
                missing_links.append(label)
        if missing_links:
            add_issue(issues, "content", "error", f"{qid} 因果链缺少：{'、'.join(missing_links)}")
        if decision_narrative and audience_mode != "external_decision_qa" and not re.search(
            r"未计入|不计入|已通过|已纳入|已列为|已设置|不影响|不构成|已在.{0,20}处理|按.{0,20}处理",
            text,
        ):
            add_issue(issues, "content", "error", f"{qid} 未说明风险已如何处理或为何不影响投资结论")
        requires_capital = planned_module in {
            "valuation_benchmark", "financial_reconciliation", "governance_ip", "policy_transaction",
        }
        if requires_capital and audience_mode != "external_decision_qa" and not CAPITAL_ACTION_PATTERN.search(text):
            add_issue(issues, "capital", "error", f"{qid} 需要交易处理但未说明付款、预留或重定价安排")
        checklist_hits = len(re.findall(r"需核验|建议取得|交割前|补充材料|待补|需提供", text))
        if checklist_hits > 2:
            add_issue(issues, "content", "error", f"{qid} 核验清单式表述达到 {checklist_hits} 处")

        used = {str(value) for value in (item.get("used_fact_ids") or [])}
        for fact_id in used:
            fact_usage.setdefault(fact_id, []).append(qid)
        missing_facts = [
            str(value) for value in planned.get("must_cover_fact_ids", [])
            if str(value) not in used
        ]
        if missing_facts:
            add_issue(issues, "content", "error", f"{qid} 未登记必用事实：{'、'.join(missing_facts)}")
        unknown_facts = sorted(used - set(fact_map))
        if unknown_facts:
            add_issue(issues, "content", "error", f"{qid} 使用不存在的事实：{'、'.join(unknown_facts)}")
        if audience_mode == "external_decision_qa":
            used_facts = [fact_map[fact_id] for fact_id in used if fact_id in fact_map]
            pending_ids = sorted(
                str(fact.get("id")) for fact in used_facts if fact.get("status") == "pending"
            )
            if pending_ids:
                add_issue(issues, "evidence", "error", f"{qid} 对外答疑使用待确认事实：{'、'.join(pending_ids)}")
        if decision_narrative:
            support_entries = item.get("claim_support") or []
            support_map: dict[str, list[dict[str, Any]]] = {}
            for support in support_entries if isinstance(support_entries, list) else []:
                if not isinstance(support, dict):
                    add_issue(issues, "claims", "error", f"{qid} claim_support 存在非对象")
                    continue
                claim_type = str(support.get("claim_type") or "")
                support_map.setdefault(claim_type, []).append(support)
                level = str(support.get("expression_level") or "")
                fact_ids = [str(value) for value in (support.get("fact_ids") or [])]
                if level not in {"confirmed", "management_stated", "conditional"} or not fact_ids:
                    add_issue(issues, "claims", "error", f"{qid} {claim_type} 的支持记录不完整")
                    continue
                support_facts = [fact_map.get(fact_id) for fact_id in fact_ids]
                if any(fact is None for fact in support_facts):
                    add_issue(issues, "claims", "error", f"{qid} {claim_type} 引用不存在的事实")
                    continue
                ranks = [fact_status_rank(str(fact.get("status"))) for fact in support_facts if fact]
                if level == "confirmed" and (not ranks or min(ranks) < fact_status_rank("verified")):
                    add_issue(issues, "claims", "error", f"{qid} {claim_type} 不得以 confirmed 表达")
                if level == "management_stated" and not re.search(r"公司称|管理层|据公司", text):
                    add_issue(issues, "claims", "error", f"{qid} {claim_type} 缺少管理层归属语")
                if level == "conditional" and not re.search(r"若|待|通过|需", text):
                    add_issue(issues, "claims", "error", f"{qid} {claim_type} 缺少条件表达")
            for claim_type, pattern in SENSITIVE_CLAIM_PATTERNS.items():
                if re.search(pattern, text) and claim_type not in support_map:
                    add_issue(issues, "claims", "error", f"{qid} 敏感断言 {claim_type} 缺少 claim_support")
        orphan = sorted(number for number in normalized_numbers(text)
                        if number not in source_numbers and not re.fullmatch(r"20[2-4]\d", number))
        if orphan:
            add_issue(issues, "content", "error", f"{qid} 数字无法溯源：{'、'.join(orphan[:8])}")
        for entry in rulings.get("blacklist", []):
            bad = entry.get("bad") if isinstance(entry, dict) else None
            if bad and str(bad) in text:
                add_issue(issues, "rulings", "error", f"{qid} 使用了否决口径“{bad}”")
        for prior_id, prior_text in seen_answers:
            ratio = SequenceMatcher(None, prior_text, text).ratio()
            if ratio > 0.82:
                add_issue(issues, "content", "error", f"{qid} 与 {prior_id} 高度重复（{ratio:.0%}）")
        seen_answers.append((qid, text))

    if decision_narrative:
        capital_heavy = sum(
            bool(CAPITAL_ACTION_PATTERN.search(answer_text(item)))
            for item in items if isinstance(item, dict)
        )
        if capital_heavy > 5:
            add_issue(issues, "narrative", "error", f"共有 {capital_heavy} 题重复展开付款/重定价，模板感过强")
        final_openings = [
            re.sub(r"[，。；：:].*$", "", str((item.get("answer_paragraphs") or [""])[-1]).strip())[:14]
            for item in items if isinstance(item, dict) and item.get("answer_paragraphs")
        ]
        for opening in set(final_openings):
            if opening in {"第一", "第二", "第三"}:
                continue
            if opening and final_openings.count(opening) >= 3:
                add_issue(issues, "narrative", "error", f"至少 3 题使用相同末段开头“{opening}”")
        if audience_mode == "external_decision_qa" and generic_closing_count > max(2, len(items) // 3):
            add_issue(issues, "narrative", "error", "过多问题使用“已形成/已验证基础”式通用结尾")
        if audience_mode == "external_decision_qa":
            maximum_reuse = max(3, len(items) // 2)
            for fact_id, qids in sorted(fact_usage.items()):
                if len(qids) > maximum_reuse:
                    add_issue(
                        issues,
                        "narrative",
                        "error",
                        f"事实 {fact_id} 在 {len(qids)} 题重复使用（{','.join(qids)}），缺少主要归属",
                    )

    raw_dimensions = scorecard.get("dimensions") if isinstance(scorecard, dict) else None
    total = 0.0
    if not isinstance(raw_dimensions, dict):
        add_issue(issues, "score", "error", "quality_scorecard.json 缺少 dimensions")
        raw_dimensions = {}
    active_dimensions = (
        EXTERNAL_SCORE_DIMENSIONS
        if audience_mode == "external_decision_qa"
        else INTERNAL_SCORE_DIMENSIONS
    )
    if mode == "legacy" and set(LEGACY_SCORE_DIMENSIONS).issubset(raw_dimensions):
        active_dimensions = LEGACY_SCORE_DIMENSIONS
    for name, maximum in active_dimensions.items():
        value = raw_dimensions.get(name, {})
        try:
            score = float(value.get("score"))
        except (AttributeError, TypeError, ValueError):
            score = -1
        if score < 0 or score > maximum:
            add_issue(issues, "score", "error", f"评分维度无效：{name}")
            score = 0
        if score < maximum * 0.6:
            add_issue(issues, "score", "error", f"评分维度低于 60%：{name} {score:g}/{maximum}")
        total += score
    if total < 85:
        add_issue(issues, "score", "error", f"语义总分 {total:g}/100 低于 85 分")
    if scorecard.get("hard_failures"):
        add_issue(issues, "score", "error", "quality_scorecard.json 仍有 hard_failures")
    if abs(float(scorecard.get("total", 0) or 0) - total) > 0.01:
        add_issue(issues, "score", "error", "quality_scorecard total 与各维度加总不一致")

    if not isinstance(revision_log, list):
        add_issue(issues, "revision", "error", "revision_log.json 必须是数组")
    else:
        for entry in revision_log:
            if not isinstance(entry, dict):
                add_issue(issues, "revision", "error", "revision_log 存在非对象记录")
                continue
            if entry.get("priority") in {"P0", "P1"} and entry.get("status") != "closed":
                add_issue(issues, "revision", "error", f"{entry.get('issue_id', '未编号')} 尚未关闭")

    errors = [issue for issue in issues if issue["level"] == "error"]
    warnings = [issue for issue in issues if issue["level"] == "warning"]
    return {
        "status": "pass" if not errors else "fail",
        "error_count": len(errors),
        "warning_count": len(warnings),
        "issues": issues,
        "summary": {
            "fact_count": len(facts),
            "question_count": len(questions),
            "content_item_count": len(items),
            "semantic_total": total,
        },
    }


def set_run_font(
    run: Any,
    name: str,
    size: float,
    bold: bool = False,
    latin_name: str | None = None,
) -> None:
    run.font.name = latin_name or name
    run.font.size = Pt(size)
    run.bold = bold
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin_name or name)
    fonts.set(qn("w:hAnsi"), latin_name or name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), latin_name or name)


def add_mixed_font_text(
    paragraph: Any,
    text: str,
    cjk_name: str,
    latin_name: str,
    size: float,
    bold: bool = False,
) -> None:
    """Split CJK and ASCII into separate runs for stable Word/LibreOffice fallback."""
    for segment in re.findall(r"[^\x00-\x7f]+|[\x00-\x7f]+", text):
        if segment.isascii():
            set_run_font(
                paragraph.add_run(segment), cjk_name, size, bold, latin_name
            )
        else:
            set_run_font(paragraph.add_run(segment), cjk_name, size, bold)


def add_lead_emphasis(paragraph: Any, text: str) -> None:
    """Bold only the first sentence so the decision remains easy to scan."""
    split_at = text.find("。")
    if split_at < 0:
        set_run_font(paragraph.add_run(text), FONT_NAME, BODY_PT, True)
        return
    lead = text[: split_at + 1]
    remainder = text[split_at + 1 :]
    set_run_font(paragraph.add_run(lead), FONT_NAME, BODY_PT, True)
    if remainder:
        set_run_font(paragraph.add_run(remainder), FONT_NAME, BODY_PT, False)


def render_docx(root: Path, output: Path, allow_failed_audit: bool = False) -> None:
    report = audit_artifacts(root)
    write_json(root / "qa_audit.json", report)
    if report["status"] != "pass" and not allow_failed_audit:
        raise SystemExit(
            f"QA 硬门禁未通过（{report['error_count']} 项）；详见 {root / 'qa_audit.json'}"
        )
    content = read_json(root / "qa_content.json")
    meta = content.get("meta", {})
    deta_profile = meta.get("format_profile") == "deta_qa_pdf"
    mode = narrative_mode({"profile": {"narrative_mode": meta.get("narrative_mode")}}, content)
    company = str(meta.get("company") or "目标公司")
    title = str(meta.get("title") or f"{company}项目 Q&A")

    cjk_font = DETA_CJK_FONT_NAME if deta_profile else FONT_NAME
    latin_font = DETA_LATIN_FONT_NAME if deta_profile else None
    title_pt = DETA_TITLE_PT if deta_profile else TITLE_PT
    body_pt = DETA_BODY_PT if deta_profile else BODY_PT
    line_pt = DETA_LINE_PT if deta_profile else 17
    margin_tb = DETA_MARGIN_TOP_BOTTOM_CM if deta_profile else MARGIN_TOP_BOTTOM_CM
    margin_lr = DETA_MARGIN_LEFT_RIGHT_CM if deta_profile else MARGIN_LEFT_RIGHT_CM

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(PAGE_WIDTH_CM)
    section.page_height = Cm(PAGE_HEIGHT_CM)
    section.top_margin = Cm(margin_tb)
    section.bottom_margin = Cm(margin_tb)
    section.left_margin = Cm(margin_lr)
    section.right_margin = Cm(margin_lr)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)
    document.core_properties.title = title
    document.core_properties.subject = "投资委员会答辩 Q&A"

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_before = Pt(0 if deta_profile else 12)
    summary_items = [] if deta_profile or mode == "dual_use" else (meta.get("investment_summary") or [])
    title_paragraph.paragraph_format.space_after = Pt(
        16 if deta_profile else (10 if summary_items or mode == "dual_use" else 18)
    )
    if deta_profile:
        title_paragraph.paragraph_format.line_spacing = Pt(20)
    title_paragraph.paragraph_format.keep_with_next = True
    if deta_profile:
        add_mixed_font_text(
            title_paragraph, title, cjk_font, latin_font, title_pt, True
        )
    else:
        set_run_font(title_paragraph.add_run(title), cjk_font, title_pt, True)

    if summary_items:
        summary_heading = document.add_paragraph()
        summary_heading.paragraph_format.space_before = Pt(0)
        summary_heading.paragraph_format.space_after = Pt(4)
        summary_heading.paragraph_format.line_spacing = Pt(17)
        summary_heading.paragraph_format.keep_with_next = True
        set_run_font(summary_heading.add_run("投资建议摘要"), FONT_NAME, BODY_PT, True)

        for summary_item in summary_items:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = Pt(17)
            paragraph.paragraph_format.keep_together = True
            if isinstance(summary_item, dict):
                label = str(summary_item.get("label") or "").strip()
                text = str(summary_item.get("text") or "").strip()
                if label:
                    set_run_font(paragraph.add_run(f"{label}："), FONT_NAME, BODY_PT, True)
                if text:
                    set_run_font(paragraph.add_run(text), FONT_NAME, BODY_PT, False)
            else:
                set_run_font(paragraph.add_run(str(summary_item)), FONT_NAME, BODY_PT, False)

    overview = content.get("overview") or {}
    if mode == "dual_use" and isinstance(overview, dict):
        overview_paragraphs = [
            str(overview.get("market_question") or "").strip(),
            *[str(value).strip() for value in (overview.get("players") or [])],
            str(overview.get("project_position") or "").strip(),
        ]
        for overview_index, overview_text in enumerate(overview_paragraphs):
            if not overview_text:
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7.8 if deta_profile else 5)
            paragraph.paragraph_format.line_spacing = Pt(line_pt)
            paragraph.paragraph_format.keep_together = True
            add_mixed_font_text(
                paragraph, overview_text, cjk_font, latin_font, body_pt, False
            )

    for index, item in enumerate(content.get("items", []), 1):
        qid = str(item.get("id") or f"Q{index}")
        question = str(item.get("question", "")).strip()
        q_paragraph = document.add_paragraph()
        q_paragraph.paragraph_format.space_before = Pt(
            (20.4 if mode == "dual_use" else 0)
            if index == 1
            else (20.4 if deta_profile else 10)
        )
        q_paragraph.paragraph_format.space_after = Pt(7.8 if deta_profile else 4)
        q_paragraph.paragraph_format.line_spacing = Pt(line_pt)
        q_paragraph.paragraph_format.keep_with_next = True
        q_paragraph.paragraph_format.keep_together = True
        if deta_profile:
            add_mixed_font_text(
                q_paragraph,
                f"{qid}：{question}",
                cjk_font,
                latin_font,
                body_pt,
                True,
            )
        else:
            set_run_font(
                q_paragraph.add_run(f"{qid}：{question}"),
                cjk_font,
                body_pt,
                True,
            )

        for paragraph_index, text in enumerate(item.get("answer_paragraphs", [])):
            text = str(text).strip()
            if not text:
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(7.8 if deta_profile else 5)
            paragraph.paragraph_format.line_spacing = Pt(line_pt)
            if not deta_profile:
                paragraph.paragraph_format.keep_together = True
            if deta_profile:
                visible_text = f"回答：{text}" if paragraph_index == 0 else text
                add_mixed_font_text(
                    paragraph,
                    visible_text,
                    cjk_font,
                    latin_font,
                    body_pt,
                    False,
                )
            elif paragraph_index == 0:
                add_lead_emphasis(paragraph, text)
            else:
                set_run_font(paragraph.add_run(text), FONT_NAME, BODY_PT, False)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"rendered {output}")


def inspect_docx(docx_path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(docx_path) as bundle:
        names = set(bundle.namelist())
        document_root = ET.fromstring(bundle.read("word/document.xml"))
        paragraphs = document_root.findall(f".//{W}p")
        text_by_paragraph = [
            "".join(node.text or "" for node in paragraph.findall(f".//{W}t"))
            for paragraph in paragraphs
        ]
        sections = document_root.findall(f".//{W}sectPr")
        tables = document_root.findall(f".//{W}tbl")
        heading_count = 0
        toc_fields = 0
        page_fields = 0
        font_names: set[str] = set()
        font_sizes: set[float] = set()
        paragraph_layouts: list[dict[str, Any]] = []
        for paragraph, paragraph_text in zip(paragraphs, text_by_paragraph):
            style = paragraph.find(f"./{W}pPr/{W}pStyle")
            style_value = style.get(f"{W}val", "") if style is not None else ""
            if style_value.lower().startswith("heading"):
                heading_count += 1
            ppr = paragraph.find(f"./{W}pPr")
            spacing = ppr.find(f"./{W}spacing") if ppr is not None else None
            alignment = ppr.find(f"./{W}jc") if ppr is not None else None
            indent = ppr.find(f"./{W}ind") if ppr is not None else None
            run_bold: list[bool] = []
            run_sizes: set[float] = set()
            for run in paragraph.findall(f"./{W}r"):
                run_text = "".join(node.text or "" for node in run.findall(f".//{W}t"))
                if not run_text:
                    continue
                rpr = run.find(f"./{W}rPr")
                bold = rpr.find(f"./{W}b") if rpr is not None else None
                bold_value = bold.get(f"{W}val", "1") if bold is not None else "0"
                run_bold.append(str(bold_value).lower() not in {"0", "false", "off"})
                size = rpr.find(f"./{W}sz") if rpr is not None else None
                if size is not None and size.get(f"{W}val"):
                    run_sizes.add(int(size.get(f"{W}val")) / 2)
            paragraph_layouts.append({
                "text": paragraph_text,
                "alignment": alignment.get(f"{W}val") if alignment is not None else "",
                "space_before": int(spacing.get(f"{W}before", "0")) if spacing is not None else 0,
                "space_after": int(spacing.get(f"{W}after", "0")) if spacing is not None else 0,
                "line": int(spacing.get(f"{W}line", "0")) if spacing is not None else 0,
                "line_rule": spacing.get(f"{W}lineRule", "") if spacing is not None else "",
                "first_line_indent": int(indent.get(f"{W}firstLine", "0")) if indent is not None else 0,
                "left_indent": int(indent.get(f"{W}left", "0")) if indent is not None else 0,
                "keep_next": ppr.find(f"./{W}keepNext") is not None if ppr is not None else False,
                "all_bold": bool(run_bold) and all(run_bold),
                "any_bold": any(run_bold),
                "font_sizes": sorted(run_sizes),
            })
        for instruction in document_root.findall(f".//{W}instrText"):
            value = instruction.text or ""
            toc_fields += int("TOC" in value)
            page_fields += int("PAGE" in value)
        for run_properties in document_root.findall(f".//{W}rPr"):
            fonts = run_properties.find(f"./{W}rFonts")
            if fonts is not None:
                for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                    value = fonts.get(f"{W}{key}")
                    if value:
                        font_names.add(value)
            size = run_properties.find(f"./{W}sz")
            if size is not None and size.get(f"{W}val"):
                font_sizes.add(int(size.get(f"{W}val")) / 2)
        page_sizes: list[dict[str, int]] = []
        margins: list[dict[str, int]] = []
        for section in sections:
            page = section.find(f"./{W}pgSz")
            margin = section.find(f"./{W}pgMar")
            page_sizes.append({
                "width": int(page.get(f"{W}w", "0")) if page is not None else 0,
                "height": int(page.get(f"{W}h", "0")) if page is not None else 0,
            })
            margins.append({
                key: int(margin.get(f"{W}{key}", "0")) if margin is not None else 0
                for key in ("top", "bottom", "left", "right")
            })
        return {
            "section_count": len(sections),
            "table_count": len(tables),
            "paragraph_count": len(paragraphs),
            "question_count": sum(bool(re.match(r"^Q\d+：", text)) for text in text_by_paragraph),
            "heading_count": heading_count,
            "toc_fields": toc_fields,
            "page_fields": page_fields,
            "header_parts": sorted(name for name in names if name.startswith("word/header")),
            "footer_parts": sorted(name for name in names if name.startswith("word/footer")),
            "page_sizes": page_sizes,
            "margins": margins,
            "font_names": sorted(font_names),
            "font_sizes": sorted(font_sizes),
            "title": text_by_paragraph[0] if text_by_paragraph else "",
            "paragraph_layouts": paragraph_layouts,
        }


def verify_docx(root: Path, docx_path: Path) -> dict[str, Any]:
    artifact_audit = audit_artifacts(root)
    package = inspect_docx(docx_path)
    content = read_json(root / "qa_content.json")
    deta_profile = content.get("meta", {}).get("format_profile") == "deta_qa_pdf"
    issues: list[str] = []
    if package["section_count"] != 1:
        issues.append(f"Q&A 必须为单一连续节，当前 {package['section_count']} 节")
    if package["table_count"]:
        issues.append("Q&A 答辩体不得包含表格")
    if package["heading_count"] or package["toc_fields"]:
        issues.append("Q&A 不得使用标题层级或目录域")
    if package["page_fields"] or package["header_parts"] or package["footer_parts"]:
        issues.append("Q&A 不得包含页眉、页脚或页码域")
    if not MIN_QUESTIONS <= package["question_count"] <= MAX_QUESTIONS:
        issues.append(f"DOCX 问题数为 {package['question_count']}，必须为 6—9 个")
    for page in package["page_sizes"]:
        if abs(page["width"] - 11906) > 12 or abs(page["height"] - 16838) > 12:
            issues.append(f"页面不是 A4：{page['width']}×{page['height']} DXA")
    for margin in package["margins"]:
        expected = (
            {"top": 1440, "bottom": 1440, "left": 1800, "right": 1800}
            if deta_profile
            else {"top": 879, "bottom": 879, "left": 1729, "right": 1729}
        )
        for key, target in expected.items():
            if abs(margin[key] - target) > 24:
                issues.append(f"{key} 页边距异常：{margin[key]} DXA")
    allowed_fonts = (
        {DETA_CJK_FONT_NAME, DETA_LATIN_FONT_NAME}
        if deta_profile
        else {FONT_NAME}
    )
    foreign_fonts = [name for name in package["font_names"] if name not in allowed_fonts]
    if foreign_fonts:
        issues.append("出现非标准字体：" + "、".join(foreign_fonts))
    required_sizes = (
        {DETA_BODY_PT, DETA_TITLE_PT} if deta_profile else {BODY_PT, TITLE_PT}
    )
    if not required_sizes.issubset(set(package["font_sizes"])):
        issues.append(f"字号不完整：需要 {sorted(required_sizes)}，实际 {package['font_sizes']}")
    if deta_profile:
        layouts = [entry for entry in package["paragraph_layouts"] if entry["text"].strip()]
        if not layouts:
            issues.append("德塔版式缺少可见段落")
        else:
            title = layouts[0]
            if title["alignment"] != "center":
                issues.append("德塔标题必须居中")
            if title["font_sizes"] != [DETA_TITLE_PT] or not title["all_bold"]:
                issues.append("德塔标题必须为16pt加粗")
            if (title["space_before"], title["space_after"], title["line"], title["line_rule"]) != (0, 320, 400, "exact"):
                issues.append("德塔标题间距必须为段前0、段后16pt、固定20pt行距")
            question_positions = [
                index for index, entry in enumerate(layouts)
                if re.match(r"^Q\d+：", entry["text"])
            ]
            if not question_positions or question_positions[0] != 1:
                issues.append("德塔标题后必须直接进入Q1")
            for number, position in enumerate(question_positions, 1):
                question = layouts[position]
                expected_before = 0 if number == 1 else 408
                if question["font_sizes"] != [DETA_BODY_PT] or not question["all_bold"]:
                    issues.append(f"Q{number} 必须为11pt加粗")
                if (
                    question["space_before"], question["space_after"],
                    question["line"], question["line_rule"]
                ) != (expected_before, 156, 288, "exact"):
                    issues.append(f"Q{number} 的段前、段后或固定行距不符合德塔模板")
                if not question["keep_next"]:
                    issues.append(f"Q{number} 未与首段回答保持同页")
                if position + 1 >= len(layouts) or not layouts[position + 1]["text"].startswith("回答："):
                    issues.append(f"Q{number} 首段回答缺少“回答：”")
            question_set = set(question_positions)
            for position, answer in enumerate(layouts[1:], 1):
                if position in question_set:
                    continue
                if answer["font_sizes"] != [DETA_BODY_PT] or answer["any_bold"]:
                    issues.append("德塔回答必须为11pt常规字重")
                    break
                if (
                    answer["space_after"], answer["line"], answer["line_rule"],
                    answer["first_line_indent"], answer["left_indent"]
                ) != (156, 288, "exact", 0, 0):
                    issues.append("德塔回答必须段后7.8pt、固定14.4pt行距且不缩进")
                    break
    status = "pass" if artifact_audit["status"] == "pass" and not issues else "fail"
    return {
        "status": status,
        "artifact_audit": artifact_audit,
        "docx_package": package,
        "render_issues": issues,
        "visual_status": "pending_manual_page_review",
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="独立投资 Q&A 报告确定性处理程序")
    subparsers = parser.add_subparsers(dest="command", required=True)

    prepare = subparsers.add_parser("prepare", help="解包材料并生成证据包和阶段文件")
    prepare.add_argument("input", type=Path)
    prepare.add_argument("--workdir", type=Path, required=True)

    audit = subparsers.add_parser("audit", help="执行事实、规划、正文和评分硬门禁")
    audit.add_argument("--artifacts", type=Path, required=True)
    audit.add_argument("--out", type=Path)

    render = subparsers.add_parser("render", help="生成单节 A4 黑白答辩体 DOCX")
    render.add_argument("--artifacts", type=Path, required=True)
    render.add_argument("--output", type=Path, required=True)
    render.add_argument("--allow-failed-audit", action="store_true")

    verify = subparsers.add_parser("verify", help="核验阶段文件和 DOCX 包结构")
    verify.add_argument("--artifacts", type=Path, required=True)
    verify.add_argument("--docx", type=Path, required=True)
    verify.add_argument("--out", type=Path, required=True)
    return parser


def main(argv: list[str] | None = None) -> None:
    args = build_parser().parse_args(argv)
    if args.command == "prepare":
        prepare_input(args.input.resolve(), args.workdir.resolve())
        return
    if args.command == "audit":
        report = audit_artifacts(args.artifacts.resolve())
        output = args.out or args.artifacts / "qa_audit.json"
        write_json(output, report)
        print(f"{report['status']}: {report['error_count']} errors, "
              f"{report['warning_count']} warnings -> {output}")
        if report["status"] != "pass":
            raise SystemExit(2)
        return
    if args.command == "render":
        render_docx(args.artifacts.resolve(), args.output.resolve(), args.allow_failed_audit)
        return
    if args.command == "verify":
        if not args.docx.is_file():
            raise SystemExit(f"DOCX 不存在：{args.docx}")
        report = verify_docx(args.artifacts.resolve(), args.docx.resolve())
        write_json(args.out.resolve(), report)
        print(f"{report['status']} -> {args.out}")
        if report["status"] != "pass":
            raise SystemExit(2)


if __name__ == "__main__":
    main()
